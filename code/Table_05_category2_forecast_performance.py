from pathlib import Path
import argparse
import warnings
import numpy as np
import pandas as pd
from functools import partial
from tqdm import tqdm as _tqdm
tqdm = partial(_tqdm, bar_format="{bar} {percentage:3.0f}%", ncols=50)
from arch.bootstrap import MCS
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, Border, Side

ROOT = Path(__file__).resolve().parents[1]
RV_FILE = ROOT / "data/raw/RV.csv"
RETURN_FILE = ROOT / "data/raw/Return.csv"
LOG_CENTRALITY_FILE = ROOT / "data/processed/centrality/P_log.csv"
QLIKE_DROP_DATES_FILE = ROOT / "data/processed/evaluation/qlike_intermediate.csv"
OUTPUT_FILE = ROOT / "results/generated/Table_05_category2_forecast_performance.xlsx"
THRESH = 1e-3
SCALE_IF_MOST_GT_THRESH = 1e-4
EPS = 1e-12
RATIO_MIN = 1e-6
RATIO_MAX = 1e6
QLIKE_CAP = 100.0
POST_SCALE_THRESHOLD = 1e4
POST_SCALE_FACTOR = 1e-8


def run_har():
    model_outputs = {}
    import pandas as pd
    import numpy as np
    from functools import partial
    from tqdm import tqdm as _tqdm
    tqdm = partial(_tqdm, bar_format="{bar} {percentage:3.0f}%", ncols=50)
    import statsmodels.formula.api as smf
    data = pd.read_csv(RV_FILE)
    num_cols = data.columns.drop('date')
    data[num_cols] = data[num_cols].apply(pd.to_numeric, errors='coerce')
    data = data.dropna()
    data['date'] = pd.to_datetime(data['date'], utc=True)
    data['date'] = data['date'].dt.tz_localize(None)
    data['date'] = data['date'].dt.normalize()
    data = data.set_index('date')
    for c in data.columns:
        data_RV = pd.DataFrame(data[c])
        data_RV.columns = ['RV_now']
        data_RV['RV_D'] = data_RV['RV_now'].shift(1)
        data_RV['RV_W'] = data_RV['RV_D'].rolling(5).mean()
        data_RV['RV_M'] = data_RV['RV_D'].rolling(22).mean()
        data_RV['RV_MA22'] = data_RV['RV_now'].rolling(window=22, center=False).mean().shift(-21)
        data_RV = data_RV.dropna().loc['2014-07-04':]
        time_start = '2018/05/02'
        time_stop = '2022-02-25'
        data_test = data_RV.loc[time_start:time_stop]
        data_predict = data_test['RV_MA22']
        train_len = len(data_RV.loc[:time_start]) - 1
        result = []
        for i in tqdm(range(len(data_test)), desc=f'HAR {str(c).strip().lstrip(".")}', leave=False):
            data_rolling = data_RV.iloc[i:i + train_len]
            params = smf.ols('RV_MA22 ~ RV_D + RV_W + RV_M', data_rolling).fit().params
            result.append(
                params.iloc[0]
                + params.iloc[1] * data_test['RV_D'].iloc[i]
                + params.iloc[2] * data_test['RV_W'].iloc[i]
                + params.iloc[3] * data_test['RV_M'].iloc[i]
            )
        result = pd.Series(result, index=data_predict.index)
        frame = pd.concat([data_predict, result], axis=1)
        frame.columns = ['RV_True', 'RV_Predict']
        frame.index = pd.to_datetime(data_test.index)
        frame['MSE'] = (frame['RV_True'] - frame['RV_Predict']) ** 2
        frame['MAE'] = np.abs(frame['RV_True'] - frame['RV_Predict'])
        frame['Qlike'] = frame['RV_True'] / frame['RV_Predict'] + np.log(frame['RV_Predict'])
        model_outputs[str(c).strip().lstrip('.')] = frame.copy()
    return model_outputs


def run_loghar():
    model_outputs = {}
    import math
    import scipy as sc
    import time
    import os
    from scipy.stats import norm
    import pandas as pd
    import numpy as np
    from functools import partial
    from tqdm import tqdm as _tqdm
    tqdm = partial(_tqdm, bar_format="{bar} {percentage:3.0f}%", ncols=50)
    import statistics
    import statsmodels.formula.api as smf
    EPS_LOG = 1e-12
    EXP_CLIP_MAX = 50
    H_LIST = [22]
    SAMPLE_START = '2014-07-04'
    time_start = '2018/05/02'
    time_stop = '2022-02-25'

    def safe_log(s: pd.Series, eps: float=EPS_LOG) -> pd.Series:
        return np.log(s.clip(lower=eps))

    def safe_exp(x, clip_max: float=EXP_CLIP_MAX):
        return np.exp(np.clip(x, -clip_max, clip_max))

    def future_mean_level(rv_now: pd.Series, H: int) -> pd.Series:
        if H == 1:
            return rv_now
        ma = rv_now.rolling(window=H, center=False).mean()
        return ma.shift(-(H - 1))
    data = pd.read_csv(RV_FILE)
    num_cols = [c for c in data.columns if c != 'date']
    data[num_cols] = data[num_cols].apply(pd.to_numeric, errors='coerce')
    data = data.dropna()
    data['date'] = pd.to_datetime(data['date'], utc=True)
    data['date'] = data['date'].dt.tz_localize(None)
    data['date'] = data['date'].dt.normalize()
    data = data.set_index('date')
    for H in H_LIST:
        error = []
        for c in data.columns:
            data_RV = data[[c]].copy()
            data_RV.columns = ['RV_now']
            data_RV['RV_D_lvl'] = data_RV['RV_now'].shift(1)
            data_RV['RV_W_lvl'] = data_RV['RV_D_lvl'].rolling(5).mean()
            data_RV['RV_M_lvl'] = data_RV['RV_D_lvl'].rolling(22).mean()
            data_RV['RV_D'] = safe_log(data_RV['RV_D_lvl'])
            data_RV['RV_W'] = safe_log(data_RV['RV_W_lvl'])
            data_RV['RV_M'] = safe_log(data_RV['RV_M_lvl'])
            data_RV[f'RV_MA{H}_lvl'] = future_mean_level(data_RV['RV_now'], H)
            data_RV[f'RV_MA{H}'] = safe_log(data_RV[f'RV_MA{H}_lvl'])
            data_RV = data_RV.dropna()
            data_RV = data_RV.loc[SAMPLE_START:]
            data_test = data_RV.loc[time_start:time_stop].copy()
            if len(data_test) == 0:
                error.append([np.nan, np.nan, np.nan])
                tmp = pd.DataFrame(index=pd.DatetimeIndex([], name='date'))
                tmp['RV_True'] = np.nan
                tmp['RV_Predict'] = np.nan
                tmp['MSE'] = np.nan
                tmp['MAE'] = np.nan
                tmp['Qlike'] = np.nan
                model_outputs[str(c).strip().lstrip('.')] = tmp.copy()
                continue
            y_true_log = data_test[f'RV_MA{H}'].copy()
            result_log = []
            n_test = len(data_RV.loc[time_start:time_stop])
            train_len = len(data_RV.loc[:time_start]) - 1
            for i in tqdm(range(n_test), desc=f'LogHAR {c} H{H}', leave=False):
                data_rolling = data_RV.iloc[i:i + train_len].copy()
                mod = smf.ols(f'RV_MA{H} ~ RV_D + RV_W + RV_M', data_rolling)
                res = mod.fit()
                params_new = res.params
                r_predict_log = params_new.iloc[0] + params_new.iloc[1] * data_test['RV_D'].iloc[i] + params_new.iloc[2] * data_test['RV_W'].iloc[i] + params_new.iloc[3] * data_test['RV_M'].iloc[i]
                result_log.append(r_predict_log)
            result_log = pd.Series(result_log, index=y_true_log.index)
            y_true = pd.Series(safe_exp(y_true_log.values), index=y_true_log.index)
            y_pred = pd.Series(safe_exp(result_log.values), index=result_log.index)
            SUM = pd.concat([y_true, y_pred], axis=1)
            SUM.columns = ['RV_True', 'RV_Predict']
            SUM_final = SUM.reset_index(drop=True)
            SUM_final.index = pd.to_datetime(data_test.index)
            SUM_final['MSE'] = (SUM_final['RV_True'] - SUM_final['RV_Predict']) ** 2
            SUM_final['MAE'] = np.abs(SUM_final['RV_True'] - SUM_final['RV_Predict'])
            pred_safe = SUM_final['RV_Predict'].clip(lower=EPS_LOG)
            true_safe = SUM_final['RV_True'].clip(lower=EPS_LOG)
            SUM_final['Qlike'] = true_safe / pred_safe + np.log(pred_safe)
            MSE = float(np.mean(SUM_final['MSE']))
            MAE = float(np.mean(SUM_final['MAE']))
            Theil = float(np.sqrt(np.mean((SUM_final['RV_True'] - SUM_final['RV_Predict']) ** 2)) / (np.sqrt(np.mean(SUM_final['RV_Predict'] ** 2)) + np.sqrt(np.mean(SUM_final['RV_True'] ** 2))))
            error.append([MSE, MAE, Theil])
            model_outputs[str(c).strip().lstrip('.')] = SUM_final.copy()
        ERROR = pd.DataFrame(error, columns=['MSE', 'MAE', 'Theil'], index=data.columns).T
    return model_outputs

def run_loghar_ci():
    model_outputs = {}
    import os
    import numpy as np
    import pandas as pd
    import statsmodels.formula.api as smf
    from functools import partial
    from tqdm import tqdm as _tqdm
    tqdm = partial(_tqdm, bar_format="{bar} {percentage:3.0f}%", ncols=50)
    RET_PATH = RETURN_FILE
    RV_PATH = RV_FILE
    P_STORE_PATH = LOG_CENTRALITY_FILE
    fit_end_date = pd.to_datetime('2018-05-01')
    pred_start = pd.to_datetime('2018-05-02')
    pred_end = pd.to_datetime('2022-02-25')
    H_LIST = [22]
    FIT_START = 1000
    FIT_STEP = 100
    FIT_MAX = 1
    L_fit_list = [FIT_START + FIT_STEP * k for k in range(FIT_MAX)]
    NET_MIN = 700
    NET_MAX = 700
    NET_STEP = 50
    L_net_list = list(range(NET_MIN, NET_MAX + 1, NET_STEP))
    EPS_LOG = 1e-12
    EXP_CLIP_MAX = 50

    def _norm_date(x):
        return pd.to_datetime(x).tz_localize(None).normalize()

    def safe_log(s: pd.Series, eps: float=EPS_LOG) -> pd.Series:
        return np.log(s.clip(lower=eps))

    def safe_exp(x: pd.Series | np.ndarray, clip_max: float=EXP_CLIP_MAX):
        return np.exp(np.clip(x, -clip_max, clip_max))

    def make_future_mean(rv_now: pd.Series, H: int) -> pd.Series:
        if H == 1:
            return rv_now
        ma = rv_now.rolling(window=H, center=False).mean()
        return ma.shift(-(H - 1))

    def make_future_mean_log(rv_now: pd.Series, H: int) -> pd.Series:
        future_mean = make_future_mean(rv_now, H)
        return safe_log(future_mean)

    def read_returns(path):
        df = pd.read_csv(path)
        df = df.loc[:, ~df.columns.str.startswith('Unnamed')]
        df['date'] = pd.to_datetime(df['date'].astype(str).str.strip().str[:10], errors='coerce')
        df['date'] = df['date'].dt.tz_localize(None).dt.normalize()
        df = df.set_index('date').dropna(how='all')
        return df

    def read_rv_no_scale(path):
        data = pd.read_csv(path)
        if 'date' not in data.columns:
            raise RuntimeError("RV file must contain a 'date' column.")
        num_cols = [c for c in data.columns if c != 'date']
        data[num_cols] = data[num_cols].apply(pd.to_numeric, errors='coerce')
        data = data.dropna(how='all')
        data['date'] = pd.to_datetime(data['date'], utc=True, errors='coerce')
        data = data.dropna(subset=['date'])
        data['date'] = data['date'].dt.tz_localize(None)
        data['date'] = data['date'].dt.normalize()
        data = data.set_index('date')
        data = data.dropna(axis=1, how='all')
        return data
    df_ret = read_returns(RET_PATH)
    data_rv = read_rv_no_scale(RV_PATH)
    markets = [c for c in data_rv.columns if c in df_ret.columns]
    P_all = pd.read_csv(P_STORE_PATH)
    need_cols = ['date', 'market', 'ci_start_date', 'ci_end_date', 'L_net', 'P_corr', 'P_kendall', 'P_partial', 'P_tail']
    miss = [c for c in need_cols if c not in P_all.columns]
    if miss:
        raise RuntimeError(f'Model execution message{miss}Model execution message{list(P_all.columns)}')
    P_all['date'] = pd.to_datetime(P_all['date'], errors='coerce').dt.tz_localize(None).dt.normalize()
    P_all['ci_start_date'] = P_all['ci_start_date'].apply(_norm_date)
    P_all['ci_end_date'] = P_all['ci_end_date'].apply(_norm_date)
    P_all['L_net'] = pd.to_numeric(P_all['L_net'], errors='coerce').astype('Int64')
    P_all = P_all.dropna(subset=['date', 'market', 'ci_start_date', 'ci_end_date', 'L_net']).copy()
    P_cache = {}

    def load_P_from_csv(L_net: int, ci_start_date, ci_end_date):
        key = (int(L_net), _norm_date(ci_start_date), _norm_date(ci_end_date))
        if key in P_cache:
            return P_cache[key]
        (ln, cs, ce) = key
        sub = P_all[(P_all['L_net'] == ln) & (P_all['ci_start_date'] == cs) & (P_all['ci_end_date'] == ce)]
        if sub.empty:
            P_cache[key] = (None, None, None, None)
            return P_cache[key]

        def _pivot(col):
            df = sub.pivot_table(index='date', columns='market', values=col, aggfunc='mean')
            return df.sort_index()
        P_corr = _pivot('P_corr')
        P_kendall = _pivot('P_kendall')
        P_partial = _pivot('P_partial')
        P_tail = _pivot('P_tail')
        P_cache[key] = (P_corr, P_kendall, P_partial, P_tail)
        return P_cache[key]

    def forecast_one_market_oneH_logharci(data_rv_all: pd.DataFrame, P_corr: pd.DataFrame, P_kendall: pd.DataFrame, P_partial: pd.DataFrame, P_tail: pd.DataFrame, market: str, fit_start_date: pd.Timestamp, H: int):
        df = data_rv_all[market].to_frame('RV_now')
        df['RV_D_lvl'] = df['RV_now'].shift(1)
        df['RV_W_lvl'] = df['RV_D_lvl'].rolling(5).mean()
        df['RV_M_lvl'] = df['RV_D_lvl'].rolling(22).mean()
        df['RV_D'] = safe_log(df['RV_D_lvl'])
        df['RV_W'] = safe_log(df['RV_W_lvl'])
        df['RV_M'] = safe_log(df['RV_M_lvl'])
        df['Y'] = make_future_mean_log(df['RV_now'], H)
        ci = pd.DataFrame({'CI_corr': P_corr[market], 'CI_kendall': P_kendall[market], 'CI_partial': P_partial[market], 'CI_tail': P_tail[market]}).reindex(df.index)
        df = df.join(ci, how='inner').dropna()
        df['CI_corr_D'] = df['CI_corr'].shift(1)
        df['CI_kendall_D'] = df['CI_kendall'].shift(1)
        df['CI_partial_D'] = df['CI_partial'].shift(1)
        df['CI_tail_D'] = df['CI_tail'].shift(1)
        df = df.dropna()
        df = df.loc[fit_start_date:]
        data_test = df.loc[pred_start:pred_end].copy()
        if len(data_test) == 0:
            return (None, np.nan, np.nan, np.nan)
        window_len = len(df.loc[:pred_start]) - 1
        if window_len <= 0:
            return (None, np.nan, np.nan, np.nan)
        y_true_log = data_test['Y'].astype(float).copy()
        idx = y_true_log.index
        preds_log = []
        for i in tqdm(range(len(idx)), desc=f'LogHAR-CI {market} H{H}', leave=False):
            train = df.iloc[i:i + window_len].copy()
            train['_Y_'] = train['Y'].astype(float).values
            res = smf.ols('_Y_ ~ RV_D + RV_W + RV_M + CI_corr_D + CI_kendall_D + CI_partial_D + CI_tail_D', train).fit()
            p = res.params
            pred_log = p.iloc[0] + p.iloc[1] * data_test['RV_D'].iloc[i] + p.iloc[2] * data_test['RV_W'].iloc[i] + p.iloc[3] * data_test['RV_M'].iloc[i] + p.iloc[4] * data_test['CI_corr_D'].iloc[i] + p.iloc[5] * data_test['CI_kendall_D'].iloc[i] + p.iloc[6] * data_test['CI_partial_D'].iloc[i] + p.iloc[7] * data_test['CI_tail_D'].iloc[i]
            preds_log.append(float(pred_log))
        y_pred_log = pd.Series(preds_log, index=idx)
        y_true = pd.Series(safe_exp(y_true_log.values), index=idx)
        y_pred = pd.Series(safe_exp(y_pred_log.values), index=idx)
        mcs = pd.concat([y_true.rename('RV_True'), y_pred.rename('RV_Predict')], axis=1)
        mcs['MSE'] = (mcs['RV_True'] - mcs['RV_Predict']) ** 2
        mcs['MAE'] = np.abs(mcs['RV_True'] - mcs['RV_Predict'])
        mse = float(np.nanmean(mcs['MSE'].values))
        mae = float(np.nanmean(mcs['MAE'].values))
        denom = np.sqrt(np.nanmean(mcs['RV_Predict'].values ** 2)) + np.sqrt(np.nanmean(mcs['RV_True'].values ** 2))
        theil = float(np.sqrt(np.nanmean((mcs['RV_True'].values - mcs['RV_Predict'].values) ** 2)) / denom) if denom != 0 else np.nan
        return (mcs, mse, mae, theil)
    L_fit = L_fit_list[0]
    L_net = L_net_list[0]
    for H in H_LIST:
        error_by_market = {}
        for mkt in markets:
            if fit_end_date not in data_rv.index:
                error_by_market[mkt] = [np.nan, np.nan, np.nan]
                continue
            fit_end_pos = data_rv.index.get_loc(fit_end_date)
            fit_start_pos = fit_end_pos - (L_fit - 1)
            if fit_start_pos < 0:
                error_by_market[mkt] = [np.nan, np.nan, np.nan]
                continue
            fit_start_date = data_rv.index[fit_start_pos]
            ret_pos_end = df_ret.index.searchsorted(fit_start_date) - 1
            if ret_pos_end < 0:
                error_by_market[mkt] = [np.nan, np.nan, np.nan]
                continue
            net_end_date = df_ret.index[ret_pos_end]
            ci_start_date = net_end_date
            ci_end_date = pred_end
            (P_corr, P_kendall, P_partial, P_tail) = load_P_from_csv(L_net=L_net, ci_start_date=ci_start_date, ci_end_date=ci_end_date)
            if P_corr is None:
                error_by_market[mkt] = [np.nan, np.nan, np.nan]
                continue
            if mkt not in P_corr.columns:
                error_by_market[mkt] = [np.nan, np.nan, np.nan]
                continue
            (mcs, mse, mae, theil) = forecast_one_market_oneH_logharci(data_rv_all=data_rv, P_corr=P_corr, P_kendall=P_kendall, P_partial=P_partial, P_tail=P_tail, market=mkt, fit_start_date=fit_start_date, H=H)
            if mcs is None:
                tmp = pd.DataFrame(index=pd.DatetimeIndex([], name='date'))
                tmp['RV_True'] = np.nan
                tmp['RV_Predict'] = np.nan
                tmp['MSE'] = np.nan
                tmp['MAE'] = np.nan
                model_outputs[str(mkt).strip().lstrip('.')] = tmp.copy()
            else:
                model_outputs[str(mkt).strip().lstrip('.')] = mcs.copy()
            error_by_market[mkt] = [mse, mae, theil]
        ERROR = pd.DataFrame(error_by_market, index=['MSE', 'MAE', 'Theil'])
    return model_outputs

def run_hexpgl():
    model_outputs = {}
    import pandas as pd
    import numpy as np
    import statsmodels.formula.api as smf
    from functools import partial
    from tqdm import tqdm as _tqdm
    tqdm = partial(_tqdm, bar_format="{bar} {percentage:3.0f}%", ncols=50)
    import os

    def get_lambda(com):
        return np.log(1 + 1 / com)

    def calculate_exp_factor(series, com, window=500):
        lam = get_lambda(com)
        weights = np.exp(-np.arange(1, window + 1) * lam)
        weights /= weights.sum()
        return series.rolling(window).apply(lambda x: np.dot(x[::-1], weights), raw=True)
    data = pd.read_csv(RV_FILE)
    data['date'] = pd.to_datetime(data['date'])
    data = data.set_index('date')
    mean_rv_all = data.mean()
    global_signal = (data / mean_rv_all).mean(axis=1)
    exp_global_5 = calculate_exp_factor(global_signal, com=5)
    H_LIST = [22]
    com_list = [1, 5, 25, 125]
    time = '2018/05/02'
    time_stop = '2022-02-25'
    error_by_H = {H: [] for H in H_LIST}
    for c in data.columns:
        asset_mean = mean_rv_all[c]
        df_base = pd.DataFrame(index=data.index)
        df_base['RV_now'] = data[c]
        for j in com_list:
            df_base[f'ExpRV_{j}'] = calculate_exp_factor(df_base['RV_now'], com=j)
        df_base['ExpGlRV_5'] = exp_global_5 * asset_mean
        features = [f'ExpRV_{j}' for j in com_list] + ['ExpGlRV_5']
        for feat in features:
            df_base[f'{feat}_lag'] = df_base[feat].shift(1)
        reg_features = []
        for feat in features:
            col_name = f'diff_{feat}'
            df_base[col_name] = df_base[f'{feat}_lag'] - asset_mean
            reg_features.append(col_name)
        for H in H_LIST:
            df_model = df_base.copy()
            target_col = f'RV_MA{H}'
            if H == 1:
                df_model[target_col] = df_model['RV_now']
            else:
                df_model[target_col] = df_model['RV_now'].rolling(window=H).mean()
                df_model[target_col] = df_model[target_col].shift(-H)
            df_model['target_diff'] = df_model[target_col] - asset_mean
            df_model = df_model.dropna()
            df_model = df_model.loc['2014-07-04':]
            data_test = df_model.loc[time:time_stop]
            data_predict = df_model.loc[time:time_stop][target_col]
            if len(data_test) == 0:
                error_by_H[H].append([np.nan, np.nan, np.nan])
                continue
            result = []
            for i in tqdm(range(len(df_model.loc[time:time_stop])), desc=f'{c}-H{H}'):
                data_rolling = df_model.iloc[i:i + (len(df_model.loc[:time]) - 1)]
                formula = 'target_diff ~ ' + ' + '.join(reg_features)
                mod = smf.ols(formula, data_rolling)
                res = mod.fit()
                params_new = res.params
                current_row = data_test.iloc[i]
                pred_diff = params_new.iloc[0]
                for (k, feat) in enumerate(reg_features, start=1):
                    pred_diff += params_new.iloc[k] * current_row[feat]
                r_predict = pred_diff + asset_mean
                result.append(r_predict)
            result = pd.Series(result, index=data_predict.index)
            SUM = pd.concat([data_predict, result], axis=1)
            SUM.columns = ['RV_True', 'RV_Predict']
            SUM_final = SUM
            SUM_final = SUM_final.reset_index(drop=True)
            SUM_final.index = pd.to_datetime(data_test.index)
            SUM_final['MSE'] = (SUM_final['RV_True'] - SUM_final['RV_Predict']) ** 2
            SUM_final['MAE'] = np.abs(SUM_final['RV_True'] - SUM_final['RV_Predict'])
            eps = 1e-12
            pred_clip = np.maximum(SUM_final['RV_Predict'], eps)
            SUM_final['Qlike'] = SUM_final['RV_True'] / pred_clip + np.log(pred_clip)
            MSE = np.mean((SUM_final['RV_True'] - SUM_final['RV_Predict']) ** 2)
            MAE = np.mean(np.abs(SUM_final['RV_True'] - SUM_final['RV_Predict']))
            Theil = np.sqrt(np.mean((SUM_final['RV_True'] - SUM_final['RV_Predict']) ** 2)) / (np.sqrt(np.mean(SUM_final['RV_Predict'] ** 2)) + np.sqrt(np.mean(SUM_final['RV_True'] ** 2)))
            error_by_H[H].append([MSE, MAE, Theil])
            model_outputs[str(c).strip().lstrip('.')] = SUM_final.copy()
    for H in H_LIST:
        ERROR = pd.DataFrame(error_by_H[H], columns=['MSE', 'MAE', 'Theil'], index=data.columns).T
    return model_outputs

def run_mvf():
    model_outputs = {}
    import os
    import numpy as np
    import pandas as pd
    import statsmodels.formula.api as smf
    from functools import partial
    from tqdm import tqdm as _tqdm
    tqdm = partial(_tqdm, bar_format="{bar} {percentage:3.0f}%", ncols=50)
    H_LIST = [22]
    time_start = '2018/05/02'
    time_stop = '2022-02-25'
    EPS = 1e-12
    SCALE_10000 = False
    data = pd.read_csv(RV_FILE)
    if 'date' not in data.columns:
        raise RuntimeError("RV file must contain a 'date' column.")
    num_cols = data.columns.drop('date')
    data[num_cols] = data[num_cols].apply(pd.to_numeric, errors='coerce')
    if SCALE_10000:
        data[num_cols] = data[num_cols] * 10000
    data = data.dropna()
    data['date'] = pd.to_datetime(data['date'])
    data = data.set_index('date')
    data[data <= 0] = EPS
    crv_raw = data.mean(axis=1).clip(lower=EPS)
    xi_raw = data.divide(crv_raw, axis=0).clip(lower=EPS)

    def get_logmean_har_features(raw_series, h=22, eps=1e-12):
        raw = pd.Series(raw_series).astype(float).clip(lower=eps)
        df = pd.DataFrame(index=raw.index)
        df['D'] = np.log(raw.shift(1).clip(lower=eps))
        df['W'] = np.log(raw.shift(1).rolling(5).mean().clip(lower=eps))
        df['M'] = np.log(raw.shift(1).rolling(22).mean().clip(lower=eps))
        future_mean = raw.rolling(window=h).mean().shift(-(h - 1)).clip(lower=eps)
        df['Target'] = np.log(future_mean)
        return df.dropna()
    for H in H_LIST:
        df_log_crv = get_logmean_har_features(crv_raw, h=H, eps=EPS)
        df_log_crv = df_log_crv.loc['2014-07-04':]
        data_test_crv = df_log_crv.loc[time_start:time_stop]
        if data_test_crv.empty:
            raise RuntimeError(f'[H={H}] data_test_crv is empty. Check dates.')
        train_len_crv = len(df_log_crv.loc[:time_start]) - H
        if train_len_crv <= 10:
            raise RuntimeError(f'[H={H}] train_len_crv too small ({train_len_crv}). Check start date or reduce H.')
        crv_log_predictions = []
        for i in tqdm(range(len(data_test_crv)), desc=f'CRV-Log-HAR H={H}'):
            train_data = df_log_crv.iloc[i:i + train_len_crv]
            model = smf.ols('Target ~ D + W + M', train_data).fit()
            row = data_test_crv.iloc[i]
            pred = model.params['Intercept'] + model.params['D'] * row['D'] + model.params['W'] * row['W'] + model.params['M'] * row['M']
            crv_log_predictions.append(pred)
        pred_log_crv_series = pd.Series(crv_log_predictions, index=data_test_crv.index)
        error_rows = []
        for c in data.columns:
            df_log_xi = get_logmean_har_features(xi_raw[c], h=H, eps=EPS)
            df_log_xi = df_log_xi.loc['2014-07-04':]
            data_test_xi = df_log_xi.loc[time_start:time_stop]
            train_len_xi = len(df_log_xi.loc[:time_start]) - H
            if train_len_xi <= 10 or data_test_xi.empty:
                error_rows.append([np.nan, np.nan, np.nan])
                model_outputs[str(c).strip().lstrip('.')] = pd.DataFrame()
                continue
            xi_log_predictions = []
            for i in tqdm(range(len(data_test_xi)), desc=f'{c}-xi-Log-HAR H={H}', leave=False):
                train_data = df_log_xi.iloc[i:i + train_len_xi]
                model = smf.ols('Target ~ D + W + M', train_data).fit()
                row = data_test_xi.iloc[i]
                pred = model.params['Intercept'] + model.params['D'] * row['D'] + model.params['W'] * row['W'] + model.params['M'] * row['M']
                xi_log_predictions.append(pred)
            pred_log_xi_series = pd.Series(xi_log_predictions, index=data_test_xi.index)
            idx = pred_log_xi_series.index.intersection(pred_log_crv_series.index)
            if len(idx) == 0:
                error_rows.append([np.nan, np.nan, np.nan])
                model_outputs[str(c).strip().lstrip('.')] = pd.DataFrame()
                continue
            pred_crv_mean = np.exp(pred_log_crv_series.loc[idx])
            pred_xi_mean = np.exp(pred_log_xi_series.loc[idx])
            rv_pred_mean = (pred_crv_mean * pred_xi_mean).clip(lower=EPS)
            rv_true_mean = data[c].rolling(H).mean().shift(-(H - 1)).loc[idx].clip(lower=EPS)
            out = pd.DataFrame({'RV_True': rv_true_mean, 'RV_Predict': rv_pred_mean}).dropna()
            if out.empty:
                error_rows.append([np.nan, np.nan, np.nan])
                model_outputs[str(c).strip().lstrip('.')] = out.copy()
                continue
            out['MSE'] = (out['RV_True'] - out['RV_Predict']) ** 2
            out['MAE'] = np.abs(out['RV_True'] - out['RV_Predict'])
            out['Qlike'] = out['RV_True'] / out['RV_Predict'].clip(lower=EPS) + np.log(out['RV_Predict'].clip(lower=EPS))
            mse_val = float(out['MSE'].mean())
            mae_val = float(out['MAE'].mean())
            theil_val = float(np.sqrt(mse_val) / (np.sqrt((out['RV_Predict'] ** 2).mean()) + np.sqrt((out['RV_True'] ** 2).mean())))
            error_rows.append([mse_val, mae_val, theil_val])
            model_outputs[str(c).strip().lstrip('.')] = out.copy()
        ERROR_DF = pd.DataFrame(error_rows, columns=['MSE', 'MAE', 'Theil'], index=data.columns).T
    return model_outputs

def run_multivariate():
    model_outputs = {}
    import os
    import numpy as np
    import pandas as pd
    from functools import partial
    from tqdm import tqdm as _tqdm
    tqdm = partial(_tqdm, bar_format="{bar} {percentage:3.0f}%", ncols=50)
    import statsmodels.api as sm
    EPS_LOG = 1e-12
    EXP_CLIP_MAX = 50
    H = 22
    SAMPLE_START = '2014-07-04'
    time_start = '2018/05/02'
    time_stop = '2022-02-25'

    def safe_log(x, eps: float=EPS_LOG):
        return np.log(np.clip(x, eps, None))

    def safe_exp(x, clip_max: float=EXP_CLIP_MAX):
        return np.exp(np.clip(x, -clip_max, clip_max))

    def future_mean_level(rv_now: pd.Series, H: int) -> pd.Series:
        if H == 1:
            return rv_now
        ma = rv_now.rolling(window=H, center=False).mean()
        return ma.shift(-(H - 1))
    data = pd.read_csv(RV_FILE)
    num_cols = [c for c in data.columns if c != 'date']
    data[num_cols] = data[num_cols].apply(pd.to_numeric, errors='coerce')
    data = data.dropna()
    data['date'] = pd.to_datetime(data['date'], utc=True)
    data['date'] = data['date'].dt.tz_localize(None)
    data['date'] = data['date'].dt.normalize()
    data = data.set_index('date').sort_index()
    data = data.fillna(method='ffill').dropna()
    data[data <= 0] = EPS_LOG
    markets = list(data.columns)
    rv_now = data.copy()
    rv_d_lvl = rv_now.shift(1)
    rv_w_lvl = rv_d_lvl.rolling(5).mean()
    rv_m_lvl = rv_d_lvl.rolling(22).mean()
    rv_d = safe_log(rv_d_lvl)
    rv_w = safe_log(rv_w_lvl)
    rv_m = safe_log(rv_m_lvl)
    X_all = pd.DataFrame(index=rv_now.index)
    for m in markets:
        X_all[f'RV_D_{m}'] = rv_d[m]
        X_all[f'RV_W_{m}'] = rv_w[m]
        X_all[f'RV_M_{m}'] = rv_m[m]
    X_cols = list(X_all.columns)
    error_rows = []
    time_start_dt = pd.to_datetime(time_start)
    time_stop_dt = pd.to_datetime(time_stop)
    for target in markets:
        y_lvl = future_mean_level(rv_now[target], H)
        y_log = pd.Series(safe_log(y_lvl), index=rv_now.index, name=f'RV_MA{H}')
        df = pd.concat([y_log, X_all], axis=1).dropna()
        df = df.loc[SAMPLE_START:].copy()
        data_test = df.loc[time_start_dt:time_stop_dt].copy()
        if len(data_test) == 0:
            tmp = pd.DataFrame(index=pd.DatetimeIndex([], name='date'))
            tmp['RV_True'] = np.nan
            tmp['RV_Predict'] = np.nan
            tmp['MSE'] = np.nan
            tmp['MAE'] = np.nan
            tmp['Qlike'] = np.nan
            model_outputs[str(target).strip().lstrip('.')] = tmp.copy()
            error_rows.append([np.nan, np.nan, np.nan])
            continue
        y_true_log = data_test[f'RV_MA{H}'].copy()
        test_index = data_test.index
        start_pos = df.index.get_loc(test_index[0])
        if start_pos <= H:
            tmp = pd.DataFrame(index=pd.DatetimeIndex([], name='date'))
            tmp['RV_True'] = np.nan
            tmp['RV_Predict'] = np.nan
            tmp['MSE'] = np.nan
            tmp['MAE'] = np.nan
            tmp['Qlike'] = np.nan
            model_outputs[str(target).strip().lstrip('.')] = tmp.copy()
            error_rows.append([np.nan, np.nan, np.nan])
            continue
        preds_log = []
        for i, forecast_date in enumerate(tqdm(test_index, desc=f'MV-LogHAR {target} H{H}', leave=False)):
            train_slice = df.iloc[:start_pos + i].copy()
            train_slice = train_slice.iloc[:-(H - 1)]
            y_tr = train_slice[f'RV_MA{H}'].values
            X_tr = train_slice[X_cols].values
            X_tr = sm.add_constant(X_tr, has_constant='add')
            res = sm.OLS(y_tr, X_tr).fit()
            x_row = data_test.loc[[forecast_date], X_cols].values
            x_row = sm.add_constant(x_row, has_constant='add')
            pred_i = float(res.predict(x_row)[0])
            preds_log.append(pred_i)
        preds_log = pd.Series(preds_log, index=y_true_log.index)
        y_true = pd.Series(safe_exp(y_true_log.values), index=y_true_log.index)
        y_pred = pd.Series(safe_exp(preds_log.values), index=preds_log.index)
        SUM = pd.concat([y_true, y_pred], axis=1)
        SUM.columns = ['RV_True', 'RV_Predict']
        SUM.index.name = 'date'
        SUM['MSE'] = (SUM['RV_True'] - SUM['RV_Predict']) ** 2
        SUM['MAE'] = np.abs(SUM['RV_True'] - SUM['RV_Predict'])
        pred_safe = SUM['RV_Predict'].clip(lower=EPS_LOG)
        true_safe = SUM['RV_True'].clip(lower=EPS_LOG)
        SUM['Qlike'] = true_safe / pred_safe + np.log(pred_safe)
        MSE = float(SUM['MSE'].mean())
        MAE = float(SUM['MAE'].mean())
        Theil = float(np.sqrt(np.mean((SUM['RV_True'] - SUM['RV_Predict']) ** 2)) / (np.sqrt(np.mean(SUM['RV_Predict'] ** 2)) + np.sqrt(np.mean(SUM['RV_True'] ** 2))))
        error_rows.append([MSE, MAE, Theil])
        model_outputs[str(target).strip().lstrip('.')] = SUM.copy()
    ERROR = pd.DataFrame(error_rows, columns=['MSE', 'MAE', 'Theil'], index=markets).T
    return model_outputs

def calc_mse(y_true: np.ndarray, y_pred: np.ndarray) -> dict:
    y_true = np.asarray(y_true, dtype=float)
    y_pred = np.asarray(y_pred, dtype=float)
    mask = np.isfinite(y_true) & np.isfinite(y_pred)
    yt = y_true[mask]
    yp = y_pred[mask]
    if len(yt) == 0:
        return {'MSE': np.nan}
    err = yp - yt
    mse = float(np.mean(err ** 2))
    return {'MSE': mse}

def qlike_with_shared_drop(y_true: pd.Series, y_pred: pd.Series, drop_dates: set) -> float:
    df = pd.concat([y_true.rename('y'), y_pred.rename('yhat')], axis=1).dropna()
    if df.empty:
        return np.nan
    if drop_dates:
        df = df.loc[~df.index.isin(drop_dates)]
    if df.empty:
        return np.nan
    y = df['y'].astype(float).values
    yhat = df['yhat'].astype(float).values
    ok_pos = (y > EPS) & (yhat > EPS)
    if not np.any(ok_pos):
        return np.nan
    ratio = y[ok_pos] / yhat[ok_pos]
    ok_ratio = np.isfinite(ratio) & (ratio >= RATIO_MIN) & (ratio <= RATIO_MAX)
    if not np.any(ok_ratio):
        return np.nan
    ratio = ratio[ok_ratio]
    q = ratio - np.log(ratio) - 1.0
    ok_q = np.isfinite(q) & (q >= 0.0) & (q <= QLIKE_CAP)
    if not np.any(ok_q):
        return np.nan
    return float(np.mean(q[ok_q]))

def add_summary_rows(df: pd.DataFrame) -> pd.DataFrame:
    mean_row = df.mean(axis=0, skipna=True)
    q1_row = df.quantile(0.25, axis=0, interpolation='linear')
    q2_row = df.quantile(0.5, axis=0, interpolation='linear')
    q3_row = df.quantile(0.75, axis=0, interpolation='linear')
    df2 = df.copy()
    df2.loc['Mean'] = mean_row
    df2.loc['Q1'] = q1_row
    df2.loc['Q2'] = q2_row
    df2.loc['Q3'] = q3_row
    return df2

def prepare_mcs_matrix_relaxed(sum_pre: pd.DataFrame, min_valid_per_model: int=10, min_rows_for_mcs: int=5):
    info = {'initial_shape': tuple(sum_pre.shape), 'dropped_low_valid_models': [], 'dropped_const_models': [], 'final_shape': None, 'reason': None}
    if sum_pre.empty or sum_pre.shape[1] == 0:
        info['reason'] = 'The loss matrix is empty'
        return (None, info)
    valid_counts = sum_pre.notna().sum(axis=0)
    keep_cols = valid_counts[valid_counts >= min_valid_per_model].index.tolist()
    info['dropped_low_valid_models'] = [c for c in sum_pre.columns if c not in keep_cols]
    if len(keep_cols) < 2:
        info['reason'] = f'Fewer than two models with at least {min_valid_per_model} valid observations'
        return (None, info)
    x = sum_pre[keep_cols].copy()
    x = x.dropna(axis=0, how='any')
    if x.empty:
        info['reason'] = 'No common valid observations'
        return (None, info)
    if x.shape[1] < 2:
        info['reason'] = 'Fewer than two models after dropping missing observations'
        return (None, info)
    if x.shape[0] < min_rows_for_mcs:
        info['reason'] = f'Insufficient common observations ({x.shape[0]} < {min_rows_for_mcs})'
        return (None, info)
    variances = x.var(axis=0, skipna=True)
    non_const_cols = variances[variances > 0].index.tolist()
    info['dropped_const_models'] = [c for c in x.columns if c not in non_const_cols]
    x = x[non_const_cols]
    if x.shape[1] < 2:
        info['reason'] = 'Fewer than two models after removing constant columns'
        return (None, info)
    if x.shape[0] < min_rows_for_mcs:
        info['reason'] = f'Insufficient common observations after removing constant columns ({x.shape[0]} < {min_rows_for_mcs})'
        return (None, info)
    info['final_shape'] = tuple(x.shape)
    return (x, info)

def choose_one_model_to_drop_for_retry(raw_sum_pre: pd.DataFrame, current_cols: list):
    if len(current_cols) <= 2:
        return None
    sub = raw_sum_pre[current_cols].copy()
    valid_counts = sub.notna().sum(axis=0)
    min_valid = valid_counts.min()
    candidates = valid_counts[valid_counts == min_valid].index.tolist()
    if len(candidates) == 1:
        return candidates[0]
    var_map = {}
    for c in candidates:
        s = pd.to_numeric(sub[c], errors='coerce').dropna()
        if len(s) <= 1:
            var_map[c] = -1.0
        else:
            v = float(s.var())
            if np.isnan(v):
                v = -1.0
            var_map[c] = v
    drop_col = sorted(var_map.items(), key=lambda kv: (kv[1], kv[0]))[0][0]
    return drop_col

def run_mcs_with_retry(sum_pre: pd.DataFrame, country_code: str, min_valid_per_model: int, min_rows_for_mcs: int, max_retry_drop: int, debug_logs: list):
    details = {'country': country_code, 'retry_count': 0, 'dropped_in_retry': [], 'fail_reason': None, 'final_shape': None, 'final_models': None}
    if sum_pre.empty or sum_pre.shape[1] < 2:
        details['fail_reason'] = 'Fewer than two available models'
        return (False, None, None, None, details)
    current_cols = list(sum_pre.columns)
    for _ in range(max_retry_drop + 1):
        tmp_sum = sum_pre[current_cols].copy()
        (mcs_df, prep_info) = prepare_mcs_matrix_relaxed(tmp_sum, min_valid_per_model=min_valid_per_model, min_rows_for_mcs=min_rows_for_mcs)
        if mcs_df is None:
            if len(current_cols) <= 2:
                details['fail_reason'] = f"Preparation failed with too few models to retry: {prep_info.get('reason')}"
                return (False, None, None, None, details)
            drop_col = choose_one_model_to_drop_for_retry(sum_pre, current_cols)
            if drop_col is None:
                details['fail_reason'] = f"Preparation failed and no model can be removed: {prep_info.get('reason')}"
                return (False, None, None, None, details)
            current_cols.remove(drop_col)
            details['retry_count'] += 1
            details['dropped_in_retry'].append(drop_col)
            debug_logs.append(f"[{country_code}] Preparation failed; retry without model: {drop_col} | Reason: {prep_info.get('reason')}")
            continue
        try:
            mcs = MCS(losses=mcs_df, size=0.9, method='R', reps=20000, block_size=5)
            mcs.compute()
            details['final_shape'] = tuple(mcs_df.shape)
            details['final_models'] = list(mcs_df.columns)
            return (True, mcs.pvalues, mcs.included, mcs_df, details)
        except Exception as e:
            if len(current_cols) <= 2:
                details['fail_reason'] = f'MCS failed with too few models to retry: {str(e)}'
                return (False, None, None, None, details)
            drop_col = choose_one_model_to_drop_for_retry(sum_pre, current_cols)
            if drop_col is None:
                details['fail_reason'] = f'MCS failed and no model can be removed: {str(e)}'
                return (False, None, None, None, details)
            current_cols.remove(drop_col)
            details['retry_count'] += 1
            details['dropped_in_retry'].append(drop_col)
            debug_logs.append(f'[{country_code}] MCS failed; retry without model: {drop_col} | Error: {str(e)}')
            continue
    details['fail_reason'] = f'Maximum retry count exceeded: {max_retry_drop}'
    return (False, None, None, None, details)

def qlike_series_with_rules(y_true: np.ndarray, y_pred: np.ndarray) -> np.ndarray:
    y_true = np.asarray(y_true, dtype=float)
    y_pred = np.asarray(y_pred, dtype=float)
    q = np.full(shape=len(y_true), fill_value=np.nan, dtype=float)
    mask = np.isfinite(y_true) & np.isfinite(y_pred)
    if not np.any(mask):
        return q
    yt = y_true[mask]
    yp = y_pred[mask]
    ok_pos = (yt > EPS) & (yp > EPS)
    ratio = np.full_like(yt, np.nan, dtype=float)
    ratio[ok_pos] = yt[ok_pos] / yp[ok_pos]
    ok_ratio = np.isfinite(ratio) & (ratio >= RATIO_MIN) & (ratio <= RATIO_MAX)
    q_sub = np.full_like(yt, np.nan, dtype=float)
    q_sub[ok_ratio] = ratio[ok_ratio] - np.log(ratio[ok_ratio]) - 1.0
    ok_q = np.isfinite(q_sub) & (q_sub >= 0.0) & (q_sub <= QLIKE_CAP)
    q_masked = np.full_like(yt, np.nan, dtype=float)
    q_masked[ok_q] = q_sub[ok_q]
    q[mask] = q_masked
    return q


MODEL_RUNNERS = [
    ("LogHAR", "LogHAR", run_loghar),
    ("LogHAR-CI", "LogHAR-CI", run_loghar_ci),
    ("HExp-GlRV", "HExpGl", run_hexpgl),
    ("MVF-Log", "MVF", run_mvf),
    ("Multivariate-LogHAR", "Multivariate", run_multivariate),
]

AUXILIARY_RUNNERS = [
    ("HAR-RV", "HAR", run_har),
]


def load_prediction_pair(frame):
    dates = pd.to_datetime(frame.index, errors="coerce")
    y_true = pd.Series(pd.to_numeric(frame["RV_True"], errors="coerce").to_numpy(), index=dates)
    y_pred = pd.Series(pd.to_numeric(frame["RV_Predict"], errors="coerce").to_numpy(), index=dates)
    valid = dates.notna() & np.isfinite(y_true.to_numpy()) & np.isfinite(y_pred.to_numpy())
    y_true, y_pred = y_true[valid], y_pred[valid]
    if len(y_true) and np.mean(y_true.to_numpy() > THRESH) > 0.5:
        y_true = y_true * SCALE_IF_MOST_GT_THRESH
        y_pred = y_pred * SCALE_IF_MOST_GT_THRESH
    y_true = y_true[~y_true.index.duplicated(keep="last")].sort_index()
    y_pred = y_pred[~y_pred.index.duplicated(keep="last")].sort_index()
    return y_true, y_pred


def load_shared_drop_dates(markets):
    data = pd.read_csv(QLIKE_DROP_DATES_FILE)
    required = {"market", "date"}
    if not required.issubset(data.columns):
        raise ValueError(f"Missing QLIKE exclusion columns: {sorted(required - set(data.columns))}")
    data = data[["market", "date"]].copy()
    data["market"] = data["market"].astype(str).str.strip().str.lstrip(".")
    data["date"] = pd.to_datetime(data["date"], errors="coerce").dt.normalize()
    if data["date"].isna().any():
        raise ValueError("The QLIKE exclusion file contains invalid dates")
    if data.duplicated(["market", "date"]).any():
        raise ValueError("The QLIKE exclusion file contains duplicate market-date rows")
    file_markets = set(data["market"])
    expected_markets = set(markets)
    if file_markets != expected_markets:
        difference = sorted(file_markets.symmetric_difference(expected_markets))
        raise ValueError(f"QLIKE exclusion markets do not match the forecast markets: {difference}")
    return {market: set(group["date"]) for market, group in data.groupby("market")}


def summarize_losses(predictions):
    models = [key for key, _, _ in MODEL_RUNNERS]
    all_models = [key for key, _, _ in AUXILIARY_RUNNERS] + models
    markets = sorted(predictions["LogHAR"])
    cache = {market: {model: load_prediction_pair(predictions[model][market]) for model in all_models} for market in markets}
    shared_drop_dates = load_shared_drop_dates(markets)
    raw = {metric: pd.DataFrame(index=markets, columns=all_models, dtype=float) for metric in ("MSE", "QLIKE")}
    for market in markets:
        for model in all_models:
            y_true, y_pred = cache[market][model]
            raw["MSE"].loc[market, model] = calc_mse(y_true.to_numpy(), y_pred.to_numpy())["MSE"]
            raw["QLIKE"].loc[market, model] = qlike_with_shared_drop(y_true, y_pred, shared_drop_dates[market])
    summaries = {}
    for metric, values in raw.items():
        values = values.dropna(how="all").dropna(subset=["HAR-RV", "LogHAR"], how="any")
        har_relative = values.div(values["HAR-RV"].replace(0.0, np.nan), axis=0)
        if np.any(np.isfinite(har_relative.to_numpy()) & (np.abs(har_relative.to_numpy()) > POST_SCALE_THRESHOLD)):
            warnings.warn("The original relative-loss scaling rule was triggered for " + metric, RuntimeWarning)
            har_relative = har_relative.where(~(np.abs(har_relative) > POST_SCALE_THRESHOLD), har_relative * POST_SCALE_FACTOR)
        statistics = add_summary_rows(har_relative).loc[["Mean", "Q1", "Q2"]]
        normalized = statistics.div(statistics["LogHAR"].replace(0.0, np.nan), axis=0)
        summaries[metric] = normalized[models].T
    return summaries, shared_drop_dates


def qlike_loss_series(frame, drop_dates):
    data = pd.DataFrame({
        "date": pd.to_datetime(frame.index, errors="coerce"),
        "y_true": pd.to_numeric(frame["RV_True"], errors="coerce").to_numpy(),
        "y_pred": pd.to_numeric(frame["RV_Predict"], errors="coerce").to_numpy(),
    }).dropna(subset=["date"])
    if data["date"].dt.tz is not None:
        data["date"] = data["date"].dt.tz_localize(None)
    data["date"] = data["date"].dt.normalize()
    data = data.sort_values("date").drop_duplicates(subset=["date"], keep="last").reset_index(drop=True)
    data = data.loc[~data["date"].isin(drop_dates)]
    return pd.Series(qlike_series_with_rules(data["y_true"].to_numpy(), data["y_pred"].to_numpy()), index=data["date"])


def count_mcs_memberships(predictions, shared_drop_dates):
    models = [key for key, _, _ in MODEL_RUNNERS]
    markets = sorted(predictions["LogHAR"])
    counts = {metric: dict.fromkeys(models, 0) for metric in ("MSE", "QLIKE")}
    for metric in counts:
        for market in tqdm(markets, desc="MCS " + metric):
            losses = pd.DataFrame()
            for model in models:
                frame = predictions[model][market]
                if metric == "MSE":
                    series = pd.to_numeric(frame["MSE"], errors="coerce").reset_index(drop=True)
                else:
                    series = qlike_loss_series(frame, shared_drop_dates[market])
                if series.notna().sum():
                    losses[model] = series
            logs = []
            success, _, included, used, details = run_mcs_with_retry(losses, market, 10, 5, 50, logs)
            if not success:
                raise RuntimeError(f"MCS {metric} failed for {market}: {details['fail_reason']}")
            omitted = [model for model in models if model not in used.columns]
            if omitted:
                warnings.warn(f"MCS {metric}, {market}: excluded from the test: {', '.join(omitted)}", RuntimeWarning)
            for model in included:
                counts[metric][model] += 1
    return counts


def save_table(summaries, counts, output_path=OUTPUT_FILE, bold_mode="loghar-ci"):
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Table 5"
    sheet.merge_cells("A1:E1")
    sheet.merge_cells("G1:K1")
    sheet["A1"] = "MSE"
    sheet["G1"] = "QLIKE"
    for start in (1, 7):
        for offset, label in enumerate(("Model", "Mean", "Q1", "Median", "MCS(#)")):
            sheet.cell(2, start + offset, label)
    for row, (key, label, _) in enumerate(MODEL_RUNNERS, start=3):
        for start, metric in ((1, "MSE"), (7, "QLIKE")):
            sheet.cell(row, start, label)
            for offset, statistic in enumerate(("Mean", "Q1", "Q2"), start=1):
                value = float(summaries[metric].loc[key, statistic])
                if not np.isfinite(value):
                    raise ValueError(f"Cannot export a missing {metric} {statistic} for {key}")
                sheet.cell(row, start + offset, value).number_format = "0.000"
            sheet.cell(row, start + 4, int(counts[metric][key])).number_format = "0"
    thin = Side(style="thin", color="000000")
    last_row = 2 + len(MODEL_RUNNERS)
    for row in sheet:
        for cell in row:
            cell.font = Font(name="Times New Roman", size=11, bold=False)
            cell.alignment = Alignment(horizontal="left" if cell.column in (1, 7) and cell.row > 1 else "center", vertical="center")
            if cell.column != 6:
                if cell.row == 1:
                    cell.border = Border(top=thin, bottom=thin)
                elif cell.row in (2, last_row):
                    cell.border = Border(bottom=thin)
    for start, metric in ((1, "MSE"), (7, "QLIKE")):
        for row, (key, _, _) in enumerate(MODEL_RUNNERS, start=3):
            for offset, stat in enumerate(("Mean", "Q1", "Q2", "MCS(#)"), start=1):
                bold = bold_mode == "loghar-ci" and key == "LogHAR-CI"
                if bold_mode == "best":
                    bold = counts[metric][key] == max(counts[metric].values()) if offset == 4 else np.isclose(summaries[metric].loc[key, stat], summaries[metric][stat].min(), rtol=0, atol=1e-12)
                if bold:
                    sheet.cell(row, start + offset).font = Font(name="Times New Roman", size=11, bold=True)
    for column in ("A", "G"):
        sheet.column_dimensions[column].width = 22
    for column in ("B", "C", "D", "E", "H", "I", "J", "K"):
        sheet.column_dimensions[column].width = 12
    sheet.column_dimensions["F"].width = 3
    sheet.freeze_panes = "B3"
    sheet.sheet_view.showGridLines = False
    sheet.print_options.horizontalCentered = True
    sheet.page_setup.orientation = "landscape"
    sheet.page_setup.paperSize = sheet.PAPERSIZE_A4
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 1
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    sheet.print_area = f"A1:K{last_row}"
    output_path = Path(output_path)
    output_path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(output_path)


def check_inputs():
    for path in (RV_FILE, RETURN_FILE, LOG_CENTRALITY_FILE, QLIKE_DROP_DATES_FILE):
        if not path.is_file():
            raise FileNotFoundError(f"Required input not found: {path}")
    markets = [str(c).strip().lstrip(".") for c in pd.read_csv(RV_FILE, nrows=0).columns if c != "date"]
    if len(markets) != 18 or len(set(markets)) != 18:
        raise ValueError("Table 5 requires 18 distinct markets in RV.csv")
    centrality = pd.read_csv(LOG_CENTRALITY_FILE, usecols=["market", "L_net"])
    centrality["market"] = centrality["market"].astype(str).str.strip().str.lstrip(".")
    windows = set(pd.to_numeric(centrality["L_net"], errors="coerce").dropna().astype(int))
    if 700 not in windows:
        raise ValueError("P_log.csv does not contain the 700-day network window required by LogHAR-CI")
    if set(centrality["market"].dropna()) != set(markets):
        raise ValueError("P_log.csv markets do not match RV.csv")
    load_shared_drop_dates(markets)
    return set(markets)


def main():
    global LOG_CENTRALITY_FILE
    parser = argparse.ArgumentParser(description="Generate Table 5 from the original Category 2 model specifications.")
    parser.add_argument("--recomputed-centrality", action="store_true")
    parser.add_argument("--bold", choices=("loghar-ci", "best"), default="loghar-ci")
    args = parser.parse_args()
    if args.recomputed_centrality:
        LOG_CENTRALITY_FILE = ROOT / "results/generated/centrality/P_log.csv"
    markets = check_inputs()
    predictions = {}
    for key, label, runner in tqdm(AUXILIARY_RUNNERS, desc="Table 5 auxiliary benchmark"):
        result = runner()
        if set(result) != markets:
            raise RuntimeError(f"{label}: missing or unexpected markets: {sorted(markets.symmetric_difference(result))}")
        predictions[key] = result
    for key, label, runner in tqdm(MODEL_RUNNERS, desc="Table 5 models"):
        result = runner()
        if set(result) != markets:
            raise RuntimeError(f"{label}: missing or unexpected markets: {sorted(markets.symmetric_difference(result))}")
        for market, frame in result.items():
            if frame.empty or not {"RV_True", "RV_Predict", "MSE"}.issubset(frame.columns):
                raise RuntimeError(f"{label}: no usable forecast output for {market}")
        predictions[key] = result
    summaries, shared_drop_dates = summarize_losses(predictions)
    counts = count_mcs_memberships(predictions, shared_drop_dates)
    save_table(summaries, counts, bold_mode=args.bold)


if __name__ == "__main__":
    main()

