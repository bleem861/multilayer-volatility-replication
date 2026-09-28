from pathlib import Path
import argparse
import warnings
import numpy as np
import pandas as pd
from functools import partial
from tqdm import tqdm as _tqdm
tqdm = partial(_tqdm, bar_format="{bar} {percentage:3.0f}%", ncols=50)
from arch.bootstrap import MCS
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, Border, Side

ROOT = Path(__file__).resolve().parents[1]
RV_FILE = ROOT / "data/raw/RV.csv"
RETURN_FILE = ROOT / "data/raw/Return.csv"
CENTRALITY_FILE = ROOT / "data/processed/centrality/P.csv"
QLIKE_DROP_DATES_FILE = ROOT / "data/processed/evaluation/qlike_intermediate.csv"
OUTPUT_FILE = ROOT / "results/generated/Table_06_category3_forecast_performance.xlsx"
THRESH = 1e-3
SCALE_IF_MOST_GT_THRESH = 1e-4
EPS = 1e-12
RATIO_MIN = 1e-6
RATIO_MAX = 1e6
QLIKE_CAP = 100.0
POST_SCALE_THRESHOLD = 1e4
POST_SCALE_FACTOR = 1e-8


def run_har():
    model_outputs = {}
    import math
    import scipy as sc
    import time
    import os
    from scipy.stats import norm
    import pandas as pd
    import numpy as np
    from functools import partial
    from tqdm import tqdm as _tqdm
    tqdm = partial(_tqdm, bar_format="{bar} {percentage:3.0f}%", ncols=50)
    import statistics
    import statsmodels.formula.api as smf
    import os
    data = pd.read_csv(RV_FILE)
    num_cols = data.columns.drop('date')
    data[num_cols] = data[num_cols].apply(pd.to_numeric, errors='coerce')
    data = data.dropna()
    data['date'] = pd.to_datetime(data['date'], utc=True)
    data['date'] = data['date'].dt.tz_localize(None)
    data['date'] = data['date'].dt.normalize()
    data = data.set_index('date')
    error = []
    H = 22
    for c in data.columns:
        data_RV = data['%s' % c]
        data_RV = pd.DataFrame(data_RV)
        data_RV.columns = ['RV_now']
        data_RV['RV_D'] = data_RV['RV_now'].shift(1)
        data_RV['RV_W'] = data_RV['RV_D'].rolling(5).mean()
        data_RV['RV_M'] = data_RV['RV_D'].rolling(22).mean()
        data_RV['RV_MA22'] = data_RV['RV_now'].rolling(window=22, center=False).mean()
        data_RV['RV_MA22'] = data_RV['RV_MA22'].shift(-21)
        data_RV = data_RV.dropna()
        data_RV = data_RV.loc['2014-07-04':]
        time = '2018/05/02'
        time_stop = '2022-02-25'
        Loss_ALL = []
        data_test = data_RV.loc[time:time_stop]
        data_predict = data_RV.loc[time:time_stop]['RV_MA22']
        result = []
        for i in tqdm(range(len(data_RV.loc[time:time_stop]))):
            data_rolling = data_RV.iloc[i:i + (len(data_RV.loc[:time]) - 1)]
            mod = smf.ols('RV_MA22 ~ RV_D + RV_W + RV_M', data_rolling)
            res = mod.fit()
            params_new = res.params
            r_predict = params_new.iloc[0] + params_new.iloc[1] * data_test['RV_D'].iloc[i] + params_new.iloc[2] * data_test['RV_W'].iloc[i] + params_new.iloc[3] * data_test['RV_M'].iloc[i]
            result.append(r_predict)
        result = pd.Series(result, index=data_predict.index)
        SUM = pd.concat([data_predict, result], axis=1)
        SUM.columns = ['RV_True', 'RV_Predict']
        SUM_final = SUM
        SUM_final = SUM_final.reset_index(drop=True)
        SUM_final.index = pd.to_datetime(data_test.index)
        SUM_final['MSE'] = (SUM_final['RV_True'] - SUM_final['RV_Predict']) ** 2
        SUM_final['MAE'] = np.abs(SUM_final['RV_True'] - SUM_final['RV_Predict'])
        SUM_final['Qlike'] = SUM_final['RV_True'] / SUM_final['RV_Predict'] + np.log(SUM_final['RV_Predict'])
        MSE = np.mean((SUM_final['RV_True'] - SUM_final['RV_Predict']) ** 2)
        MAE = np.mean(np.abs(SUM_final['RV_True'] - SUM_final['RV_Predict']))
        Theil = np.sqrt(np.mean((SUM_final['RV_True'] - SUM_final['RV_Predict']) ** 2)) / (np.sqrt(np.mean(SUM_final['RV_Predict'] ** 2)) + np.sqrt(np.mean(SUM_final['RV_True'] ** 2)))
        Loss_ALL = [MSE, MAE, Theil]
        error.append(Loss_ALL)
        model_outputs[str(c).strip().lstrip('.')] = SUM_final.copy()
    ERROR = pd.DataFrame(error, columns=['MSE', 'MAE', 'Theil'], index=data.columns).T
    return model_outputs

def run_har_ci():
    model_outputs = {}
    import os
    import numpy as np
    import pandas as pd
    import statsmodels.formula.api as smf
    from functools import partial
    from tqdm import tqdm as _tqdm
    tqdm = partial(_tqdm, bar_format="{bar} {percentage:3.0f}%", ncols=50)
    RET_PATH = RETURN_FILE
    RV_PATH = RV_FILE
    P_STORE_PATH = CENTRALITY_FILE
    PLOT = False
    fit_end_date = pd.to_datetime('2018-05-01')
    pred_start = pd.to_datetime('2018-05-02')
    pred_end = pd.to_datetime('2022-02-25')
    H_LIST = [22]
    FIT_START = 1000
    FIT_STEP = 100
    FIT_MAX = 1
    L_fit_list = [FIT_START + FIT_STEP * k for k in range(FIT_MAX)]
    NET_MIN = 1000
    NET_MAX = 1000
    NET_STEP = 50
    L_net_list = list(range(NET_MIN, NET_MAX + 1, NET_STEP))

    def _norm_date(x):
        return pd.to_datetime(x).tz_localize(None).normalize()

    def make_future_mean(rv_now: pd.Series, H: int) -> pd.Series:
        if H == 1:
            return rv_now
        ma = rv_now.rolling(window=H, center=False).mean()
        return ma.shift(-(H - 1))

    def read_returns(path):
        df = pd.read_csv(path)
        df = df.loc[:, ~df.columns.str.startswith('Unnamed')]
        df['date'] = pd.to_datetime(df['date'].astype(str).str.strip().str[:10], errors='coerce')
        df['date'] = df['date'].dt.tz_localize(None).dt.normalize()
        df = df.set_index('date').dropna(how='all')
        return df

    def read_rv_no_scale(path):
        data = pd.read_csv(path)
        if 'date' not in data.columns:
            raise RuntimeError("RV file must contain a 'date' column.")
        num_cols = [c for c in data.columns if c != 'date']
        data[num_cols] = data[num_cols].apply(pd.to_numeric, errors='coerce')
        data = data.dropna(how='all')
        data['date'] = pd.to_datetime(data['date'], utc=True, errors='coerce')
        data = data.dropna(subset=['date'])
        data['date'] = data['date'].dt.tz_localize(None)
        data['date'] = data['date'].dt.normalize()
        data = data.set_index('date')
        data = data.dropna(axis=1, how='all')
        return data
    df_ret = read_returns(RET_PATH)
    data_rv = read_rv_no_scale(RV_PATH)
    markets = [c for c in data_rv.columns if c in df_ret.columns]
    P_all = pd.read_csv(P_STORE_PATH)
    need_cols = ['date', 'market', 'ci_start_date', 'ci_end_date', 'L_net', 'P_corr', 'P_kendall', 'P_partial', 'P_tail']
    miss = [c for c in need_cols if c not in P_all.columns]
    if miss:
        raise RuntimeError(f'P.csv is missing required columns: {miss}; available columns={list(P_all.columns)}')
    P_all['date'] = pd.to_datetime(P_all['date'], errors='coerce').dt.tz_localize(None).dt.normalize()
    P_all['ci_start_date'] = P_all['ci_start_date'].apply(_norm_date)
    P_all['ci_end_date'] = P_all['ci_end_date'].apply(_norm_date)
    P_all['L_net'] = pd.to_numeric(P_all['L_net'], errors='coerce').astype('Int64')
    P_all = P_all.dropna(subset=['date', 'market', 'ci_start_date', 'ci_end_date', 'L_net']).copy()
    P_cache = {}

    def load_P_from_csv(L_net: int, ci_start_date, ci_end_date):
        key = (int(L_net), _norm_date(ci_start_date), _norm_date(ci_end_date))
        if key in P_cache:
            return P_cache[key]
        (ln, cs, ce) = key
        sub = P_all[(P_all['L_net'] == ln) & (P_all['ci_start_date'] == cs) & (P_all['ci_end_date'] == ce)]
        if sub.empty:
            P_cache[key] = (None, None, None, None)
            return P_cache[key]

        def _pivot(col):
            df = sub.pivot_table(index='date', columns='market', values=col, aggfunc='mean')
            return df.sort_index()
        P_corr = _pivot('P_corr')
        P_kendall = _pivot('P_kendall')
        P_partial = _pivot('P_partial')
        P_tail = _pivot('P_tail')
        P_cache[key] = (P_corr, P_kendall, P_partial, P_tail)
        return P_cache[key]

    def forecast_one_market_oneH_harci(data_rv_all: pd.DataFrame, P_corr: pd.DataFrame, P_kendall: pd.DataFrame, P_partial: pd.DataFrame, P_tail: pd.DataFrame, market: str, fit_start_date: pd.Timestamp, H: int):
        df = data_rv_all[market].to_frame('RV_now')
        df['RV_D'] = df['RV_now'].shift(1)
        df['RV_W'] = df['RV_D'].rolling(5).mean()
        df['RV_M'] = df['RV_D'].rolling(22).mean()
        df['Y'] = make_future_mean(df['RV_now'], H)
        ci = pd.DataFrame({'CI_corr': P_corr[market], 'CI_kendall': P_kendall[market], 'CI_partial': P_partial[market], 'CI_tail': P_tail[market]}).reindex(df.index)
        df = df.join(ci, how='inner').dropna()
        df['CI_corr_D'] = df['CI_corr'].shift(1)
        df['CI_kendall_D'] = df['CI_kendall'].shift(1)
        df['CI_partial_D'] = df['CI_partial'].shift(1)
        df['CI_tail_D'] = df['CI_tail'].shift(1)
        df = df.dropna()
        df = df.loc[fit_start_date:]
        data_test = df.loc[pred_start:pred_end].copy()
        if len(data_test) == 0:
            return (None, np.nan, np.nan, np.nan)
        window_len = len(df.loc[:pred_start]) - 1
        if window_len <= 0:
            return (None, np.nan, np.nan, np.nan)
        y_true = data_test['Y'].astype(float).copy()
        idx = y_true.index
        preds = []
        for i in tqdm(range(len(idx)), desc=f'HAR-CI {market} H{H}', leave=False):
            train = df.iloc[i:i + window_len].copy()
            train['_Y_'] = train['Y'].astype(float).values
            res = smf.ols('_Y_ ~ RV_D + RV_W + RV_M + CI_corr_D + CI_kendall_D + CI_partial_D + CI_tail_D', train).fit()
            p = res.params
            pred = p.iloc[0] + p.iloc[1] * data_test['RV_D'].iloc[i] + p.iloc[2] * data_test['RV_W'].iloc[i] + p.iloc[3] * data_test['RV_M'].iloc[i] + p.iloc[4] * data_test['CI_corr_D'].iloc[i] + p.iloc[5] * data_test['CI_kendall_D'].iloc[i] + p.iloc[6] * data_test['CI_partial_D'].iloc[i] + p.iloc[7] * data_test['CI_tail_D'].iloc[i]
            preds.append(float(pred))
        y_pred = pd.Series(preds, index=idx)
        mcs = pd.concat([y_true.rename('RV_True'), y_pred.rename('RV_Predict')], axis=1)
        mcs['MSE'] = (mcs['RV_True'] - mcs['RV_Predict']) ** 2
        mcs['MAE'] = np.abs(mcs['RV_True'] - mcs['RV_Predict'])
        mse = float(np.nanmean(mcs['MSE'].values))
        mae = float(np.nanmean(mcs['MAE'].values))
        denom = np.sqrt(np.nanmean(mcs['RV_Predict'].values ** 2)) + np.sqrt(np.nanmean(mcs['RV_True'].values ** 2))
        theil = float(np.sqrt(np.nanmean((mcs['RV_True'].values - mcs['RV_Predict'].values) ** 2)) / denom) if denom != 0 else np.nan
        return (mcs, mse, mae, theil)
    L_fit = L_fit_list[0]
    L_net = L_net_list[0]
    for H in H_LIST:
        error_by_market = {}
        for mkt in markets:
            if fit_end_date not in data_rv.index:
                error_by_market[mkt] = [np.nan, np.nan, np.nan]
                continue
            fit_end_pos = data_rv.index.get_loc(fit_end_date)
            fit_start_pos = fit_end_pos - (L_fit - 1)
            if fit_start_pos < 0:
                error_by_market[mkt] = [np.nan, np.nan, np.nan]
                continue
            fit_start_date = data_rv.index[fit_start_pos]
            ret_pos_end = df_ret.index.searchsorted(fit_start_date) - 1
            if ret_pos_end < 0:
                error_by_market[mkt] = [np.nan, np.nan, np.nan]
                continue
            net_end_date = df_ret.index[ret_pos_end]
            ci_start_date = net_end_date
            ci_end_date = pred_end
            (P_corr, P_kendall, P_partial, P_tail) = load_P_from_csv(L_net=L_net, ci_start_date=ci_start_date, ci_end_date=ci_end_date)
            if P_corr is None:
                error_by_market[mkt] = [np.nan, np.nan, np.nan]
                continue
            if mkt not in P_corr.columns:
                error_by_market[mkt] = [np.nan, np.nan, np.nan]
                continue
            (mcs, mse, mae, theil) = forecast_one_market_oneH_harci(data_rv_all=data_rv, P_corr=P_corr, P_kendall=P_kendall, P_partial=P_partial, P_tail=P_tail, market=mkt, fit_start_date=fit_start_date, H=H)
            if mcs is None:
                tmp = pd.DataFrame(index=pd.DatetimeIndex([], name='date'))
                tmp['RV_True'] = np.nan
                tmp['RV_Predict'] = np.nan
                tmp['MSE'] = np.nan
                tmp['MAE'] = np.nan
                model_outputs[str(mkt).strip().lstrip('.')] = tmp.copy()
            else:
                model_outputs[str(mkt).strip().lstrip('.')] = mcs.copy()
            error_by_market[mkt] = [mse, mae, theil]
        ERROR = pd.DataFrame(error_by_market, index=['MSE', 'MAE', 'Theil'])
    return model_outputs

def run_bagging():
    model_outputs = {}
    from sklearn import metrics
    import pandas as pd
    from functools import partial
    from tqdm import tqdm as _tqdm
    tqdm = partial(_tqdm, bar_format="{bar} {percentage:3.0f}%", ncols=50)
    import numpy as np
    from sklearn.preprocessing import StandardScaler
    import warnings
    warnings.filterwarnings('ignore')
    from sklearn.metrics import mean_squared_error
    from sklearn.metrics import mean_absolute_error
    from sklearn.tree import DecisionTreeRegressor
    from sklearn.tree import ExtraTreeRegressor
    from sklearn.preprocessing import MinMaxScaler
    from sklearn.ensemble import RandomForestRegressor
    from sklearn.ensemble import AdaBoostRegressor
    from sklearn.ensemble import GradientBoostingRegressor
    from sklearn.ensemble import BaggingRegressor
    from sklearn.metrics import mean_absolute_percentage_error
    import os
    data = pd.read_csv(RV_FILE)
    data = data.dropna()
    data['date'] = pd.to_datetime(data['date'], utc=True)
    data['date'] = data['date'].dt.tz_localize(None)
    data['date'] = data['date'].dt.normalize()
    data = data.fillna(method='ffill')
    data = data.set_index('date')
    P_STORE_PATH = CENTRALITY_FILE
    P_all = pd.read_csv(P_STORE_PATH)
    need_cols = ['date', 'market', 'P_corr', 'P_kendall', 'P_partial', 'P_tail']
    miss = [c for c in need_cols if c not in P_all.columns]
    if miss:
        raise RuntimeError(f'Model execution message{miss}Model execution message{list(P_all.columns)}')
    P_all['date'] = pd.to_datetime(P_all['date'], errors='coerce').dt.tz_localize(None).dt.normalize()
    P_all = P_all.dropna(subset=['date', 'market']).copy()
    P_ci = P_all.groupby(['date', 'market'])[['P_corr', 'P_kendall', 'P_partial', 'P_tail']].mean().reset_index()
    error = []
    for country in tqdm(data.columns):
        data_RV = data[[country]].copy()
        data_RV['RV_now'] = data_RV[country]
        data_RV = data_RV.drop(columns=[country])
        data_RV['RV_MA22'] = data_RV['RV_now'].rolling(window=22, center=False).mean()
        data_RV['RV_MA22'] = data_RV['RV_MA22'].shift(-21)
        data_RV['RV_D'] = data_RV['RV_now'].shift(1)
        data_RV['RV_W'] = data_RV['RV_D'].rolling(5).mean()
        data_RV['RV_M'] = data_RV['RV_D'].rolling(22).mean()
        ci_one = P_ci[P_ci['market'] == country].copy()
        if ci_one.empty:
            tmp = pd.DataFrame(index=pd.DatetimeIndex([], name='date'))
            tmp['RV_True'] = np.nan
            tmp['RV_Predict'] = np.nan
            tmp['MSE'] = np.nan
            tmp['MAE'] = np.nan
            tmp['MAPE'] = np.nan
            model_outputs[str(country).strip().lstrip('.')] = tmp.copy()
            error.append([np.nan, np.nan, np.nan])
            continue
        ci_one = ci_one.set_index('date')[['P_corr', 'P_kendall', 'P_partial', 'P_tail']]
        ci_one = ci_one.rename(columns={'P_corr': 'CI_corr', 'P_kendall': 'CI_kendall', 'P_partial': 'CI_partial', 'P_tail': 'CI_tail'})
        data_RV = data_RV.join(ci_one, how='left')
        data_RV['CI_corr_D'] = data_RV['CI_corr'].shift(1)
        data_RV['CI_kendall_D'] = data_RV['CI_kendall'].shift(1)
        data_RV['CI_partial_D'] = data_RV['CI_partial'].shift(1)
        data_RV['CI_tail_D'] = data_RV['CI_tail'].shift(1)
        data_RV = data_RV.dropna()
        data_RV = data_RV.loc['2014-07-04':]
        time = '2018-05-02'
        time_stop = '2022-02-25'
        X_all = data_RV[['RV_D', 'RV_W', 'RV_M', 'CI_corr_D', 'CI_kendall_D', 'CI_partial_D', 'CI_tail_D']]
        y_all = data_RV['RV_MA22']
        train_mask = X_all.index < pd.to_datetime(time)
        test_mask = (X_all.index >= pd.to_datetime(time)) & (X_all.index <= pd.to_datetime(time_stop))
        (X_train, y_train) = (X_all[train_mask], y_all[train_mask])
        (X_test, y_test) = (X_all[test_mask], y_all[test_mask])
        if X_test.dropna().empty or y_test.dropna().empty or X_train.dropna().empty or y_train.dropna().empty:
            tmp = pd.DataFrame(index=pd.DatetimeIndex([], name='date'))
            tmp['RV_True'] = np.nan
            tmp['RV_Predict'] = np.nan
            tmp['MSE'] = np.nan
            tmp['MAE'] = np.nan
            tmp['MAPE'] = np.nan
            model_outputs[str(country).strip().lstrip('.')] = tmp.copy()
            error.append([np.nan, np.nan, np.nan])
            continue
        base_estimator = DecisionTreeRegressor(max_depth=10, random_state=42)
        try:
            Bagging = BaggingRegressor(estimator=base_estimator, n_estimators=1000, max_samples=0.8, max_features=0.8, random_state=42, n_jobs=-1)
        except TypeError:
            Bagging = BaggingRegressor(base_estimator=base_estimator, n_estimators=1000, max_samples=0.8, max_features=0.8, random_state=42, n_jobs=-1)
        Bagging.fit(X_train.dropna(), y_train.dropna())
        predictions = Bagging.predict(X_test.dropna())
        valid_idx = X_test.dropna().index
        result = pd.Series(predictions, index=valid_idx, name='RV_Predict')
        SUM = pd.concat([y_test.loc[valid_idx], result], axis=1)
        SUM.columns = ['RV_True', 'RV_Predict']
        SUM_final = SUM.dropna()
        SUM_final['MSE'] = (SUM_final['RV_True'] - SUM_final['RV_Predict']) ** 2
        SUM_final['MAE'] = np.abs(SUM_final['RV_True'] - SUM_final['RV_Predict'])
        SUM_final['MAPE'] = np.abs((SUM_final['RV_True'] - SUM_final['RV_Predict']) / SUM_final['RV_True'])
        MSE = np.mean((SUM_final['RV_True'] - SUM_final['RV_Predict']) ** 2)
        MAE = np.mean(np.abs(SUM_final['RV_True'] - SUM_final['RV_Predict']))
        Theil = np.sqrt(np.mean((SUM_final['RV_True'] - SUM_final['RV_Predict']) ** 2)) / (np.sqrt(np.mean(SUM_final['RV_Predict'] ** 2)) + np.sqrt(np.mean(SUM_final['RV_True'] ** 2)))
        error.append([MSE, MAE, Theil])
        model_outputs[str(country).strip().lstrip('.')] = SUM_final.copy()
    ERROR = pd.DataFrame(error, columns=['MSE', 'MAE', 'Theil'], index=data.columns).T
    return model_outputs

def run_gradient():
    model_outputs = {}
    from sklearn import metrics
    import pandas as pd
    from functools import partial
    from tqdm import tqdm as _tqdm
    tqdm = partial(_tqdm, bar_format="{bar} {percentage:3.0f}%", ncols=50)
    import numpy as np
    from sklearn.preprocessing import StandardScaler
    import warnings
    warnings.filterwarnings('ignore')
    from sklearn.metrics import mean_squared_error
    from sklearn.metrics import mean_absolute_error
    from sklearn.tree import DecisionTreeRegressor
    from sklearn.tree import ExtraTreeRegressor
    from sklearn.preprocessing import MinMaxScaler
    from sklearn.ensemble import RandomForestRegressor
    from sklearn.ensemble import AdaBoostRegressor
    from sklearn.ensemble import GradientBoostingRegressor
    from sklearn.ensemble import BaggingRegressor
    from sklearn.metrics import mean_absolute_percentage_error
    import os
    data = pd.read_csv(RV_FILE)
    data = data.dropna()
    data['date'] = pd.to_datetime(data['date'], utc=True)
    data['date'] = data['date'].dt.tz_localize(None)
    data['date'] = data['date'].dt.normalize()
    data = data.fillna(method='ffill')
    data = data.set_index('date')
    P_STORE_PATH = CENTRALITY_FILE
    P_all = pd.read_csv(P_STORE_PATH)
    need_cols = ['date', 'market', 'P_corr', 'P_kendall', 'P_partial', 'P_tail']
    miss = [c for c in need_cols if c not in P_all.columns]
    if miss:
        raise RuntimeError(f'Model execution message{miss}Model execution message{list(P_all.columns)}')
    P_all['date'] = pd.to_datetime(P_all['date'], errors='coerce').dt.tz_localize(None).dt.normalize()
    P_all = P_all.dropna(subset=['date', 'market']).copy()
    P_ci = P_all.groupby(['date', 'market'])[['P_corr', 'P_kendall', 'P_partial', 'P_tail']].mean().reset_index()
    error = []
    for country in tqdm(data.columns):
        data_RV = data[[country]].copy()
        data_RV['RV_now'] = data_RV[country]
        data_RV = data_RV.drop(columns=[country])
        data_RV['RV_D'] = data_RV['RV_now'].shift(1)
        data_RV['RV_W'] = data_RV['RV_D'].rolling(5).mean()
        data_RV['RV_M'] = data_RV['RV_D'].rolling(22).mean()
        data_RV['RV_MA22'] = data_RV['RV_now'].rolling(window=22, center=False).mean()
        data_RV['RV_MA22'] = data_RV['RV_MA22'].shift(-21)
        ci_one = P_ci[P_ci['market'] == country].copy()
        if ci_one.empty:
            tmp = pd.DataFrame(index=pd.DatetimeIndex([], name='date'))
            tmp['RV_True'] = np.nan
            tmp['RV_Predict'] = np.nan
            tmp['MSE'] = np.nan
            tmp['MAE'] = np.nan
            tmp['MAPE'] = np.nan
            model_outputs[str(country).strip().lstrip('.')] = tmp.copy()
            error.append([np.nan, np.nan, np.nan])
            continue
        ci_one = ci_one.set_index('date')[['P_corr', 'P_kendall', 'P_partial', 'P_tail']]
        ci_one = ci_one.rename(columns={'P_corr': 'CI_corr', 'P_kendall': 'CI_kendall', 'P_partial': 'CI_partial', 'P_tail': 'CI_tail'})
        data_RV = data_RV.join(ci_one, how='left')
        data_RV['CI_corr_D'] = data_RV['CI_corr'].shift(1)
        data_RV['CI_kendall_D'] = data_RV['CI_kendall'].shift(1)
        data_RV['CI_partial_D'] = data_RV['CI_partial'].shift(1)
        data_RV['CI_tail_D'] = data_RV['CI_tail'].shift(1)
        data_RV = data_RV.dropna()
        data_RV = data_RV.loc['2014-07-04':]
        time = '2018-05-02'
        time_stop = '2022-02-25'
        X_all = data_RV[['RV_D', 'RV_W', 'RV_M', 'CI_corr_D', 'CI_kendall_D', 'CI_partial_D', 'CI_tail_D']]
        y_all = data_RV['RV_MA22']
        train_mask = X_all.index < pd.to_datetime(time)
        test_mask = (X_all.index >= pd.to_datetime(time)) & (X_all.index <= pd.to_datetime(time_stop))
        (X_train, y_train) = (X_all[train_mask], y_all[train_mask])
        (X_test, y_test) = (X_all[test_mask], y_all[test_mask])
        if X_test.dropna().empty or y_test.dropna().empty or X_train.dropna().empty or y_train.dropna().empty:
            tmp = pd.DataFrame(index=pd.DatetimeIndex([], name='date'))
            tmp['RV_True'] = np.nan
            tmp['RV_Predict'] = np.nan
            tmp['MSE'] = np.nan
            tmp['MAE'] = np.nan
            tmp['MAPE'] = np.nan
            model_outputs[str(country).strip().lstrip('.')] = tmp.copy()
            error.append([np.nan, np.nan, np.nan])
            continue
        gb = GradientBoostingRegressor(n_estimators=1000, learning_rate=0.02, max_depth=10, subsample=0.8, min_samples_split=50, random_state=32, validation_fraction=0.1, n_iter_no_change=10)
        gb.fit(X_train.dropna(), y_train.dropna())
        predictions = gb.predict(X_test.dropna())
        valid_idx = X_test.dropna().index
        result = pd.Series(predictions, index=valid_idx, name='RV_Predict')
        SUM = pd.concat([y_test.loc[valid_idx], result], axis=1)
        SUM.columns = ['RV_True', 'RV_Predict']
        SUM_final = SUM.dropna()
        SUM_final['MSE'] = (SUM_final['RV_True'] - SUM_final['RV_Predict']) ** 2
        SUM_final['MAE'] = np.abs(SUM_final['RV_True'] - SUM_final['RV_Predict'])
        SUM_final['MAPE'] = np.abs((SUM_final['RV_True'] - SUM_final['RV_Predict']) / SUM_final['RV_True'])
        MSE = np.mean((SUM_final['RV_True'] - SUM_final['RV_Predict']) ** 2)
        MAE = np.mean(np.abs(SUM_final['RV_True'] - SUM_final['RV_Predict']))
        Theil = np.sqrt(np.mean((SUM_final['RV_True'] - SUM_final['RV_Predict']) ** 2)) / (np.sqrt(np.mean(SUM_final['RV_Predict'] ** 2)) + np.sqrt(np.mean(SUM_final['RV_True'] ** 2)))
        error.append([MSE, MAE, Theil])
        model_outputs[str(country).strip().lstrip('.')] = SUM_final.copy()
    ERROR = pd.DataFrame(error, columns=['MSE', 'MAE', 'Theil'], index=data.columns).T
    return model_outputs

def run_random_forest():
    model_outputs = {}
    from sklearn import metrics
    import pandas as pd
    from functools import partial
    from tqdm import tqdm as _tqdm
    tqdm = partial(_tqdm, bar_format="{bar} {percentage:3.0f}%", ncols=50)
    import numpy as np
    from sklearn.preprocessing import StandardScaler
    import warnings
    warnings.filterwarnings('ignore')
    from sklearn.metrics import mean_squared_error
    from sklearn.metrics import mean_absolute_error
    from sklearn.tree import DecisionTreeRegressor
    from sklearn.tree import ExtraTreeRegressor
    from sklearn.preprocessing import MinMaxScaler
    from sklearn.ensemble import RandomForestRegressor
    from sklearn.ensemble import AdaBoostRegressor
    from sklearn.ensemble import GradientBoostingRegressor
    from sklearn.ensemble import BaggingRegressor
    from sklearn.metrics import mean_absolute_percentage_error
    import os
    H = 22
    data = pd.read_csv(RV_FILE)
    data = data.dropna()
    data['date'] = pd.to_datetime(data['date'], utc=True)
    data['date'] = data['date'].dt.tz_localize(None)
    data['date'] = data['date'].dt.normalize()
    data = data.fillna(method='ffill')
    data = data.set_index('date')
    P_STORE_PATH = CENTRALITY_FILE
    P_all = pd.read_csv(P_STORE_PATH)
    need_cols = ['date', 'market', 'P_corr', 'P_kendall', 'P_partial', 'P_tail']
    miss = [c for c in need_cols if c not in P_all.columns]
    if miss:
        raise RuntimeError(f'Model execution message{miss}Model execution message{list(P_all.columns)}')
    P_all['date'] = pd.to_datetime(P_all['date'], errors='coerce').dt.tz_localize(None).dt.normalize()
    P_all = P_all.dropna(subset=['date', 'market']).copy()
    P_ci = P_all.groupby(['date', 'market'])[['P_corr', 'P_kendall', 'P_partial', 'P_tail']].mean().reset_index()
    error = []
    for country in data.columns:
        data_RV = data[[country]].copy()
        data_RV['RV_now'] = data_RV[country]
        data_RV = data_RV.drop(columns=[country])
        data_RV['RV_D'] = data_RV['RV_now'].shift(1)
        data_RV['RV_W'] = data_RV['RV_D'].rolling(5).mean()
        data_RV['RV_M'] = data_RV['RV_D'].rolling(22).mean()
        data_RV['RV_MA22'] = data_RV['RV_now'].rolling(window=H, center=False).mean()
        data_RV['RV_MA22'] = data_RV['RV_MA22'].shift(-(H - 1))
        ci_one = P_ci[P_ci['market'] == country].copy()
        if ci_one.empty:
            tmp = pd.DataFrame(index=pd.DatetimeIndex([], name='date'))
            tmp['RV_True'] = np.nan
            tmp['RV_Predict'] = np.nan
            tmp['MSE'] = np.nan
            tmp['MAE'] = np.nan
            model_outputs[str(country).strip().lstrip('.')] = tmp.copy()
            error.append([np.nan, np.nan, np.nan])
            continue
        ci_one = ci_one.set_index('date')[['P_corr', 'P_kendall', 'P_partial', 'P_tail']]
        ci_one = ci_one.rename(columns={'P_corr': 'CI_corr', 'P_kendall': 'CI_kendall', 'P_partial': 'CI_partial', 'P_tail': 'CI_tail'})
        data_RV = data_RV.join(ci_one, how='left')
        data_RV['CI_corr_D'] = data_RV['CI_corr'].shift(1)
        data_RV['CI_kendall_D'] = data_RV['CI_kendall'].shift(1)
        data_RV['CI_partial_D'] = data_RV['CI_partial'].shift(1)
        data_RV['CI_tail_D'] = data_RV['CI_tail'].shift(1)
        data_RV = data_RV.dropna()
        data_RV = data_RV.loc['2014-07-04':]
        time = '2018-05-02'
        time_stop = '2022-02-25'
        X_all = data_RV[['RV_D', 'RV_W', 'RV_M', 'CI_corr_D', 'CI_kendall_D', 'CI_partial_D', 'CI_tail_D']].copy()
        y_all = data_RV['RV_MA22'].copy()
        time_ts = pd.to_datetime(time)
        time_stop_ts = pd.to_datetime(time_stop)
        if len(data_RV.loc[time:time_stop]) == 0:
            tmp = pd.DataFrame(index=pd.DatetimeIndex([], name='date'))
            tmp['RV_True'] = np.nan
            tmp['RV_Predict'] = np.nan
            tmp['MSE'] = np.nan
            tmp['MAE'] = np.nan
            model_outputs[str(country).strip().lstrip('.')] = tmp.copy()
            error.append([np.nan, np.nan, np.nan])
            continue
        if time_ts in data_RV.index:
            train_idx = int(data_RV.index.get_loc(time_ts))
        else:
            train_idx = int(data_RV.index.searchsorted(time_ts))
        test_len = len(data_RV.loc[time:time_stop])
        X_test = X_all.iloc[train_idx:train_idx + test_len]
        y_test = y_all.iloc[train_idx:train_idx + test_len]
        if X_test.empty or y_test.empty:
            tmp = pd.DataFrame(index=pd.DatetimeIndex([], name='date'))
            tmp['RV_True'] = np.nan
            tmp['RV_Predict'] = np.nan
            tmp['MSE'] = np.nan
            tmp['MAE'] = np.nan
            model_outputs[str(country).strip().lstrip('.')] = tmp.copy()
            error.append([np.nan, np.nan, np.nan])
            continue
        rf = RandomForestRegressor(n_estimators=200, max_depth=1, random_state=42, n_jobs=-1)
        predictions = []
        for i in tqdm(range(len(X_test)), desc=f'RF {country}', leave=False):
            pos = train_idx + i
            train_end_exclusive = pos - H + 1
            train_start = i
            if train_end_exclusive <= train_start:
                predictions.append(np.nan)
                continue
            current_train_X = X_all.iloc[train_start:train_end_exclusive].copy()
            current_train_y = y_all.iloc[train_start:train_end_exclusive].copy()
            current_train_X = current_train_X.dropna()
            current_train_y = current_train_y.loc[current_train_X.index].dropna()
            current_train_X = current_train_X.loc[current_train_y.index]
            if current_train_X.empty or current_train_y.empty:
                predictions.append(np.nan)
                continue
            rf.fit(current_train_X, current_train_y)
            x_one = X_test.iloc[[i]].dropna()
            if x_one.empty:
                predictions.append(np.nan)
                continue
            pred = rf.predict(x_one)
            predictions.append(float(pred[0]) if len(pred) > 0 else np.nan)
        result = pd.Series(predictions, index=y_test.index, name='RV_Predict')
        SUM = pd.concat([y_test.rename('RV_True'), result], axis=1)
        SUM_final = SUM.dropna()
        if len(SUM_final) == 0:
            tmp = pd.DataFrame(index=pd.DatetimeIndex([], name='date'))
            tmp['RV_True'] = np.nan
            tmp['RV_Predict'] = np.nan
            tmp['MSE'] = np.nan
            tmp['MAE'] = np.nan
            model_outputs[str(country).strip().lstrip('.')] = tmp.copy()
            error.append([np.nan, np.nan, np.nan])
            continue
        SUM_final['MSE'] = (SUM_final['RV_True'] - SUM_final['RV_Predict']) ** 2
        SUM_final['MAE'] = np.abs(SUM_final['RV_True'] - SUM_final['RV_Predict'])
        MSE = float(np.mean((SUM_final['RV_True'] - SUM_final['RV_Predict']) ** 2))
        MAE = float(np.mean(np.abs(SUM_final['RV_True'] - SUM_final['RV_Predict'])))
        denom = np.sqrt(np.mean(SUM_final['RV_Predict'] ** 2)) + np.sqrt(np.mean(SUM_final['RV_True'] ** 2))
        Theil = float(np.sqrt(np.mean((SUM_final['RV_True'] - SUM_final['RV_Predict']) ** 2)) / denom) if denom != 0 else np.nan
        error.append([MSE, MAE, Theil])
        model_outputs[str(country).strip().lstrip('.')] = SUM_final.copy()
    ERROR = pd.DataFrame(error, columns=['MSE', 'MAE', 'Theil'], index=data.columns).T
    return model_outputs

def run_deep_learning():
    deep_outputs = {'LSTM': {}, 'CNN': {}, 'Transformer': {}}
    import os
    import warnings
    import numpy as np
    import pandas as pd
    import tensorflow as tf
    from functools import partial
    from tqdm import tqdm as _tqdm
    tqdm = partial(_tqdm, bar_format="{bar} {percentage:3.0f}%", ncols=50)
    from sklearn.preprocessing import MinMaxScaler
    from sklearn.metrics import mean_squared_error, mean_absolute_error
    from tensorflow.keras.models import Sequential
    from tensorflow.keras.layers import LSTM, Dense, Conv1D, MaxPooling1D, Flatten, MultiHeadAttention, Dropout, Input, GlobalAveragePooling1D
    from tensorflow.keras.callbacks import EarlyStopping
    warnings.filterwarnings('ignore')
    np.random.seed(42)
    tf.random.set_seed(42)
    RV_PATH = RV_FILE
    P_STORE_PATH = CENTRALITY_FILE
    OUT_LSTM = 'LSTM'
    OUT_CNN = 'CNN'
    OUT_TR = 'Transformer'
    PRED_START = pd.to_datetime('2018-05-02')
    PRED_END = pd.to_datetime('2022-02-25')
    FIT_START = pd.to_datetime('2014-07-04')
    n_steps = 1
    data = pd.read_csv(RV_PATH)
    data = data.dropna()
    data['date'] = pd.to_datetime(data['date'], utc=True)
    data['date'] = data['date'].dt.tz_localize(None)
    data['date'] = data['date'].dt.normalize()
    data = data.fillna(method='ffill')
    data = data.set_index('date')
    data = data.loc[FIT_START:]
    P_all = pd.read_csv(P_STORE_PATH)
    need_cols = ['date', 'market', 'P_corr', 'P_kendall', 'P_partial', 'P_tail']
    miss = [c for c in need_cols if c not in P_all.columns]
    if miss:
        raise RuntimeError(f'Model execution message{miss}Model execution message{list(P_all.columns)}')
    P_all['date'] = pd.to_datetime(P_all['date'], errors='coerce').dt.tz_localize(None).dt.normalize()
    P_all = P_all.dropna(subset=['date', 'market']).copy()
    P_ci = P_all.groupby(['date', 'market'])[['P_corr', 'P_kendall', 'P_partial', 'P_tail']].mean().reset_index()

    def build_lstm_model(input_shape):
        model = Sequential([LSTM(64, activation='relu', input_shape=input_shape), Dense(1)])
        model.compile(optimizer='adam', loss='mse')
        return model

    def build_cnn_model(input_shape):
        model = Sequential([Conv1D(filters=64, kernel_size=1, activation='relu', input_shape=input_shape), MaxPooling1D(pool_size=1), Flatten(), Dense(50, activation='relu'), Dense(1)])
        model.compile(optimizer='adam', loss='mse')
        return model

    def transformer_encoder(inputs, head_size, num_heads, ff_dim, dropout=0.0):
        x = MultiHeadAttention(key_dim=head_size, num_heads=num_heads, dropout=dropout)(inputs, inputs)
        x = Dropout(dropout)(x)
        res = x + inputs
        x = Dense(ff_dim, activation='relu')(res)
        x = Dense(inputs.shape[-1])(x)
        return x + res

    def build_transformer_model(input_shape):
        inputs = Input(shape=input_shape)
        x = transformer_encoder(inputs, head_size=32, num_heads=4, ff_dim=64, dropout=0.0)
        x = GlobalAveragePooling1D()(x)
        outputs = Dense(1)(x)
        model = tf.keras.Model(inputs, outputs)
        model.compile(optimizer='adam', loss='mse')
        return model

    def make_h22_target(rv_now: pd.Series) -> pd.Series:
        y = rv_now.rolling(window=22, center=False).mean().shift(-21)
        return y

    def create_dataset_with_dates(X_df: pd.DataFrame, y_s: pd.Series, time_steps=1):
        (Xs, ys, y_dates) = ([], [], [])
        for i in range(len(X_df) - time_steps):
            Xs.append(X_df.iloc[i:i + time_steps].values)
            ys.append(y_s.iloc[i + time_steps])
            y_dates.append(y_s.index[i + time_steps])
        return (np.array(Xs), np.array(ys), pd.DatetimeIndex(y_dates))

    def calc_theil(y_true, y_pred):
        y_true = np.asarray(y_true).reshape(-1, 1)
        y_pred = np.asarray(y_pred).reshape(-1, 1)
        denom = np.sqrt(np.mean(y_pred ** 2)) + np.sqrt(np.mean(y_true ** 2))
        if denom == 0:
            return np.nan
        return float(np.sqrt(np.mean((y_true - y_pred) ** 2)) / denom)

    def save_mcs_and_metrics(model, X_test, y_test_scaled, scaler_y, date_index, out_dir, country):
        y_pred_scaled = model.predict(X_test, verbose=0)
        y_pred = scaler_y.inverse_transform(y_pred_scaled).reshape(-1)
        y_true = scaler_y.inverse_transform(y_test_scaled.reshape(-1, 1)).reshape(-1)
        mcs = pd.DataFrame({'RV_True': y_true, 'RV_Predict': y_pred}, index=pd.DatetimeIndex(date_index, name='date'))
        mcs = mcs.dropna()
        if len(mcs) == 0:
            tmp = pd.DataFrame(index=pd.DatetimeIndex([], name='date'))
            tmp['RV_True'] = np.nan
            tmp['RV_Predict'] = np.nan
            tmp['MSE'] = np.nan
            tmp['MAE'] = np.nan
            deep_outputs[out_dir][str(country).strip().lstrip('.')] = tmp.copy()
            return [np.nan, np.nan, np.nan]
        mcs['MSE'] = (mcs['RV_True'] - mcs['RV_Predict']) ** 2
        mcs['MAE'] = np.abs(mcs['RV_True'] - mcs['RV_Predict'])
        eps = 1e-12
        pred_safe = np.clip(mcs['RV_Predict'].values, eps, None)
        true_safe = np.clip(mcs['RV_True'].values, eps, None)
        mcs['Qlike'] = true_safe / pred_safe + np.log(pred_safe)
        deep_outputs[out_dir][str(country).strip().lstrip('.')] = mcs.copy()
        mse = float(mean_squared_error(mcs['RV_True'], mcs['RV_Predict']))
        mae = float(mean_absolute_error(mcs['RV_True'], mcs['RV_Predict']))
        theil = calc_theil(mcs['RV_True'].values, mcs['RV_Predict'].values)
        return [mse, mae, theil]
    error_lstm = []
    error_cnn = []
    error_transformer = []
    for country in tqdm(data.columns, desc='Markets'):
        data_RV = data[[country]].copy()
        data_RV['RV_now'] = data_RV[country]
        data_RV = data_RV.drop(columns=[country])
        data_RV['RV_D'] = data_RV['RV_now'].shift(1)
        data_RV['RV_W'] = data_RV['RV_D'].rolling(5).mean()
        data_RV['RV_M'] = data_RV['RV_D'].rolling(22).mean()
        data_RV['RV_MA22'] = make_h22_target(data_RV['RV_now'])
        ci_one = P_ci[P_ci['market'] == country].copy()
        if ci_one.empty:
            for out_dir in [OUT_LSTM, OUT_CNN, OUT_TR]:
                tmp = pd.DataFrame(index=pd.DatetimeIndex([], name='date'))
                tmp['RV_True'] = np.nan
                tmp['RV_Predict'] = np.nan
                tmp['MSE'] = np.nan
                tmp['MAE'] = np.nan
                tmp['Qlike'] = np.nan
                deep_outputs[out_dir][str(country).strip().lstrip('.')] = tmp.copy()
            error_lstm.append([np.nan, np.nan, np.nan])
            error_cnn.append([np.nan, np.nan, np.nan])
            error_transformer.append([np.nan, np.nan, np.nan])
            continue
        ci_one = ci_one.set_index('date')[['P_corr', 'P_kendall', 'P_partial', 'P_tail']]
        ci_one = ci_one.rename(columns={'P_corr': 'CI_corr', 'P_kendall': 'CI_kendall', 'P_partial': 'CI_partial', 'P_tail': 'CI_tail'})
        data_RV = data_RV.join(ci_one, how='left')
        data_RV['CI_corr_D'] = data_RV['CI_corr'].shift(1)
        data_RV['CI_kendall_D'] = data_RV['CI_kendall'].shift(1)
        data_RV['CI_partial_D'] = data_RV['CI_partial'].shift(1)
        data_RV['CI_tail_D'] = data_RV['CI_tail'].shift(1)
        data_RV = data_RV.dropna()
        data_RV = data_RV.loc[FIT_START:]
        X = data_RV[['RV_D', 'RV_W', 'RV_M', 'CI_corr_D', 'CI_kendall_D', 'CI_partial_D', 'CI_tail_D']].copy()
        y = data_RV['RV_MA22'].copy()
        if len(X) <= n_steps + 10:
            for out_dir in [OUT_LSTM, OUT_CNN, OUT_TR]:
                tmp = pd.DataFrame(index=pd.DatetimeIndex([], name='date'))
                tmp['RV_True'] = np.nan
                tmp['RV_Predict'] = np.nan
                tmp['MSE'] = np.nan
                tmp['MAE'] = np.nan
                tmp['Qlike'] = np.nan
                deep_outputs[out_dir][str(country).strip().lstrip('.')] = tmp.copy()
            error_lstm.append([np.nan, np.nan, np.nan])
            error_cnn.append([np.nan, np.nan, np.nan])
            error_transformer.append([np.nan, np.nan, np.nan])
            continue
        train_mask = X.index < PRED_START
        test_mask = (X.index >= PRED_START) & (X.index <= PRED_END)
        X_train_df = X.loc[train_mask].copy()
        y_train_s = y.loc[train_mask].copy()
        X_test_df = X.loc[test_mask].copy()
        y_test_s = y.loc[test_mask].copy()
        if len(X_test_df) == 0 or len(X_train_df) == 0:
            for out_dir in [OUT_LSTM, OUT_CNN, OUT_TR]:
                tmp = pd.DataFrame(index=pd.DatetimeIndex([], name='date'))
                tmp['RV_True'] = np.nan
                tmp['RV_Predict'] = np.nan
                tmp['MSE'] = np.nan
                tmp['MAE'] = np.nan
                tmp['Qlike'] = np.nan
                deep_outputs[out_dir][str(country).strip().lstrip('.')] = tmp.copy()
            error_lstm.append([np.nan, np.nan, np.nan])
            error_cnn.append([np.nan, np.nan, np.nan])
            error_transformer.append([np.nan, np.nan, np.nan])
            continue
        X_context = X.iloc[max(0, X.index.get_loc(X_test_df.index[0]) - n_steps):X.index.get_loc(X_test_df.index[0])]
        y_context = y.loc[X_context.index]
        X_test_for_seq = pd.concat([X_context, X_test_df], axis=0)
        y_test_for_seq = pd.concat([y_context, y_test_s], axis=0)
        scaler_X = MinMaxScaler()
        scaler_y = MinMaxScaler()
        X_train_scaled = pd.DataFrame(scaler_X.fit_transform(X_train_df), index=X_train_df.index, columns=X_train_df.columns)
        y_train_scaled = pd.Series(scaler_y.fit_transform(y_train_s.values.reshape(-1, 1)).reshape(-1), index=y_train_s.index)
        X_test_scaled = pd.DataFrame(scaler_X.transform(X_test_for_seq), index=X_test_for_seq.index, columns=X_test_for_seq.columns)
        y_test_scaled = pd.Series(scaler_y.transform(y_test_for_seq.values.reshape(-1, 1)).reshape(-1), index=y_test_for_seq.index)
        (X_train_seq, y_train_seq, train_dates) = create_dataset_with_dates(X_train_scaled, y_train_scaled, n_steps)
        (X_test_seq, y_test_seq, test_dates) = create_dataset_with_dates(X_test_scaled, y_test_scaled, n_steps)
        keep_test = (test_dates >= PRED_START) & (test_dates <= PRED_END)
        X_test_seq = X_test_seq[keep_test]
        y_test_seq = y_test_seq[keep_test]
        test_dates = test_dates[keep_test]
        if len(X_train_seq) == 0 or len(X_test_seq) == 0:
            for out_dir in [OUT_LSTM, OUT_CNN, OUT_TR]:
                tmp = pd.DataFrame(index=pd.DatetimeIndex([], name='date'))
                tmp['RV_True'] = np.nan
                tmp['RV_Predict'] = np.nan
                tmp['MSE'] = np.nan
                tmp['MAE'] = np.nan
                tmp['Qlike'] = np.nan
                deep_outputs[out_dir][str(country).strip().lstrip('.')] = tmp.copy()
            error_lstm.append([np.nan, np.nan, np.nan])
            error_cnn.append([np.nan, np.nan, np.nan])
            error_transformer.append([np.nan, np.nan, np.nan])
            continue
        early_stopping = EarlyStopping(monitor='val_loss', patience=10, verbose=0, restore_best_weights=True)
        lstm_model = build_lstm_model((n_steps, X_train_seq.shape[2]))
        lstm_model.fit(X_train_seq, y_train_seq, epochs=200, batch_size=32, validation_split=0.1, callbacks=[early_stopping], verbose=0)
        lstm_pred_scaled = lstm_model.predict(X_test_seq, verbose=0)
        lstm_metrics_real = save_mcs_and_metrics(lstm_model, X_test_seq, y_test_seq, scaler_y, test_dates, OUT_LSTM, country)
        cnn_model = build_cnn_model((n_steps, X_train_seq.shape[2]))
        cnn_model.fit(X_train_seq, y_train_seq, epochs=200, batch_size=32, validation_split=0.1, callbacks=[early_stopping], verbose=0)
        cnn_pred_scaled = cnn_model.predict(X_test_seq, verbose=0)
        cnn_metrics_real = save_mcs_and_metrics(cnn_model, X_test_seq, y_test_seq, scaler_y, test_dates, OUT_CNN, country)
        transformer_model = build_transformer_model((n_steps, X_train_seq.shape[2]))
        transformer_model.fit(X_train_seq, y_train_seq, epochs=200, batch_size=32, validation_split=0.1, callbacks=[early_stopping], verbose=0)
        tr_pred_scaled = transformer_model.predict(X_test_seq, verbose=0)
        transformer_metrics_real = save_mcs_and_metrics(transformer_model, X_test_seq, y_test_seq, scaler_y, test_dates, OUT_TR, country)
        error_lstm.append(lstm_metrics_real)
        error_cnn.append(cnn_metrics_real)
        error_transformer.append(transformer_metrics_real)
    ERROR_lstm = pd.DataFrame(error_lstm, columns=['MSE', 'MAE', 'Theil'], index=data.columns).T
    ERROR_cnn = pd.DataFrame(error_cnn, columns=['MSE', 'MAE', 'Theil'], index=data.columns).T
    ERROR_transformer = pd.DataFrame(error_transformer, columns=['MSE', 'MAE', 'Theil'], index=data.columns).T
    return deep_outputs

def calc_mse(y_true: np.ndarray, y_pred: np.ndarray) -> dict:
    y_true = np.asarray(y_true, dtype=float)
    y_pred = np.asarray(y_pred, dtype=float)
    mask = np.isfinite(y_true) & np.isfinite(y_pred)
    yt = y_true[mask]
    yp = y_pred[mask]
    if len(yt) == 0:
        return {'MSE': np.nan}
    err = yp - yt
    mse = float(np.mean(err ** 2))
    return {'MSE': mse}

def qlike_with_shared_drop(y_true: pd.Series, y_pred: pd.Series, drop_dates: set) -> float:
    df = pd.concat([y_true.rename('y'), y_pred.rename('yhat')], axis=1).dropna()
    if df.empty:
        return np.nan
    if drop_dates:
        df = df.loc[~df.index.isin(drop_dates)]
    if df.empty:
        return np.nan
    y = df['y'].astype(float).values
    yhat = df['yhat'].astype(float).values
    ok_pos = (y > EPS) & (yhat > EPS)
    if not np.any(ok_pos):
        return np.nan
    ratio = y[ok_pos] / yhat[ok_pos]
    ok_ratio = np.isfinite(ratio) & (ratio >= RATIO_MIN) & (ratio <= RATIO_MAX)
    if not np.any(ok_ratio):
        return np.nan
    ratio = ratio[ok_ratio]
    q = ratio - np.log(ratio) - 1.0
    ok_q = np.isfinite(q) & (q >= 0.0) & (q <= QLIKE_CAP)
    if not np.any(ok_q):
        return np.nan
    return float(np.mean(q[ok_q]))

def add_summary_rows(df: pd.DataFrame) -> pd.DataFrame:
    mean_row = df.mean(axis=0, skipna=True)
    q1_row = df.quantile(0.25, axis=0, interpolation='linear')
    q2_row = df.quantile(0.5, axis=0, interpolation='linear')
    q3_row = df.quantile(0.75, axis=0, interpolation='linear')
    df2 = df.copy()
    df2.loc['Mean'] = mean_row
    df2.loc['Q1'] = q1_row
    df2.loc['Q2'] = q2_row
    df2.loc['Q3'] = q3_row
    return df2

def prepare_mcs_matrix_relaxed(sum_pre: pd.DataFrame, min_valid_per_model: int=10, min_rows_for_mcs: int=5):
    info = {'initial_shape': tuple(sum_pre.shape), 'dropped_low_valid_models': [], 'dropped_const_models': [], 'final_shape': None, 'reason': None}
    if sum_pre.empty or sum_pre.shape[1] == 0:
        info['reason'] = 'The loss matrix is empty'
        return (None, info)
    valid_counts = sum_pre.notna().sum(axis=0)
    keep_cols = valid_counts[valid_counts >= min_valid_per_model].index.tolist()
    info['dropped_low_valid_models'] = [c for c in sum_pre.columns if c not in keep_cols]
    if len(keep_cols) < 2:
        info['reason'] = f'Fewer than two models with at least {min_valid_per_model} valid observations'
        return (None, info)
    x = sum_pre[keep_cols].copy()
    x = x.dropna(axis=0, how='any')
    if x.empty:
        info['reason'] = 'No common valid observations'
        return (None, info)
    if x.shape[1] < 2:
        info['reason'] = 'Fewer than two models after dropping missing observations'
        return (None, info)
    if x.shape[0] < min_rows_for_mcs:
        info['reason'] = f'Insufficient common observations ({x.shape[0]} < {min_rows_for_mcs})'
        return (None, info)
    variances = x.var(axis=0, skipna=True)
    non_const_cols = variances[variances > 0].index.tolist()
    info['dropped_const_models'] = [c for c in x.columns if c not in non_const_cols]
    x = x[non_const_cols]
    if x.shape[1] < 2:
        info['reason'] = 'Fewer than two models after removing constant columns'
        return (None, info)
    if x.shape[0] < min_rows_for_mcs:
        info['reason'] = f'Insufficient common observations after removing constant columns ({x.shape[0]} < {min_rows_for_mcs})'
        return (None, info)
    info['final_shape'] = tuple(x.shape)
    return (x, info)

def choose_one_model_to_drop_for_retry(raw_sum_pre: pd.DataFrame, current_cols: list):
    if len(current_cols) <= 2:
        return None
    sub = raw_sum_pre[current_cols].copy()
    valid_counts = sub.notna().sum(axis=0)
    min_valid = valid_counts.min()
    candidates = valid_counts[valid_counts == min_valid].index.tolist()
    if len(candidates) == 1:
        return candidates[0]
    var_map = {}
    for c in candidates:
        s = pd.to_numeric(sub[c], errors='coerce').dropna()
        if len(s) <= 1:
            var_map[c] = -1.0
        else:
            v = float(s.var())
            if np.isnan(v):
                v = -1.0
            var_map[c] = v
    drop_col = sorted(var_map.items(), key=lambda kv: (kv[1], kv[0]))[0][0]
    return drop_col

def run_mcs_with_retry(sum_pre: pd.DataFrame, country_code: str, min_valid_per_model: int, min_rows_for_mcs: int, max_retry_drop: int, debug_logs: list):
    details = {'country': country_code, 'retry_count': 0, 'dropped_in_retry': [], 'fail_reason': None, 'final_shape': None, 'final_models': None}
    if sum_pre.empty or sum_pre.shape[1] < 2:
        details['fail_reason'] = 'Fewer than two available models'
        return (False, None, None, None, details)
    current_cols = list(sum_pre.columns)
    for _ in range(max_retry_drop + 1):
        tmp_sum = sum_pre[current_cols].copy()
        (mcs_df, prep_info) = prepare_mcs_matrix_relaxed(tmp_sum, min_valid_per_model=min_valid_per_model, min_rows_for_mcs=min_rows_for_mcs)
        if mcs_df is None:
            if len(current_cols) <= 2:
                details['fail_reason'] = f"Preparation failed with too few models to retry: {prep_info.get('reason')}"
                return (False, None, None, None, details)
            drop_col = choose_one_model_to_drop_for_retry(sum_pre, current_cols)
            if drop_col is None:
                details['fail_reason'] = f"Preparation failed and no model can be removed: {prep_info.get('reason')}"
                return (False, None, None, None, details)
            current_cols.remove(drop_col)
            details['retry_count'] += 1
            details['dropped_in_retry'].append(drop_col)
            debug_logs.append(f"[{country_code}] Preparation failed; retry without model: {drop_col} | Reason: {prep_info.get('reason')}")
            continue
        try:
            mcs = MCS(losses=mcs_df, size=0.9, method='R', reps=20000, block_size=5)
            mcs.compute()
            details['final_shape'] = tuple(mcs_df.shape)
            details['final_models'] = list(mcs_df.columns)
            return (True, mcs.pvalues, mcs.included, mcs_df, details)
        except Exception as e:
            if len(current_cols) <= 2:
                details['fail_reason'] = f'MCS failed with too few models to retry: {str(e)}'
                return (False, None, None, None, details)
            drop_col = choose_one_model_to_drop_for_retry(sum_pre, current_cols)
            if drop_col is None:
                details['fail_reason'] = f'MCS failed and no model can be removed: {str(e)}'
                return (False, None, None, None, details)
            current_cols.remove(drop_col)
            details['retry_count'] += 1
            details['dropped_in_retry'].append(drop_col)
            debug_logs.append(f'[{country_code}] MCS failed; retry without model: {drop_col} | Error: {str(e)}')
            continue
    details['fail_reason'] = f'Maximum retry count exceeded: {max_retry_drop}'
    return (False, None, None, None, details)

def qlike_series_with_rules(y_true: np.ndarray, y_pred: np.ndarray) -> np.ndarray:
    y_true = np.asarray(y_true, dtype=float)
    y_pred = np.asarray(y_pred, dtype=float)
    q = np.full(shape=len(y_true), fill_value=np.nan, dtype=float)
    mask = np.isfinite(y_true) & np.isfinite(y_pred)
    if not np.any(mask):
        return q
    yt = y_true[mask]
    yp = y_pred[mask]
    ok_pos = (yt > EPS) & (yp > EPS)
    ratio = np.full_like(yt, np.nan, dtype=float)
    ratio[ok_pos] = yt[ok_pos] / yp[ok_pos]
    ok_ratio = np.isfinite(ratio) & (ratio >= RATIO_MIN) & (ratio <= RATIO_MAX)
    q_sub = np.full_like(yt, np.nan, dtype=float)
    q_sub[ok_ratio] = ratio[ok_ratio] - np.log(ratio[ok_ratio]) - 1.0
    ok_q = np.isfinite(q_sub) & (q_sub >= 0.0) & (q_sub <= QLIKE_CAP)
    q_masked = np.full_like(yt, np.nan, dtype=float)
    q_masked[ok_q] = q_sub[ok_q]
    q[mask] = q_masked
    return q


MODEL_LABELS = [
    ("HAR-RV", "HAR"),
    ("HAR-CI", "HAR-CI"),
    ("bagging", "Bagging"),
    ("gradient", "Gradient"),
    ("RF", "RF"),
    ("CNN", "CNN"),
    ("LSTM", "LSTM"),
    ("Transformer", "Transformer"),
]

STANDARD_RUNNERS = [
    ("HAR-RV", "HAR", run_har),
    ("HAR-CI", "HAR-CI", run_har_ci),
    ("bagging", "Bagging", run_bagging),
    ("gradient", "Gradient", run_gradient),
    ("RF", "RF", run_random_forest),
]


def load_prediction_pair(frame):
    dates = pd.to_datetime(frame.index, errors="coerce")
    y_true = pd.Series(pd.to_numeric(frame["RV_True"], errors="coerce").to_numpy(), index=dates)
    y_pred = pd.Series(pd.to_numeric(frame["RV_Predict"], errors="coerce").to_numpy(), index=dates)
    valid = dates.notna() & np.isfinite(y_true.to_numpy()) & np.isfinite(y_pred.to_numpy())
    y_true, y_pred = y_true[valid], y_pred[valid]
    if len(y_true) and np.mean(y_true.to_numpy() > THRESH) > 0.5:
        y_true = y_true * SCALE_IF_MOST_GT_THRESH
        y_pred = y_pred * SCALE_IF_MOST_GT_THRESH
    y_true = y_true[~y_true.index.duplicated(keep="last")].sort_index()
    y_pred = y_pred[~y_pred.index.duplicated(keep="last")].sort_index()
    return y_true, y_pred


def load_shared_drop_dates(markets):
    data = pd.read_csv(QLIKE_DROP_DATES_FILE)
    required = {"market", "date"}
    if not required.issubset(data.columns):
        raise ValueError(f"Missing QLIKE exclusion columns: {sorted(required - set(data.columns))}")
    data = data[["market", "date"]].copy()
    data["market"] = data["market"].astype(str).str.strip().str.lstrip(".")
    data["date"] = pd.to_datetime(data["date"], errors="coerce").dt.normalize()
    if data["date"].isna().any():
        raise ValueError("The QLIKE exclusion file contains invalid dates")
    if data.duplicated(["market", "date"]).any():
        raise ValueError("The QLIKE exclusion file contains duplicate market-date rows")
    file_markets = set(data["market"])
    expected_markets = set(markets)
    if file_markets != expected_markets:
        difference = sorted(file_markets.symmetric_difference(expected_markets))
        raise ValueError(f"QLIKE exclusion markets do not match the forecast markets: {difference}")
    return {market: set(group["date"]) for market, group in data.groupby("market")}


def summarize_losses(predictions):
    models = [key for key, _ in MODEL_LABELS]
    markets = sorted(predictions["HAR-RV"])
    cache = {market: {model: load_prediction_pair(predictions[model][market]) for model in models} for market in markets}
    shared_drop_dates = load_shared_drop_dates(markets)
    raw = {metric: pd.DataFrame(index=markets, columns=models, dtype=float) for metric in ("MSE", "QLIKE")}
    for market in markets:
        for model in models:
            y_true, y_pred = cache[market][model]
            raw["MSE"].loc[market, model] = calc_mse(y_true.to_numpy(), y_pred.to_numpy())["MSE"]
            raw["QLIKE"].loc[market, model] = qlike_with_shared_drop(y_true, y_pred, shared_drop_dates[market])
    summaries = {}
    for metric, values in raw.items():
        values = values.dropna(how="all").dropna(subset=["HAR-RV"], how="all")
        relative = values.div(values["HAR-RV"].replace(0.0, np.nan), axis=0)
        if np.any(np.isfinite(relative.to_numpy()) & (np.abs(relative.to_numpy()) > POST_SCALE_THRESHOLD)):
            warnings.warn("The original relative-loss scaling rule was triggered for " + metric, RuntimeWarning)
            relative = relative.where(~(np.abs(relative) > POST_SCALE_THRESHOLD), relative * POST_SCALE_FACTOR)
        summaries[metric] = add_summary_rows(relative).loc[["Mean", "Q1", "Q2"]].T
    return summaries, shared_drop_dates


def qlike_loss_series(frame, drop_dates):
    data = pd.DataFrame({
        "date": pd.to_datetime(frame.index, errors="coerce"),
        "y_true": pd.to_numeric(frame["RV_True"], errors="coerce").to_numpy(),
        "y_pred": pd.to_numeric(frame["RV_Predict"], errors="coerce").to_numpy(),
    }).dropna(subset=["date"])
    if data["date"].dt.tz is not None:
        data["date"] = data["date"].dt.tz_localize(None)
    data["date"] = data["date"].dt.normalize()
    data = data.sort_values("date").drop_duplicates(subset=["date"], keep="last").reset_index(drop=True)
    data = data.loc[~data["date"].isin(drop_dates)]
    return pd.Series(qlike_series_with_rules(data["y_true"].to_numpy(), data["y_pred"].to_numpy()), index=data["date"])


def count_mcs_memberships(predictions, shared_drop_dates):
    models = [key for key, _ in MODEL_LABELS]
    markets = sorted(predictions["HAR-RV"])
    counts = {metric: dict.fromkeys(models, 0) for metric in ("MSE", "QLIKE")}
    for metric in counts:
        for market in tqdm(markets, desc="MCS " + metric):
            losses = pd.DataFrame()
            for model in models:
                frame = predictions[model][market]
                if metric == "MSE":
                    series = pd.to_numeric(frame["MSE"], errors="coerce").reset_index(drop=True)
                else:
                    series = qlike_loss_series(frame, shared_drop_dates[market])
                if series.notna().sum():
                    losses[model] = series
            logs = []
            success, _, included, used, details = run_mcs_with_retry(losses, market, 10, 5, 50, logs)
            if not success:
                raise RuntimeError(f"MCS {metric} failed for {market}: {details['fail_reason']}")
            omitted = [model for model in models if model not in used.columns]
            if omitted:
                warnings.warn(f"MCS {metric}, {market}: excluded from the test: {', '.join(omitted)}", RuntimeWarning)
            for model in included:
                counts[metric][model] += 1
    return counts


def save_table(summaries, counts, output_path=OUTPUT_FILE, bold_mode="har-ci"):
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Table 6"
    sheet.merge_cells("A1:E1")
    sheet.merge_cells("G1:K1")
    sheet["A1"] = "MSE"
    sheet["G1"] = "QLIKE"
    for start in (1, 7):
        for offset, label in enumerate(("Model", "Mean", "Q1", "Median", "MCS(#)")):
            sheet.cell(2, start + offset, label)
    for row, (key, label) in enumerate(MODEL_LABELS, start=3):
        for start, metric in ((1, "MSE"), (7, "QLIKE")):
            sheet.cell(row, start, label)
            for offset, statistic in enumerate(("Mean", "Q1", "Q2"), start=1):
                value = float(summaries[metric].loc[key, statistic])
                if not np.isfinite(value):
                    raise ValueError(f"Cannot export a missing {metric} {statistic} for {key}")
                sheet.cell(row, start + offset, value).number_format = "0.000"
            sheet.cell(row, start + 4, int(counts[metric][key])).number_format = "0"
    thin = Side(style="thin", color="000000")
    last_row = 2 + len(MODEL_LABELS)
    for row in sheet:
        for cell in row:
            cell.font = Font(name="Times New Roman", size=11, bold=False)
            cell.alignment = Alignment(horizontal="left" if cell.column in (1, 7) and cell.row > 1 else "center", vertical="center")
            if cell.column != 6:
                if cell.row == 1:
                    cell.border = Border(top=thin, bottom=thin)
                elif cell.row in (2, last_row):
                    cell.border = Border(bottom=thin)
    for start, metric in ((1, "MSE"), (7, "QLIKE")):
        for row, (key, _) in enumerate(MODEL_LABELS, start=3):
            for offset, stat in enumerate(("Mean", "Q1", "Q2", "MCS(#)"), start=1):
                bold = bold_mode == "har-ci" and key == "HAR-CI"
                if bold_mode == "best":
                    bold = counts[metric][key] == max(counts[metric].values()) if offset == 4 else np.isclose(summaries[metric].loc[key, stat], summaries[metric][stat].min(), rtol=0, atol=1e-12)
                if bold:
                    sheet.cell(row, start + offset).font = Font(name="Times New Roman", size=11, bold=True)
    for column in ("A", "G"):
        sheet.column_dimensions[column].width = 18
    for column in ("B", "C", "D", "E", "H", "I", "J", "K"):
        sheet.column_dimensions[column].width = 12
    sheet.column_dimensions["F"].width = 3
    sheet.freeze_panes = "B3"
    sheet.sheet_view.showGridLines = False
    sheet.print_options.horizontalCentered = True
    sheet.page_setup.orientation = "landscape"
    sheet.page_setup.paperSize = sheet.PAPERSIZE_A4
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 1
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    sheet.print_area = f"A1:K{last_row}"
    output_path = Path(output_path)
    output_path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(output_path)


def check_inputs():
    for path in (RV_FILE, RETURN_FILE, CENTRALITY_FILE, QLIKE_DROP_DATES_FILE):
        if not path.is_file():
            raise FileNotFoundError(f"Required input not found: {path}")
    markets = [str(c).strip().lstrip(".") for c in pd.read_csv(RV_FILE, nrows=0).columns if c != "date"]
    if len(markets) != 18 or len(set(markets)) != 18:
        raise ValueError("Table 6 requires 18 distinct markets in RV.csv")
    centrality = pd.read_csv(CENTRALITY_FILE, usecols=["market", "L_net"])
    centrality["market"] = centrality["market"].astype(str).str.strip().str.lstrip(".")
    windows = set(pd.to_numeric(centrality["L_net"], errors="coerce").dropna().astype(int))
    if 1000 not in windows:
        raise ValueError("P.csv does not contain the 1000-day network window required by Table 6")
    if set(centrality["market"].dropna()) != set(markets):
        raise ValueError("P.csv markets do not match RV.csv")
    load_shared_drop_dates(markets)
    return set(markets)


def validate_model_output(label, result, markets):
    if set(result) != markets:
        raise RuntimeError(f"{label}: missing or unexpected markets: {sorted(markets.symmetric_difference(result))}")
    for market, frame in result.items():
        if frame.empty or not {"RV_True", "RV_Predict", "MSE"}.issubset(frame.columns):
            raise RuntimeError(f"{label}: no usable forecast output for {market}")


def main():
    parser = argparse.ArgumentParser(description="Generate Table 6 from the original Category 3 model specifications.")
    parser.add_argument("--bold", choices=("har-ci", "best"), default="har-ci")
    args = parser.parse_args()
    markets = check_inputs()
    predictions = {}
    for key, label, runner in tqdm(STANDARD_RUNNERS, desc="Table 6 models"):
        result = runner()
        validate_model_output(label, result, markets)
        predictions[key] = result
    deep_outputs = run_deep_learning()
    for key in ("CNN", "LSTM", "Transformer"):
        result = deep_outputs[key]
        validate_model_output(key, result, markets)
        predictions[key] = result
    summaries, shared_drop_dates = summarize_losses(predictions)
    counts = count_mcs_memberships(predictions, shared_drop_dates)
    save_table(summaries, counts, bold_mode=args.bold)


if __name__ == "__main__":
    main()

