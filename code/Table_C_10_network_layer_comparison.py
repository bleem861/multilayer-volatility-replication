from pathlib import Path
import argparse
import warnings
import numpy as np
import pandas as pd
from functools import partial
from tqdm import tqdm as _tqdm
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, Border, Side

tqdm = partial(_tqdm, bar_format="{bar} {percentage:3.0f}%", ncols=50)
ROOT = Path(__file__).resolve().parents[1]
RV_FILE = ROOT / "data/raw/RV.csv"
RETURN_FILE = ROOT / "data/raw/Return.csv"
CENTRALITY_FILE = ROOT / "data/processed/centrality/P.csv"
QLIKE_DROP_DATES_FILE = ROOT / "data/processed/evaluation/qlike_intermediate.csv"
OUTPUT_FILE = ROOT / "results/generated/Table_C_10_network_layer_comparison.xlsx"
THRESH = 1e-3
SCALE_IF_MOST_GT_THRESH = 1e-4
EPS = 1e-12
RATIO_MIN = 1e-6
RATIO_MAX = 1e6
QLIKE_CAP = 100.0
POST_SCALE_THRESHOLD = 1e4
POST_SCALE_FACTOR = 1e-8


def run_har():
    model_outputs = {}
    import math
    import scipy as sc
    import time
    import os
    from scipy.stats import norm
    import pandas as pd
    import numpy as np
    from functools import partial
    from tqdm import tqdm as _tqdm
    tqdm = partial(_tqdm, bar_format='{bar} {percentage:3.0f}%', ncols=50)
    import statistics
    import statsmodels.formula.api as smf
    import os
    data = pd.read_csv(RV_FILE)
    num_cols = data.columns.drop('date')
    data[num_cols] = data[num_cols].apply(pd.to_numeric, errors='coerce')
    data = data.dropna()
    data['date'] = pd.to_datetime(data['date'], utc=True)
    data['date'] = data['date'].dt.tz_localize(None)
    data['date'] = data['date'].dt.normalize()
    data = data.set_index('date')
    error = []
    H = 22
    for c in data.columns:
        data_RV = data['%s' % c]
        data_RV = pd.DataFrame(data_RV)
        data_RV.columns = ['RV_now']
        data_RV['RV_D'] = data_RV['RV_now'].shift(1)
        data_RV['RV_W'] = data_RV['RV_D'].rolling(5).mean()
        data_RV['RV_M'] = data_RV['RV_D'].rolling(22).mean()
        data_RV['RV_MA22'] = data_RV['RV_now'].rolling(window=22, center=False).mean()
        data_RV['RV_MA22'] = data_RV['RV_MA22'].shift(-21)
        data_RV = data_RV.dropna()
        data_RV = data_RV.loc['2014-07-04':]
        time = '2018/05/02'
        time_stop = '2022-02-25'
        Loss_ALL = []
        data_test = data_RV.loc[time:time_stop]
        data_predict = data_RV.loc[time:time_stop]['RV_MA22']
        result = []
        for i in tqdm(range(len(data_RV.loc[time:time_stop]))):
            data_rolling = data_RV.iloc[i:i + (len(data_RV.loc[:time]) - 1)]
            mod = smf.ols('RV_MA22 ~ RV_D + RV_W + RV_M', data_rolling)
            res = mod.fit()
            params_new = res.params
            r_predict = params_new.iloc[0] + params_new.iloc[1] * data_test['RV_D'].iloc[i] + params_new.iloc[2] * data_test['RV_W'].iloc[i] + params_new.iloc[3] * data_test['RV_M'].iloc[i]
            result.append(r_predict)
        result = pd.Series(result, index=data_predict.index)
        SUM = pd.concat([data_predict, result], axis=1)
        SUM.columns = ['RV_True', 'RV_Predict']
        SUM_final = SUM
        SUM_final = SUM_final.reset_index(drop=True)
        SUM_final.index = pd.to_datetime(data_test.index)
        SUM_final['MSE'] = (SUM_final['RV_True'] - SUM_final['RV_Predict']) ** 2
        SUM_final['MAE'] = np.abs(SUM_final['RV_True'] - SUM_final['RV_Predict'])
        SUM_final['Qlike'] = SUM_final['RV_True'] / SUM_final['RV_Predict'] + np.log(SUM_final['RV_Predict'])
        MSE = np.mean((SUM_final['RV_True'] - SUM_final['RV_Predict']) ** 2)
        MAE = np.mean(np.abs(SUM_final['RV_True'] - SUM_final['RV_Predict']))
        Theil = np.sqrt(np.mean((SUM_final['RV_True'] - SUM_final['RV_Predict']) ** 2)) / (np.sqrt(np.mean(SUM_final['RV_Predict'] ** 2)) + np.sqrt(np.mean(SUM_final['RV_True'] ** 2)))
        Loss_ALL = [MSE, MAE, Theil]
        error.append(Loss_ALL)
        model_outputs[str(c).strip().lstrip('.')] = SUM_final.copy()
    ERROR = pd.DataFrame(error, columns=['MSE', 'MAE', 'Theil'], index=data.columns).T
    return model_outputs

def run_har_ci():
    model_outputs = {}
    import os
    import numpy as np
    import pandas as pd
    import statsmodels.formula.api as smf
    from functools import partial
    from tqdm import tqdm as _tqdm
    tqdm = partial(_tqdm, bar_format='{bar} {percentage:3.0f}%', ncols=50)
    RET_PATH = RETURN_FILE
    RV_PATH = RV_FILE
    P_STORE_PATH = CENTRALITY_FILE
    PLOT = False
    fit_end_date = pd.to_datetime('2018-05-01')
    pred_start = pd.to_datetime('2018-05-02')
    pred_end = pd.to_datetime('2022-02-25')
    H_LIST = [22]
    FIT_START = 1000
    FIT_STEP = 100
    FIT_MAX = 1
    L_fit_list = [FIT_START + FIT_STEP * k for k in range(FIT_MAX)]
    NET_MIN = 1000
    NET_MAX = 1000
    NET_STEP = 50
    L_net_list = list(range(NET_MIN, NET_MAX + 1, NET_STEP))

    def _norm_date(x):
        return pd.to_datetime(x).tz_localize(None).normalize()

    def make_future_mean(rv_now: pd.Series, H: int) -> pd.Series:
        if H == 1:
            return rv_now
        ma = rv_now.rolling(window=H, center=False).mean()
        return ma.shift(-(H - 1))

    def read_returns(path):
        df = pd.read_csv(path)
        df = df.loc[:, ~df.columns.str.startswith('Unnamed')]
        df['date'] = pd.to_datetime(df['date'].astype(str).str.strip().str[:10], errors='coerce')
        df['date'] = df['date'].dt.tz_localize(None).dt.normalize()
        df = df.set_index('date').dropna(how='all')
        return df

    def read_rv_no_scale(path):
        data = pd.read_csv(path)
        if 'date' not in data.columns:
            raise RuntimeError("RV file must contain a 'date' column.")
        num_cols = [c for c in data.columns if c != 'date']
        data[num_cols] = data[num_cols].apply(pd.to_numeric, errors='coerce')
        data = data.dropna(how='all')
        data['date'] = pd.to_datetime(data['date'], utc=True, errors='coerce')
        data = data.dropna(subset=['date'])
        data['date'] = data['date'].dt.tz_localize(None)
        data['date'] = data['date'].dt.normalize()
        data = data.set_index('date')
        data = data.dropna(axis=1, how='all')
        return data
    df_ret = read_returns(RET_PATH)
    data_rv = read_rv_no_scale(RV_PATH)
    markets = [c for c in data_rv.columns if c in df_ret.columns]
    P_all = pd.read_csv(P_STORE_PATH)
    need_cols = ['date', 'market', 'ci_start_date', 'ci_end_date', 'L_net', 'P_corr', 'P_kendall', 'P_partial', 'P_tail']
    miss = [c for c in need_cols if c not in P_all.columns]
    if miss:
        raise RuntimeError(f'P.csv is missing required columns: {miss}; available columns={list(P_all.columns)}')
    P_all['date'] = pd.to_datetime(P_all['date'], errors='coerce').dt.tz_localize(None).dt.normalize()
    P_all['ci_start_date'] = P_all['ci_start_date'].apply(_norm_date)
    P_all['ci_end_date'] = P_all['ci_end_date'].apply(_norm_date)
    P_all['L_net'] = pd.to_numeric(P_all['L_net'], errors='coerce').astype('Int64')
    P_all = P_all.dropna(subset=['date', 'market', 'ci_start_date', 'ci_end_date', 'L_net']).copy()
    P_cache = {}

    def load_P_from_csv(L_net: int, ci_start_date, ci_end_date):
        key = (int(L_net), _norm_date(ci_start_date), _norm_date(ci_end_date))
        if key in P_cache:
            return P_cache[key]
        ln, cs, ce = key
        sub = P_all[(P_all['L_net'] == ln) & (P_all['ci_start_date'] == cs) & (P_all['ci_end_date'] == ce)]
        if sub.empty:
            P_cache[key] = (None, None, None, None)
            return P_cache[key]

        def _pivot(col):
            df = sub.pivot_table(index='date', columns='market', values=col, aggfunc='mean')
            return df.sort_index()
        P_corr = _pivot('P_corr')
        P_kendall = _pivot('P_kendall')
        P_partial = _pivot('P_partial')
        P_tail = _pivot('P_tail')
        P_cache[key] = (P_corr, P_kendall, P_partial, P_tail)
        return P_cache[key]

    def forecast_one_market_oneH_harci(data_rv_all: pd.DataFrame, P_corr: pd.DataFrame, P_kendall: pd.DataFrame, P_partial: pd.DataFrame, P_tail: pd.DataFrame, market: str, fit_start_date: pd.Timestamp, H: int):
        df = data_rv_all[market].to_frame('RV_now')
        df['RV_D'] = df['RV_now'].shift(1)
        df['RV_W'] = df['RV_D'].rolling(5).mean()
        df['RV_M'] = df['RV_D'].rolling(22).mean()
        df['Y'] = make_future_mean(df['RV_now'], H)
        ci = pd.DataFrame({'CI_corr': P_corr[market], 'CI_kendall': P_kendall[market], 'CI_partial': P_partial[market], 'CI_tail': P_tail[market]}).reindex(df.index)
        df = df.join(ci, how='inner').dropna()
        df['CI_corr_D'] = df['CI_corr'].shift(1)
        df['CI_kendall_D'] = df['CI_kendall'].shift(1)
        df['CI_partial_D'] = df['CI_partial'].shift(1)
        df['CI_tail_D'] = df['CI_tail'].shift(1)
        df = df.dropna()
        df = df.loc[fit_start_date:]
        data_test = df.loc[pred_start:pred_end].copy()
        if len(data_test) == 0:
            return (None, np.nan, np.nan, np.nan)
        window_len = len(df.loc[:pred_start]) - 1
        if window_len <= 0:
            return (None, np.nan, np.nan, np.nan)
        y_true = data_test['Y'].astype(float).copy()
        idx = y_true.index
        preds = []
        for i in tqdm(range(len(idx)), desc=f'HAR-CI {market} H{H}', leave=False):
            train = df.iloc[i:i + window_len].copy()
            train['_Y_'] = train['Y'].astype(float).values
            res = smf.ols('_Y_ ~ RV_D + RV_W + RV_M + CI_corr_D + CI_kendall_D + CI_partial_D + CI_tail_D', train).fit()
            p = res.params
            pred = p.iloc[0] + p.iloc[1] * data_test['RV_D'].iloc[i] + p.iloc[2] * data_test['RV_W'].iloc[i] + p.iloc[3] * data_test['RV_M'].iloc[i] + p.iloc[4] * data_test['CI_corr_D'].iloc[i] + p.iloc[5] * data_test['CI_kendall_D'].iloc[i] + p.iloc[6] * data_test['CI_partial_D'].iloc[i] + p.iloc[7] * data_test['CI_tail_D'].iloc[i]
            preds.append(float(pred))
        y_pred = pd.Series(preds, index=idx)
        mcs = pd.concat([y_true.rename('RV_True'), y_pred.rename('RV_Predict')], axis=1)
        mcs['MSE'] = (mcs['RV_True'] - mcs['RV_Predict']) ** 2
        mcs['MAE'] = np.abs(mcs['RV_True'] - mcs['RV_Predict'])
        mse = float(np.nanmean(mcs['MSE'].values))
        mae = float(np.nanmean(mcs['MAE'].values))
        denom = np.sqrt(np.nanmean(mcs['RV_Predict'].values ** 2)) + np.sqrt(np.nanmean(mcs['RV_True'].values ** 2))
        theil = float(np.sqrt(np.nanmean((mcs['RV_True'].values - mcs['RV_Predict'].values) ** 2)) / denom) if denom != 0 else np.nan
        return (mcs, mse, mae, theil)
    L_fit = L_fit_list[0]
    L_net = L_net_list[0]
    for H in H_LIST:
        error_by_market = {}
        for mkt in markets:
            if fit_end_date not in data_rv.index:
                error_by_market[mkt] = [np.nan, np.nan, np.nan]
                continue
            fit_end_pos = data_rv.index.get_loc(fit_end_date)
            fit_start_pos = fit_end_pos - (L_fit - 1)
            if fit_start_pos < 0:
                error_by_market[mkt] = [np.nan, np.nan, np.nan]
                continue
            fit_start_date = data_rv.index[fit_start_pos]
            ret_pos_end = df_ret.index.searchsorted(fit_start_date) - 1
            if ret_pos_end < 0:
                error_by_market[mkt] = [np.nan, np.nan, np.nan]
                continue
            net_end_date = df_ret.index[ret_pos_end]
            ci_start_date = net_end_date
            ci_end_date = pred_end
            P_corr, P_kendall, P_partial, P_tail = load_P_from_csv(L_net=L_net, ci_start_date=ci_start_date, ci_end_date=ci_end_date)
            if P_corr is None:
                error_by_market[mkt] = [np.nan, np.nan, np.nan]
                continue
            if mkt not in P_corr.columns:
                error_by_market[mkt] = [np.nan, np.nan, np.nan]
                continue
            mcs, mse, mae, theil = forecast_one_market_oneH_harci(data_rv_all=data_rv, P_corr=P_corr, P_kendall=P_kendall, P_partial=P_partial, P_tail=P_tail, market=mkt, fit_start_date=fit_start_date, H=H)
            if mcs is None:
                tmp = pd.DataFrame(index=pd.DatetimeIndex([], name='date'))
                tmp['RV_True'] = np.nan
                tmp['RV_Predict'] = np.nan
                tmp['MSE'] = np.nan
                tmp['MAE'] = np.nan
                model_outputs[str(mkt).strip().lstrip('.')] = tmp.copy()
            else:
                model_outputs[str(mkt).strip().lstrip('.')] = mcs.copy()
            error_by_market[mkt] = [mse, mae, theil]
        ERROR = pd.DataFrame(error_by_market, index=['MSE', 'MAE', 'Theil'])
    return model_outputs

def calc_mse(y_true: np.ndarray, y_pred: np.ndarray) -> dict:
    y_true = np.asarray(y_true, dtype=float)
    y_pred = np.asarray(y_pred, dtype=float)
    mask = np.isfinite(y_true) & np.isfinite(y_pred)
    yt = y_true[mask]
    yp = y_pred[mask]
    if len(yt) == 0:
        return {'MSE': np.nan}
    err = yp - yt
    mse = float(np.mean(err ** 2))
    return {'MSE': mse}

def qlike_with_shared_drop(y_true: pd.Series, y_pred: pd.Series, drop_dates: set) -> float:
    df = pd.concat([y_true.rename('y'), y_pred.rename('yhat')], axis=1).dropna()
    if df.empty:
        return np.nan
    if drop_dates:
        df = df.loc[~df.index.isin(drop_dates)]
    if df.empty:
        return np.nan
    y = df['y'].astype(float).values
    yhat = df['yhat'].astype(float).values
    ok_pos = (y > EPS) & (yhat > EPS)
    if not np.any(ok_pos):
        return np.nan
    ratio = y[ok_pos] / yhat[ok_pos]
    ok_ratio = np.isfinite(ratio) & (ratio >= RATIO_MIN) & (ratio <= RATIO_MAX)
    if not np.any(ok_ratio):
        return np.nan
    ratio = ratio[ok_ratio]
    q = ratio - np.log(ratio) - 1.0
    ok_q = np.isfinite(q) & (q >= 0.0) & (q <= QLIKE_CAP)
    if not np.any(ok_q):
        return np.nan
    return float(np.mean(q[ok_q]))

def add_summary_rows(df: pd.DataFrame) -> pd.DataFrame:
    mean_row = df.mean(axis=0, skipna=True)
    q1_row = df.quantile(0.25, axis=0, interpolation='linear')
    q2_row = df.quantile(0.5, axis=0, interpolation='linear')
    q3_row = df.quantile(0.75, axis=0, interpolation='linear')
    df2 = df.copy()
    df2.loc['Mean'] = mean_row
    df2.loc['Q1'] = q1_row
    df2.loc['Q2'] = q2_row
    df2.loc['Q3'] = q3_row
    return df2

def load_prediction_pair(frame):
    dates = pd.to_datetime(frame.index, errors='coerce')
    y_true = pd.Series(pd.to_numeric(frame['RV_True'], errors='coerce').to_numpy(), index=dates)
    y_pred = pd.Series(pd.to_numeric(frame['RV_Predict'], errors='coerce').to_numpy(), index=dates)
    valid = dates.notna() & np.isfinite(y_true.to_numpy()) & np.isfinite(y_pred.to_numpy())
    y_true, y_pred = (y_true[valid], y_pred[valid])
    if len(y_true) and np.mean(y_true.to_numpy() > THRESH) > 0.5:
        y_true = y_true * SCALE_IF_MOST_GT_THRESH
        y_pred = y_pred * SCALE_IF_MOST_GT_THRESH
    y_true = y_true[~y_true.index.duplicated(keep='last')].sort_index()
    y_pred = y_pred[~y_pred.index.duplicated(keep='last')].sort_index()
    return (y_true, y_pred)

def load_shared_drop_dates(markets):
    data = pd.read_csv(QLIKE_DROP_DATES_FILE)
    required = {'market', 'date'}
    if not required.issubset(data.columns):
        raise ValueError(f'Missing QLIKE exclusion columns: {sorted(required - set(data.columns))}')
    data = data[['market', 'date']].copy()
    data['market'] = data['market'].astype(str).str.strip().str.lstrip('.')
    data['date'] = pd.to_datetime(data['date'], errors='coerce').dt.normalize()
    if data['date'].isna().any():
        raise ValueError('The QLIKE exclusion file contains invalid dates')
    if data.duplicated(['market', 'date']).any():
        raise ValueError('The QLIKE exclusion file contains duplicate market-date rows')
    file_markets = set(data['market'])
    expected_markets = set(markets)
    if file_markets != expected_markets:
        difference = sorted(file_markets.symmetric_difference(expected_markets))
        raise ValueError(f'QLIKE exclusion markets do not match the forecast markets: {difference}')
    return {market: set(group['date']) for market, group in data.groupby('market')}

def summarize_losses(predictions):
    models = [key for key, _, _ in MODEL_RUNNERS]
    markets = sorted(predictions['HAR-RV'])
    cache = {market: {model: load_prediction_pair(predictions[model][market]) for model in models} for market in markets}
    shared_drop_dates = load_shared_drop_dates(markets)
    raw = {metric: pd.DataFrame(index=markets, columns=models, dtype=float) for metric in ('MSE', 'QLIKE')}
    for market in markets:
        for model in models:
            y_true, y_pred = cache[market][model]
            raw['MSE'].loc[market, model] = calc_mse(y_true.to_numpy(), y_pred.to_numpy())['MSE']
            raw['QLIKE'].loc[market, model] = qlike_with_shared_drop(y_true, y_pred, shared_drop_dates[market])
    summaries = {}
    for metric, values in raw.items():
        values = values.dropna(how='all').dropna(subset=['HAR-RV'], how='all')
        relative = values.div(values['HAR-RV'].replace(0.0, np.nan), axis=0)
        if np.any(np.isfinite(relative.to_numpy()) & (np.abs(relative.to_numpy()) > POST_SCALE_THRESHOLD)):
            warnings.warn('The original relative-loss scaling rule was triggered for ' + metric, RuntimeWarning)
            relative = relative.where(~(np.abs(relative) > POST_SCALE_THRESHOLD), relative * POST_SCALE_FACTOR)
        summaries[metric] = add_summary_rows(relative).loc[['Mean', 'Q1', 'Q2']].T
    return (summaries, shared_drop_dates)

def run_layer_corr():
    model_outputs = {}
    import os
    import numpy as np
    import pandas as pd
    import statsmodels.formula.api as smf
    from functools import partial
    from tqdm import tqdm as _tqdm
    tqdm = partial(_tqdm, bar_format='{bar} {percentage:3.0f}%', ncols=50)
    RET_PATH = RETURN_FILE
    RV_PATH = RV_FILE
    P_STORE_PATH = CENTRALITY_FILE
    PLOT = False
    fit_end_date = pd.to_datetime('2018-05-01')
    pred_start = pd.to_datetime('2018-05-02')
    pred_end = pd.to_datetime('2022-02-25')
    H_LIST = [22]
    FIT_START = 1000
    FIT_STEP = 100
    FIT_MAX = 1
    L_fit_list = [FIT_START + FIT_STEP * k for k in range(FIT_MAX)]
    NET_MIN = 1000
    NET_MAX = 1000
    NET_STEP = 50
    L_net_list = list(range(NET_MIN, NET_MAX + 1, NET_STEP))

    def _norm_date(x):
        return pd.to_datetime(x).tz_localize(None).normalize()

    def make_future_mean(rv_now: pd.Series, H: int) -> pd.Series:
        if H == 1:
            return rv_now
        ma = rv_now.rolling(window=H, center=False).mean()
        return ma.shift(-(H - 1))

    def read_returns(path):
        df = pd.read_csv(path)
        df = df.loc[:, ~df.columns.str.startswith('Unnamed')]
        df['date'] = pd.to_datetime(df['date'].astype(str).str.strip().str[:10], errors='coerce')
        df['date'] = df['date'].dt.tz_localize(None).dt.normalize()
        df = df.set_index('date').dropna(how='all')
        return df

    def read_rv_no_scale(path):
        data = pd.read_csv(path)
        if 'date' not in data.columns:
            raise RuntimeError("RV file must contain a 'date' column.")
        num_cols = [c for c in data.columns if c != 'date']
        data[num_cols] = data[num_cols].apply(pd.to_numeric, errors='coerce')
        data = data.dropna(how='all')
        data['date'] = pd.to_datetime(data['date'], utc=True, errors='coerce')
        data = data.dropna(subset=['date'])
        data['date'] = data['date'].dt.tz_localize(None)
        data['date'] = data['date'].dt.normalize()
        data = data.set_index('date')
        data = data.dropna(axis=1, how='all')
        return data
    df_ret = read_returns(RET_PATH)
    data_rv = read_rv_no_scale(RV_PATH)
    markets = [c for c in data_rv.columns if c in df_ret.columns]
    P_all = pd.read_csv(P_STORE_PATH)
    need_cols = ['date', 'market', 'ci_start_date', 'ci_end_date', 'L_net', 'P_corr', 'P_kendall', 'P_partial', 'P_tail']
    miss = [c for c in need_cols if c not in P_all.columns]
    if miss:
        raise RuntimeError(f'P.csv is missing required columns: {miss}; available columns={list(P_all.columns)}')
    P_all['date'] = pd.to_datetime(P_all['date'], errors='coerce').dt.tz_localize(None).dt.normalize()
    P_all['ci_start_date'] = P_all['ci_start_date'].apply(_norm_date)
    P_all['ci_end_date'] = P_all['ci_end_date'].apply(_norm_date)
    P_all['L_net'] = pd.to_numeric(P_all['L_net'], errors='coerce').astype('Int64')
    P_all = P_all.dropna(subset=['date', 'market', 'ci_start_date', 'ci_end_date', 'L_net']).copy()
    P_cache = {}

    def load_P_from_csv(L_net: int, ci_start_date, ci_end_date):
        key = (int(L_net), _norm_date(ci_start_date), _norm_date(ci_end_date))
        if key in P_cache:
            return P_cache[key]
        ln, cs, ce = key
        sub = P_all[(P_all['L_net'] == ln) & (P_all['ci_start_date'] == cs) & (P_all['ci_end_date'] == ce)]
        if sub.empty:
            P_cache[key] = (None, None, None, None)
            return P_cache[key]

        def _pivot(col):
            df = sub.pivot_table(index='date', columns='market', values=col, aggfunc='mean')
            return df.sort_index()
        P_corr = _pivot('P_corr')
        P_kendall = _pivot('P_kendall')
        P_partial = _pivot('P_partial')
        P_tail = _pivot('P_tail')
        P_cache[key] = (P_corr, P_kendall, P_partial, P_tail)
        return P_cache[key]

    def forecast_one_market_oneH_harci(data_rv_all: pd.DataFrame, P_corr: pd.DataFrame, P_kendall: pd.DataFrame, P_partial: pd.DataFrame, P_tail: pd.DataFrame, market: str, fit_start_date: pd.Timestamp, H: int):
        df = data_rv_all[market].to_frame('RV_now')
        df['RV_D'] = df['RV_now'].shift(1)
        df['RV_W'] = df['RV_D'].rolling(5).mean()
        df['RV_M'] = df['RV_D'].rolling(22).mean()
        df['Y'] = make_future_mean(df['RV_now'], H)
        ci = pd.DataFrame({'CI_corr': P_corr[market], 'CI_kendall': P_kendall[market], 'CI_partial': P_partial[market], 'CI_tail': P_tail[market]}).reindex(df.index)
        df = df.join(ci, how='inner').dropna()
        df['CI_corr_D'] = df['CI_corr'].shift(1)
        df['CI_kendall_D'] = df['CI_kendall'].shift(1)
        df['CI_partial_D'] = df['CI_partial'].shift(1)
        df['CI_tail_D'] = df['CI_tail'].shift(1)
        df = df.dropna()
        df = df.loc[fit_start_date:]
        data_test = df.loc[pred_start:pred_end].copy()
        if len(data_test) == 0:
            return (None, np.nan, np.nan, np.nan)
        window_len = len(df.loc[:pred_start]) - 1
        if window_len <= 0:
            return (None, np.nan, np.nan, np.nan)
        y_true = data_test['Y'].astype(float).copy()
        idx = y_true.index
        preds = []
        for i in tqdm(range(len(idx)), desc=f'HAR-CI {market} H{H}', leave=False):
            train = df.iloc[i:i + window_len].copy()
            train['_Y_'] = train['Y'].astype(float).values
            res = smf.ols('_Y_ ~ RV_D + RV_W + RV_M + CI_corr_D', train).fit()
            p = res.params
            pred = p.iloc[0] + p.iloc[1] * data_test['RV_D'].iloc[i] + p.iloc[2] * data_test['RV_W'].iloc[i] + p.iloc[3] * data_test['RV_M'].iloc[i] + p.iloc[4] * data_test['CI_corr_D'].iloc[i]
            preds.append(float(pred))
        y_pred = pd.Series(preds, index=idx)
        mcs = pd.concat([y_true.rename('RV_True'), y_pred.rename('RV_Predict')], axis=1)
        mcs['MSE'] = (mcs['RV_True'] - mcs['RV_Predict']) ** 2
        mcs['MAE'] = np.abs(mcs['RV_True'] - mcs['RV_Predict'])
        mse = float(np.nanmean(mcs['MSE'].values))
        mae = float(np.nanmean(mcs['MAE'].values))
        denom = np.sqrt(np.nanmean(mcs['RV_Predict'].values ** 2)) + np.sqrt(np.nanmean(mcs['RV_True'].values ** 2))
        theil = float(np.sqrt(np.nanmean((mcs['RV_True'].values - mcs['RV_Predict'].values) ** 2)) / denom) if denom != 0 else np.nan
        return (mcs, mse, mae, theil)
    L_fit = L_fit_list[0]
    L_net = L_net_list[0]
    for H in H_LIST:
        error_by_market = {}
        for mkt in markets:
            if fit_end_date not in data_rv.index:
                error_by_market[mkt] = [np.nan, np.nan, np.nan]
                continue
            fit_end_pos = data_rv.index.get_loc(fit_end_date)
            fit_start_pos = fit_end_pos - (L_fit - 1)
            if fit_start_pos < 0:
                error_by_market[mkt] = [np.nan, np.nan, np.nan]
                continue
            fit_start_date = data_rv.index[fit_start_pos]
            ret_pos_end = df_ret.index.searchsorted(fit_start_date) - 1
            if ret_pos_end < 0:
                error_by_market[mkt] = [np.nan, np.nan, np.nan]
                continue
            net_end_date = df_ret.index[ret_pos_end]
            ci_start_date = net_end_date
            ci_end_date = pred_end
            P_corr, P_kendall, P_partial, P_tail = load_P_from_csv(L_net=L_net, ci_start_date=ci_start_date, ci_end_date=ci_end_date)
            if P_corr is None:
                error_by_market[mkt] = [np.nan, np.nan, np.nan]
                continue
            if mkt not in P_corr.columns:
                error_by_market[mkt] = [np.nan, np.nan, np.nan]
                continue
            mcs, mse, mae, theil = forecast_one_market_oneH_harci(data_rv_all=data_rv, P_corr=P_corr, P_kendall=P_kendall, P_partial=P_partial, P_tail=P_tail, market=mkt, fit_start_date=fit_start_date, H=H)
            if mcs is None:
                tmp = pd.DataFrame(index=pd.DatetimeIndex([], name='date'))
                tmp['RV_True'] = np.nan
                tmp['RV_Predict'] = np.nan
                tmp['MSE'] = np.nan
                tmp['MAE'] = np.nan
                model_outputs[str(mkt).strip().lstrip('.')] = tmp.copy()
            else:
                model_outputs[str(mkt).strip().lstrip('.')] = mcs.copy()
            error_by_market[mkt] = [mse, mae, theil]
        ERROR = pd.DataFrame(error_by_market, index=['MSE', 'MAE', 'Theil'])
    return model_outputs

def run_layer_kendall():
    model_outputs = {}
    import os
    import numpy as np
    import pandas as pd
    import statsmodels.formula.api as smf
    from functools import partial
    from tqdm import tqdm as _tqdm
    tqdm = partial(_tqdm, bar_format='{bar} {percentage:3.0f}%', ncols=50)
    RET_PATH = RETURN_FILE
    RV_PATH = RV_FILE
    P_STORE_PATH = CENTRALITY_FILE
    PLOT = False
    fit_end_date = pd.to_datetime('2018-05-01')
    pred_start = pd.to_datetime('2018-05-02')
    pred_end = pd.to_datetime('2022-02-25')
    H_LIST = [22]
    FIT_START = 1000
    FIT_STEP = 100
    FIT_MAX = 1
    L_fit_list = [FIT_START + FIT_STEP * k for k in range(FIT_MAX)]
    NET_MIN = 1000
    NET_MAX = 1000
    NET_STEP = 50
    L_net_list = list(range(NET_MIN, NET_MAX + 1, NET_STEP))

    def _norm_date(x):
        return pd.to_datetime(x).tz_localize(None).normalize()

    def make_future_mean(rv_now: pd.Series, H: int) -> pd.Series:
        if H == 1:
            return rv_now
        ma = rv_now.rolling(window=H, center=False).mean()
        return ma.shift(-(H - 1))

    def read_returns(path):
        df = pd.read_csv(path)
        df = df.loc[:, ~df.columns.str.startswith('Unnamed')]
        df['date'] = pd.to_datetime(df['date'].astype(str).str.strip().str[:10], errors='coerce')
        df['date'] = df['date'].dt.tz_localize(None).dt.normalize()
        df = df.set_index('date').dropna(how='all')
        return df

    def read_rv_no_scale(path):
        data = pd.read_csv(path)
        if 'date' not in data.columns:
            raise RuntimeError("RV file must contain a 'date' column.")
        num_cols = [c for c in data.columns if c != 'date']
        data[num_cols] = data[num_cols].apply(pd.to_numeric, errors='coerce')
        data = data.dropna(how='all')
        data['date'] = pd.to_datetime(data['date'], utc=True, errors='coerce')
        data = data.dropna(subset=['date'])
        data['date'] = data['date'].dt.tz_localize(None)
        data['date'] = data['date'].dt.normalize()
        data = data.set_index('date')
        data = data.dropna(axis=1, how='all')
        return data
    df_ret = read_returns(RET_PATH)
    data_rv = read_rv_no_scale(RV_PATH)
    markets = [c for c in data_rv.columns if c in df_ret.columns]
    P_all = pd.read_csv(P_STORE_PATH)
    need_cols = ['date', 'market', 'ci_start_date', 'ci_end_date', 'L_net', 'P_corr', 'P_kendall', 'P_partial', 'P_tail']
    miss = [c for c in need_cols if c not in P_all.columns]
    if miss:
        raise RuntimeError(f'P.csv is missing required columns: {miss}; available columns={list(P_all.columns)}')
    P_all['date'] = pd.to_datetime(P_all['date'], errors='coerce').dt.tz_localize(None).dt.normalize()
    P_all['ci_start_date'] = P_all['ci_start_date'].apply(_norm_date)
    P_all['ci_end_date'] = P_all['ci_end_date'].apply(_norm_date)
    P_all['L_net'] = pd.to_numeric(P_all['L_net'], errors='coerce').astype('Int64')
    P_all = P_all.dropna(subset=['date', 'market', 'ci_start_date', 'ci_end_date', 'L_net']).copy()
    P_cache = {}

    def load_P_from_csv(L_net: int, ci_start_date, ci_end_date):
        key = (int(L_net), _norm_date(ci_start_date), _norm_date(ci_end_date))
        if key in P_cache:
            return P_cache[key]
        ln, cs, ce = key
        sub = P_all[(P_all['L_net'] == ln) & (P_all['ci_start_date'] == cs) & (P_all['ci_end_date'] == ce)]
        if sub.empty:
            P_cache[key] = (None, None, None, None)
            return P_cache[key]

        def _pivot(col):
            df = sub.pivot_table(index='date', columns='market', values=col, aggfunc='mean')
            return df.sort_index()
        P_corr = _pivot('P_corr')
        P_kendall = _pivot('P_kendall')
        P_partial = _pivot('P_partial')
        P_tail = _pivot('P_tail')
        P_cache[key] = (P_corr, P_kendall, P_partial, P_tail)
        return P_cache[key]

    def forecast_one_market_oneH_harci(data_rv_all: pd.DataFrame, P_corr: pd.DataFrame, P_kendall: pd.DataFrame, P_partial: pd.DataFrame, P_tail: pd.DataFrame, market: str, fit_start_date: pd.Timestamp, H: int):
        df = data_rv_all[market].to_frame('RV_now')
        df['RV_D'] = df['RV_now'].shift(1)
        df['RV_W'] = df['RV_D'].rolling(5).mean()
        df['RV_M'] = df['RV_D'].rolling(22).mean()
        df['Y'] = make_future_mean(df['RV_now'], H)
        ci = pd.DataFrame({'CI_corr': P_corr[market], 'CI_kendall': P_kendall[market], 'CI_partial': P_partial[market], 'CI_tail': P_tail[market]}).reindex(df.index)
        df = df.join(ci, how='inner').dropna()
        df['CI_corr_D'] = df['CI_corr'].shift(1)
        df['CI_kendall_D'] = df['CI_kendall'].shift(1)
        df['CI_partial_D'] = df['CI_partial'].shift(1)
        df['CI_tail_D'] = df['CI_tail'].shift(1)
        df = df.dropna()
        df = df.loc[fit_start_date:]
        data_test = df.loc[pred_start:pred_end].copy()
        if len(data_test) == 0:
            return (None, np.nan, np.nan, np.nan)
        window_len = len(df.loc[:pred_start]) - 1
        if window_len <= 0:
            return (None, np.nan, np.nan, np.nan)
        y_true = data_test['Y'].astype(float).copy()
        idx = y_true.index
        preds = []
        for i in tqdm(range(len(idx)), desc=f'HAR-CI {market} H{H}', leave=False):
            train = df.iloc[i:i + window_len].copy()
            train['_Y_'] = train['Y'].astype(float).values
            res = smf.ols('_Y_ ~ RV_D + RV_W + RV_M + CI_kendall_D', train).fit()
            p = res.params
            pred = p.iloc[0] + p.iloc[1] * data_test['RV_D'].iloc[i] + p.iloc[2] * data_test['RV_W'].iloc[i] + p.iloc[3] * data_test['RV_M'].iloc[i] + p.iloc[4] * data_test['CI_kendall_D'].iloc[i]
            preds.append(float(pred))
        y_pred = pd.Series(preds, index=idx)
        mcs = pd.concat([y_true.rename('RV_True'), y_pred.rename('RV_Predict')], axis=1)
        mcs['MSE'] = (mcs['RV_True'] - mcs['RV_Predict']) ** 2
        mcs['MAE'] = np.abs(mcs['RV_True'] - mcs['RV_Predict'])
        mse = float(np.nanmean(mcs['MSE'].values))
        mae = float(np.nanmean(mcs['MAE'].values))
        denom = np.sqrt(np.nanmean(mcs['RV_Predict'].values ** 2)) + np.sqrt(np.nanmean(mcs['RV_True'].values ** 2))
        theil = float(np.sqrt(np.nanmean((mcs['RV_True'].values - mcs['RV_Predict'].values) ** 2)) / denom) if denom != 0 else np.nan
        return (mcs, mse, mae, theil)
    L_fit = L_fit_list[0]
    L_net = L_net_list[0]
    for H in H_LIST:
        error_by_market = {}
        for mkt in markets:
            if fit_end_date not in data_rv.index:
                error_by_market[mkt] = [np.nan, np.nan, np.nan]
                continue
            fit_end_pos = data_rv.index.get_loc(fit_end_date)
            fit_start_pos = fit_end_pos - (L_fit - 1)
            if fit_start_pos < 0:
                error_by_market[mkt] = [np.nan, np.nan, np.nan]
                continue
            fit_start_date = data_rv.index[fit_start_pos]
            ret_pos_end = df_ret.index.searchsorted(fit_start_date) - 1
            if ret_pos_end < 0:
                error_by_market[mkt] = [np.nan, np.nan, np.nan]
                continue
            net_end_date = df_ret.index[ret_pos_end]
            ci_start_date = net_end_date
            ci_end_date = pred_end
            P_corr, P_kendall, P_partial, P_tail = load_P_from_csv(L_net=L_net, ci_start_date=ci_start_date, ci_end_date=ci_end_date)
            if P_corr is None:
                error_by_market[mkt] = [np.nan, np.nan, np.nan]
                continue
            if mkt not in P_corr.columns:
                error_by_market[mkt] = [np.nan, np.nan, np.nan]
                continue
            mcs, mse, mae, theil = forecast_one_market_oneH_harci(data_rv_all=data_rv, P_corr=P_corr, P_kendall=P_kendall, P_partial=P_partial, P_tail=P_tail, market=mkt, fit_start_date=fit_start_date, H=H)
            if mcs is None:
                tmp = pd.DataFrame(index=pd.DatetimeIndex([], name='date'))
                tmp['RV_True'] = np.nan
                tmp['RV_Predict'] = np.nan
                tmp['MSE'] = np.nan
                tmp['MAE'] = np.nan
                model_outputs[str(mkt).strip().lstrip('.')] = tmp.copy()
            else:
                model_outputs[str(mkt).strip().lstrip('.')] = mcs.copy()
            error_by_market[mkt] = [mse, mae, theil]
        ERROR = pd.DataFrame(error_by_market, index=['MSE', 'MAE', 'Theil'])
    return model_outputs

def run_layer_partial():
    model_outputs = {}
    import os
    import numpy as np
    import pandas as pd
    import statsmodels.formula.api as smf
    from functools import partial
    from tqdm import tqdm as _tqdm
    tqdm = partial(_tqdm, bar_format='{bar} {percentage:3.0f}%', ncols=50)
    RET_PATH = RETURN_FILE
    RV_PATH = RV_FILE
    P_STORE_PATH = CENTRALITY_FILE
    PLOT = False
    fit_end_date = pd.to_datetime('2018-05-01')
    pred_start = pd.to_datetime('2018-05-02')
    pred_end = pd.to_datetime('2022-02-25')
    H_LIST = [22]
    FIT_START = 1000
    FIT_STEP = 100
    FIT_MAX = 1
    L_fit_list = [FIT_START + FIT_STEP * k for k in range(FIT_MAX)]
    NET_MIN = 1000
    NET_MAX = 1000
    NET_STEP = 50
    L_net_list = list(range(NET_MIN, NET_MAX + 1, NET_STEP))

    def _norm_date(x):
        return pd.to_datetime(x).tz_localize(None).normalize()

    def make_future_mean(rv_now: pd.Series, H: int) -> pd.Series:
        if H == 1:
            return rv_now
        ma = rv_now.rolling(window=H, center=False).mean()
        return ma.shift(-(H - 1))

    def read_returns(path):
        df = pd.read_csv(path)
        df = df.loc[:, ~df.columns.str.startswith('Unnamed')]
        df['date'] = pd.to_datetime(df['date'].astype(str).str.strip().str[:10], errors='coerce')
        df['date'] = df['date'].dt.tz_localize(None).dt.normalize()
        df = df.set_index('date').dropna(how='all')
        return df

    def read_rv_no_scale(path):
        data = pd.read_csv(path)
        if 'date' not in data.columns:
            raise RuntimeError("RV file must contain a 'date' column.")
        num_cols = [c for c in data.columns if c != 'date']
        data[num_cols] = data[num_cols].apply(pd.to_numeric, errors='coerce')
        data = data.dropna(how='all')
        data['date'] = pd.to_datetime(data['date'], utc=True, errors='coerce')
        data = data.dropna(subset=['date'])
        data['date'] = data['date'].dt.tz_localize(None)
        data['date'] = data['date'].dt.normalize()
        data = data.set_index('date')
        data = data.dropna(axis=1, how='all')
        return data
    df_ret = read_returns(RET_PATH)
    data_rv = read_rv_no_scale(RV_PATH)
    markets = [c for c in data_rv.columns if c in df_ret.columns]
    P_all = pd.read_csv(P_STORE_PATH)
    need_cols = ['date', 'market', 'ci_start_date', 'ci_end_date', 'L_net', 'P_corr', 'P_kendall', 'P_partial', 'P_tail']
    miss = [c for c in need_cols if c not in P_all.columns]
    if miss:
        raise RuntimeError(f'P.csv is missing required columns: {miss}; available columns={list(P_all.columns)}')
    P_all['date'] = pd.to_datetime(P_all['date'], errors='coerce').dt.tz_localize(None).dt.normalize()
    P_all['ci_start_date'] = P_all['ci_start_date'].apply(_norm_date)
    P_all['ci_end_date'] = P_all['ci_end_date'].apply(_norm_date)
    P_all['L_net'] = pd.to_numeric(P_all['L_net'], errors='coerce').astype('Int64')
    P_all = P_all.dropna(subset=['date', 'market', 'ci_start_date', 'ci_end_date', 'L_net']).copy()
    P_cache = {}

    def load_P_from_csv(L_net: int, ci_start_date, ci_end_date):
        key = (int(L_net), _norm_date(ci_start_date), _norm_date(ci_end_date))
        if key in P_cache:
            return P_cache[key]
        ln, cs, ce = key
        sub = P_all[(P_all['L_net'] == ln) & (P_all['ci_start_date'] == cs) & (P_all['ci_end_date'] == ce)]
        if sub.empty:
            P_cache[key] = (None, None, None, None)
            return P_cache[key]

        def _pivot(col):
            df = sub.pivot_table(index='date', columns='market', values=col, aggfunc='mean')
            return df.sort_index()
        P_corr = _pivot('P_corr')
        P_kendall = _pivot('P_kendall')
        P_partial = _pivot('P_partial')
        P_tail = _pivot('P_tail')
        P_cache[key] = (P_corr, P_kendall, P_partial, P_tail)
        return P_cache[key]

    def forecast_one_market_oneH_harci(data_rv_all: pd.DataFrame, P_corr: pd.DataFrame, P_kendall: pd.DataFrame, P_partial: pd.DataFrame, P_tail: pd.DataFrame, market: str, fit_start_date: pd.Timestamp, H: int):
        df = data_rv_all[market].to_frame('RV_now')
        df['RV_D'] = df['RV_now'].shift(1)
        df['RV_W'] = df['RV_D'].rolling(5).mean()
        df['RV_M'] = df['RV_D'].rolling(22).mean()
        df['Y'] = make_future_mean(df['RV_now'], H)
        ci = pd.DataFrame({'CI_corr': P_corr[market], 'CI_kendall': P_kendall[market], 'CI_partial': P_partial[market], 'CI_tail': P_tail[market]}).reindex(df.index)
        df = df.join(ci, how='inner').dropna()
        df['CI_corr_D'] = df['CI_corr'].shift(1)
        df['CI_kendall_D'] = df['CI_kendall'].shift(1)
        df['CI_partial_D'] = df['CI_partial'].shift(1)
        df['CI_tail_D'] = df['CI_tail'].shift(1)
        df = df.dropna()
        df = df.loc[fit_start_date:]
        data_test = df.loc[pred_start:pred_end].copy()
        if len(data_test) == 0:
            return (None, np.nan, np.nan, np.nan)
        window_len = len(df.loc[:pred_start]) - 1
        if window_len <= 0:
            return (None, np.nan, np.nan, np.nan)
        y_true = data_test['Y'].astype(float).copy()
        idx = y_true.index
        preds = []
        for i in tqdm(range(len(idx)), desc=f'HAR-CI {market} H{H}', leave=False):
            train = df.iloc[i:i + window_len].copy()
            train['_Y_'] = train['Y'].astype(float).values
            res = smf.ols('_Y_ ~ RV_D + RV_W + RV_M + CI_partial_D', train).fit()
            p = res.params
            pred = p.iloc[0] + p.iloc[1] * data_test['RV_D'].iloc[i] + p.iloc[2] * data_test['RV_W'].iloc[i] + p.iloc[3] * data_test['RV_M'].iloc[i] + p.iloc[4] * data_test['CI_partial_D'].iloc[i]
            preds.append(float(pred))
        y_pred = pd.Series(preds, index=idx)
        mcs = pd.concat([y_true.rename('RV_True'), y_pred.rename('RV_Predict')], axis=1)
        mcs['MSE'] = (mcs['RV_True'] - mcs['RV_Predict']) ** 2
        mcs['MAE'] = np.abs(mcs['RV_True'] - mcs['RV_Predict'])
        mse = float(np.nanmean(mcs['MSE'].values))
        mae = float(np.nanmean(mcs['MAE'].values))
        denom = np.sqrt(np.nanmean(mcs['RV_Predict'].values ** 2)) + np.sqrt(np.nanmean(mcs['RV_True'].values ** 2))
        theil = float(np.sqrt(np.nanmean((mcs['RV_True'].values - mcs['RV_Predict'].values) ** 2)) / denom) if denom != 0 else np.nan
        return (mcs, mse, mae, theil)
    L_fit = L_fit_list[0]
    L_net = L_net_list[0]
    for H in H_LIST:
        error_by_market = {}
        for mkt in markets:
            if fit_end_date not in data_rv.index:
                error_by_market[mkt] = [np.nan, np.nan, np.nan]
                continue
            fit_end_pos = data_rv.index.get_loc(fit_end_date)
            fit_start_pos = fit_end_pos - (L_fit - 1)
            if fit_start_pos < 0:
                error_by_market[mkt] = [np.nan, np.nan, np.nan]
                continue
            fit_start_date = data_rv.index[fit_start_pos]
            ret_pos_end = df_ret.index.searchsorted(fit_start_date) - 1
            if ret_pos_end < 0:
                error_by_market[mkt] = [np.nan, np.nan, np.nan]
                continue
            net_end_date = df_ret.index[ret_pos_end]
            ci_start_date = net_end_date
            ci_end_date = pred_end
            P_corr, P_kendall, P_partial, P_tail = load_P_from_csv(L_net=L_net, ci_start_date=ci_start_date, ci_end_date=ci_end_date)
            if P_corr is None:
                error_by_market[mkt] = [np.nan, np.nan, np.nan]
                continue
            if mkt not in P_corr.columns:
                error_by_market[mkt] = [np.nan, np.nan, np.nan]
                continue
            mcs, mse, mae, theil = forecast_one_market_oneH_harci(data_rv_all=data_rv, P_corr=P_corr, P_kendall=P_kendall, P_partial=P_partial, P_tail=P_tail, market=mkt, fit_start_date=fit_start_date, H=H)
            if mcs is None:
                tmp = pd.DataFrame(index=pd.DatetimeIndex([], name='date'))
                tmp['RV_True'] = np.nan
                tmp['RV_Predict'] = np.nan
                tmp['MSE'] = np.nan
                tmp['MAE'] = np.nan
                model_outputs[str(mkt).strip().lstrip('.')] = tmp.copy()
            else:
                model_outputs[str(mkt).strip().lstrip('.')] = mcs.copy()
            error_by_market[mkt] = [mse, mae, theil]
        ERROR = pd.DataFrame(error_by_market, index=['MSE', 'MAE', 'Theil'])
    return model_outputs

def run_layer_tail():
    model_outputs = {}
    import os
    import numpy as np
    import pandas as pd
    import statsmodels.formula.api as smf
    from functools import partial
    from tqdm import tqdm as _tqdm
    tqdm = partial(_tqdm, bar_format='{bar} {percentage:3.0f}%', ncols=50)
    RET_PATH = RETURN_FILE
    RV_PATH = RV_FILE
    P_STORE_PATH = CENTRALITY_FILE
    PLOT = False
    fit_end_date = pd.to_datetime('2018-05-01')
    pred_start = pd.to_datetime('2018-05-02')
    pred_end = pd.to_datetime('2022-02-25')
    H_LIST = [22]
    FIT_START = 1000
    FIT_STEP = 100
    FIT_MAX = 1
    L_fit_list = [FIT_START + FIT_STEP * k for k in range(FIT_MAX)]
    NET_MIN = 1000
    NET_MAX = 1000
    NET_STEP = 50
    L_net_list = list(range(NET_MIN, NET_MAX + 1, NET_STEP))

    def _norm_date(x):
        return pd.to_datetime(x).tz_localize(None).normalize()

    def make_future_mean(rv_now: pd.Series, H: int) -> pd.Series:
        if H == 1:
            return rv_now
        ma = rv_now.rolling(window=H, center=False).mean()
        return ma.shift(-(H - 1))

    def read_returns(path):
        df = pd.read_csv(path)
        df = df.loc[:, ~df.columns.str.startswith('Unnamed')]
        df['date'] = pd.to_datetime(df['date'].astype(str).str.strip().str[:10], errors='coerce')
        df['date'] = df['date'].dt.tz_localize(None).dt.normalize()
        df = df.set_index('date').dropna(how='all')
        return df

    def read_rv_no_scale(path):
        data = pd.read_csv(path)
        if 'date' not in data.columns:
            raise RuntimeError("RV file must contain a 'date' column.")
        num_cols = [c for c in data.columns if c != 'date']
        data[num_cols] = data[num_cols].apply(pd.to_numeric, errors='coerce')
        data = data.dropna(how='all')
        data['date'] = pd.to_datetime(data['date'], utc=True, errors='coerce')
        data = data.dropna(subset=['date'])
        data['date'] = data['date'].dt.tz_localize(None)
        data['date'] = data['date'].dt.normalize()
        data = data.set_index('date')
        data = data.dropna(axis=1, how='all')
        return data
    df_ret = read_returns(RET_PATH)
    data_rv = read_rv_no_scale(RV_PATH)
    markets = [c for c in data_rv.columns if c in df_ret.columns]
    P_all = pd.read_csv(P_STORE_PATH)
    need_cols = ['date', 'market', 'ci_start_date', 'ci_end_date', 'L_net', 'P_corr', 'P_kendall', 'P_partial', 'P_tail']
    miss = [c for c in need_cols if c not in P_all.columns]
    if miss:
        raise RuntimeError(f'P.csv is missing required columns: {miss}; available columns={list(P_all.columns)}')
    P_all['date'] = pd.to_datetime(P_all['date'], errors='coerce').dt.tz_localize(None).dt.normalize()
    P_all['ci_start_date'] = P_all['ci_start_date'].apply(_norm_date)
    P_all['ci_end_date'] = P_all['ci_end_date'].apply(_norm_date)
    P_all['L_net'] = pd.to_numeric(P_all['L_net'], errors='coerce').astype('Int64')
    P_all = P_all.dropna(subset=['date', 'market', 'ci_start_date', 'ci_end_date', 'L_net']).copy()
    P_cache = {}

    def load_P_from_csv(L_net: int, ci_start_date, ci_end_date):
        key = (int(L_net), _norm_date(ci_start_date), _norm_date(ci_end_date))
        if key in P_cache:
            return P_cache[key]
        ln, cs, ce = key
        sub = P_all[(P_all['L_net'] == ln) & (P_all['ci_start_date'] == cs) & (P_all['ci_end_date'] == ce)]
        if sub.empty:
            P_cache[key] = (None, None, None, None)
            return P_cache[key]

        def _pivot(col):
            df = sub.pivot_table(index='date', columns='market', values=col, aggfunc='mean')
            return df.sort_index()
        P_corr = _pivot('P_corr')
        P_kendall = _pivot('P_kendall')
        P_partial = _pivot('P_partial')
        P_tail = _pivot('P_tail')
        P_cache[key] = (P_corr, P_kendall, P_partial, P_tail)
        return P_cache[key]

    def forecast_one_market_oneH_harci(data_rv_all: pd.DataFrame, P_corr: pd.DataFrame, P_kendall: pd.DataFrame, P_partial: pd.DataFrame, P_tail: pd.DataFrame, market: str, fit_start_date: pd.Timestamp, H: int):
        df = data_rv_all[market].to_frame('RV_now')
        df['RV_D'] = df['RV_now'].shift(1)
        df['RV_W'] = df['RV_D'].rolling(5).mean()
        df['RV_M'] = df['RV_D'].rolling(22).mean()
        df['Y'] = make_future_mean(df['RV_now'], H)
        ci = pd.DataFrame({'CI_corr': P_corr[market], 'CI_kendall': P_kendall[market], 'CI_partial': P_partial[market], 'CI_tail': P_tail[market]}).reindex(df.index)
        df = df.join(ci, how='inner').dropna()
        df['CI_corr_D'] = df['CI_corr'].shift(1)
        df['CI_kendall_D'] = df['CI_kendall'].shift(1)
        df['CI_partial_D'] = df['CI_partial'].shift(1)
        df['CI_tail_D'] = df['CI_tail'].shift(1)
        df = df.dropna()
        df = df.loc[fit_start_date:]
        data_test = df.loc[pred_start:pred_end].copy()
        if len(data_test) == 0:
            return (None, np.nan, np.nan, np.nan)
        window_len = len(df.loc[:pred_start]) - 1
        if window_len <= 0:
            return (None, np.nan, np.nan, np.nan)
        y_true = data_test['Y'].astype(float).copy()
        idx = y_true.index
        preds = []
        for i in tqdm(range(len(idx)), desc=f'HAR-CI {market} H{H}', leave=False):
            train = df.iloc[i:i + window_len].copy()
            train['_Y_'] = train['Y'].astype(float).values
            res = smf.ols('_Y_ ~ RV_D + RV_W + RV_M + CI_tail_D', train).fit()
            p = res.params
            pred = p.iloc[0] + p.iloc[1] * data_test['RV_D'].iloc[i] + p.iloc[2] * data_test['RV_W'].iloc[i] + p.iloc[3] * data_test['RV_M'].iloc[i] + p.iloc[4] * data_test['CI_tail_D'].iloc[i]
            preds.append(float(pred))
        y_pred = pd.Series(preds, index=idx)
        mcs = pd.concat([y_true.rename('RV_True'), y_pred.rename('RV_Predict')], axis=1)
        mcs['MSE'] = (mcs['RV_True'] - mcs['RV_Predict']) ** 2
        mcs['MAE'] = np.abs(mcs['RV_True'] - mcs['RV_Predict'])
        mse = float(np.nanmean(mcs['MSE'].values))
        mae = float(np.nanmean(mcs['MAE'].values))
        denom = np.sqrt(np.nanmean(mcs['RV_Predict'].values ** 2)) + np.sqrt(np.nanmean(mcs['RV_True'].values ** 2))
        theil = float(np.sqrt(np.nanmean((mcs['RV_True'].values - mcs['RV_Predict'].values) ** 2)) / denom) if denom != 0 else np.nan
        return (mcs, mse, mae, theil)
    L_fit = L_fit_list[0]
    L_net = L_net_list[0]
    for H in H_LIST:
        error_by_market = {}
        for mkt in markets:
            if fit_end_date not in data_rv.index:
                error_by_market[mkt] = [np.nan, np.nan, np.nan]
                continue
            fit_end_pos = data_rv.index.get_loc(fit_end_date)
            fit_start_pos = fit_end_pos - (L_fit - 1)
            if fit_start_pos < 0:
                error_by_market[mkt] = [np.nan, np.nan, np.nan]
                continue
            fit_start_date = data_rv.index[fit_start_pos]
            ret_pos_end = df_ret.index.searchsorted(fit_start_date) - 1
            if ret_pos_end < 0:
                error_by_market[mkt] = [np.nan, np.nan, np.nan]
                continue
            net_end_date = df_ret.index[ret_pos_end]
            ci_start_date = net_end_date
            ci_end_date = pred_end
            P_corr, P_kendall, P_partial, P_tail = load_P_from_csv(L_net=L_net, ci_start_date=ci_start_date, ci_end_date=ci_end_date)
            if P_corr is None:
                error_by_market[mkt] = [np.nan, np.nan, np.nan]
                continue
            if mkt not in P_corr.columns:
                error_by_market[mkt] = [np.nan, np.nan, np.nan]
                continue
            mcs, mse, mae, theil = forecast_one_market_oneH_harci(data_rv_all=data_rv, P_corr=P_corr, P_kendall=P_kendall, P_partial=P_partial, P_tail=P_tail, market=mkt, fit_start_date=fit_start_date, H=H)
            if mcs is None:
                tmp = pd.DataFrame(index=pd.DatetimeIndex([], name='date'))
                tmp['RV_True'] = np.nan
                tmp['RV_Predict'] = np.nan
                tmp['MSE'] = np.nan
                tmp['MAE'] = np.nan
                model_outputs[str(mkt).strip().lstrip('.')] = tmp.copy()
            else:
                model_outputs[str(mkt).strip().lstrip('.')] = mcs.copy()
            error_by_market[mkt] = [mse, mae, theil]
        ERROR = pd.DataFrame(error_by_market, index=['MSE', 'MAE', 'Theil'])
    return model_outputs

def run_layer_corr_kendall_partial():
    model_outputs = {}
    import os
    import numpy as np
    import pandas as pd
    import statsmodels.formula.api as smf
    from functools import partial
    from tqdm import tqdm as _tqdm
    tqdm = partial(_tqdm, bar_format='{bar} {percentage:3.0f}%', ncols=50)
    RET_PATH = RETURN_FILE
    RV_PATH = RV_FILE
    P_STORE_PATH = CENTRALITY_FILE
    PLOT = False
    fit_end_date = pd.to_datetime('2018-05-01')
    pred_start = pd.to_datetime('2018-05-02')
    pred_end = pd.to_datetime('2022-02-25')
    H_LIST = [22]
    FIT_START = 1000
    FIT_STEP = 100
    FIT_MAX = 1
    L_fit_list = [FIT_START + FIT_STEP * k for k in range(FIT_MAX)]
    NET_MIN = 1000
    NET_MAX = 1000
    NET_STEP = 50
    L_net_list = list(range(NET_MIN, NET_MAX + 1, NET_STEP))

    def _norm_date(x):
        return pd.to_datetime(x).tz_localize(None).normalize()

    def make_future_mean(rv_now: pd.Series, H: int) -> pd.Series:
        if H == 1:
            return rv_now
        ma = rv_now.rolling(window=H, center=False).mean()
        return ma.shift(-(H - 1))

    def read_returns(path):
        df = pd.read_csv(path)
        df = df.loc[:, ~df.columns.str.startswith('Unnamed')]
        df['date'] = pd.to_datetime(df['date'].astype(str).str.strip().str[:10], errors='coerce')
        df['date'] = df['date'].dt.tz_localize(None).dt.normalize()
        df = df.set_index('date').dropna(how='all')
        return df

    def read_rv_no_scale(path):
        data = pd.read_csv(path)
        if 'date' not in data.columns:
            raise RuntimeError("RV file must contain a 'date' column.")
        num_cols = [c for c in data.columns if c != 'date']
        data[num_cols] = data[num_cols].apply(pd.to_numeric, errors='coerce')
        data = data.dropna(how='all')
        data['date'] = pd.to_datetime(data['date'], utc=True, errors='coerce')
        data = data.dropna(subset=['date'])
        data['date'] = data['date'].dt.tz_localize(None)
        data['date'] = data['date'].dt.normalize()
        data = data.set_index('date')
        data = data.dropna(axis=1, how='all')
        return data
    df_ret = read_returns(RET_PATH)
    data_rv = read_rv_no_scale(RV_PATH)
    markets = [c for c in data_rv.columns if c in df_ret.columns]
    P_all = pd.read_csv(P_STORE_PATH)
    need_cols = ['date', 'market', 'ci_start_date', 'ci_end_date', 'L_net', 'P_corr', 'P_kendall', 'P_partial', 'P_tail']
    miss = [c for c in need_cols if c not in P_all.columns]
    if miss:
        raise RuntimeError(f'P.csv is missing required columns: {miss}; available columns={list(P_all.columns)}')
    P_all['date'] = pd.to_datetime(P_all['date'], errors='coerce').dt.tz_localize(None).dt.normalize()
    P_all['ci_start_date'] = P_all['ci_start_date'].apply(_norm_date)
    P_all['ci_end_date'] = P_all['ci_end_date'].apply(_norm_date)
    P_all['L_net'] = pd.to_numeric(P_all['L_net'], errors='coerce').astype('Int64')
    P_all = P_all.dropna(subset=['date', 'market', 'ci_start_date', 'ci_end_date', 'L_net']).copy()
    P_cache = {}

    def load_P_from_csv(L_net: int, ci_start_date, ci_end_date):
        key = (int(L_net), _norm_date(ci_start_date), _norm_date(ci_end_date))
        if key in P_cache:
            return P_cache[key]
        ln, cs, ce = key
        sub = P_all[(P_all['L_net'] == ln) & (P_all['ci_start_date'] == cs) & (P_all['ci_end_date'] == ce)]
        if sub.empty:
            P_cache[key] = (None, None, None, None)
            return P_cache[key]

        def _pivot(col):
            df = sub.pivot_table(index='date', columns='market', values=col, aggfunc='mean')
            return df.sort_index()
        P_corr = _pivot('P_corr')
        P_kendall = _pivot('P_kendall')
        P_partial = _pivot('P_partial')
        P_tail = _pivot('P_tail')
        P_cache[key] = (P_corr, P_kendall, P_partial, P_tail)
        return P_cache[key]

    def forecast_one_market_oneH_harci(data_rv_all: pd.DataFrame, P_corr: pd.DataFrame, P_kendall: pd.DataFrame, P_partial: pd.DataFrame, P_tail: pd.DataFrame, market: str, fit_start_date: pd.Timestamp, H: int):
        df = data_rv_all[market].to_frame('RV_now')
        df['RV_D'] = df['RV_now'].shift(1)
        df['RV_W'] = df['RV_D'].rolling(5).mean()
        df['RV_M'] = df['RV_D'].rolling(22).mean()
        df['Y'] = make_future_mean(df['RV_now'], H)
        ci = pd.DataFrame({'CI_corr': P_corr[market], 'CI_kendall': P_kendall[market], 'CI_partial': P_partial[market], 'CI_tail': P_tail[market]}).reindex(df.index)
        df = df.join(ci, how='inner').dropna()
        df['CI_corr_D'] = df['CI_corr'].shift(1)
        df['CI_kendall_D'] = df['CI_kendall'].shift(1)
        df['CI_partial_D'] = df['CI_partial'].shift(1)
        df['CI_tail_D'] = df['CI_tail'].shift(1)
        df = df.dropna()
        df = df.loc[fit_start_date:]
        data_test = df.loc[pred_start:pred_end].copy()
        if len(data_test) == 0:
            return (None, np.nan, np.nan, np.nan)
        window_len = len(df.loc[:pred_start]) - 1
        if window_len <= 0:
            return (None, np.nan, np.nan, np.nan)
        y_true = data_test['Y'].astype(float).copy()
        idx = y_true.index
        preds = []
        for i in tqdm(range(len(idx)), desc=f'HAR-CI {market} H{H}', leave=False):
            train = df.iloc[i:i + window_len].copy()
            train['_Y_'] = train['Y'].astype(float).values
            res = smf.ols('_Y_ ~ RV_D + RV_W + RV_M + CI_corr_D + CI_kendall_D + CI_partial_D', train).fit()
            p = res.params
            pred = p.iloc[0] + p.iloc[1] * data_test['RV_D'].iloc[i] + p.iloc[2] * data_test['RV_W'].iloc[i] + p.iloc[3] * data_test['RV_M'].iloc[i] + p.iloc[4] * data_test['CI_corr_D'].iloc[i] + p.iloc[5] * data_test['CI_kendall_D'].iloc[i] + p.iloc[6] * data_test['CI_partial_D'].iloc[i]
            preds.append(float(pred))
        y_pred = pd.Series(preds, index=idx)
        mcs = pd.concat([y_true.rename('RV_True'), y_pred.rename('RV_Predict')], axis=1)
        mcs['MSE'] = (mcs['RV_True'] - mcs['RV_Predict']) ** 2
        mcs['MAE'] = np.abs(mcs['RV_True'] - mcs['RV_Predict'])
        mse = float(np.nanmean(mcs['MSE'].values))
        mae = float(np.nanmean(mcs['MAE'].values))
        denom = np.sqrt(np.nanmean(mcs['RV_Predict'].values ** 2)) + np.sqrt(np.nanmean(mcs['RV_True'].values ** 2))
        theil = float(np.sqrt(np.nanmean((mcs['RV_True'].values - mcs['RV_Predict'].values) ** 2)) / denom) if denom != 0 else np.nan
        return (mcs, mse, mae, theil)
    L_fit = L_fit_list[0]
    L_net = L_net_list[0]
    for H in H_LIST:
        error_by_market = {}
        for mkt in markets:
            if fit_end_date not in data_rv.index:
                error_by_market[mkt] = [np.nan, np.nan, np.nan]
                continue
            fit_end_pos = data_rv.index.get_loc(fit_end_date)
            fit_start_pos = fit_end_pos - (L_fit - 1)
            if fit_start_pos < 0:
                error_by_market[mkt] = [np.nan, np.nan, np.nan]
                continue
            fit_start_date = data_rv.index[fit_start_pos]
            ret_pos_end = df_ret.index.searchsorted(fit_start_date) - 1
            if ret_pos_end < 0:
                error_by_market[mkt] = [np.nan, np.nan, np.nan]
                continue
            net_end_date = df_ret.index[ret_pos_end]
            ci_start_date = net_end_date
            ci_end_date = pred_end
            P_corr, P_kendall, P_partial, P_tail = load_P_from_csv(L_net=L_net, ci_start_date=ci_start_date, ci_end_date=ci_end_date)
            if P_corr is None:
                error_by_market[mkt] = [np.nan, np.nan, np.nan]
                continue
            if mkt not in P_corr.columns:
                error_by_market[mkt] = [np.nan, np.nan, np.nan]
                continue
            mcs, mse, mae, theil = forecast_one_market_oneH_harci(data_rv_all=data_rv, P_corr=P_corr, P_kendall=P_kendall, P_partial=P_partial, P_tail=P_tail, market=mkt, fit_start_date=fit_start_date, H=H)
            if mcs is None:
                tmp = pd.DataFrame(index=pd.DatetimeIndex([], name='date'))
                tmp['RV_True'] = np.nan
                tmp['RV_Predict'] = np.nan
                tmp['MSE'] = np.nan
                tmp['MAE'] = np.nan
                model_outputs[str(mkt).strip().lstrip('.')] = tmp.copy()
            else:
                model_outputs[str(mkt).strip().lstrip('.')] = mcs.copy()
            error_by_market[mkt] = [mse, mae, theil]
        ERROR = pd.DataFrame(error_by_market, index=['MSE', 'MAE', 'Theil'])
    return model_outputs

def run_layer_corr_kendall_tail():
    model_outputs = {}
    import os
    import numpy as np
    import pandas as pd
    import statsmodels.formula.api as smf
    from functools import partial
    from tqdm import tqdm as _tqdm
    tqdm = partial(_tqdm, bar_format='{bar} {percentage:3.0f}%', ncols=50)
    RET_PATH = RETURN_FILE
    RV_PATH = RV_FILE
    P_STORE_PATH = CENTRALITY_FILE
    PLOT = False
    fit_end_date = pd.to_datetime('2018-05-01')
    pred_start = pd.to_datetime('2018-05-02')
    pred_end = pd.to_datetime('2022-02-25')
    H_LIST = [22]
    FIT_START = 1000
    FIT_STEP = 100
    FIT_MAX = 1
    L_fit_list = [FIT_START + FIT_STEP * k for k in range(FIT_MAX)]
    NET_MIN = 1000
    NET_MAX = 1000
    NET_STEP = 50
    L_net_list = list(range(NET_MIN, NET_MAX + 1, NET_STEP))

    def _norm_date(x):
        return pd.to_datetime(x).tz_localize(None).normalize()

    def make_future_mean(rv_now: pd.Series, H: int) -> pd.Series:
        if H == 1:
            return rv_now
        ma = rv_now.rolling(window=H, center=False).mean()
        return ma.shift(-(H - 1))

    def read_returns(path):
        df = pd.read_csv(path)
        df = df.loc[:, ~df.columns.str.startswith('Unnamed')]
        df['date'] = pd.to_datetime(df['date'].astype(str).str.strip().str[:10], errors='coerce')
        df['date'] = df['date'].dt.tz_localize(None).dt.normalize()
        df = df.set_index('date').dropna(how='all')
        return df

    def read_rv_no_scale(path):
        data = pd.read_csv(path)
        if 'date' not in data.columns:
            raise RuntimeError("RV file must contain a 'date' column.")
        num_cols = [c for c in data.columns if c != 'date']
        data[num_cols] = data[num_cols].apply(pd.to_numeric, errors='coerce')
        data = data.dropna(how='all')
        data['date'] = pd.to_datetime(data['date'], utc=True, errors='coerce')
        data = data.dropna(subset=['date'])
        data['date'] = data['date'].dt.tz_localize(None)
        data['date'] = data['date'].dt.normalize()
        data = data.set_index('date')
        data = data.dropna(axis=1, how='all')
        return data
    df_ret = read_returns(RET_PATH)
    data_rv = read_rv_no_scale(RV_PATH)
    markets = [c for c in data_rv.columns if c in df_ret.columns]
    P_all = pd.read_csv(P_STORE_PATH)
    need_cols = ['date', 'market', 'ci_start_date', 'ci_end_date', 'L_net', 'P_corr', 'P_kendall', 'P_partial', 'P_tail']
    miss = [c for c in need_cols if c not in P_all.columns]
    if miss:
        raise RuntimeError(f'P.csv is missing required columns: {miss}; available columns={list(P_all.columns)}')
    P_all['date'] = pd.to_datetime(P_all['date'], errors='coerce').dt.tz_localize(None).dt.normalize()
    P_all['ci_start_date'] = P_all['ci_start_date'].apply(_norm_date)
    P_all['ci_end_date'] = P_all['ci_end_date'].apply(_norm_date)
    P_all['L_net'] = pd.to_numeric(P_all['L_net'], errors='coerce').astype('Int64')
    P_all = P_all.dropna(subset=['date', 'market', 'ci_start_date', 'ci_end_date', 'L_net']).copy()
    P_cache = {}

    def load_P_from_csv(L_net: int, ci_start_date, ci_end_date):
        key = (int(L_net), _norm_date(ci_start_date), _norm_date(ci_end_date))
        if key in P_cache:
            return P_cache[key]
        ln, cs, ce = key
        sub = P_all[(P_all['L_net'] == ln) & (P_all['ci_start_date'] == cs) & (P_all['ci_end_date'] == ce)]
        if sub.empty:
            P_cache[key] = (None, None, None, None)
            return P_cache[key]

        def _pivot(col):
            df = sub.pivot_table(index='date', columns='market', values=col, aggfunc='mean')
            return df.sort_index()
        P_corr = _pivot('P_corr')
        P_kendall = _pivot('P_kendall')
        P_partial = _pivot('P_partial')
        P_tail = _pivot('P_tail')
        P_cache[key] = (P_corr, P_kendall, P_partial, P_tail)
        return P_cache[key]

    def forecast_one_market_oneH_harci(data_rv_all: pd.DataFrame, P_corr: pd.DataFrame, P_kendall: pd.DataFrame, P_partial: pd.DataFrame, P_tail: pd.DataFrame, market: str, fit_start_date: pd.Timestamp, H: int):
        df = data_rv_all[market].to_frame('RV_now')
        df['RV_D'] = df['RV_now'].shift(1)
        df['RV_W'] = df['RV_D'].rolling(5).mean()
        df['RV_M'] = df['RV_D'].rolling(22).mean()
        df['Y'] = make_future_mean(df['RV_now'], H)
        ci = pd.DataFrame({'CI_corr': P_corr[market], 'CI_kendall': P_kendall[market], 'CI_partial': P_partial[market], 'CI_tail': P_tail[market]}).reindex(df.index)
        df = df.join(ci, how='inner').dropna()
        df['CI_corr_D'] = df['CI_corr'].shift(1)
        df['CI_kendall_D'] = df['CI_kendall'].shift(1)
        df['CI_partial_D'] = df['CI_partial'].shift(1)
        df['CI_tail_D'] = df['CI_tail'].shift(1)
        df = df.dropna()
        df = df.loc[fit_start_date:]
        data_test = df.loc[pred_start:pred_end].copy()
        if len(data_test) == 0:
            return (None, np.nan, np.nan, np.nan)
        window_len = len(df.loc[:pred_start]) - 1
        if window_len <= 0:
            return (None, np.nan, np.nan, np.nan)
        y_true = data_test['Y'].astype(float).copy()
        idx = y_true.index
        preds = []
        for i in tqdm(range(len(idx)), desc=f'HAR-CI {market} H{H}', leave=False):
            train = df.iloc[i:i + window_len].copy()
            train['_Y_'] = train['Y'].astype(float).values
            res = smf.ols('_Y_ ~ RV_D + RV_W + RV_M + CI_corr_D + CI_kendall_D + CI_tail_D', train).fit()
            p = res.params
            pred = p.iloc[0] + p.iloc[1] * data_test['RV_D'].iloc[i] + p.iloc[2] * data_test['RV_W'].iloc[i] + p.iloc[3] * data_test['RV_M'].iloc[i] + p.iloc[4] * data_test['CI_corr_D'].iloc[i] + p.iloc[5] * data_test['CI_kendall_D'].iloc[i] + p.iloc[6] * data_test['CI_tail_D'].iloc[i]
            preds.append(float(pred))
        y_pred = pd.Series(preds, index=idx)
        mcs = pd.concat([y_true.rename('RV_True'), y_pred.rename('RV_Predict')], axis=1)
        mcs['MSE'] = (mcs['RV_True'] - mcs['RV_Predict']) ** 2
        mcs['MAE'] = np.abs(mcs['RV_True'] - mcs['RV_Predict'])
        mse = float(np.nanmean(mcs['MSE'].values))
        mae = float(np.nanmean(mcs['MAE'].values))
        denom = np.sqrt(np.nanmean(mcs['RV_Predict'].values ** 2)) + np.sqrt(np.nanmean(mcs['RV_True'].values ** 2))
        theil = float(np.sqrt(np.nanmean((mcs['RV_True'].values - mcs['RV_Predict'].values) ** 2)) / denom) if denom != 0 else np.nan
        return (mcs, mse, mae, theil)
    L_fit = L_fit_list[0]
    L_net = L_net_list[0]
    for H in H_LIST:
        error_by_market = {}
        for mkt in markets:
            if fit_end_date not in data_rv.index:
                error_by_market[mkt] = [np.nan, np.nan, np.nan]
                continue
            fit_end_pos = data_rv.index.get_loc(fit_end_date)
            fit_start_pos = fit_end_pos - (L_fit - 1)
            if fit_start_pos < 0:
                error_by_market[mkt] = [np.nan, np.nan, np.nan]
                continue
            fit_start_date = data_rv.index[fit_start_pos]
            ret_pos_end = df_ret.index.searchsorted(fit_start_date) - 1
            if ret_pos_end < 0:
                error_by_market[mkt] = [np.nan, np.nan, np.nan]
                continue
            net_end_date = df_ret.index[ret_pos_end]
            ci_start_date = net_end_date
            ci_end_date = pred_end
            P_corr, P_kendall, P_partial, P_tail = load_P_from_csv(L_net=L_net, ci_start_date=ci_start_date, ci_end_date=ci_end_date)
            if P_corr is None:
                error_by_market[mkt] = [np.nan, np.nan, np.nan]
                continue
            if mkt not in P_corr.columns:
                error_by_market[mkt] = [np.nan, np.nan, np.nan]
                continue
            mcs, mse, mae, theil = forecast_one_market_oneH_harci(data_rv_all=data_rv, P_corr=P_corr, P_kendall=P_kendall, P_partial=P_partial, P_tail=P_tail, market=mkt, fit_start_date=fit_start_date, H=H)
            if mcs is None:
                tmp = pd.DataFrame(index=pd.DatetimeIndex([], name='date'))
                tmp['RV_True'] = np.nan
                tmp['RV_Predict'] = np.nan
                tmp['MSE'] = np.nan
                tmp['MAE'] = np.nan
                model_outputs[str(mkt).strip().lstrip('.')] = tmp.copy()
            else:
                model_outputs[str(mkt).strip().lstrip('.')] = mcs.copy()
            error_by_market[mkt] = [mse, mae, theil]
        ERROR = pd.DataFrame(error_by_market, index=['MSE', 'MAE', 'Theil'])
    return model_outputs

def run_layer_corr_tail_partial():
    model_outputs = {}
    import os
    import numpy as np
    import pandas as pd
    import statsmodels.formula.api as smf
    from functools import partial
    from tqdm import tqdm as _tqdm
    tqdm = partial(_tqdm, bar_format='{bar} {percentage:3.0f}%', ncols=50)
    RET_PATH = RETURN_FILE
    RV_PATH = RV_FILE
    P_STORE_PATH = CENTRALITY_FILE
    PLOT = False
    fit_end_date = pd.to_datetime('2018-05-01')
    pred_start = pd.to_datetime('2018-05-02')
    pred_end = pd.to_datetime('2022-02-25')
    H_LIST = [22]
    FIT_START = 1000
    FIT_STEP = 100
    FIT_MAX = 1
    L_fit_list = [FIT_START + FIT_STEP * k for k in range(FIT_MAX)]
    NET_MIN = 1000
    NET_MAX = 1000
    NET_STEP = 50
    L_net_list = list(range(NET_MIN, NET_MAX + 1, NET_STEP))

    def _norm_date(x):
        return pd.to_datetime(x).tz_localize(None).normalize()

    def make_future_mean(rv_now: pd.Series, H: int) -> pd.Series:
        if H == 1:
            return rv_now
        ma = rv_now.rolling(window=H, center=False).mean()
        return ma.shift(-(H - 1))

    def read_returns(path):
        df = pd.read_csv(path)
        df = df.loc[:, ~df.columns.str.startswith('Unnamed')]
        df['date'] = pd.to_datetime(df['date'].astype(str).str.strip().str[:10], errors='coerce')
        df['date'] = df['date'].dt.tz_localize(None).dt.normalize()
        df = df.set_index('date').dropna(how='all')
        return df

    def read_rv_no_scale(path):
        data = pd.read_csv(path)
        if 'date' not in data.columns:
            raise RuntimeError("RV file must contain a 'date' column.")
        num_cols = [c for c in data.columns if c != 'date']
        data[num_cols] = data[num_cols].apply(pd.to_numeric, errors='coerce')
        data = data.dropna(how='all')
        data['date'] = pd.to_datetime(data['date'], utc=True, errors='coerce')
        data = data.dropna(subset=['date'])
        data['date'] = data['date'].dt.tz_localize(None)
        data['date'] = data['date'].dt.normalize()
        data = data.set_index('date')
        data = data.dropna(axis=1, how='all')
        return data
    df_ret = read_returns(RET_PATH)
    data_rv = read_rv_no_scale(RV_PATH)
    markets = [c for c in data_rv.columns if c in df_ret.columns]
    P_all = pd.read_csv(P_STORE_PATH)
    need_cols = ['date', 'market', 'ci_start_date', 'ci_end_date', 'L_net', 'P_corr', 'P_kendall', 'P_partial', 'P_tail']
    miss = [c for c in need_cols if c not in P_all.columns]
    if miss:
        raise RuntimeError(f'P.csv is missing required columns: {miss}; available columns={list(P_all.columns)}')
    P_all['date'] = pd.to_datetime(P_all['date'], errors='coerce').dt.tz_localize(None).dt.normalize()
    P_all['ci_start_date'] = P_all['ci_start_date'].apply(_norm_date)
    P_all['ci_end_date'] = P_all['ci_end_date'].apply(_norm_date)
    P_all['L_net'] = pd.to_numeric(P_all['L_net'], errors='coerce').astype('Int64')
    P_all = P_all.dropna(subset=['date', 'market', 'ci_start_date', 'ci_end_date', 'L_net']).copy()
    P_cache = {}

    def load_P_from_csv(L_net: int, ci_start_date, ci_end_date):
        key = (int(L_net), _norm_date(ci_start_date), _norm_date(ci_end_date))
        if key in P_cache:
            return P_cache[key]
        ln, cs, ce = key
        sub = P_all[(P_all['L_net'] == ln) & (P_all['ci_start_date'] == cs) & (P_all['ci_end_date'] == ce)]
        if sub.empty:
            P_cache[key] = (None, None, None, None)
            return P_cache[key]

        def _pivot(col):
            df = sub.pivot_table(index='date', columns='market', values=col, aggfunc='mean')
            return df.sort_index()
        P_corr = _pivot('P_corr')
        P_kendall = _pivot('P_kendall')
        P_partial = _pivot('P_partial')
        P_tail = _pivot('P_tail')
        P_cache[key] = (P_corr, P_kendall, P_partial, P_tail)
        return P_cache[key]

    def forecast_one_market_oneH_harci(data_rv_all: pd.DataFrame, P_corr: pd.DataFrame, P_kendall: pd.DataFrame, P_partial: pd.DataFrame, P_tail: pd.DataFrame, market: str, fit_start_date: pd.Timestamp, H: int):
        df = data_rv_all[market].to_frame('RV_now')
        df['RV_D'] = df['RV_now'].shift(1)
        df['RV_W'] = df['RV_D'].rolling(5).mean()
        df['RV_M'] = df['RV_D'].rolling(22).mean()
        df['Y'] = make_future_mean(df['RV_now'], H)
        ci = pd.DataFrame({'CI_corr': P_corr[market], 'CI_kendall': P_kendall[market], 'CI_partial': P_partial[market], 'CI_tail': P_tail[market]}).reindex(df.index)
        df = df.join(ci, how='inner').dropna()
        df['CI_corr_D'] = df['CI_corr'].shift(1)
        df['CI_kendall_D'] = df['CI_kendall'].shift(1)
        df['CI_partial_D'] = df['CI_partial'].shift(1)
        df['CI_tail_D'] = df['CI_tail'].shift(1)
        df = df.dropna()
        df = df.loc[fit_start_date:]
        data_test = df.loc[pred_start:pred_end].copy()
        if len(data_test) == 0:
            return (None, np.nan, np.nan, np.nan)
        window_len = len(df.loc[:pred_start]) - 1
        if window_len <= 0:
            return (None, np.nan, np.nan, np.nan)
        y_true = data_test['Y'].astype(float).copy()
        idx = y_true.index
        preds = []
        for i in tqdm(range(len(idx)), desc=f'HAR-CI {market} H{H}', leave=False):
            train = df.iloc[i:i + window_len].copy()
            train['_Y_'] = train['Y'].astype(float).values
            res = smf.ols('_Y_ ~ RV_D + RV_W + RV_M + CI_corr_D + CI_tail_D + CI_partial_D', train).fit()
            p = res.params
            pred = p.iloc[0] + p.iloc[1] * data_test['RV_D'].iloc[i] + p.iloc[2] * data_test['RV_W'].iloc[i] + p.iloc[3] * data_test['RV_M'].iloc[i] + p.iloc[4] * data_test['CI_corr_D'].iloc[i] + p.iloc[5] * data_test['CI_tail_D'].iloc[i] + p.iloc[6] * data_test['CI_partial_D'].iloc[i]
            preds.append(float(pred))
        y_pred = pd.Series(preds, index=idx)
        mcs = pd.concat([y_true.rename('RV_True'), y_pred.rename('RV_Predict')], axis=1)
        mcs['MSE'] = (mcs['RV_True'] - mcs['RV_Predict']) ** 2
        mcs['MAE'] = np.abs(mcs['RV_True'] - mcs['RV_Predict'])
        mse = float(np.nanmean(mcs['MSE'].values))
        mae = float(np.nanmean(mcs['MAE'].values))
        denom = np.sqrt(np.nanmean(mcs['RV_Predict'].values ** 2)) + np.sqrt(np.nanmean(mcs['RV_True'].values ** 2))
        theil = float(np.sqrt(np.nanmean((mcs['RV_True'].values - mcs['RV_Predict'].values) ** 2)) / denom) if denom != 0 else np.nan
        return (mcs, mse, mae, theil)
    L_fit = L_fit_list[0]
    L_net = L_net_list[0]
    for H in H_LIST:
        error_by_market = {}
        for mkt in markets:
            if fit_end_date not in data_rv.index:
                error_by_market[mkt] = [np.nan, np.nan, np.nan]
                continue
            fit_end_pos = data_rv.index.get_loc(fit_end_date)
            fit_start_pos = fit_end_pos - (L_fit - 1)
            if fit_start_pos < 0:
                error_by_market[mkt] = [np.nan, np.nan, np.nan]
                continue
            fit_start_date = data_rv.index[fit_start_pos]
            ret_pos_end = df_ret.index.searchsorted(fit_start_date) - 1
            if ret_pos_end < 0:
                error_by_market[mkt] = [np.nan, np.nan, np.nan]
                continue
            net_end_date = df_ret.index[ret_pos_end]
            ci_start_date = net_end_date
            ci_end_date = pred_end
            P_corr, P_kendall, P_partial, P_tail = load_P_from_csv(L_net=L_net, ci_start_date=ci_start_date, ci_end_date=ci_end_date)
            if P_corr is None:
                error_by_market[mkt] = [np.nan, np.nan, np.nan]
                continue
            if mkt not in P_corr.columns:
                error_by_market[mkt] = [np.nan, np.nan, np.nan]
                continue
            mcs, mse, mae, theil = forecast_one_market_oneH_harci(data_rv_all=data_rv, P_corr=P_corr, P_kendall=P_kendall, P_partial=P_partial, P_tail=P_tail, market=mkt, fit_start_date=fit_start_date, H=H)
            if mcs is None:
                tmp = pd.DataFrame(index=pd.DatetimeIndex([], name='date'))
                tmp['RV_True'] = np.nan
                tmp['RV_Predict'] = np.nan
                tmp['MSE'] = np.nan
                tmp['MAE'] = np.nan
                model_outputs[str(mkt).strip().lstrip('.')] = tmp.copy()
            else:
                model_outputs[str(mkt).strip().lstrip('.')] = mcs.copy()
            error_by_market[mkt] = [mse, mae, theil]
        ERROR = pd.DataFrame(error_by_market, index=['MSE', 'MAE', 'Theil'])
    return model_outputs

def run_layer_kendall_tail_partial():
    model_outputs = {}
    import os
    import numpy as np
    import pandas as pd
    import statsmodels.formula.api as smf
    from functools import partial
    from tqdm import tqdm as _tqdm
    tqdm = partial(_tqdm, bar_format='{bar} {percentage:3.0f}%', ncols=50)
    RET_PATH = RETURN_FILE
    RV_PATH = RV_FILE
    P_STORE_PATH = CENTRALITY_FILE
    PLOT = False
    fit_end_date = pd.to_datetime('2018-05-01')
    pred_start = pd.to_datetime('2018-05-02')
    pred_end = pd.to_datetime('2022-02-25')
    H_LIST = [22]
    FIT_START = 1000
    FIT_STEP = 100
    FIT_MAX = 1
    L_fit_list = [FIT_START + FIT_STEP * k for k in range(FIT_MAX)]
    NET_MIN = 1000
    NET_MAX = 1000
    NET_STEP = 50
    L_net_list = list(range(NET_MIN, NET_MAX + 1, NET_STEP))

    def _norm_date(x):
        return pd.to_datetime(x).tz_localize(None).normalize()

    def make_future_mean(rv_now: pd.Series, H: int) -> pd.Series:
        if H == 1:
            return rv_now
        ma = rv_now.rolling(window=H, center=False).mean()
        return ma.shift(-(H - 1))

    def read_returns(path):
        df = pd.read_csv(path)
        df = df.loc[:, ~df.columns.str.startswith('Unnamed')]
        df['date'] = pd.to_datetime(df['date'].astype(str).str.strip().str[:10], errors='coerce')
        df['date'] = df['date'].dt.tz_localize(None).dt.normalize()
        df = df.set_index('date').dropna(how='all')
        return df

    def read_rv_no_scale(path):
        data = pd.read_csv(path)
        if 'date' not in data.columns:
            raise RuntimeError("RV file must contain a 'date' column.")
        num_cols = [c for c in data.columns if c != 'date']
        data[num_cols] = data[num_cols].apply(pd.to_numeric, errors='coerce')
        data = data.dropna(how='all')
        data['date'] = pd.to_datetime(data['date'], utc=True, errors='coerce')
        data = data.dropna(subset=['date'])
        data['date'] = data['date'].dt.tz_localize(None)
        data['date'] = data['date'].dt.normalize()
        data = data.set_index('date')
        data = data.dropna(axis=1, how='all')
        return data
    df_ret = read_returns(RET_PATH)
    data_rv = read_rv_no_scale(RV_PATH)
    markets = [c for c in data_rv.columns if c in df_ret.columns]
    P_all = pd.read_csv(P_STORE_PATH)
    need_cols = ['date', 'market', 'ci_start_date', 'ci_end_date', 'L_net', 'P_corr', 'P_kendall', 'P_partial', 'P_tail']
    miss = [c for c in need_cols if c not in P_all.columns]
    if miss:
        raise RuntimeError(f'P.csv is missing required columns: {miss}; available columns={list(P_all.columns)}')
    P_all['date'] = pd.to_datetime(P_all['date'], errors='coerce').dt.tz_localize(None).dt.normalize()
    P_all['ci_start_date'] = P_all['ci_start_date'].apply(_norm_date)
    P_all['ci_end_date'] = P_all['ci_end_date'].apply(_norm_date)
    P_all['L_net'] = pd.to_numeric(P_all['L_net'], errors='coerce').astype('Int64')
    P_all = P_all.dropna(subset=['date', 'market', 'ci_start_date', 'ci_end_date', 'L_net']).copy()
    P_cache = {}

    def load_P_from_csv(L_net: int, ci_start_date, ci_end_date):
        key = (int(L_net), _norm_date(ci_start_date), _norm_date(ci_end_date))
        if key in P_cache:
            return P_cache[key]
        ln, cs, ce = key
        sub = P_all[(P_all['L_net'] == ln) & (P_all['ci_start_date'] == cs) & (P_all['ci_end_date'] == ce)]
        if sub.empty:
            P_cache[key] = (None, None, None, None)
            return P_cache[key]

        def _pivot(col):
            df = sub.pivot_table(index='date', columns='market', values=col, aggfunc='mean')
            return df.sort_index()
        P_corr = _pivot('P_corr')
        P_kendall = _pivot('P_kendall')
        P_partial = _pivot('P_partial')
        P_tail = _pivot('P_tail')
        P_cache[key] = (P_corr, P_kendall, P_partial, P_tail)
        return P_cache[key]

    def forecast_one_market_oneH_harci(data_rv_all: pd.DataFrame, P_corr: pd.DataFrame, P_kendall: pd.DataFrame, P_partial: pd.DataFrame, P_tail: pd.DataFrame, market: str, fit_start_date: pd.Timestamp, H: int):
        df = data_rv_all[market].to_frame('RV_now')
        df['RV_D'] = df['RV_now'].shift(1)
        df['RV_W'] = df['RV_D'].rolling(5).mean()
        df['RV_M'] = df['RV_D'].rolling(22).mean()
        df['Y'] = make_future_mean(df['RV_now'], H)
        ci = pd.DataFrame({'CI_corr': P_corr[market], 'CI_kendall': P_kendall[market], 'CI_partial': P_partial[market], 'CI_tail': P_tail[market]}).reindex(df.index)
        df = df.join(ci, how='inner').dropna()
        df['CI_corr_D'] = df['CI_corr'].shift(1)
        df['CI_kendall_D'] = df['CI_kendall'].shift(1)
        df['CI_partial_D'] = df['CI_partial'].shift(1)
        df['CI_tail_D'] = df['CI_tail'].shift(1)
        df = df.dropna()
        df = df.loc[fit_start_date:]
        data_test = df.loc[pred_start:pred_end].copy()
        if len(data_test) == 0:
            return (None, np.nan, np.nan, np.nan)
        window_len = len(df.loc[:pred_start]) - 1
        if window_len <= 0:
            return (None, np.nan, np.nan, np.nan)
        y_true = data_test['Y'].astype(float).copy()
        idx = y_true.index
        preds = []
        for i in tqdm(range(len(idx)), desc=f'HAR-CI {market} H{H}', leave=False):
            train = df.iloc[i:i + window_len].copy()
            train['_Y_'] = train['Y'].astype(float).values
            res = smf.ols('_Y_ ~ RV_D + RV_W + RV_M + CI_kendall_D + CI_tail_D + CI_partial_D', train).fit()
            p = res.params
            pred = p.iloc[0] + p.iloc[1] * data_test['RV_D'].iloc[i] + p.iloc[2] * data_test['RV_W'].iloc[i] + p.iloc[3] * data_test['RV_M'].iloc[i] + p.iloc[4] * data_test['CI_kendall_D'].iloc[i] + p.iloc[5] * data_test['CI_tail_D'].iloc[i] + p.iloc[6] * data_test['CI_partial_D'].iloc[i]
            preds.append(float(pred))
        y_pred = pd.Series(preds, index=idx)
        mcs = pd.concat([y_true.rename('RV_True'), y_pred.rename('RV_Predict')], axis=1)
        mcs['MSE'] = (mcs['RV_True'] - mcs['RV_Predict']) ** 2
        mcs['MAE'] = np.abs(mcs['RV_True'] - mcs['RV_Predict'])
        mse = float(np.nanmean(mcs['MSE'].values))
        mae = float(np.nanmean(mcs['MAE'].values))
        denom = np.sqrt(np.nanmean(mcs['RV_Predict'].values ** 2)) + np.sqrt(np.nanmean(mcs['RV_True'].values ** 2))
        theil = float(np.sqrt(np.nanmean((mcs['RV_True'].values - mcs['RV_Predict'].values) ** 2)) / denom) if denom != 0 else np.nan
        return (mcs, mse, mae, theil)
    L_fit = L_fit_list[0]
    L_net = L_net_list[0]
    for H in H_LIST:
        error_by_market = {}
        for mkt in markets:
            if fit_end_date not in data_rv.index:
                error_by_market[mkt] = [np.nan, np.nan, np.nan]
                continue
            fit_end_pos = data_rv.index.get_loc(fit_end_date)
            fit_start_pos = fit_end_pos - (L_fit - 1)
            if fit_start_pos < 0:
                error_by_market[mkt] = [np.nan, np.nan, np.nan]
                continue
            fit_start_date = data_rv.index[fit_start_pos]
            ret_pos_end = df_ret.index.searchsorted(fit_start_date) - 1
            if ret_pos_end < 0:
                error_by_market[mkt] = [np.nan, np.nan, np.nan]
                continue
            net_end_date = df_ret.index[ret_pos_end]
            ci_start_date = net_end_date
            ci_end_date = pred_end
            P_corr, P_kendall, P_partial, P_tail = load_P_from_csv(L_net=L_net, ci_start_date=ci_start_date, ci_end_date=ci_end_date)
            if P_corr is None:
                error_by_market[mkt] = [np.nan, np.nan, np.nan]
                continue
            if mkt not in P_corr.columns:
                error_by_market[mkt] = [np.nan, np.nan, np.nan]
                continue
            mcs, mse, mae, theil = forecast_one_market_oneH_harci(data_rv_all=data_rv, P_corr=P_corr, P_kendall=P_kendall, P_partial=P_partial, P_tail=P_tail, market=mkt, fit_start_date=fit_start_date, H=H)
            if mcs is None:
                tmp = pd.DataFrame(index=pd.DatetimeIndex([], name='date'))
                tmp['RV_True'] = np.nan
                tmp['RV_Predict'] = np.nan
                tmp['MSE'] = np.nan
                tmp['MAE'] = np.nan
                model_outputs[str(mkt).strip().lstrip('.')] = tmp.copy()
            else:
                model_outputs[str(mkt).strip().lstrip('.')] = mcs.copy()
            error_by_market[mkt] = [mse, mae, theil]
        ERROR = pd.DataFrame(error_by_market, index=['MSE', 'MAE', 'Theil'])
    return model_outputs

MODEL_RUNNERS = [
    ("HAR-RV", "HAR", run_har),
    ("HAR-CI", "HAR-CI", run_har_ci),
    ("HAR-CI-corr", "HAR-P_pearson", run_layer_corr),
    ("HAR-CI-kendall", "HAR-P_kendall", run_layer_kendall),
    ("HAR-CI-partial", "HAR-P_partial", run_layer_partial),
    ("HAR-CI-tail", "HAR-P_tail", run_layer_tail),
    ("HAR-CI-kendall_tail_partial", "HAR-drop P_pearson", run_layer_kendall_tail_partial),
    ("HAR-CI-corr_tail_partial", "HAR-drop P_kendall", run_layer_corr_tail_partial),
    ("HAR-CI-corr_kendall_tail", "HAR-drop P_partial", run_layer_corr_kendall_tail),
    ("HAR-CI-corr_kendall_partial", "HAR-drop P_tail", run_layer_corr_kendall_partial),
]
LAYER_SPECIFICATIONS = [(), ("pearson", "kendall", "partial", "tail"), ("pearson",), ("kendall",), ("partial",), ("tail",), ("kendall", "partial", "tail"), ("pearson", "partial", "tail"), ("pearson", "kendall", "tail"), ("pearson", "kendall", "partial")]


def save_table(summaries, output_path=OUTPUT_FILE):
    book = Workbook()
    sheet = book.active
    sheet.title = "Table C.10"
    for row, label in ((1, "Panel A: Model specifications"), (10, "Panel B: Summary statistics of MSE for monthly realized volatility forecasts"), (14, "Panel C: Summary statistics of QLIKE for monthly realized volatility forecasts")):
        sheet.merge_cells(start_row=row, start_column=1, end_row=row, end_column=11)
        sheet.cell(row, 1, label)
    sheet.cell(2, 1, "Predictor\\Model")
    for col, (_, label, _) in enumerate(MODEL_RUNNERS, 2):
        sheet.cell(2, col, label)
    for row, label in enumerate(("RV^(d)", "RV^(w)", "RV^(m)", "P_pearson", "P_kendall", "P_partial", "P_tail"), 3):
        sheet.cell(row, 1, label)
        for col, layers in enumerate(LAYER_SPECIFICATIONS, 2):
            sheet.cell(row, col, "\u2713" if row <= 5 or label[2:] in layers else None)
    for start, metric in ((11, "MSE"), (15, "QLIKE")):
        for row, (stat, label) in enumerate((("Mean", "Mean"), ("Q1", "Q1"), ("Q2", "Median")), start):
            sheet.cell(row, 1, label)
            for col, (key, _, _) in enumerate(MODEL_RUNNERS, 2):
                sheet.cell(row, col, float(summaries[metric].loc[key, stat])).number_format = "0.000"
    thin = Side(style="thin", color="000000")
    for row in sheet:
        for cell in row:
            cell.font = Font(name="Times New Roman", size=11, bold=False)
            if cell.value == "\u2713":
                cell.font = Font(name="DejaVu Sans", size=11, bold=False)
            cell.alignment = Alignment(horizontal="left" if cell.column == 1 and cell.row not in (1, 10, 14) else "center", vertical="center", wrap_text=cell.row == 2)
            cell.border = Border(top=thin if cell.row in (1, 10, 14) else Side(), bottom=thin if cell.row in (1, 2, 9, 10, 13, 14, 17) else Side())
    sheet.column_dimensions["A"].width = 22
    for col in "BCDEFGHIJK":
        sheet.column_dimensions[col].width = 22
    sheet.row_dimensions[2].height = 33
    sheet.sheet_view.showGridLines = False
    sheet.freeze_panes = "B3"
    sheet.page_setup.orientation = "landscape"
    sheet.page_setup.paperSize = sheet.PAPERSIZE_A3
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 1
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    sheet.print_area = "A1:K17"
    output_path = Path(output_path)
    output_path.parent.mkdir(parents=True, exist_ok=True)
    book.save(output_path)


def main():
    global CENTRALITY_FILE
    parser = argparse.ArgumentParser()
    parser.add_argument("--recomputed-centrality", action="store_true")
    args, _ = parser.parse_known_args()
    if args.recomputed_centrality:
        CENTRALITY_FILE = ROOT / "results/generated/centrality/P.csv"
    for path in (RV_FILE, RETURN_FILE, CENTRALITY_FILE, QLIKE_DROP_DATES_FILE):
        if not path.is_file():
            raise FileNotFoundError(f"Required input not found: {path}")
    predictions = {}
    for key, _, runner in MODEL_RUNNERS:
        predictions[key] = runner()
        if not predictions[key]:
            raise ValueError(f"No forecasts generated for {key}")
    summaries, _ = summarize_losses(predictions)
    save_table(summaries)


if __name__ == "__main__":
    main()
