from pathlib import Path
from types import SimpleNamespace
import numpy as np
from functools import partial
from tqdm import tqdm as _tqdm
tqdm = partial(_tqdm, bar_format="{bar} {percentage:3.0f}%", ncols=50)



def build_table_04():
    from pathlib import Path

    import argparse

    import warnings

    import numpy as np

    import pandas as pd

    from functools import partial

    from tqdm import tqdm as _tqdm

    tqdm = partial(_tqdm, bar_format="{bar} {percentage:3.0f}%", ncols=50)

    from arch.bootstrap import MCS

    from openpyxl import Workbook

    from openpyxl.styles import Font, Alignment, Border, Side

    ROOT = Path(__file__).resolve().parents[1]

    RV_FILE = ROOT / "data/raw/RV.csv"

    RETURN_FILE = ROOT / "data/raw/Return.csv"

    CENTRALITY_FILE = ROOT / "data/processed/centrality/P.csv"

    TCI_FILE = ROOT / "data/processed/spillover/TCI_data_BK-TVP.csv"

    OUTPUT_FILE = ROOT / "results/generated/Table_04_category1_forecast_performance.xlsx"

    THRESH = 1e-3

    SCALE_IF_MOST_GT_THRESH = 1e-4

    EPS = 1e-12

    RATIO_MIN = 1e-6

    RATIO_MAX = 1e6

    QLIKE_CAP = 100.0

    POST_SCALE_THRESHOLD = 1e4

    POST_SCALE_FACTOR = 1e-8

    def run_har():
        model_outputs = {}
        import math
        import scipy as sc
        import time
        import os
        from scipy.stats import norm
        import pandas as pd
        import numpy as np
        from functools import partial
        from tqdm import tqdm as _tqdm
        tqdm = partial(_tqdm, bar_format="{bar} {percentage:3.0f}%", ncols=50)
        import statistics
        import statsmodels.formula.api as smf
        import os
        data = pd.read_csv(RV_FILE)
        num_cols = data.columns.drop('date')
        data[num_cols] = data[num_cols].apply(pd.to_numeric, errors='coerce')
        data = data.dropna()
        data['date'] = pd.to_datetime(data['date'], utc=True)
        data['date'] = data['date'].dt.tz_localize(None)
        data['date'] = data['date'].dt.normalize()
        data = data.set_index('date')
        error = []
        H = 22
        for c in data.columns:
            data_RV = data['%s' % c]
            data_RV = pd.DataFrame(data_RV)
            data_RV.columns = ['RV_now']
            data_RV['RV_D'] = data_RV['RV_now'].shift(1)
            data_RV['RV_W'] = data_RV['RV_D'].rolling(5).mean()
            data_RV['RV_M'] = data_RV['RV_D'].rolling(22).mean()
            data_RV['RV_MA22'] = data_RV['RV_now'].rolling(window=22, center=False).mean()
            data_RV['RV_MA22'] = data_RV['RV_MA22'].shift(-21)
            data_RV = data_RV.dropna()
            data_RV = data_RV.loc['2014-07-04':]
            time = '2018/05/02'
            time_stop = '2022-02-25'
            Loss_ALL = []
            data_test = data_RV.loc[time:time_stop]
            data_predict = data_RV.loc[time:time_stop]['RV_MA22']
            result = []
            for i in tqdm(range(len(data_RV.loc[time:time_stop]))):
                data_rolling = data_RV.iloc[i:i + (len(data_RV.loc[:time]) - 1)]
                mod = smf.ols('RV_MA22 ~ RV_D + RV_W + RV_M', data_rolling)
                res = mod.fit()
                params_new = res.params
                r_predict = params_new.iloc[0] + params_new.iloc[1] * data_test['RV_D'].iloc[i] + params_new.iloc[2] * data_test['RV_W'].iloc[i] + params_new.iloc[3] * data_test['RV_M'].iloc[i]
                result.append(r_predict)
            result = pd.Series(result, index=data_predict.index)
            SUM = pd.concat([data_predict, result], axis=1)
            SUM.columns = ['RV_True', 'RV_Predict']
            SUM_final = SUM
            SUM_final = SUM_final.reset_index(drop=True)
            SUM_final.index = pd.to_datetime(data_test.index)
            SUM_final['MSE'] = (SUM_final['RV_True'] - SUM_final['RV_Predict']) ** 2
            SUM_final['MAE'] = np.abs(SUM_final['RV_True'] - SUM_final['RV_Predict'])
            SUM_final['Qlike'] = SUM_final['RV_True'] / SUM_final['RV_Predict'] + np.log(SUM_final['RV_Predict'])
            MSE = np.mean((SUM_final['RV_True'] - SUM_final['RV_Predict']) ** 2)
            MAE = np.mean(np.abs(SUM_final['RV_True'] - SUM_final['RV_Predict']))
            Theil = np.sqrt(np.mean((SUM_final['RV_True'] - SUM_final['RV_Predict']) ** 2)) / (np.sqrt(np.mean(SUM_final['RV_Predict'] ** 2)) + np.sqrt(np.mean(SUM_final['RV_True'] ** 2)))
            Loss_ALL = [MSE, MAE, Theil]
            error.append(Loss_ALL)
            model_outputs[str(c).strip().lstrip('.')] = SUM_final.copy()
        ERROR = pd.DataFrame(error, columns=['MSE', 'MAE', 'Theil'], index=data.columns).T
        return model_outputs

    def run_har_ci():
        model_outputs = {}
        import os
        import numpy as np
        import pandas as pd
        import statsmodels.formula.api as smf
        from functools import partial
        from tqdm import tqdm as _tqdm
        tqdm = partial(_tqdm, bar_format="{bar} {percentage:3.0f}%", ncols=50)
        RET_PATH = RETURN_FILE
        RV_PATH = RV_FILE
        P_STORE_PATH = CENTRALITY_FILE
        PLOT = False
        fit_end_date = pd.to_datetime('2018-05-01')
        pred_start = pd.to_datetime('2018-05-02')
        pred_end = pd.to_datetime('2022-02-25')
        H_LIST = [22]
        FIT_START = 1000
        FIT_STEP = 100
        FIT_MAX = 1
        L_fit_list = [FIT_START + FIT_STEP * k for k in range(FIT_MAX)]
        NET_MIN = 1000
        NET_MAX = 1000
        NET_STEP = 50
        L_net_list = list(range(NET_MIN, NET_MAX + 1, NET_STEP))

        def _norm_date(x):
            return pd.to_datetime(x).tz_localize(None).normalize()

        def make_future_mean(rv_now: pd.Series, H: int) -> pd.Series:
            if H == 1:
                return rv_now
            ma = rv_now.rolling(window=H, center=False).mean()
            return ma.shift(-(H - 1))

        def read_returns(path):
            df = pd.read_csv(path)
            df = df.loc[:, ~df.columns.str.startswith('Unnamed')]
            df['date'] = pd.to_datetime(df['date'].astype(str).str.strip().str[:10], errors='coerce')
            df['date'] = df['date'].dt.tz_localize(None).dt.normalize()
            df = df.set_index('date').dropna(how='all')
            return df

        def read_rv_no_scale(path):
            data = pd.read_csv(path)
            if 'date' not in data.columns:
                raise RuntimeError("RV file must contain a 'date' column.")
            num_cols = [c for c in data.columns if c != 'date']
            data[num_cols] = data[num_cols].apply(pd.to_numeric, errors='coerce')
            data = data.dropna(how='all')
            data['date'] = pd.to_datetime(data['date'], utc=True, errors='coerce')
            data = data.dropna(subset=['date'])
            data['date'] = data['date'].dt.tz_localize(None)
            data['date'] = data['date'].dt.normalize()
            data = data.set_index('date')
            data = data.dropna(axis=1, how='all')
            return data
        df_ret = read_returns(RET_PATH)
        data_rv = read_rv_no_scale(RV_PATH)
        markets = [c for c in data_rv.columns if c in df_ret.columns]
        P_all = pd.read_csv(P_STORE_PATH)
        need_cols = ['date', 'market', 'ci_start_date', 'ci_end_date', 'L_net', 'P_corr', 'P_kendall', 'P_partial', 'P_tail']
        miss = [c for c in need_cols if c not in P_all.columns]
        if miss:
            raise RuntimeError(f'P.csv is missing required columns: {miss}; available columns={list(P_all.columns)}')
        P_all['date'] = pd.to_datetime(P_all['date'], errors='coerce').dt.tz_localize(None).dt.normalize()
        P_all['ci_start_date'] = P_all['ci_start_date'].apply(_norm_date)
        P_all['ci_end_date'] = P_all['ci_end_date'].apply(_norm_date)
        P_all['L_net'] = pd.to_numeric(P_all['L_net'], errors='coerce').astype('Int64')
        P_all = P_all.dropna(subset=['date', 'market', 'ci_start_date', 'ci_end_date', 'L_net']).copy()
        P_cache = {}

        def load_P_from_csv(L_net: int, ci_start_date, ci_end_date):
            key = (int(L_net), _norm_date(ci_start_date), _norm_date(ci_end_date))
            if key in P_cache:
                return P_cache[key]
            ln, cs, ce = key
            sub = P_all[(P_all['L_net'] == ln) & (P_all['ci_start_date'] == cs) & (P_all['ci_end_date'] == ce)]
            if sub.empty:
                P_cache[key] = (None, None, None, None)
                return P_cache[key]

            def _pivot(col):
                df = sub.pivot_table(index='date', columns='market', values=col, aggfunc='mean')
                return df.sort_index()
            P_corr = _pivot('P_corr')
            P_kendall = _pivot('P_kendall')
            P_partial = _pivot('P_partial')
            P_tail = _pivot('P_tail')
            P_cache[key] = (P_corr, P_kendall, P_partial, P_tail)
            return P_cache[key]

        def forecast_one_market_oneH_harci(data_rv_all: pd.DataFrame, P_corr: pd.DataFrame, P_kendall: pd.DataFrame, P_partial: pd.DataFrame, P_tail: pd.DataFrame, market: str, fit_start_date: pd.Timestamp, H: int):
            df = data_rv_all[market].to_frame('RV_now')
            df['RV_D'] = df['RV_now'].shift(1)
            df['RV_W'] = df['RV_D'].rolling(5).mean()
            df['RV_M'] = df['RV_D'].rolling(22).mean()
            df['Y'] = make_future_mean(df['RV_now'], H)
            ci = pd.DataFrame({'CI_corr': P_corr[market], 'CI_kendall': P_kendall[market], 'CI_partial': P_partial[market], 'CI_tail': P_tail[market]}).reindex(df.index)
            df = df.join(ci, how='inner').dropna()
            df['CI_corr_D'] = df['CI_corr'].shift(1)
            df['CI_kendall_D'] = df['CI_kendall'].shift(1)
            df['CI_partial_D'] = df['CI_partial'].shift(1)
            df['CI_tail_D'] = df['CI_tail'].shift(1)
            df = df.dropna()
            df = df.loc[fit_start_date:]
            data_test = df.loc[pred_start:pred_end].copy()
            if len(data_test) == 0:
                return (None, np.nan, np.nan, np.nan)
            window_len = len(df.loc[:pred_start]) - 1
            if window_len <= 0:
                return (None, np.nan, np.nan, np.nan)
            y_true = data_test['Y'].astype(float).copy()
            idx = y_true.index
            preds = []
            for i in tqdm(range(len(idx)), desc=f'HAR-CI {market} H{H}', leave=False):
                train = df.iloc[i:i + window_len].copy()
                train['_Y_'] = train['Y'].astype(float).values
                res = smf.ols('_Y_ ~ RV_D + RV_W + RV_M + CI_corr_D + CI_kendall_D + CI_partial_D + CI_tail_D', train).fit()
                p = res.params
                pred = p.iloc[0] + p.iloc[1] * data_test['RV_D'].iloc[i] + p.iloc[2] * data_test['RV_W'].iloc[i] + p.iloc[3] * data_test['RV_M'].iloc[i] + p.iloc[4] * data_test['CI_corr_D'].iloc[i] + p.iloc[5] * data_test['CI_kendall_D'].iloc[i] + p.iloc[6] * data_test['CI_partial_D'].iloc[i] + p.iloc[7] * data_test['CI_tail_D'].iloc[i]
                preds.append(float(pred))
            y_pred = pd.Series(preds, index=idx)
            mcs = pd.concat([y_true.rename('RV_True'), y_pred.rename('RV_Predict')], axis=1)
            mcs['MSE'] = (mcs['RV_True'] - mcs['RV_Predict']) ** 2
            mcs['MAE'] = np.abs(mcs['RV_True'] - mcs['RV_Predict'])
            mse = float(np.nanmean(mcs['MSE'].values))
            mae = float(np.nanmean(mcs['MAE'].values))
            denom = np.sqrt(np.nanmean(mcs['RV_Predict'].values ** 2)) + np.sqrt(np.nanmean(mcs['RV_True'].values ** 2))
            theil = float(np.sqrt(np.nanmean((mcs['RV_True'].values - mcs['RV_Predict'].values) ** 2)) / denom) if denom != 0 else np.nan
            return (mcs, mse, mae, theil)
        L_fit = L_fit_list[0]
        L_net = L_net_list[0]
        for H in H_LIST:
            error_by_market = {}
            for mkt in markets:
                if fit_end_date not in data_rv.index:
                    error_by_market[mkt] = [np.nan, np.nan, np.nan]
                    continue
                fit_end_pos = data_rv.index.get_loc(fit_end_date)
                fit_start_pos = fit_end_pos - (L_fit - 1)
                if fit_start_pos < 0:
                    error_by_market[mkt] = [np.nan, np.nan, np.nan]
                    continue
                fit_start_date = data_rv.index[fit_start_pos]
                ret_pos_end = df_ret.index.searchsorted(fit_start_date) - 1
                if ret_pos_end < 0:
                    error_by_market[mkt] = [np.nan, np.nan, np.nan]
                    continue
                net_end_date = df_ret.index[ret_pos_end]
                ci_start_date = net_end_date
                ci_end_date = pred_end
                P_corr, P_kendall, P_partial, P_tail = load_P_from_csv(L_net=L_net, ci_start_date=ci_start_date, ci_end_date=ci_end_date)
                if P_corr is None:
                    error_by_market[mkt] = [np.nan, np.nan, np.nan]
                    continue
                if mkt not in P_corr.columns:
                    error_by_market[mkt] = [np.nan, np.nan, np.nan]
                    continue
                mcs, mse, mae, theil = forecast_one_market_oneH_harci(data_rv_all=data_rv, P_corr=P_corr, P_kendall=P_kendall, P_partial=P_partial, P_tail=P_tail, market=mkt, fit_start_date=fit_start_date, H=H)
                if mcs is None:
                    tmp = pd.DataFrame(index=pd.DatetimeIndex([], name='date'))
                    tmp['RV_True'] = np.nan
                    tmp['RV_Predict'] = np.nan
                    tmp['MSE'] = np.nan
                    tmp['MAE'] = np.nan
                    model_outputs[str(mkt).strip().lstrip('.')] = tmp.copy()
                else:
                    model_outputs[str(mkt).strip().lstrip('.')] = mcs.copy()
                error_by_market[mkt] = [mse, mae, theil]
            ERROR = pd.DataFrame(error_by_market, index=['MSE', 'MAE', 'Theil'])
        return model_outputs

    def run_har_ks():
        model_outputs = {}
        from sklearn.linear_model import LinearRegression
        data = pd.read_csv(RV_FILE)
        data = data.dropna()
        data['date'] = pd.to_datetime(data['date'], utc=True)
        data['date'] = data['date'].dt.tz_localize(None)
        data['date'] = data['date'].dt.normalize()
        data = data.ffill()
        data = data.set_index('date')
        data = data.dropna()
        time = '2018-05-02'
        time_stop = '2022-02-25'
        error = []
        for country in data.columns:
            data_RV = data.copy()
            data_RV.loc[:, data_RV.columns != country] = data_RV.loc[:, data_RV.columns != country].shift(1)
            data_RV['RV_now'] = data_RV[country]
            data_RV['RV_D'] = data_RV['RV_now'].shift(1)
            data_RV['RV_W'] = data_RV['RV_D'].rolling(5).mean()
            data_RV['RV_M'] = data_RV['RV_D'].rolling(22).mean()
            data_RV['RV_MA22'] = data_RV['RV_now'].rolling(window=22, center=False).mean()
            data_RV['RV_MA22'] = data_RV['RV_MA22'].shift(-21)
            data_RV = data_RV.drop(columns=[country, 'RV_now'])
            data_RV = data_RV.loc['2014-07-04':]
            data_test = data_RV.loc[time:time_stop]
            data_predict = data_test['RV_MA22']
            result = []
            ols = LinearRegression()
            window_len = len(data_RV.loc[:time]) - 1
            for i in tqdm(range(len(data_test)), desc=str(country)):
                data_rolling = data_RV.iloc[i:i + window_len].dropna()
                X_train = data_rolling.drop(columns=['RV_MA22'])
                y_train = data_rolling['RV_MA22']
                ols.fit(X_train, y_train)
                X_test = np.array([data_test.drop(columns=['RV_MA22']).iloc[i].values])
                r_predict = ols.predict(X_test)[0]
                result.append(r_predict)
            result = pd.Series(result, index=data_predict.index)
            SUM = pd.concat([data_predict, result], axis=1)
            SUM.columns = ['RV_True', 'RV_Predict']
            SUM_final = SUM.reset_index(drop=True)
            SUM_final.index = pd.to_datetime(data_test.index)
            SUM_final['MSE'] = (SUM_final['RV_True'] - SUM_final['RV_Predict']) ** 2
            SUM_final['MAE'] = np.abs(SUM_final['RV_True'] - SUM_final['RV_Predict'])
            SUM_final['Qlike'] = SUM_final['RV_True'] / SUM_final['RV_Predict'] + np.log(SUM_final['RV_Predict'])
            MSE = np.mean((SUM_final['RV_True'] - SUM_final['RV_Predict']) ** 2)
            MAE = np.mean(np.abs(SUM_final['RV_True'] - SUM_final['RV_Predict']))
            Theil = np.sqrt(np.mean((SUM_final['RV_True'] - SUM_final['RV_Predict']) ** 2)) / (np.sqrt(np.mean(SUM_final['RV_Predict'] ** 2)) + np.sqrt(np.mean(SUM_final['RV_True'] ** 2)))
            error.append([MSE, MAE, Theil])
            model_outputs[str(country).strip().lstrip('.')] = SUM_final.copy()
        ERROR = pd.DataFrame(error, columns=['MSE', 'MAE', 'Theil'], index=data.columns).T
        return model_outputs

    def run_mc():
        model_outputs = {}
        import math
        import scipy as sc
        import time
        import os
        from scipy.stats import norm
        import pandas as pd
        import numpy as np
        from functools import partial
        from tqdm import tqdm as _tqdm
        tqdm = partial(_tqdm, bar_format="{bar} {percentage:3.0f}%", ncols=50)
        import statistics
        import statsmodels.formula.api as smf
        from sklearn.linear_model import LinearRegression
        import statsmodels.api as sm
        data = pd.read_csv(RV_FILE)
        data = data.dropna()
        data['date'] = pd.to_datetime(data['date'], utc=True)
        data['date'] = data['date'].dt.tz_localize(None)
        data['date'] = data['date'].dt.normalize()
        data = data.fillna(method='ffill')
        data = data.set_index('date')
        data = data.dropna()
        H = 22
        error_H = []
        for country in data.columns:
            data_RV_base = data.copy()
            data_RV_base.loc[:, data_RV_base.columns != '%s' % country] = data_RV_base.loc[:, data_RV_base.columns != '%s' % country].shift(1)
            data_RV_base['RV_now'] = data_RV_base['%s' % country]
            data_RV_base['RV_D'] = data_RV_base['RV_now'].shift(1)
            data_RV_base['RV_W'] = data_RV_base['RV_D'].rolling(5).mean()
            data_RV_base['RV_M'] = data_RV_base['RV_D'].rolling(22).mean()
            time_start = '2018-05-02'
            time_stop = '2022-02-25'
            data_RV = data_RV_base.copy()
            target_col = f'RV_MA{H}'
            data_RV[target_col] = data_RV['RV_now'].rolling(window=H, center=False).mean()
            data_RV[target_col] = data_RV[target_col].shift(-22)
            data_RV = data_RV.drop(columns=['%s' % country])
            data_RV = data_RV.drop(columns=['RV_now'])
            data_RV = data_RV.dropna()
            data_RV = data_RV.loc['2014-07-04':]
            data_test = data_RV.loc[time_start:time_stop]
            data_predict = data_RV.loc[time_start:time_stop][target_col]
            result = []
            ols = LinearRegression()
            window_len_fit = len(data_RV.loc[:time_start]) - 1
            for i in tqdm(range(len(data_RV.loc[time_start:time_stop])), desc=f'{country} | H{H}'):
                data_rolling = data_RV.iloc[i:i + window_len_fit]
                data_rolling = data_rolling.dropna()
                MEAN_pre = []
                for other_country in data_RV.columns.drop([target_col, 'RV_D', 'RV_W', 'RV_M']):
                    X_train = data_rolling[['RV_D', 'RV_W', 'RV_M', '%s' % other_country]]
                    y_train = data_rolling[target_col]
                    ols.fit(X_train, y_train)
                    X_test = np.array([data_test[['RV_D', 'RV_W', 'RV_M', '%s' % other_country]].iloc[i].values])
                    r_predict = ols.predict(X_test)[0]
                    MEAN_pre.append(r_predict)
                result.append(np.mean(MEAN_pre))
            result = pd.Series(result)
            result.index = data_predict.index
            SUM = pd.concat([data_predict, result], axis=1)
            SUM.columns = ['RV_True', 'RV_Predict']
            SUM_final = SUM
            SUM_final = SUM_final.reset_index(drop=True)
            SUM_final.index = pd.to_datetime(data_test.index)
            SUM_final['MSE'] = (SUM_final['RV_True'] - SUM_final['RV_Predict']) ** 2
            SUM_final['MAE'] = np.abs(SUM_final['RV_True'] - SUM_final['RV_Predict'])
            SUM_final['Qlike'] = SUM_final['RV_True'] / SUM_final['RV_Predict'] + np.log(SUM_final['RV_Predict'])
            MSE = np.mean((SUM_final['RV_True'] - SUM_final['RV_Predict']) ** 2)
            MAE = np.mean(np.abs(SUM_final['RV_True'] - SUM_final['RV_Predict']))
            Theil = np.sqrt(np.mean((SUM_final['RV_True'] - SUM_final['RV_Predict']) ** 2)) / (np.sqrt(np.mean(SUM_final['RV_Predict'] ** 2)) + np.sqrt(np.mean(SUM_final['RV_True'] ** 2)))
            error_H.append([MSE, MAE, Theil])
            model_outputs[str(country).strip().lstrip('.')] = SUM_final.copy()
        ERROR = pd.DataFrame(error_H, columns=['MSE', 'MAE', 'Theil'], index=data.columns).T
        return model_outputs

    def run_mdc():
        model_outputs = {}
        import math
        import scipy as sc
        import time
        import os
        from scipy.stats import norm
        import pandas as pd
        import numpy as np
        from functools import partial
        from tqdm import tqdm as _tqdm
        tqdm = partial(_tqdm, bar_format="{bar} {percentage:3.0f}%", ncols=50)
        import statistics
        import statsmodels.formula.api as smf
        from sklearn.linear_model import LinearRegression
        import statsmodels.api as sm
        data = pd.read_csv(RV_FILE)
        data = data.dropna()
        data['date'] = pd.to_datetime(data['date'], utc=True)
        data['date'] = data['date'].dt.tz_localize(None)
        data['date'] = data['date'].dt.normalize()
        data = data.fillna(method='ffill')
        data = data.set_index('date')
        data = data.dropna()
        H_LIST = [22]
        time = '2018-05-02'
        time_stop = '2022-02-25'
        error_by_H = {H: [] for H in H_LIST}
        for country in data.columns:
            data_RV_base = data.copy()
            data_RV_base.loc[:, data_RV_base.columns != f'{country}'] = data_RV_base.loc[:, data_RV_base.columns != f'{country}'].shift(1)
            data_RV_base['RV_now'] = data_RV_base[f'{country}']
            data_RV_base['RV_D'] = data_RV_base['RV_now'].shift(1)
            data_RV_base['RV_W'] = data_RV_base['RV_D'].rolling(5).mean()
            data_RV_base['RV_M'] = data_RV_base['RV_D'].rolling(22).mean()
            for H in H_LIST:
                target_col = f'RV_MA{H}'
                data_RV = data_RV_base.copy()
                if H == 1:
                    data_RV[target_col] = data_RV['RV_now']
                else:
                    data_RV[target_col] = data_RV['RV_now'].rolling(window=H, center=False).mean()
                    data_RV[target_col] = data_RV[target_col].shift(-22)
                data_RV = data_RV.drop(columns=[f'{country}'])
                data_RV = data_RV.drop(columns=['RV_now'])
                data_RV = data_RV.dropna()
                data_RV = data_RV.loc['2014-07-04':]
                data_test = data_RV.loc[time:time_stop]
                data_predict = data_test[target_col]
                if len(data_test) == 0:
                    error_by_H[H].append([np.nan, np.nan, np.nan])
                    continue
                window_len = len(data_RV.loc[:time]) - 1
                if window_len <= 0:
                    error_by_H[H].append([np.nan, np.nan, np.nan])
                    continue
                result = []
                ols = LinearRegression()
                for i in tqdm(range(len(data_RV.loc[time:time_stop])), desc=f'{country}-H{H}'):
                    data_rolling = data_RV.iloc[i:i + window_len]
                    data_rolling = data_rolling.dropna()
                    MEAN_pre = []
                    other_cols = data_RV.columns.drop([target_col, 'RV_D', 'RV_W', 'RV_M'])
                    for other_country in other_cols:
                        X_train = data_rolling[['RV_D', 'RV_W', 'RV_M', f'{other_country}']]
                        y_train = data_rolling[target_col]
                        ols.fit(X_train, y_train)
                        X_test = np.array([data_test[['RV_D', 'RV_W', 'RV_M', f'{other_country}']].iloc[i].values])
                        r_predict = ols.predict(X_test)[0]
                        MEAN_pre.append(r_predict)
                    result.append(np.median(MEAN_pre))
                result = pd.Series(result, index=data_predict.index)
                SUM = pd.concat([data_predict, result], axis=1)
                SUM.columns = ['RV_True', 'RV_Predict']
                SUM_final = SUM
                SUM_final = SUM_final.reset_index(drop=True)
                SUM_final.index = pd.to_datetime(data_test.index)
                SUM_final['MSE'] = (SUM_final['RV_True'] - SUM_final['RV_Predict']) ** 2
                SUM_final['MAE'] = np.abs(SUM_final['RV_True'] - SUM_final['RV_Predict'])
                SUM_final['Qlike'] = SUM_final['RV_True'] / SUM_final['RV_Predict'] + np.log(SUM_final['RV_Predict'])
                MSE = np.mean((SUM_final['RV_True'] - SUM_final['RV_Predict']) ** 2)
                MAE = np.mean(np.abs(SUM_final['RV_True'] - SUM_final['RV_Predict']))
                Theil = np.sqrt(np.mean((SUM_final['RV_True'] - SUM_final['RV_Predict']) ** 2)) / (np.sqrt(np.mean(SUM_final['RV_Predict'] ** 2)) + np.sqrt(np.mean(SUM_final['RV_True'] ** 2)))
                error_by_H[H].append([MSE, MAE, Theil])
                model_outputs[str(country).strip().lstrip('.')] = SUM_final.copy()
        for H in H_LIST:
            ERROR = pd.DataFrame(error_by_H[H], columns=['MSE', 'MAE', 'Theil'], index=data.columns).T
        return model_outputs

    def run_tmc():
        model_outputs = {}
        import math
        import scipy as sc
        import time
        import os
        from scipy.stats import norm
        import pandas as pd
        import numpy as np
        from functools import partial
        from tqdm import tqdm as _tqdm
        tqdm = partial(_tqdm, bar_format="{bar} {percentage:3.0f}%", ncols=50)
        import statistics
        import statsmodels.formula.api as smf
        from sklearn.linear_model import LinearRegression
        import statsmodels.api as sm
        data = pd.read_csv(RV_FILE)
        data = data.dropna()
        data['date'] = pd.to_datetime(data['date'], utc=True)
        data['date'] = data['date'].dt.tz_localize(None)
        data['date'] = data['date'].dt.normalize()
        data = data.fillna(method='ffill')
        data = data.set_index('date')
        data = data.dropna()
        H = 22
        time = '2018-05-02'
        time_stop = '2022-02-25'
        error_list = []
        for country in data.columns:
            data_RV_base = data.copy()
            data_RV_base.loc[:, data_RV_base.columns != '%s' % country] = data_RV_base.loc[:, data_RV_base.columns != '%s' % country].shift(1)
            data_RV_base['RV_now'] = data_RV_base['%s' % country]
            data_RV_base['RV_D'] = data_RV_base['RV_now'].shift(1)
            data_RV_base['RV_W'] = data_RV_base['RV_D'].rolling(5).mean()
            data_RV_base['RV_M'] = data_RV_base['RV_D'].rolling(22).mean()
            target_col = 'RV_MA22'
            data_RV = data_RV_base.copy()
            data_RV[target_col] = data_RV['RV_now'].rolling(window=H, center=False).mean()
            data_RV[target_col] = data_RV[target_col].shift(-22)
            data_RV = data_RV.drop(columns=['RV_now'])
            data_RV = data_RV.drop(columns=['%s' % country])
            data_RV = data_RV.dropna()
            data_RV = data_RV.loc['2014-07-04':]
            data_test = data_RV.loc[time:time_stop]
            data_predict = data_RV.loc[time:time_stop, target_col]
            if len(data_test) == 0:
                error_list.append([np.nan, np.nan, np.nan])
                continue
            window_len = len(data_RV.loc[:time]) - 1
            if window_len <= 0:
                error_list.append([np.nan, np.nan, np.nan])
                continue
            result = []
            ols = LinearRegression()
            other_markets = data_RV.columns.drop([target_col, 'RV_D', 'RV_W', 'RV_M'])
            for i in tqdm(range(len(data_RV.loc[time:time_stop])), desc=f'{country}-H{H}'):
                data_rolling = data_RV.iloc[i:i + window_len]
                data_rolling = data_rolling.dropna()
                MEAN_pre = []
                for other_country in other_markets:
                    X_train = data_rolling[['RV_D', 'RV_W', 'RV_M', other_country]]
                    y_train = data_rolling[target_col]
                    ols.fit(X_train, y_train)
                    X_test = np.array([data_test[['RV_D', 'RV_W', 'RV_M', other_country]].iloc[i].values])
                    r_predict = ols.predict(X_test)[0]
                    MEAN_pre.append(r_predict)
                result.append(np.mean(np.sort(MEAN_pre)[1:-1]))
            result = pd.Series(result, index=data_predict.index)
            SUM = pd.concat([data_predict, result], axis=1)
            SUM.columns = ['RV_True', 'RV_Predict']
            SUM_final = SUM.copy()
            SUM_final = SUM_final.reset_index(drop=True)
            SUM_final.index = pd.to_datetime(data_test.index)
            SUM_final['MSE'] = (SUM_final['RV_True'] - SUM_final['RV_Predict']) ** 2
            SUM_final['MAE'] = np.abs(SUM_final['RV_True'] - SUM_final['RV_Predict'])
            SUM_final['MAPE'] = np.abs((SUM_final['RV_True'] - SUM_final['RV_Predict']) / SUM_final['RV_True'])
            SUM_final['Qlike'] = SUM_final['RV_True'] / SUM_final['RV_Predict'] + np.log(SUM_final['RV_Predict'])
            MSE = np.mean((SUM_final['RV_True'] - SUM_final['RV_Predict']) ** 2)
            MAE = np.mean(np.abs(SUM_final['RV_True'] - SUM_final['RV_Predict']))
            Theil = np.sqrt(np.mean((SUM_final['RV_True'] - SUM_final['RV_Predict']) ** 2)) / (np.sqrt(np.mean(SUM_final['RV_Predict'] ** 2)) + np.sqrt(np.mean(SUM_final['RV_True'] ** 2)))
            error_list.append([MSE, MAE, Theil])
            model_outputs[str(country).strip().lstrip('.')] = SUM_final.copy()
        ERROR = pd.DataFrame(error_list, columns=['MSE', 'MAE', 'Theil'], index=data.columns).T
        return model_outputs

    def run_dmspe_09():
        model_outputs = {}
        from sklearn.linear_model import LinearRegression
        data = pd.read_csv(RV_FILE)
        num_cols = data.columns.drop('date')
        data[num_cols] = data[num_cols].apply(pd.to_numeric, errors='coerce') * 10000
        data = data.dropna()
        data['date'] = pd.to_datetime(data['date'], utc=True)
        data['date'] = data['date'].dt.tz_localize(None)
        data['date'] = data['date'].dt.normalize()
        data = data.bfill()
        data = data.set_index('date')

        def make_target(rv_now: pd.Series, H: int) -> pd.Series:
            if H == 1:
                return rv_now
            y = rv_now.rolling(window=H, center=False).mean()
            y = y.shift(-22)
            return y
        H_LIST = [22]
        for H in H_LIST:
            error = []
            for country in data.columns:
                data_name = data.loc[:, data.columns != '%s' % country]
                data_RV = data.copy()
                data_RV.loc[:, data_RV.columns != '%s' % country] = data_RV.loc[:, data_RV.columns != '%s' % country].shift(1)
                data_RV['RV_now'] = data_RV['%s' % country]
                data_RV['RV_D'] = data_RV['RV_now'].shift(1)
                data_RV['RV_W'] = data_RV['RV_D'].rolling(5).mean()
                data_RV['RV_M'] = data_RV['RV_D'].rolling(22).mean()
                data_RV[f'Y_H{H}'] = make_target(data_RV['RV_now'], H)
                data_RV = data_RV.drop(columns=['RV_now'])
                data_RV = data_RV.drop(columns=['%s' % country])
                data_RV = data_RV.dropna()
                data_RV = data_RV.loc['2014-07-04':]
                time = '2018-05-02'
                time_stop = '2022-02-25'
                time_holdout = '2016-5-09'
                theta = 0.9
                Loss_ALL = []
                data_test = data_RV.loc[time:time_stop]
                data_test_holdout = data_RV.loc[time_holdout:time]
                data_predict = data_test[f'Y_H{H}']
                data_predict_holdout = data_test_holdout[f'Y_H{H}']
                result = []
                ols = LinearRegression()
                predictions = pd.DataFrame(columns=data_name.columns)
                other_list = data_RV.columns.drop([f'Y_H{H}', 'RV_D', 'RV_W', 'RV_M'])
                for j in tqdm(range(len(data_test_holdout)), desc=f'{country} | H{H} | holdout'):
                    data_rolling = data_RV.loc['2014-07-04':pd.to_datetime(time_holdout) + pd.DateOffset(days=j)].iloc[:-1]
                    for other_country in other_list:
                        X_train_holdout = data_rolling[['RV_D', 'RV_W', 'RV_M', other_country]]
                        y_train_holdout = data_rolling[f'Y_H{H}']
                        ols.fit(X_train_holdout, y_train_holdout)
                        X_test_holdout = np.array([data_test_holdout[['RV_D', 'RV_W', 'RV_M', other_country]].iloc[j].values])
                        pred = ols.predict(X_test_holdout)[0]
                        predictions.loc[j, '%s' % other_country] = pred
                for i in tqdm(range(len(data_test)), desc=f'{country} | H{H} | test'):
                    current_time = data_test.index[i]
                    data_rolling = data_RV.loc[time_holdout:current_time].iloc[:-1]
                    for other_country in other_list:
                        X_train = data_rolling[['RV_D', 'RV_W', 'RV_M', other_country]]
                        y_train = data_rolling[f'Y_H{H}']
                        ols.fit(X_train, y_train)
                        X_test = np.array([data_test[['RV_D', 'RV_W', 'RV_M', other_country]].iloc[i].values])
                        pred = ols.predict(X_test)[0]
                        predictions.loc[len(data_test_holdout) + i, '%s' % other_country] = pred
                for m in tqdm(range(len(data_test)), desc=f'{country} | H{H} | dmspe'):
                    current_time = data_test.index[m]
                    y_train_sum = data_RV.loc[pd.to_datetime(time_holdout):current_time, f'Y_H{H}']
                    y_train_sum = y_train_sum.iloc[:-1]
                    y_train_sum.index = range(len(y_train_sum))
                    dmspe_errors = []
                    for other_country in other_list:
                        pred_hist = predictions.loc[:len(y_train_sum) - 1, other_country].astype(float)
                        errors = (y_train_sum - pred_hist) ** 2
                        time_weights = theta ** np.arange(len(errors) - 1, -1, -1)
                        dmspe_errors.append(np.sum(errors * time_weights) + 1e-12)
                    dmspe_errors = np.array(dmspe_errors)
                    weights = 1 / dmspe_errors
                    weights /= weights.sum()
                    dmspe_prediction = np.sum(predictions.loc[len(data_test_holdout) + m, :].astype(float).values * weights)
                    result.append(dmspe_prediction)
                result = pd.Series(result, index=data_predict.index)
                SUM = pd.concat([data_predict, result], axis=1)
                SUM.columns = ['RV_True', 'RV_Predict']
                SUM_final = SUM.reset_index(drop=True)
                SUM_final.index = pd.to_datetime(data_test.index)
                SUM_final['MSE'] = (SUM_final['RV_True'] - SUM_final['RV_Predict']) ** 2
                SUM_final['MAE'] = np.abs(SUM_final['RV_True'] - SUM_final['RV_Predict'])
                SUM_final['Qlike'] = SUM_final['RV_True'] / SUM_final['RV_Predict'] + np.log(SUM_final['RV_Predict'])
                MSE = np.mean((SUM_final['RV_True'] - SUM_final['RV_Predict']) ** 2)
                MAE = np.mean(np.abs(SUM_final['RV_True'] - SUM_final['RV_Predict']))
                Theil = np.sqrt(np.mean((SUM_final['RV_True'] - SUM_final['RV_Predict']) ** 2)) / (np.sqrt(np.mean(SUM_final['RV_Predict'] ** 2)) + np.sqrt(np.mean(SUM_final['RV_True'] ** 2)))
                Loss_ALL = [MSE, MAE, Theil]
                error.append(Loss_ALL)
                model_outputs[str(country).strip().lstrip('.')] = SUM_final.copy()
            ERROR = pd.DataFrame(error, columns=['MSE', 'MAE', 'Theil'], index=data.columns).T
        return model_outputs

    def run_dmspe_1():
        model_outputs = {}
        import math
        import scipy as sc
        import time
        import os
        from scipy.stats import norm
        import pandas as pd
        import numpy as np
        from functools import partial
        from tqdm import tqdm as _tqdm
        tqdm = partial(_tqdm, bar_format="{bar} {percentage:3.0f}%", ncols=50)
        import statistics
        import statsmodels.formula.api as smf
        from sklearn.linear_model import LinearRegression
        import statsmodels.api as sm
        data = pd.read_csv(RV_FILE)
        num_cols = data.columns.drop('date')
        data[num_cols] = data[num_cols].apply(pd.to_numeric, errors='coerce')
        data = data.dropna()
        data['date'] = pd.to_datetime(data['date'], utc=True)
        data['date'] = data['date'].dt.tz_localize(None)
        data['date'] = data['date'].dt.normalize()
        data = data.fillna(method='bfill')
        data = data.set_index('date')

        def make_target(rv_now: pd.Series, H: int) -> pd.Series:
            if H == 1:
                return rv_now
            y = rv_now.rolling(window=H, center=False).mean()
            y = y.shift(-22)
            return y
        H_LIST = [22]
        for H in H_LIST:
            error = []
            for country in data.columns:
                data_name = data.loc[:, data.columns != '%s' % country]
                data_RV = data.copy()
                data_RV.loc[:, data_RV.columns != '%s' % country] = data_RV.loc[:, data_RV.columns != '%s' % country].shift(1)
                data_RV['RV_now'] = data_RV['%s' % country]
                data_RV['RV_D'] = data_RV['RV_now'].shift(1)
                data_RV['RV_W'] = data_RV['RV_D'].rolling(5).mean()
                data_RV['RV_M'] = data_RV['RV_D'].rolling(22).mean()
                data_RV[f'Y_H{H}'] = make_target(data_RV['RV_now'], H)
                data_RV = data_RV.drop(columns=['RV_now'])
                data_RV = data_RV.drop(columns=['%s' % country])
                data_RV = data_RV.dropna()
                data_RV = data_RV.loc['2014-07-04':]
                time = '2018-05-02'
                time_stop = '2022-02-25'
                time_holdout = '2016-5-09'
                theta = 1
                Loss_ALL = []
                data_test = data_RV.loc[time:time_stop]
                data_test_holdout = data_RV.loc[time_holdout:time]
                data_predict = data_test[f'Y_H{H}']
                data_predict_holdout = data_test_holdout[f'Y_H{H}']
                result = []
                ols = LinearRegression()
                predictions = pd.DataFrame(columns=data_name.columns)
                other_list = data_RV.columns.drop([f'Y_H{H}', 'RV_D', 'RV_W', 'RV_M'])
                for j in tqdm(range(len(data_test_holdout)), desc=f'{country} | H{H} | holdout'):
                    data_rolling = data_RV.loc['2014-07-04':pd.to_datetime(time_holdout) + pd.DateOffset(days=j)].iloc[:-1]
                    for other_country in other_list:
                        X_train_holdout = data_rolling[['RV_D', 'RV_W', 'RV_M', other_country]]
                        y_train_holdout = data_rolling[f'Y_H{H}']
                        ols.fit(X_train_holdout, y_train_holdout)
                        X_test_holdout = np.array([data_test_holdout[['RV_D', 'RV_W', 'RV_M', other_country]].iloc[j].values])
                        pred = ols.predict(X_test_holdout)[0]
                        predictions.loc[j, '%s' % other_country] = pred
                for i in tqdm(range(len(data_test)), desc=f'{country} | H{H} | test'):
                    current_time = data_test.index[i]
                    data_rolling = data_RV.loc[time_holdout:current_time].iloc[:-1]
                    for other_country in other_list:
                        X_train = data_rolling[['RV_D', 'RV_W', 'RV_M', other_country]]
                        y_train = data_rolling[f'Y_H{H}']
                        ols.fit(X_train, y_train)
                        X_test = np.array([data_test[['RV_D', 'RV_W', 'RV_M', other_country]].iloc[i].values])
                        pred = ols.predict(X_test)[0]
                        predictions.loc[len(data_test_holdout) + i, '%s' % other_country] = pred
                for m in tqdm(range(len(data_test)), desc=f'{country} | H{H} | dmspe'):
                    current_time = data_test.index[m]
                    y_train_sum = data_RV.loc[pd.to_datetime(time_holdout):current_time, f'Y_H{H}']
                    y_train_sum = y_train_sum.iloc[:-1]
                    y_train_sum.index = range(len(y_train_sum))
                    dmspe_errors = []
                    for other_country in other_list:
                        pred_hist = predictions.loc[:len(y_train_sum) - 1, other_country].astype(float)
                        errors = (y_train_sum - pred_hist) ** 2
                        time_weights = theta ** np.arange(len(errors) - 1, -1, -1)
                        dmspe_errors.append(np.sum(errors * time_weights) + 1e-12)
                    dmspe_errors = np.array(dmspe_errors)
                    weights = 1 / dmspe_errors
                    weights /= weights.sum()
                    dmspe_prediction = np.sum(predictions.loc[len(data_test_holdout) + m, :].astype(float).values * weights)
                    result.append(dmspe_prediction)
                result = pd.Series(result, index=data_predict.index)
                SUM = pd.concat([data_predict, result], axis=1)
                SUM.columns = ['RV_True', 'RV_Predict']
                SUM_final = SUM.reset_index(drop=True)
                SUM_final.index = pd.to_datetime(data_test.index)
                SUM_final['MSE'] = (SUM_final['RV_True'] - SUM_final['RV_Predict']) ** 2
                SUM_final['MAE'] = np.abs(SUM_final['RV_True'] - SUM_final['RV_Predict'])
                SUM_final['Qlike'] = SUM_final['RV_True'] / SUM_final['RV_Predict'] + np.log(SUM_final['RV_Predict'])
                MSE = np.mean((SUM_final['RV_True'] - SUM_final['RV_Predict']) ** 2)
                MAE = np.mean(np.abs(SUM_final['RV_True'] - SUM_final['RV_Predict']))
                Theil = np.sqrt(np.mean((SUM_final['RV_True'] - SUM_final['RV_Predict']) ** 2)) / (np.sqrt(np.mean(SUM_final['RV_Predict'] ** 2)) + np.sqrt(np.mean(SUM_final['RV_True'] ** 2)))
                Loss_ALL = [MSE, MAE, Theil]
                error.append(Loss_ALL)
                model_outputs[str(country).strip().lstrip('.')] = SUM_final.copy()
            ERROR = pd.DataFrame(error, columns=['MSE', 'MAE', 'Theil'], index=data.columns).T
        return model_outputs

    def run_pca1():
        model_outputs = {}
        import math
        import scipy as sc
        import time
        import os
        from scipy.stats import norm
        import pandas as pd
        import numpy as np
        from functools import partial
        from tqdm import tqdm as _tqdm
        tqdm = partial(_tqdm, bar_format="{bar} {percentage:3.0f}%", ncols=50)
        import statistics
        import statsmodels.formula.api as smf
        from sklearn.linear_model import LinearRegression
        import statsmodels.api as sm
        from sklearn.decomposition import PCA
        data = pd.read_csv(RV_FILE)
        data = data.dropna()
        num_cols = data.columns.drop('date')
        data[num_cols] = data[num_cols].apply(pd.to_numeric, errors='coerce')
        data['date'] = pd.to_datetime(data['date'], utc=True)
        data['date'] = data['date'].dt.tz_localize(None)
        data['date'] = data['date'].dt.normalize()
        data = data.fillna(method='ffill')
        data = data.set_index('date')
        data = data.dropna()
        H_LIST = [22]
        time = '2018-05-02'
        time_stop = '2022-02-25'
        error_by_H = {H: [] for H in H_LIST}
        for country in data.columns:
            data_RV_base = data.copy()
            data_RV_base.loc[:, data_RV_base.columns != f'{country}'] = data_RV_base.loc[:, data_RV_base.columns != f'{country}'].shift(1)
            data_RV_base['RV_now'] = data_RV_base[f'{country}']
            data_RV_base['RV_D'] = data_RV_base['RV_now'].shift(1)
            data_RV_base['RV_W'] = data_RV_base['RV_D'].rolling(5).mean()
            data_RV_base['RV_M'] = data_RV_base['RV_D'].rolling(22).mean()
            for H in H_LIST:
                target_col = f'RV_MA{H}'
                data_RV = data_RV_base.copy()
                data_RV[target_col] = data_RV['RV_now'].rolling(window=H, center=False).mean()
                data_RV[target_col] = data_RV[target_col].shift(-(H - 1))
                data_RV = data_RV.drop(columns=[f'{country}'])
                data_RV = data_RV.drop(columns=['RV_now'])
                data_RV = data_RV.dropna()
                data_RV = data_RV.loc['2014-07-04':]
                data_test = data_RV.loc[time:time_stop]
                data_predict = data_test[target_col]
                if len(data_test) == 0:
                    error_by_H[H].append([np.nan, np.nan, np.nan])
                    continue
                result = []
                window_len = len(data_RV.loc[:time]) - 1
                for i in tqdm(range(len(data_RV.loc[time:time_stop])), desc=f'{country}-H{H}'):
                    data_rolling = data_RV.iloc[i:i + window_len]
                    data_rolling = data_rolling.dropna()
                    pca = PCA(n_components=1)
                    pca.fit(data_rolling.drop(columns=[target_col, 'RV_D', 'RV_W', 'RV_M']))
                    data_rolling['PCA_1'] = pca.fit_transform(data_rolling.drop(columns=[target_col, 'RV_D', 'RV_W', 'RV_M']))
                    mod = smf.ols(f'{target_col} ~ RV_D + RV_W + RV_M + PCA_1', data_rolling)
                    res = mod.fit()
                    params_new = res.params
                    r_predict = params_new[0] + params_new[1] * data_test['RV_D'].iloc[i] + params_new[2] * data_test['RV_W'].iloc[i] + params_new[3] * data_test['RV_M'].iloc[i] + params_new[4] * data_rolling['PCA_1'].iloc[-1]
                    result.append(r_predict)
                result = pd.Series(result)
                result.index = data_predict.index
                SUM = pd.concat([data_predict, result], axis=1)
                SUM.columns = ['RV_True', 'RV_Predict']
                SUM_final = SUM
                SUM_final = SUM_final.reset_index(drop=True)
                SUM_final.index = pd.to_datetime(data_test.index)
                SUM_final['MSE'] = (SUM_final['RV_True'] - SUM_final['RV_Predict']) ** 2
                SUM_final['MAE'] = np.abs(SUM_final['RV_True'] - SUM_final['RV_Predict'])
                SUM_final['Qlike'] = SUM_final['RV_True'] / SUM_final['RV_Predict'] + np.log(SUM_final['RV_Predict'])
                MSE = np.mean((SUM_final['RV_True'] - SUM_final['RV_Predict']) ** 2)
                MAE = np.mean(np.abs(SUM_final['RV_True'] - SUM_final['RV_Predict']))
                Theil = np.sqrt(np.mean((SUM_final['RV_True'] - SUM_final['RV_Predict']) ** 2)) / (np.sqrt(np.mean(SUM_final['RV_Predict'] ** 2)) + np.sqrt(np.mean(SUM_final['RV_True'] ** 2)))
                error_by_H[H].append([MSE, MAE, Theil])
                model_outputs[str(country).strip().lstrip('.')] = SUM_final.copy()
        for H in H_LIST:
            ERROR = pd.DataFrame(error_by_H[H], columns=['MSE', 'MAE', 'Theil'], index=data.columns).T
        return model_outputs

    def run_pca_bic():
        model_outputs = {}
        import math
        import scipy as sc
        import time
        import os
        from scipy.stats import norm
        import pandas as pd
        import numpy as np
        from functools import partial
        from tqdm import tqdm as _tqdm
        tqdm = partial(_tqdm, bar_format="{bar} {percentage:3.0f}%", ncols=50)
        import statistics
        import statsmodels.formula.api as smf
        from sklearn.linear_model import LinearRegression
        import statsmodels.api as sm
        from sklearn.decomposition import PCA
        data = pd.read_csv(RV_FILE)
        num_cols = data.columns.drop('date')
        data[num_cols] = data[num_cols].apply(pd.to_numeric, errors='coerce')
        data = data.dropna()
        data['date'] = pd.to_datetime(data['date'], utc=True)
        data['date'] = data['date'].dt.tz_localize(None)
        data['date'] = data['date'].dt.normalize()
        data = data.fillna(method='ffill')
        data = data.set_index('date')
        data = data.dropna()
        H_LIST = [22]
        time = '2018-05-02'
        time_stop = '2022-02-25'
        error_by_H = {H: [] for H in H_LIST}
        bics_by_H = {H: [] for H in H_LIST}
        for country in data.columns:
            data_RV_base = data.copy()
            data_RV_base.loc[:, data_RV_base.columns != f'{country}'] = data_RV_base.loc[:, data_RV_base.columns != f'{country}'].shift(1)
            data_RV_base['RV_now'] = data_RV_base[f'{country}']
            data_RV_base['RV_D'] = data_RV_base['RV_now'].shift(1)
            data_RV_base['RV_W'] = data_RV_base['RV_D'].rolling(5).mean()
            data_RV_base['RV_M'] = data_RV_base['RV_D'].rolling(22).mean()
            for H in H_LIST:
                target_col = f'RV_MA{H}'
                data_RV = data_RV_base.copy()
                if H == 1:
                    data_RV[target_col] = data_RV['RV_now']
                else:
                    data_RV[target_col] = data_RV['RV_now'].rolling(window=H, center=False).mean()
                    data_RV[target_col] = data_RV[target_col].shift(-(H - 1))
                data_RV = data_RV.drop(columns=[f'{country}'])
                data_RV = data_RV.drop(columns=['RV_now'])
                data_RV = data_RV.dropna()
                data_RV = data_RV.loc['2014-07-04':]
                data_test = data_RV.loc[time:time_stop]
                data_predict = data_test[target_col]
                if len(data_test) == 0:
                    error_by_H[H].append([np.nan, np.nan, np.nan])
                    bics_by_H[H].append(np.nan)
                    continue
                window_len = len(data_RV.loc[:time]) - 1
                if window_len <= 0:
                    error_by_H[H].append([np.nan, np.nan, np.nan])
                    bics_by_H[H].append(np.nan)
                    continue
                bic_values = []
                for n_pca in range(1, 17):
                    result = []
                    n_samples = len(data_RV.iloc[0:0 + window_len])
                    k = 4 + n_pca
                    for i in tqdm(range(len(data_RV.loc[time:time_stop])), desc=f'{country}-H{H}-BIC'):
                        variables = []
                        data_rolling = data_RV.iloc[i:i + window_len].copy()
                        data_rolling = data_rolling.dropna()
                        data_rolling = data_rolling.iloc[:-21]
                        if len(data_rolling) <= max(n_pca, 5):
                            result.append(np.nan)
                            continue
                        pca = PCA(n_components=n_pca)
                        X_other = data_rolling.drop(columns=[target_col, 'RV_D', 'RV_W', 'RV_M'])
                        pca.fit(X_other)
                        principal_components = pca.transform(X_other)
                        columns_names = [f'pca_{name}' for name in range(1, n_pca + 1)]
                        principal_df = pd.DataFrame(data=principal_components, columns=columns_names, index=data_rolling.index)
                        variables.append(data_rolling[target_col])
                        variables.append(data_rolling['RV_D'])
                        variables.append(data_rolling['RV_W'])
                        variables.append(data_rolling['RV_M'])
                        for colu in principal_df.columns:
                            variables.append(principal_df[colu])
                        variables_df = pd.concat(variables, axis=1)
                        pca_columns = [f'pca_{j}' for j in range(1, n_pca + 1)]
                        formula = f'{target_col} ~ RV_D + RV_W + RV_M + ' + ' + '.join(pca_columns)
                        mod = smf.ols(formula, data=variables_df)
                        res = mod.fit()
                        params_new = res.params
                        r_predict = params_new[0] + params_new[1] * data_test['RV_D'].iloc[i] + params_new[2] * data_test['RV_W'].iloc[i] + params_new[3] * data_test['RV_M'].iloc[i] + sum((params_new[4 + j] * principal_df.iloc[-1, j] for j in range(n_pca)))
                        result.append(r_predict)
                    result = pd.Series(result, index=data_predict.index)
                    SUM = pd.concat([data_predict, result], axis=1)
                    SUM.columns = ['RV_True', 'RV_Predict']
                    SUM_final = SUM.reset_index(drop=True)
                    SUM_final.index = pd.to_datetime(data_test.index)
                    mse_tmp = np.mean((SUM_final['RV_True'] - SUM_final['RV_Predict']) ** 2)
                    if not np.isfinite(mse_tmp) or mse_tmp <= 0:
                        bic_values.append(np.inf)
                        continue
                    log_likelihood = -0.5 * n_samples * (np.log(2 * np.pi) + np.log(mse_tmp) + 1)
                    bic = -2 * log_likelihood + k * np.log(n_samples)
                    bic_values.append(bic)
                bic_range = []
                for l in range(1, 18):
                    bic_range.append(l)
                best_pca_number = bic_range[int(np.argmin(bic_values))]
                bics_by_H[H].append(best_pca_number)
                result_final = []
                for m in tqdm(range(len(data_RV.loc[time:time_stop])), desc=f'{country}-H{H}-FINAL'):
                    variables = []
                    data_rolling = data_RV.iloc[m:m + window_len]
                    data_rolling = data_rolling.dropna()
                    if len(data_rolling) <= max(best_pca_number, 5):
                        result_final.append(np.nan)
                        continue
                    pca = PCA(n_components=best_pca_number)
                    X_other = data_rolling.drop(columns=[target_col, 'RV_D', 'RV_W', 'RV_M'])
                    pca.fit(X_other)
                    principal_components = pca.transform(X_other)
                    columns_names = [f'pca_{name}' for name in range(1, best_pca_number + 1)]
                    principal_df = pd.DataFrame(data=principal_components, columns=columns_names, index=data_rolling.index)
                    variables.append(data_rolling[target_col])
                    variables.append(data_rolling['RV_D'])
                    variables.append(data_rolling['RV_W'])
                    variables.append(data_rolling['RV_M'])
                    for colu in principal_df.columns:
                        variables.append(principal_df[colu])
                    variables_df = pd.concat(variables, axis=1)
                    pca_columns = [f'pca_{j}' for j in range(1, best_pca_number + 1)]
                    formula = f'{target_col} ~ RV_D + RV_W + RV_M + ' + ' + '.join(pca_columns)
                    mod = smf.ols(formula, data=variables_df)
                    res = mod.fit()
                    params_new = res.params
                    r_predict = params_new[0] + params_new[1] * data_test['RV_D'].iloc[m] + params_new[2] * data_test['RV_W'].iloc[m] + params_new[3] * data_test['RV_M'].iloc[m] + sum((params_new[4 + j] * principal_df.iloc[-1, j] for j in range(best_pca_number)))
                    result_final.append(r_predict)
                result_final = pd.Series(result_final, index=data_predict.index)
                SUM = pd.concat([data_predict, result_final], axis=1)
                SUM.columns = ['RV_True', 'RV_Predict']
                SUM_final = SUM.reset_index(drop=True)
                SUM_final.index = pd.to_datetime(data_test.index)
                SUM_final['MSE'] = (SUM_final['RV_True'] - SUM_final['RV_Predict']) ** 2
                SUM_final['MAE'] = np.abs(SUM_final['RV_True'] - SUM_final['RV_Predict'])
                SUM_final['Qlike'] = SUM_final['RV_True'] / SUM_final['RV_Predict'] + np.log(SUM_final['RV_Predict'])
                MSE = np.mean((SUM_final['RV_True'] - SUM_final['RV_Predict']) ** 2)
                MAE = np.mean(np.abs(SUM_final['RV_True'] - SUM_final['RV_Predict']))
                Theil = np.sqrt(np.mean((SUM_final['RV_True'] - SUM_final['RV_Predict']) ** 2)) / (np.sqrt(np.mean(SUM_final['RV_Predict'] ** 2)) + np.sqrt(np.mean(SUM_final['RV_True'] ** 2)))
                error_by_H[H].append([MSE, MAE, Theil])
                model_outputs[str(country).strip().lstrip('.')] = SUM_final.copy()
        for H in H_LIST:
            ERROR = pd.DataFrame(error_by_H[H], columns=['MSE', 'MAE', 'Theil'], index=data.columns).T
        return model_outputs

    def run_pls():
        model_outputs = {}
        import math
        import scipy as sc
        import time
        import os
        from scipy.stats import norm
        import pandas as pd
        import numpy as np
        from functools import partial
        from tqdm import tqdm as _tqdm
        tqdm = partial(_tqdm, bar_format="{bar} {percentage:3.0f}%", ncols=50)
        import statistics
        import statsmodels.formula.api as smf
        from sklearn.linear_model import LinearRegression
        import statsmodels.api as sm
        from sklearn.cross_decomposition import PLSRegression
        data = pd.read_csv(RV_FILE)
        data = data.dropna()
        data['date'] = pd.to_datetime(data['date'], utc=True)
        data['date'] = data['date'].dt.tz_localize(None)
        data['date'] = data['date'].dt.normalize()
        data = data.fillna(method='ffill')
        data = data.set_index('date')
        H = 22
        time_start = '2018-05-02'
        time_stop = '2022-02-25'
        error_H22 = []
        for country in data.columns:
            data_RV_base = data.copy()
            data_RV_base.loc[:, data_RV_base.columns != f'{country}'] = data_RV_base.loc[:, data_RV_base.columns != f'{country}'].shift(1)
            data_RV_base['RV_now'] = data_RV_base[f'{country}']
            data_RV_base['RV_D'] = data_RV_base['RV_now'].shift(1)
            data_RV_base['RV_W'] = data_RV_base['RV_D'].rolling(5).mean()
            data_RV_base['RV_M'] = data_RV_base['RV_D'].rolling(22).mean()
            target_col = f'RV_MA{H}'
            data_RV = data_RV_base.copy()
            data_RV[target_col] = data_RV['RV_now'].rolling(window=H, center=False).mean()
            data_RV[target_col] = data_RV[target_col].shift(-22)
            data_RV = data_RV.drop(columns=[f'{country}'])
            data_RV = data_RV.drop(columns=['RV_now'])
            data_RV = data_RV.dropna()
            data_RV = data_RV.loc['2014-07-04':]
            data_predict = data_RV.loc[time_start:time_stop][target_col]
            if len(data_predict) == 0:
                error_H22.append([np.nan, np.nan, np.nan])
                continue
            result = []
            ols = LinearRegression()
            Sensitivity = []
            for other_country in data_RV.columns.drop([target_col, 'RV_D', 'RV_W', 'RV_M']):
                X_train = data_RV.loc[:time_start][target_col].values.reshape(-1, 1)
                y_train = data_RV.loc[:time_start][f'{other_country}'].values
                ols.fit(X_train, y_train)
                Sensitivity.append(ols.coef_[0])
            data_RV_second_step = data_RV.copy()
            data_RV['FPLS'] = 0.0
            for t in tqdm(range(len(data_RV)), desc=f'{country}-H{H}-FPLS'):
                X_train = np.array(Sensitivity).reshape(-1, 1)
                y_train = data_RV_second_step.drop(columns=[target_col, 'RV_D', 'RV_W', 'RV_M']).iloc[t]
                ols.fit(X_train, y_train)
                data_RV.loc[data_RV.index[t], 'FPLS'] = ols.coef_[0]
            data_test = data_RV.loc[time_start:time_stop]
            window_len = len(data_RV.loc[:time_start]) - 1
            for i in tqdm(range(len(data_RV.loc[time_start:time_stop])), desc=f'{country}-H{H}-ROLL'):
                data_rolling = data_RV.iloc[i:i + window_len].copy()
                data_rolling = data_rolling.dropna()
                mod = smf.ols(f'{target_col} ~ RV_D + RV_W + RV_M + FPLS', data_rolling)
                res = mod.fit()
                params_new = res.params
                r_predict = params_new.iloc[0] + params_new.iloc[1] * data_test['RV_D'].iloc[i] + params_new.iloc[2] * data_test['RV_W'].iloc[i] + params_new.iloc[3] * data_test['RV_M'].iloc[i] + params_new.iloc[4] * data_test['FPLS'].iloc[-1]
                result.append(r_predict)
            result = pd.Series(result, index=data_predict.index)
            SUM = pd.concat([data_predict, result], axis=1)
            SUM.columns = ['RV_True', 'RV_Predict']
            SUM_final = SUM.reset_index(drop=True)
            SUM_final.index = pd.to_datetime(data_test.index)
            SUM_final['MSE'] = (SUM_final['RV_True'] - SUM_final['RV_Predict']) ** 2
            SUM_final['MAE'] = np.abs(SUM_final['RV_True'] - SUM_final['RV_Predict'])
            SUM_final['Qlike'] = SUM_final['RV_True'] / SUM_final['RV_Predict'] + np.log(SUM_final['RV_Predict'])
            MSE = np.mean((SUM_final['RV_True'] - SUM_final['RV_Predict']) ** 2)
            MAE = np.mean(np.abs(SUM_final['RV_True'] - SUM_final['RV_Predict']))
            Theil = np.sqrt(np.mean((SUM_final['RV_True'] - SUM_final['RV_Predict']) ** 2)) / (np.sqrt(np.mean(SUM_final['RV_Predict'] ** 2)) + np.sqrt(np.mean(SUM_final['RV_True'] ** 2)))
            error_H22.append([MSE, MAE, Theil])
            model_outputs[str(country).strip().lstrip('.')] = SUM_final.copy()
        ERROR = pd.DataFrame(error_H22, columns=['MSE', 'MAE', 'Theil'], index=data.columns).T
        return model_outputs

    def run_ridge():
        model_outputs = {}
        import math
        import scipy as sc
        import time
        import os
        from scipy.stats import norm
        import pandas as pd
        import numpy as np
        from functools import partial
        from tqdm import tqdm as _tqdm
        tqdm = partial(_tqdm, bar_format="{bar} {percentage:3.0f}%", ncols=50)
        import statistics
        import statsmodels.formula.api as smf
        from sklearn.linear_model import RidgeCV
        data = pd.read_csv(RV_FILE)
        data = data.dropna()
        data['date'] = pd.to_datetime(data['date'], utc=True)
        data['date'] = data['date'].dt.tz_localize(None)
        data['date'] = data['date'].dt.normalize()
        data = data.fillna(method='ffill')
        data = data.set_index('date')
        data = data.dropna()
        H_LIST = [22]
        time_start = '2018-05-02'
        time_stop = '2022-02-25'
        error_by_H = {H: [] for H in H_LIST}
        for country in data.columns:
            data_RV_base = data.copy()
            data_RV_base.loc[:, data_RV_base.columns != f'{country}'] = data_RV_base.loc[:, data_RV_base.columns != f'{country}'].shift(1)
            data_RV_base['RV_now'] = data_RV_base[f'{country}']
            data_RV_base['RV_D'] = data_RV_base['RV_now'].shift(1)
            data_RV_base['RV_W'] = data_RV_base['RV_D'].rolling(5).mean()
            data_RV_base['RV_M'] = data_RV_base['RV_D'].rolling(22).mean()
            for H in H_LIST:
                target_col = f'RV_MA{H}'
                data_RV = data_RV_base.copy()
                if H == 1:
                    data_RV[target_col] = data_RV['RV_now']
                else:
                    data_RV[target_col] = data_RV['RV_now'].rolling(window=H, center=False).mean()
                    data_RV[target_col] = data_RV[target_col].shift(-(H - 1))
                data_RV = data_RV.drop(columns=[f'{country}', 'RV_now'])
                data_RV = data_RV.dropna()
                data_RV = data_RV.loc['2014-07-04':]
                data_test = data_RV.loc[time_start:time_stop]
                if len(data_test) == 0:
                    error_by_H[H].append([np.nan, np.nan, np.nan])
                    continue
                test_start_date = data_test.index[0]
                window_len = len(data_RV.loc[:test_start_date]) - 1
                if window_len <= 10:
                    error_by_H[H].append([np.nan, np.nan, np.nan])
                    continue
                X_all = data_RV.drop(columns=[target_col])
                y_all = data_RV[target_col]
                all_index = data_RV.index
                test_index = data_test.index
                start_pos = all_index.get_loc(test_start_date)
                result = []
                for s in tqdm(range(len(test_index)), desc=f'{country}-H{H}'):
                    test_pos = start_pos + s
                    train_end_pos = test_pos - 1
                    train_start_pos = train_end_pos - (window_len - 1)
                    if train_start_pos < 0:
                        result.append(np.nan)
                        continue
                    data_rolling = data_RV.iloc[train_start_pos:train_end_pos + 1].copy()
                    if H > 1:
                        if len(data_rolling) <= H - 1:
                            result.append(np.nan)
                            continue
                        data_rolling = data_rolling.iloc[:-(H - 1)]
                    data_rolling = data_rolling.dropna()
                    if len(data_rolling) < 30:
                        result.append(np.nan)
                        continue
                    X_train = data_rolling.drop(columns=[target_col])
                    y_train = data_rolling[target_col]
                    ridge_cv = RidgeCV(alphas=np.logspace(-4, 1, 10), cv=3)
                    ridge_cv.fit(X_train, y_train)
                    X_test_row = X_all.iloc[test_pos].values.reshape(1, -1)
                    r_predict = ridge_cv.predict(X_test_row)[0]
                    result.append(r_predict)
                result = pd.Series(result, index=test_index)
                SUM = pd.concat([y_all.loc[test_index], result], axis=1)
                SUM.columns = ['RV_True', 'RV_Predict']
                SUM_final = SUM.reset_index(drop=True)
                SUM_final.index = pd.to_datetime(test_index)
                eps = 1e-12
                SUM_final['RV_Predict_clip'] = np.maximum(SUM_final['RV_Predict'], eps)
                SUM_final['MSE'] = (SUM_final['RV_True'] - SUM_final['RV_Predict']) ** 2
                SUM_final['MAE'] = np.abs(SUM_final['RV_True'] - SUM_final['RV_Predict'])
                SUM_final['Qlike'] = SUM_final['RV_True'] / SUM_final['RV_Predict_clip'] + np.log(SUM_final['RV_Predict_clip'])
                MSE = np.nanmean(SUM_final['MSE'])
                MAE = np.nanmean(SUM_final['MAE'])
                Theil = np.sqrt(np.nanmean((SUM_final['RV_True'] - SUM_final['RV_Predict']) ** 2)) / (np.sqrt(np.nanmean(SUM_final['RV_Predict'] ** 2)) + np.sqrt(np.nanmean(SUM_final['RV_True'] ** 2)))
                error_by_H[H].append([MSE, MAE, Theil])
                model_outputs[str(country).strip().lstrip('.')] = SUM_final.drop(columns=['RV_Predict_clip']).copy()
        for H in H_LIST:
            ERROR = pd.DataFrame(error_by_H[H], columns=['MSE', 'MAE', 'Theil'], index=data.columns).T
        return model_outputs

    def run_lasso():
        model_outputs = {}
        import math
        import scipy as sc
        import time
        import os
        from scipy.stats import norm
        import pandas as pd
        import numpy as np
        from functools import partial
        from tqdm import tqdm as _tqdm
        tqdm = partial(_tqdm, bar_format="{bar} {percentage:3.0f}%", ncols=50)
        import statistics
        import statsmodels.formula.api as smf
        from sklearn.linear_model import Lasso
        from sklearn.linear_model import LassoCV
        data = pd.read_csv(RV_FILE)
        data = data.dropna()
        data['date'] = pd.to_datetime(data['date'], utc=True)
        data['date'] = data['date'].dt.tz_localize(None)
        data['date'] = data['date'].dt.normalize()
        data = data.fillna(method='ffill')
        data = data.set_index('date')
        data = data.dropna()
        H_LIST = [22]
        time = '2018-05-02'
        time_stop = '2022-02-25'
        error_by_H = {H: [] for H in H_LIST}
        for country in data.columns:
            data_RV_base = data.copy()
            data_RV_base.loc[:, data_RV_base.columns != '%s' % country] = data_RV_base.loc[:, data_RV_base.columns != '%s' % country].shift(1)
            data_RV_base['RV_now'] = data_RV_base['%s' % country]
            data_RV_base['RV_D'] = data_RV_base['RV_now'].shift(1)
            data_RV_base['RV_W'] = data_RV_base['RV_D'].rolling(5).mean()
            data_RV_base['RV_M'] = data_RV_base['RV_D'].rolling(22).mean()
            data_RV_base = data_RV_base.drop(columns=['%s' % country])
            data_RV_base = data_RV_base.drop(columns=['RV_now'])
            data_RV_base = data_RV_base.dropna()
            data_RV_base = data_RV_base.loc['2014-07-04':]
            for H in H_LIST:
                target_col = f'RV_MA{H}'
                data_RV = data_RV_base.copy()
                rv_now = data[country].rename('RV_now')
                data_RV = pd.concat([data_RV, rv_now], axis=1)
                if H == 1:
                    data_RV[target_col] = data_RV['RV_now']
                else:
                    data_RV[target_col] = data_RV['RV_now'].rolling(window=H, center=False).mean()
                    data_RV[target_col] = data_RV[target_col].shift(-(H - 1))
                data_RV = data_RV.drop(columns=['RV_now'])
                data_RV = data_RV.dropna()
                data_RV = data_RV.loc['2014-07-04':]
                data_test = data_RV.loc[time:time_stop]
                data_predict = data_RV.loc[time:time_stop][target_col]
                result = []
                if len(data_test) == 0:
                    error_by_H[H].append([np.nan, np.nan, np.nan])
                    continue
                window_len = len(data_RV.loc[:time]) - 1
                if window_len <= 0:
                    error_by_H[H].append([np.nan, np.nan, np.nan])
                    continue
                for i in tqdm(range(len(data_RV.loc[time:time_stop])), desc=f'{country}-H{H}'):
                    data_rolling = data_RV.iloc[i:i + window_len]
                    data_rolling = data_rolling.dropna()
                    X_train = data_rolling.drop(columns=[target_col])
                    y_train = data_rolling[target_col]
                    lasso_cv = LassoCV(alphas=np.logspace(-4, 1, 50), cv=5)
                    lasso_cv.fit(X_train, y_train)
                    X_test = np.array([data_test.drop(columns=[target_col]).iloc[i].values])
                    r_predict = lasso_cv.predict(X_test)[0]
                    result.append(r_predict)
                result = pd.Series(result)
                result.index = data_predict.index
                SUM = pd.concat([data_predict, result], axis=1)
                SUM.columns = ['RV_True', 'RV_Predict']
                SUM_final = SUM
                SUM_final = SUM_final.reset_index(drop=True)
                SUM_final.index = pd.to_datetime(data_test.index)
                SUM_final['MSE'] = (SUM_final['RV_True'] - SUM_final['RV_Predict']) ** 2
                SUM_final['MAE'] = np.abs(SUM_final['RV_True'] - SUM_final['RV_Predict'])
                SUM_final['Qlike'] = SUM_final['RV_True'] / SUM_final['RV_Predict'] + np.log(SUM_final['RV_Predict'])
                MSE = np.mean((SUM_final['RV_True'] - SUM_final['RV_Predict']) ** 2)
                MAE = np.mean(np.abs(SUM_final['RV_True'] - SUM_final['RV_Predict']))
                Theil = np.sqrt(np.mean((SUM_final['RV_True'] - SUM_final['RV_Predict']) ** 2)) / (np.sqrt(np.mean(SUM_final['RV_Predict'] ** 2)) + np.sqrt(np.mean(SUM_final['RV_True'] ** 2)))
                error_by_H[H].append([MSE, MAE, Theil])
                model_outputs[str(country).strip().lstrip('.')] = SUM_final.copy()
        for H in H_LIST:
            ERROR = pd.DataFrame(error_by_H[H], columns=['MSE', 'MAE', 'Theil'], index=data.columns).T
        return model_outputs

    def run_enet():
        model_outputs = {}
        import math
        import scipy as sc
        import time
        import os
        from scipy.stats import norm
        import pandas as pd
        import numpy as np
        from functools import partial
        from tqdm import tqdm as _tqdm
        tqdm = partial(_tqdm, bar_format="{bar} {percentage:3.0f}%", ncols=50)
        import statistics
        import statsmodels.formula.api as smf
        from sklearn.linear_model import ElasticNetCV
        import warnings
        warnings.simplefilter('ignore', FutureWarning)
        data = pd.read_csv(RV_FILE)
        data = data.dropna()
        num_cols = data.columns.drop('date')
        data[num_cols] = data[num_cols].apply(pd.to_numeric, errors='coerce')
        data['date'] = pd.to_datetime(data['date'], utc=True)
        data['date'] = data['date'].dt.tz_localize(None)
        data['date'] = data['date'].dt.normalize()
        data = data.fillna(method='ffill')
        data = data.set_index('date')
        data = data.dropna()
        H_LIST = [22]
        time_start = '2018-05-02'
        time_stop = '2022-02-25'
        error_by_H = {H: [] for H in H_LIST}
        for country in data.columns:
            df = data.copy()
            other_cols = [c for c in df.columns if c != country]
            df.loc[:, other_cols] = df.loc[:, other_cols].shift(1)
            df['RV_now'] = df[country]
            df['RV_D'] = df['RV_now'].shift(1)
            df['RV_W'] = df['RV_D'].rolling(5).mean()
            df['RV_M'] = df['RV_D'].rolling(22).mean()
            df = df.drop(columns=[country])
            feature_base = df.drop(columns=['RV_now'])
            feature_base = feature_base.loc['2014-07-04':].dropna()
            for H in H_LIST:
                target_col = f'RV_MA{H}'
                df_model = feature_base.copy()
                rv_now_aligned = df.loc[feature_base.index, 'RV_now']
                if H == 1:
                    df_model[target_col] = rv_now_aligned.shift(-1)
                else:
                    df_model[target_col] = rv_now_aligned.rolling(window=H).mean().shift(-(H - 1))
                df_model = df_model.dropna()
                data_test = df_model.loc[time_start:time_stop]
                if len(data_test) == 0:
                    error_by_H[H].append([np.nan, np.nan, np.nan])
                    continue
                y_true = data_test[target_col]
                window_len = len(df_model.loc[:time_start]) - 1
                if window_len <= 0:
                    error_by_H[H].append([np.nan, np.nan, np.nan])
                    continue
                feature_cols = [c for c in df_model.columns if c != target_col]
                preds = []
                test_index = data_test.index
                for i in tqdm(range(len(test_index)), desc=f'{country}-H{H}'):
                    train_slice = df_model.iloc[i:i + window_len].dropna()
                    if len(train_slice) < 50:
                        preds.append(np.nan)
                        continue
                    X_train = train_slice[feature_cols].values
                    y_train = train_slice[target_col].values
                    elastic_net_cv = ElasticNetCV(alphas=None, l1_ratio=np.linspace(0.1, 1, 10), cv=5)
                    elastic_net_cv.fit(X_train, y_train)
                    X_test = data_test.iloc[i][feature_cols].values.reshape(1, -1)
                    preds.append(float(elastic_net_cv.predict(X_test)[0]))
                y_pred = pd.Series(preds, index=test_index)
                SUM_final = pd.DataFrame({'RV_True': y_true, 'RV_Predict': y_pred}).dropna()
                if SUM_final.empty:
                    error_by_H[H].append([np.nan, np.nan, np.nan])
                    model_outputs[str(country).strip().lstrip('.')] = SUM_final.copy()
                    continue
                SUM_final['MSE'] = (SUM_final['RV_True'] - SUM_final['RV_Predict']) ** 2
                SUM_final['MAE'] = np.abs(SUM_final['RV_True'] - SUM_final['RV_Predict'])
                SUM_final['Qlike'] = SUM_final['RV_True'] / SUM_final['RV_Predict'] + np.log(SUM_final['RV_Predict'])
                MSE = float(SUM_final['MSE'].mean())
                MAE = float(SUM_final['MAE'].mean())
                Theil = float(np.sqrt(MSE) / (np.sqrt(np.mean(SUM_final['RV_Predict'] ** 2)) + np.sqrt(np.mean(SUM_final['RV_True'] ** 2))))
                error_by_H[H].append([MSE, MAE, Theil])
                model_outputs[str(country).strip().lstrip('.')] = SUM_final.copy()
        for H in H_LIST:
            ERROR = pd.DataFrame(error_by_H[H], columns=['MSE', 'MAE', 'Theil'], index=data.columns).T
        return model_outputs

    def run_spillover_time():
        model_outputs = {}
        import pandas as pd
        import numpy as np
        import math
        import scipy as sc
        import time
        from functools import partial
        from tqdm import tqdm as _tqdm
        tqdm = partial(_tqdm, bar_format="{bar} {percentage:3.0f}%", ncols=50)
        import os
        from scipy.stats import norm
        import statsmodels.formula.api as smf
        import warnings
        warnings.simplefilter('ignore', FutureWarning)
        data = pd.read_csv(RV_FILE)
        data = data.dropna()
        data['date'] = pd.to_datetime(data['date'], utc=True)
        data['date'] = data['date'].dt.tz_localize(None)
        data['date'] = data['date'].dt.normalize()
        data = data.fillna(method='bfill')
        data = data.set_index('date')
        data_CI_time = pd.read_csv(TCI_FILE)
        data_CI_time['date'] = pd.to_datetime(data_CI_time['Unnamed: 0'], format='mixed')
        data_CI_time = data_CI_time.set_index('date')
        H_LIST = [22]
        time = '2018-05-02'
        time_stop = '2022-02-25'
        error_by_H = {H: [] for H in H_LIST}
        for country in data.columns:
            data_RV_base = pd.concat([data['%s' % country], data_CI_time['TCI.Total']], axis=1)
            data_RV_base = data_RV_base.dropna()
            data_RV_base = data_RV_base.loc['2014-07-04':]
            data_RV_base.columns = ['RV_now', 'CI_time']
            decay_factor = np.exp(-np.arange(len(data_RV_base)) / 7 * math.pi)
            CI_time_ED = np.tril(np.outer(data_RV_base['CI_time'], decay_factor))
            data_RV_base['CI_time_ED'] = CI_time_ED.sum(axis=1)
            data_RV_base['CI_time_D'] = data_RV_base['CI_time_ED'].shift(1)
            data_RV_base['RV_D'] = data_RV_base['RV_now'].shift(1)
            data_RV_base['RV_W'] = data_RV_base['RV_D'].rolling(5).mean()
            data_RV_base['RV_M'] = data_RV_base['RV_D'].rolling(22).mean()
            for H in H_LIST:
                target_col = f'RV_MA{H}'
                data_RV = data_RV_base.copy()
                if H == 1:
                    data_RV[target_col] = data_RV['RV_now']
                else:
                    data_RV[target_col] = data_RV['RV_now'].rolling(window=H, center=False).mean()
                    data_RV[target_col] = data_RV[target_col].shift(-(H - 1))
                data_RV = data_RV.dropna()
                mod = smf.ols(f'{target_col} ~ RV_D + RV_W + RV_M + CI_time_D', data_RV)
                res = mod.fit()
                data_test = data_RV.loc[time:time_stop]
                data_predict = data_RV.loc[time:time_stop][target_col]
                if len(data_test) == 0:
                    error_by_H[H].append([np.nan, np.nan, np.nan])
                    continue
                result = []
                window_len = len(data_RV.loc[:time]) - 1
                if window_len <= 0:
                    error_by_H[H].append([np.nan, np.nan, np.nan])
                    continue
                for i in tqdm(range(len(data_RV.loc[time:time_stop])), desc=f'{country}-H{H}'):
                    data_rolling = data_RV.iloc[i:i + window_len]
                    mod = smf.ols(f'{target_col} ~ RV_D + RV_W + RV_M + CI_time_D', data_rolling)
                    res = mod.fit()
                    params_new = res.params
                    r_predict = params_new[0] + params_new[1] * data_test['RV_D'].iloc[i] + params_new[2] * data_test['RV_W'].iloc[i] + params_new[3] * data_test['RV_M'].iloc[i] + params_new[4] * data_test['CI_time_D'].iloc[i]
                    result.append(r_predict)
                result = pd.Series(result)
                result.index = data_predict.index
                SUM = pd.concat([data_predict, result], axis=1)
                SUM.columns = ['RV_True', 'RV_Predict']
                SUM_final = SUM
                SUM_final = SUM_final.reset_index(drop=True)
                SUM_final.index = pd.to_datetime(data_test.index)
                SUM_final['MSE'] = (SUM_final['RV_True'] - SUM_final['RV_Predict']) ** 2
                SUM_final['MAE'] = np.abs(SUM_final['RV_True'] - SUM_final['RV_Predict'])
                SUM_final['Qlike'] = SUM_final['RV_True'] / SUM_final['RV_Predict'] + np.log(SUM_final['RV_Predict'])
                MSE = np.mean((SUM_final['RV_True'] - SUM_final['RV_Predict']) ** 2)
                MAE = np.mean(np.abs(SUM_final['RV_True'] - SUM_final['RV_Predict']))
                Theil = np.sqrt(np.mean((SUM_final['RV_True'] - SUM_final['RV_Predict']) ** 2)) / (np.sqrt(np.mean(SUM_final['RV_Predict'] ** 2)) + np.sqrt(np.mean(SUM_final['RV_True'] ** 2)))
                error_by_H[H].append([MSE, MAE, Theil])
                model_outputs[str(country).strip().lstrip('.')] = SUM_final.copy()
        for H in H_LIST:
            ERROR = pd.DataFrame(error_by_H[H], columns=['MSE', 'MAE', 'Theil'], index=data.columns).T
        return model_outputs

    def run_spillover_frequency():
        model_outputs = {}
        import pandas as pd
        import numpy as np
        import math
        import scipy as sc
        import time
        from functools import partial
        from tqdm import tqdm as _tqdm
        tqdm = partial(_tqdm, bar_format="{bar} {percentage:3.0f}%", ncols=50)
        import os
        from scipy.stats import norm
        import statsmodels.formula.api as smf
        import warnings
        warnings.simplefilter('ignore', FutureWarning)
        data = pd.read_csv(RV_FILE)
        data = data.dropna()
        data['date'] = pd.to_datetime(data['date'], utc=True)
        data['date'] = data['date'].dt.tz_localize(None)
        data['date'] = data['date'].dt.normalize()
        data = data.fillna(method='ffill')
        data = data.set_index('date')
        data_CI = pd.read_csv(TCI_FILE)
        data_CI['date'] = pd.to_datetime(data_CI['Unnamed: 0'], format='mixed')
        data_CI = data_CI.set_index('date')
        H = 22
        target_col = f'RV_MA{H}'
        time = '2018-05-02'
        time_stop = '2022-02-25'
        error_H22 = []
        for country in data.columns:
            data_RV_base = pd.concat([data[f'{country}'], data_CI['TCI.20.Inf'], data_CI['TCI.5.20'], data_CI['TCI.1.5']], axis=1)
            data_RV_base.columns = ['RV_now', 'CI_long', 'CI_mid', 'CI_short']
            data_RV_base = data_RV_base.fillna(method='bfill')
            data_RV_base['CI_long_D'] = data_RV_base['CI_long'].shift(1)
            data_RV_base['CI_mid_D'] = data_RV_base['CI_mid'].shift(1)
            data_RV_base['CI_short_D'] = data_RV_base['CI_short'].shift(1)
            data_RV_base['RV_D'] = data_RV_base['RV_now'].shift(1)
            data_RV_base['RV_W'] = data_RV_base['RV_D'].rolling(5).mean()
            data_RV_base['RV_M'] = data_RV_base['RV_D'].rolling(22).mean()
            data_RV = data_RV_base.copy()
            data_RV[target_col] = data_RV['RV_now'].rolling(window=H, center=False).mean()
            data_RV[target_col] = data_RV[target_col].shift(-22)
            data_RV = data_RV.loc['2014-07-04':]
            data_RV = data_RV.dropna()
            mod = smf.ols(f'{target_col} ~ RV_D + RV_W + RV_M + CI_long_D + CI_mid_D + CI_short_D', data_RV)
            res = mod.fit()
            data_test = data_RV.loc[time:time_stop]
            data_predict = data_RV.loc[time:time_stop][target_col]
            if len(data_test) == 0:
                error_H22.append([np.nan, np.nan, np.nan])
                continue
            result = []
            window_len = len(data_RV.loc[:time]) - 1
            if window_len <= 0:
                error_H22.append([np.nan, np.nan, np.nan])
                continue
            for i in tqdm(range(len(data_RV.loc[time:time_stop])), desc=f'{country}-H{H}'):
                data_rolling = data_RV.iloc[i:i + window_len]
                mod = smf.ols(f'{target_col} ~ RV_D + RV_W + RV_M + CI_long_D + CI_mid_D + CI_short_D', data_rolling)
                res = mod.fit()
                params_new = res.params
                r_predict = params_new[0] + params_new[1] * data_test['RV_D'].iloc[i] + params_new[2] * data_test['RV_W'].iloc[i] + params_new[3] * data_test['RV_M'].iloc[i] + params_new[4] * data_test['CI_long_D'].iloc[i] + params_new[5] * data_test['CI_mid_D'].iloc[i] + params_new[6] * data_test['CI_short_D'].iloc[i]
                result.append(r_predict)
            result = pd.Series(result)
            result.index = data_predict.index
            SUM = pd.concat([data_predict, result], axis=1)
            SUM.columns = ['RV_True', 'RV_Predict']
            SUM_final = SUM
            SUM_final = SUM_final.reset_index(drop=True)
            SUM_final.index = pd.to_datetime(data_test.index)
            SUM_final['MSE'] = (SUM_final['RV_True'] - SUM_final['RV_Predict']) ** 2
            SUM_final['MAE'] = np.abs(SUM_final['RV_True'] - SUM_final['RV_Predict'])
            SUM_final['Qlike'] = SUM_final['RV_True'] / SUM_final['RV_Predict'] + np.log(SUM_final['RV_Predict'])
            MSE = np.mean((SUM_final['RV_True'] - SUM_final['RV_Predict']) ** 2)
            MAE = np.mean(np.abs(SUM_final['RV_True'] - SUM_final['RV_Predict']))
            Theil = np.sqrt(np.mean((SUM_final['RV_True'] - SUM_final['RV_Predict']) ** 2)) / (np.sqrt(np.mean(SUM_final['RV_Predict'] ** 2)) + np.sqrt(np.mean(SUM_final['RV_True'] ** 2)))
            error_H22.append([MSE, MAE, Theil])
            model_outputs[str(country).strip().lstrip('.')] = SUM_final.copy()
        ERROR = pd.DataFrame(error_H22, columns=['MSE', 'MAE', 'Theil'], index=data.columns).T
        return model_outputs

    def calc_mse(y_true: np.ndarray, y_pred: np.ndarray) -> dict:
        y_true = np.asarray(y_true, dtype=float)
        y_pred = np.asarray(y_pred, dtype=float)
        mask = np.isfinite(y_true) & np.isfinite(y_pred)
        yt = y_true[mask]
        yp = y_pred[mask]
        if len(yt) == 0:
            return {'MSE': np.nan}
        err = yp - yt
        mse = float(np.mean(err ** 2))
        return {'MSE': mse}

    def qlike_with_shared_drop(y_true: pd.Series, y_pred: pd.Series, drop_dates: set) -> float:
        df = pd.concat([y_true.rename('y'), y_pred.rename('yhat')], axis=1).dropna()
        if df.empty:
            return np.nan
        if drop_dates:
            df = df.loc[~df.index.isin(drop_dates)]
        if df.empty:
            return np.nan
        y = df['y'].astype(float).values
        yhat = df['yhat'].astype(float).values
        ok_pos = (y > EPS) & (yhat > EPS)
        if not np.any(ok_pos):
            return np.nan
        ratio = y[ok_pos] / yhat[ok_pos]
        ok_ratio = np.isfinite(ratio) & (ratio >= RATIO_MIN) & (ratio <= RATIO_MAX)
        if not np.any(ok_ratio):
            return np.nan
        ratio = ratio[ok_ratio]
        q = ratio - np.log(ratio) - 1.0
        ok_q = np.isfinite(q) & (q >= 0.0) & (q <= QLIKE_CAP)
        if not np.any(ok_q):
            return np.nan
        return float(np.mean(q[ok_q]))

    def add_summary_rows(df: pd.DataFrame) -> pd.DataFrame:
        mean_row = df.mean(axis=0, skipna=True)
        q1_row = df.quantile(0.25, axis=0, interpolation='linear')
        q2_row = df.quantile(0.5, axis=0, interpolation='linear')
        q3_row = df.quantile(0.75, axis=0, interpolation='linear')
        df2 = df.copy()
        df2.loc['Mean'] = mean_row
        df2.loc['Q1'] = q1_row
        df2.loc['Q2'] = q2_row
        df2.loc['Q3'] = q3_row
        return df2

    def prepare_mcs_matrix_relaxed(sum_pre: pd.DataFrame, min_valid_per_model: int=10, min_rows_for_mcs: int=5):
        info = {'initial_shape': tuple(sum_pre.shape), 'dropped_low_valid_models': [], 'dropped_const_models': [], 'final_shape': None, 'reason': None}
        if sum_pre.empty or sum_pre.shape[1] == 0:
            info['reason'] = 'The loss matrix is empty'
            return (None, info)
        valid_counts = sum_pre.notna().sum(axis=0)
        keep_cols = valid_counts[valid_counts >= min_valid_per_model].index.tolist()
        info['dropped_low_valid_models'] = [c for c in sum_pre.columns if c not in keep_cols]
        if len(keep_cols) < 2:
            info['reason'] = f'Fewer than two models with at least {min_valid_per_model} valid observations'
            return (None, info)
        x = sum_pre[keep_cols].copy()
        x = x.dropna(axis=0, how='any')
        if x.empty:
            info['reason'] = 'No common valid observations'
            return (None, info)
        if x.shape[1] < 2:
            info['reason'] = 'Fewer than two models after dropping missing observations'
            return (None, info)
        if x.shape[0] < min_rows_for_mcs:
            info['reason'] = f'Insufficient common observations ({x.shape[0]} < {min_rows_for_mcs})'
            return (None, info)
        variances = x.var(axis=0, skipna=True)
        non_const_cols = variances[variances > 0].index.tolist()
        info['dropped_const_models'] = [c for c in x.columns if c not in non_const_cols]
        x = x[non_const_cols]
        if x.shape[1] < 2:
            info['reason'] = 'Fewer than two models after removing constant columns'
            return (None, info)
        if x.shape[0] < min_rows_for_mcs:
            info['reason'] = f'Insufficient common observations after removing constant columns ({x.shape[0]} < {min_rows_for_mcs})'
            return (None, info)
        info['final_shape'] = tuple(x.shape)
        return (x, info)

    def choose_one_model_to_drop_for_retry(raw_sum_pre: pd.DataFrame, current_cols: list):
        if len(current_cols) <= 2:
            return None
        sub = raw_sum_pre[current_cols].copy()
        valid_counts = sub.notna().sum(axis=0)
        min_valid = valid_counts.min()
        candidates = valid_counts[valid_counts == min_valid].index.tolist()
        if len(candidates) == 1:
            return candidates[0]
        var_map = {}
        for c in candidates:
            s = pd.to_numeric(sub[c], errors='coerce').dropna()
            if len(s) <= 1:
                var_map[c] = -1.0
            else:
                v = float(s.var())
                if np.isnan(v):
                    v = -1.0
                var_map[c] = v
        drop_col = sorted(var_map.items(), key=lambda kv: (kv[1], kv[0]))[0][0]
        return drop_col

    def run_mcs_with_retry(sum_pre: pd.DataFrame, country_code: str, min_valid_per_model: int, min_rows_for_mcs: int, max_retry_drop: int, debug_logs: list):
        details = {'country': country_code, 'retry_count': 0, 'dropped_in_retry': [], 'fail_reason': None, 'final_shape': None, 'final_models': None}
        if sum_pre.empty or sum_pre.shape[1] < 2:
            details['fail_reason'] = 'Fewer than two available models'
            return (False, None, None, None, details)
        current_cols = list(sum_pre.columns)
        for _ in range(max_retry_drop + 1):
            tmp_sum = sum_pre[current_cols].copy()
            mcs_df, prep_info = prepare_mcs_matrix_relaxed(tmp_sum, min_valid_per_model=min_valid_per_model, min_rows_for_mcs=min_rows_for_mcs)
            if mcs_df is None:
                if len(current_cols) <= 2:
                    details['fail_reason'] = f"Preparation failed with too few models to retry: {prep_info.get('reason')}"
                    return (False, None, None, None, details)
                drop_col = choose_one_model_to_drop_for_retry(sum_pre, current_cols)
                if drop_col is None:
                    details['fail_reason'] = f"Preparation failed and no model can be removed: {prep_info.get('reason')}"
                    return (False, None, None, None, details)
                current_cols.remove(drop_col)
                details['retry_count'] += 1
                details['dropped_in_retry'].append(drop_col)
                debug_logs.append(f"[{country_code}] Preparation failed; retry without model: {drop_col} | Reason: {prep_info.get('reason')}")
                continue
            try:
                mcs = MCS(losses=mcs_df, size=0.9, method='R', reps=20000, block_size=5)
                mcs.compute()
                details['final_shape'] = tuple(mcs_df.shape)
                details['final_models'] = list(mcs_df.columns)
                return (True, mcs.pvalues, mcs.included, mcs_df, details)
            except Exception as e:
                if len(current_cols) <= 2:
                    details['fail_reason'] = f'MCS failed with too few models to retry: {str(e)}'
                    return (False, None, None, None, details)
                drop_col = choose_one_model_to_drop_for_retry(sum_pre, current_cols)
                if drop_col is None:
                    details['fail_reason'] = f'MCS failed and no model can be removed: {str(e)}'
                    return (False, None, None, None, details)
                current_cols.remove(drop_col)
                details['retry_count'] += 1
                details['dropped_in_retry'].append(drop_col)
                debug_logs.append(f'[{country_code}] MCS failed; retry without model: {drop_col} | Error: {str(e)}')
                continue
        details['fail_reason'] = f'Maximum retry count exceeded: {max_retry_drop}'
        return (False, None, None, None, details)

    def qlike_series_with_rules(y_true: np.ndarray, y_pred: np.ndarray) -> np.ndarray:
        y_true = np.asarray(y_true, dtype=float)
        y_pred = np.asarray(y_pred, dtype=float)
        q = np.full(shape=len(y_true), fill_value=np.nan, dtype=float)
        mask = np.isfinite(y_true) & np.isfinite(y_pred)
        if not np.any(mask):
            return q
        yt = y_true[mask]
        yp = y_pred[mask]
        ok_pos = (yt > EPS) & (yp > EPS)
        ratio = np.full_like(yt, np.nan, dtype=float)
        ratio[ok_pos] = yt[ok_pos] / yp[ok_pos]
        ok_ratio = np.isfinite(ratio) & (ratio >= RATIO_MIN) & (ratio <= RATIO_MAX)
        q_sub = np.full_like(yt, np.nan, dtype=float)
        q_sub[ok_ratio] = ratio[ok_ratio] - np.log(ratio[ok_ratio]) - 1.0
        ok_q = np.isfinite(q_sub) & (q_sub >= 0.0) & (q_sub <= QLIKE_CAP)
        q_masked = np.full_like(yt, np.nan, dtype=float)
        q_masked[ok_q] = q_sub[ok_q]
        q[mask] = q_masked
        return q

    MODEL_RUNNERS = [
        ('HAR-RV', 'HAR', run_har),
        ('HAR-CI-new', 'HAR-CI', run_har_ci),
        ('ks', 'HAR-KS', run_har_ks),
        ('Mean', 'MC', run_mc),
        ('median', 'MDC', run_mdc),
        ('tmean', 'TMC', run_tmc),
        ('dm0.9', 'DMSPE(0.9)', run_dmspe_09),
        ('dm1', 'DMSPE(1)', run_dmspe_1),
        ('pca1', 'PCA1', run_pca1),
        ('pca-bic', 'PCABIC', run_pca_bic),
        ('pls', 'PLS', run_pls),
        ('Ridge', 'Ridge', run_ridge),
        ('Lasso', 'Lasso', run_lasso),
        ('ENET', 'ENET', run_enet),
        ('spillover-t', 'Spillover-Time', run_spillover_time),
        ('spillover-f', 'Spillover-Frequency', run_spillover_frequency),
    ]

    def load_prediction_pair(frame):
        dates = pd.to_datetime(frame.index, errors="coerce")
        y_true = pd.Series(pd.to_numeric(frame["RV_True"], errors="coerce").to_numpy(), index=dates)
        y_pred = pd.Series(pd.to_numeric(frame["RV_Predict"], errors="coerce").to_numpy(), index=dates)
        valid = dates.notna() & np.isfinite(y_true.to_numpy()) & np.isfinite(y_pred.to_numpy())
        y_true, y_pred = y_true[valid], y_pred[valid]
        if len(y_true) and np.mean(y_true.to_numpy() > THRESH) > 0.5:
            y_true = y_true * SCALE_IF_MOST_GT_THRESH
            y_pred = y_pred * SCALE_IF_MOST_GT_THRESH
        y_true = y_true[~y_true.index.duplicated(keep="last")].sort_index()
        y_pred = y_pred[~y_pred.index.duplicated(keep="last")].sort_index()
        return y_true, y_pred

    def summarize_losses(predictions, shared_drop_dates):
        models = [key for key, _, _ in MODEL_RUNNERS]
        markets = sorted(predictions["HAR-RV"])
        cache = {market: {model: load_prediction_pair(predictions[model][market]) for model in models} for market in markets}
        raw = {metric: pd.DataFrame(index=markets, columns=models, dtype=float) for metric in ("MSE", "QLIKE")}
        for market in markets:
            for model in models:
                y_true, y_pred = cache[market][model]
                raw["MSE"].loc[market, model] = calc_mse(y_true.to_numpy(), y_pred.to_numpy())["MSE"]
                raw["QLIKE"].loc[market, model] = qlike_with_shared_drop(y_true, y_pred, shared_drop_dates[market])
        summaries = {}
        for metric, values in raw.items():
            values = values.dropna(how="all").dropna(subset=["HAR-RV"], how="all")
            relative = values.div(values["HAR-RV"].replace(0.0, np.nan), axis=0)
            if np.any(np.isfinite(relative.to_numpy()) & (np.abs(relative.to_numpy()) > POST_SCALE_THRESHOLD)):
                warnings.warn("The original relative-loss scaling rule was triggered for " + metric, RuntimeWarning)
                relative = relative.where(~(np.abs(relative) > POST_SCALE_THRESHOLD), relative * POST_SCALE_FACTOR)
            summaries[metric] = add_summary_rows(relative).loc[["Mean", "Q1", "Q2"]].T
        return summaries, shared_drop_dates

    def qlike_loss_series(frame, drop_dates):
        data = pd.DataFrame({
            "date": pd.to_datetime(frame.index, errors="coerce"),
            "y_true": pd.to_numeric(frame["RV_True"], errors="coerce").to_numpy(),
            "y_pred": pd.to_numeric(frame["RV_Predict"], errors="coerce").to_numpy(),
        }).dropna(subset=["date"])
        if data["date"].dt.tz is not None:
            data["date"] = data["date"].dt.tz_localize(None)
        data["date"] = data["date"].dt.normalize()
        data = data.sort_values("date").drop_duplicates(subset=["date"], keep="last").reset_index(drop=True)
        data = data.loc[~data["date"].isin(drop_dates)]
        return pd.Series(qlike_series_with_rules(data["y_true"].to_numpy(), data["y_pred"].to_numpy()), index=data["date"])

    def count_mcs_memberships(predictions, shared_drop_dates):
        models = [key for key, _, _ in MODEL_RUNNERS]
        markets = sorted(predictions["HAR-RV"])
        counts = {metric: dict.fromkeys(models, 0) for metric in ("MSE", "QLIKE")}
        for metric in counts:
            for market in tqdm(markets, desc="MCS " + metric):
                losses = pd.DataFrame()
                for model in models:
                    frame = predictions[model][market]
                    if metric == "MSE":
                        series = pd.to_numeric(frame["MSE"], errors="coerce").reset_index(drop=True)
                    else:
                        series = qlike_loss_series(frame, shared_drop_dates[market])
                    if series.notna().sum():
                        losses[model] = series
                logs = []
                success, _, included, used, details = run_mcs_with_retry(losses, market, 10, 5, 50, logs)
                if not success:
                    raise RuntimeError(f"MCS {metric} failed for {market}: {details['fail_reason']}")
                omitted = [model for model in models if model not in used.columns]
                if omitted:
                    warnings.warn(f"MCS {metric}, {market}: excluded from the test: {', '.join(omitted)}", RuntimeWarning)
                for model in included:
                    counts[metric][model] += 1
        return counts

    def save_table(summaries, counts, output_path=OUTPUT_FILE, bold_mode="har-ci"):
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "Table 4"
        sheet.merge_cells("A1:E1")
        sheet.merge_cells("G1:K1")
        sheet["A1"] = "MSE"
        sheet["G1"] = "QLIKE"
        for start in (1, 7):
            for offset, label in enumerate(("Model", "Mean", "Q1", "Median", "MCS(#)")):
                sheet.cell(2, start + offset, label)
        for row, (key, label, _) in enumerate(MODEL_RUNNERS, start=3):
            for start, metric in ((1, "MSE"), (7, "QLIKE")):
                sheet.cell(row, start, label)
                for offset, statistic in enumerate(("Mean", "Q1", "Q2"), start=1):
                    value = float(summaries[metric].loc[key, statistic])
                    if not np.isfinite(value):
                        raise ValueError(f"Cannot export a missing {metric} {statistic} for {key}")
                    sheet.cell(row, start + offset, value).number_format = "0.000"
                sheet.cell(row, start + 4, int(counts[metric][key])).number_format = "0"
        thin = Side(style="thin", color="000000")
        for row in sheet:
            for cell in row:
                cell.font = Font(name="Times New Roman", size=11, bold=False)
                cell.alignment = Alignment(horizontal="left" if cell.column in (1, 7) and cell.row > 1 else "center", vertical="center")
                if cell.column != 6:
                    if cell.row == 1:
                        cell.border = Border(top=thin, bottom=thin)
                    elif cell.row in (2, 18):
                        cell.border = Border(bottom=thin)
        for start, metric in ((1, "MSE"), (7, "QLIKE")):
            for row, (key, _, _) in enumerate(MODEL_RUNNERS, start=3):
                for offset, stat in enumerate(("Mean", "Q1", "Q2", "MCS(#)"), start=1):
                    bold = bold_mode == "har-ci" and key == "HAR-CI-new"
                    if bold_mode == "best":
                        bold = counts[metric][key] == max(counts[metric].values()) if offset == 4 else np.isclose(summaries[metric].loc[key, stat], summaries[metric][stat].min(), rtol=0, atol=1e-12)
                    if bold:
                        sheet.cell(row, start + offset).font = Font(name="Times New Roman", size=11, bold=True)
        for column in ("A", "G"):
            sheet.column_dimensions[column].width = 23
        for column in ("B", "C", "D", "E", "H", "I", "J", "K"):
            sheet.column_dimensions[column].width = 12
        sheet.column_dimensions["F"].width = 3
        sheet.freeze_panes = "B3"
        sheet.sheet_view.showGridLines = False
        sheet.print_options.horizontalCentered = True
        sheet.page_setup.orientation = "landscape"
        sheet.page_setup.paperSize = sheet.PAPERSIZE_A4
        sheet.page_setup.fitToWidth = 1
        sheet.page_setup.fitToHeight = 1
        sheet.sheet_properties.pageSetUpPr.fitToPage = True
        sheet.print_area = "A1:K18"
        output_path = Path(output_path)
        output_path.parent.mkdir(parents=True, exist_ok=True)
        workbook.save(output_path)

    def check_inputs():
        for path in (RV_FILE, RETURN_FILE, CENTRALITY_FILE, TCI_FILE):
            if not path.is_file():
                raise FileNotFoundError(f"Required input not found: {path}")
        markets = [str(c).strip().lstrip(".") for c in pd.read_csv(RV_FILE, nrows=0).columns if c != "date"]
        if len(markets) != 18 or len(set(markets)) != 18:
            raise ValueError("Table 4 requires 18 distinct markets in RV.csv")
        tci_columns = set(pd.read_csv(TCI_FILE, nrows=0).columns)
        required = {"Unnamed: 0", "TCI.Total", "TCI.20.Inf", "TCI.5.20", "TCI.1.5"}
        if not required.issubset(tci_columns):
            raise ValueError(f"Missing TCI columns: {sorted(required - tci_columns)}")
        return set(markets)

    def run_layer_corr():
        model_outputs = {}
        import os
        import numpy as np
        import pandas as pd
        import statsmodels.formula.api as smf
        from functools import partial
        from tqdm import tqdm as _tqdm
        tqdm = partial(_tqdm, bar_format="{bar} {percentage:3.0f}%", ncols=50)
        RET_PATH = RETURN_FILE
        RV_PATH = RV_FILE
        P_STORE_PATH = CENTRALITY_FILE
        PLOT = False
        fit_end_date = pd.to_datetime('2018-05-01')
        pred_start = pd.to_datetime('2018-05-02')
        pred_end = pd.to_datetime('2022-02-25')
        H_LIST = [22]
        FIT_START = 1000
        FIT_STEP = 100
        FIT_MAX = 1
        L_fit_list = [FIT_START + FIT_STEP * k for k in range(FIT_MAX)]
        NET_MIN = 1000
        NET_MAX = 1000
        NET_STEP = 50
        L_net_list = list(range(NET_MIN, NET_MAX + 1, NET_STEP))

        def _norm_date(x):
            return pd.to_datetime(x).tz_localize(None).normalize()

        def make_future_mean(rv_now: pd.Series, H: int) -> pd.Series:
            if H == 1:
                return rv_now
            ma = rv_now.rolling(window=H, center=False).mean()
            return ma.shift(-(H - 1))

        def read_returns(path):
            df = pd.read_csv(path)
            df = df.loc[:, ~df.columns.str.startswith('Unnamed')]
            df['date'] = pd.to_datetime(df['date'].astype(str).str.strip().str[:10], errors='coerce')
            df['date'] = df['date'].dt.tz_localize(None).dt.normalize()
            df = df.set_index('date').dropna(how='all')
            return df

        def read_rv_no_scale(path):
            data = pd.read_csv(path)
            if 'date' not in data.columns:
                raise RuntimeError("RV file must contain a 'date' column.")
            num_cols = [c for c in data.columns if c != 'date']
            data[num_cols] = data[num_cols].apply(pd.to_numeric, errors='coerce')
            data = data.dropna(how='all')
            data['date'] = pd.to_datetime(data['date'], utc=True, errors='coerce')
            data = data.dropna(subset=['date'])
            data['date'] = data['date'].dt.tz_localize(None)
            data['date'] = data['date'].dt.normalize()
            data = data.set_index('date')
            data = data.dropna(axis=1, how='all')
            return data
        df_ret = read_returns(RET_PATH)
        data_rv = read_rv_no_scale(RV_PATH)
        markets = [c for c in data_rv.columns if c in df_ret.columns]
        P_all = pd.read_csv(P_STORE_PATH)
        need_cols = ['date', 'market', 'ci_start_date', 'ci_end_date', 'L_net', 'P_corr', 'P_kendall', 'P_partial', 'P_tail']
        miss = [c for c in need_cols if c not in P_all.columns]
        if miss:
            raise RuntimeError(f'P.csv is missing required columns: {miss}; available columns={list(P_all.columns)}')
        P_all['date'] = pd.to_datetime(P_all['date'], errors='coerce').dt.tz_localize(None).dt.normalize()
        P_all['ci_start_date'] = P_all['ci_start_date'].apply(_norm_date)
        P_all['ci_end_date'] = P_all['ci_end_date'].apply(_norm_date)
        P_all['L_net'] = pd.to_numeric(P_all['L_net'], errors='coerce').astype('Int64')
        P_all = P_all.dropna(subset=['date', 'market', 'ci_start_date', 'ci_end_date', 'L_net']).copy()
        P_cache = {}

        def load_P_from_csv(L_net: int, ci_start_date, ci_end_date):
            key = (int(L_net), _norm_date(ci_start_date), _norm_date(ci_end_date))
            if key in P_cache:
                return P_cache[key]
            ln, cs, ce = key
            sub = P_all[(P_all['L_net'] == ln) & (P_all['ci_start_date'] == cs) & (P_all['ci_end_date'] == ce)]
            if sub.empty:
                P_cache[key] = (None, None, None, None)
                return P_cache[key]

            def _pivot(col):
                df = sub.pivot_table(index='date', columns='market', values=col, aggfunc='mean')
                return df.sort_index()
            P_corr = _pivot('P_corr')
            P_kendall = _pivot('P_kendall')
            P_partial = _pivot('P_partial')
            P_tail = _pivot('P_tail')
            P_cache[key] = (P_corr, P_kendall, P_partial, P_tail)
            return P_cache[key]

        def forecast_one_market_oneH_harci(data_rv_all: pd.DataFrame, P_corr: pd.DataFrame, P_kendall: pd.DataFrame, P_partial: pd.DataFrame, P_tail: pd.DataFrame, market: str, fit_start_date: pd.Timestamp, H: int):
            df = data_rv_all[market].to_frame('RV_now')
            df['RV_D'] = df['RV_now'].shift(1)
            df['RV_W'] = df['RV_D'].rolling(5).mean()
            df['RV_M'] = df['RV_D'].rolling(22).mean()
            df['Y'] = make_future_mean(df['RV_now'], H)
            ci = pd.DataFrame({'CI_corr': P_corr[market], 'CI_kendall': P_kendall[market], 'CI_partial': P_partial[market], 'CI_tail': P_tail[market]}).reindex(df.index)
            df = df.join(ci, how='inner').dropna()
            df['CI_corr_D'] = df['CI_corr'].shift(1)
            df['CI_kendall_D'] = df['CI_kendall'].shift(1)
            df['CI_partial_D'] = df['CI_partial'].shift(1)
            df['CI_tail_D'] = df['CI_tail'].shift(1)
            df = df.dropna()
            df = df.loc[fit_start_date:]
            data_test = df.loc[pred_start:pred_end].copy()
            if len(data_test) == 0:
                return (None, np.nan, np.nan, np.nan)
            window_len = len(df.loc[:pred_start]) - 1
            if window_len <= 0:
                return (None, np.nan, np.nan, np.nan)
            y_true = data_test['Y'].astype(float).copy()
            idx = y_true.index
            preds = []
            for i in tqdm(range(len(idx)), desc=f'HAR-CI {market} H{H}', leave=False):
                train = df.iloc[i:i + window_len].copy()
                train['_Y_'] = train['Y'].astype(float).values
                res = smf.ols('_Y_ ~ RV_D + RV_W + RV_M + CI_corr_D', train).fit()
                p = res.params
                pred = p.iloc[0] + p.iloc[1] * data_test['RV_D'].iloc[i] + p.iloc[2] * data_test['RV_W'].iloc[i] + p.iloc[3] * data_test['RV_M'].iloc[i] + p.iloc[4] * data_test['CI_corr_D'].iloc[i]
                preds.append(float(pred))
            y_pred = pd.Series(preds, index=idx)
            mcs = pd.concat([y_true.rename('RV_True'), y_pred.rename('RV_Predict')], axis=1)
            mcs['MSE'] = (mcs['RV_True'] - mcs['RV_Predict']) ** 2
            mcs['MAE'] = np.abs(mcs['RV_True'] - mcs['RV_Predict'])
            mse = float(np.nanmean(mcs['MSE'].values))
            mae = float(np.nanmean(mcs['MAE'].values))
            denom = np.sqrt(np.nanmean(mcs['RV_Predict'].values ** 2)) + np.sqrt(np.nanmean(mcs['RV_True'].values ** 2))
            theil = float(np.sqrt(np.nanmean((mcs['RV_True'].values - mcs['RV_Predict'].values) ** 2)) / denom) if denom != 0 else np.nan
            return (mcs, mse, mae, theil)
        L_fit = L_fit_list[0]
        L_net = L_net_list[0]
        for H in H_LIST:
            error_by_market = {}
            for mkt in markets:
                if fit_end_date not in data_rv.index:
                    error_by_market[mkt] = [np.nan, np.nan, np.nan]
                    continue
                fit_end_pos = data_rv.index.get_loc(fit_end_date)
                fit_start_pos = fit_end_pos - (L_fit - 1)
                if fit_start_pos < 0:
                    error_by_market[mkt] = [np.nan, np.nan, np.nan]
                    continue
                fit_start_date = data_rv.index[fit_start_pos]
                ret_pos_end = df_ret.index.searchsorted(fit_start_date) - 1
                if ret_pos_end < 0:
                    error_by_market[mkt] = [np.nan, np.nan, np.nan]
                    continue
                net_end_date = df_ret.index[ret_pos_end]
                ci_start_date = net_end_date
                ci_end_date = pred_end
                P_corr, P_kendall, P_partial, P_tail = load_P_from_csv(L_net=L_net, ci_start_date=ci_start_date, ci_end_date=ci_end_date)
                if P_corr is None:
                    error_by_market[mkt] = [np.nan, np.nan, np.nan]
                    continue
                if mkt not in P_corr.columns:
                    error_by_market[mkt] = [np.nan, np.nan, np.nan]
                    continue
                mcs, mse, mae, theil = forecast_one_market_oneH_harci(data_rv_all=data_rv, P_corr=P_corr, P_kendall=P_kendall, P_partial=P_partial, P_tail=P_tail, market=mkt, fit_start_date=fit_start_date, H=H)
                if mcs is None:
                    tmp = pd.DataFrame(index=pd.DatetimeIndex([], name='date'))
                    tmp['RV_True'] = np.nan
                    tmp['RV_Predict'] = np.nan
                    tmp['MSE'] = np.nan
                    tmp['MAE'] = np.nan
                    model_outputs[str(mkt).strip().lstrip('.')] = tmp.copy()
                else:
                    model_outputs[str(mkt).strip().lstrip('.')] = mcs.copy()
                error_by_market[mkt] = [mse, mae, theil]
            ERROR = pd.DataFrame(error_by_market, index=['MSE', 'MAE', 'Theil'])
        return model_outputs

    def run_layer_kendall():
        model_outputs = {}
        import os
        import numpy as np
        import pandas as pd
        import statsmodels.formula.api as smf
        from functools import partial
        from tqdm import tqdm as _tqdm
        tqdm = partial(_tqdm, bar_format="{bar} {percentage:3.0f}%", ncols=50)
        RET_PATH = RETURN_FILE
        RV_PATH = RV_FILE
        P_STORE_PATH = CENTRALITY_FILE
        PLOT = False
        fit_end_date = pd.to_datetime('2018-05-01')
        pred_start = pd.to_datetime('2018-05-02')
        pred_end = pd.to_datetime('2022-02-25')
        H_LIST = [22]
        FIT_START = 1000
        FIT_STEP = 100
        FIT_MAX = 1
        L_fit_list = [FIT_START + FIT_STEP * k for k in range(FIT_MAX)]
        NET_MIN = 1000
        NET_MAX = 1000
        NET_STEP = 50
        L_net_list = list(range(NET_MIN, NET_MAX + 1, NET_STEP))

        def _norm_date(x):
            return pd.to_datetime(x).tz_localize(None).normalize()

        def make_future_mean(rv_now: pd.Series, H: int) -> pd.Series:
            if H == 1:
                return rv_now
            ma = rv_now.rolling(window=H, center=False).mean()
            return ma.shift(-(H - 1))

        def read_returns(path):
            df = pd.read_csv(path)
            df = df.loc[:, ~df.columns.str.startswith('Unnamed')]
            df['date'] = pd.to_datetime(df['date'].astype(str).str.strip().str[:10], errors='coerce')
            df['date'] = df['date'].dt.tz_localize(None).dt.normalize()
            df = df.set_index('date').dropna(how='all')
            return df

        def read_rv_no_scale(path):
            data = pd.read_csv(path)
            if 'date' not in data.columns:
                raise RuntimeError("RV file must contain a 'date' column.")
            num_cols = [c for c in data.columns if c != 'date']
            data[num_cols] = data[num_cols].apply(pd.to_numeric, errors='coerce')
            data = data.dropna(how='all')
            data['date'] = pd.to_datetime(data['date'], utc=True, errors='coerce')
            data = data.dropna(subset=['date'])
            data['date'] = data['date'].dt.tz_localize(None)
            data['date'] = data['date'].dt.normalize()
            data = data.set_index('date')
            data = data.dropna(axis=1, how='all')
            return data
        df_ret = read_returns(RET_PATH)
        data_rv = read_rv_no_scale(RV_PATH)
        markets = [c for c in data_rv.columns if c in df_ret.columns]
        P_all = pd.read_csv(P_STORE_PATH)
        need_cols = ['date', 'market', 'ci_start_date', 'ci_end_date', 'L_net', 'P_corr', 'P_kendall', 'P_partial', 'P_tail']
        miss = [c for c in need_cols if c not in P_all.columns]
        if miss:
            raise RuntimeError(f'P.csv is missing required columns: {miss}; available columns={list(P_all.columns)}')
        P_all['date'] = pd.to_datetime(P_all['date'], errors='coerce').dt.tz_localize(None).dt.normalize()
        P_all['ci_start_date'] = P_all['ci_start_date'].apply(_norm_date)
        P_all['ci_end_date'] = P_all['ci_end_date'].apply(_norm_date)
        P_all['L_net'] = pd.to_numeric(P_all['L_net'], errors='coerce').astype('Int64')
        P_all = P_all.dropna(subset=['date', 'market', 'ci_start_date', 'ci_end_date', 'L_net']).copy()
        P_cache = {}

        def load_P_from_csv(L_net: int, ci_start_date, ci_end_date):
            key = (int(L_net), _norm_date(ci_start_date), _norm_date(ci_end_date))
            if key in P_cache:
                return P_cache[key]
            ln, cs, ce = key
            sub = P_all[(P_all['L_net'] == ln) & (P_all['ci_start_date'] == cs) & (P_all['ci_end_date'] == ce)]
            if sub.empty:
                P_cache[key] = (None, None, None, None)
                return P_cache[key]

            def _pivot(col):
                df = sub.pivot_table(index='date', columns='market', values=col, aggfunc='mean')
                return df.sort_index()
            P_corr = _pivot('P_corr')
            P_kendall = _pivot('P_kendall')
            P_partial = _pivot('P_partial')
            P_tail = _pivot('P_tail')
            P_cache[key] = (P_corr, P_kendall, P_partial, P_tail)
            return P_cache[key]

        def forecast_one_market_oneH_harci(data_rv_all: pd.DataFrame, P_corr: pd.DataFrame, P_kendall: pd.DataFrame, P_partial: pd.DataFrame, P_tail: pd.DataFrame, market: str, fit_start_date: pd.Timestamp, H: int):
            df = data_rv_all[market].to_frame('RV_now')
            df['RV_D'] = df['RV_now'].shift(1)
            df['RV_W'] = df['RV_D'].rolling(5).mean()
            df['RV_M'] = df['RV_D'].rolling(22).mean()
            df['Y'] = make_future_mean(df['RV_now'], H)
            ci = pd.DataFrame({'CI_corr': P_corr[market], 'CI_kendall': P_kendall[market], 'CI_partial': P_partial[market], 'CI_tail': P_tail[market]}).reindex(df.index)
            df = df.join(ci, how='inner').dropna()
            df['CI_corr_D'] = df['CI_corr'].shift(1)
            df['CI_kendall_D'] = df['CI_kendall'].shift(1)
            df['CI_partial_D'] = df['CI_partial'].shift(1)
            df['CI_tail_D'] = df['CI_tail'].shift(1)
            df = df.dropna()
            df = df.loc[fit_start_date:]
            data_test = df.loc[pred_start:pred_end].copy()
            if len(data_test) == 0:
                return (None, np.nan, np.nan, np.nan)
            window_len = len(df.loc[:pred_start]) - 1
            if window_len <= 0:
                return (None, np.nan, np.nan, np.nan)
            y_true = data_test['Y'].astype(float).copy()
            idx = y_true.index
            preds = []
            for i in tqdm(range(len(idx)), desc=f'HAR-CI {market} H{H}', leave=False):
                train = df.iloc[i:i + window_len].copy()
                train['_Y_'] = train['Y'].astype(float).values
                res = smf.ols('_Y_ ~ RV_D + RV_W + RV_M + CI_kendall_D', train).fit()
                p = res.params
                pred = p.iloc[0] + p.iloc[1] * data_test['RV_D'].iloc[i] + p.iloc[2] * data_test['RV_W'].iloc[i] + p.iloc[3] * data_test['RV_M'].iloc[i] + p.iloc[4] * data_test['CI_kendall_D'].iloc[i]
                preds.append(float(pred))
            y_pred = pd.Series(preds, index=idx)
            mcs = pd.concat([y_true.rename('RV_True'), y_pred.rename('RV_Predict')], axis=1)
            mcs['MSE'] = (mcs['RV_True'] - mcs['RV_Predict']) ** 2
            mcs['MAE'] = np.abs(mcs['RV_True'] - mcs['RV_Predict'])
            mse = float(np.nanmean(mcs['MSE'].values))
            mae = float(np.nanmean(mcs['MAE'].values))
            denom = np.sqrt(np.nanmean(mcs['RV_Predict'].values ** 2)) + np.sqrt(np.nanmean(mcs['RV_True'].values ** 2))
            theil = float(np.sqrt(np.nanmean((mcs['RV_True'].values - mcs['RV_Predict'].values) ** 2)) / denom) if denom != 0 else np.nan
            return (mcs, mse, mae, theil)
        L_fit = L_fit_list[0]
        L_net = L_net_list[0]
        for H in H_LIST:
            error_by_market = {}
            for mkt in markets:
                if fit_end_date not in data_rv.index:
                    error_by_market[mkt] = [np.nan, np.nan, np.nan]
                    continue
                fit_end_pos = data_rv.index.get_loc(fit_end_date)
                fit_start_pos = fit_end_pos - (L_fit - 1)
                if fit_start_pos < 0:
                    error_by_market[mkt] = [np.nan, np.nan, np.nan]
                    continue
                fit_start_date = data_rv.index[fit_start_pos]
                ret_pos_end = df_ret.index.searchsorted(fit_start_date) - 1
                if ret_pos_end < 0:
                    error_by_market[mkt] = [np.nan, np.nan, np.nan]
                    continue
                net_end_date = df_ret.index[ret_pos_end]
                ci_start_date = net_end_date
                ci_end_date = pred_end
                P_corr, P_kendall, P_partial, P_tail = load_P_from_csv(L_net=L_net, ci_start_date=ci_start_date, ci_end_date=ci_end_date)
                if P_corr is None:
                    error_by_market[mkt] = [np.nan, np.nan, np.nan]
                    continue
                if mkt not in P_corr.columns:
                    error_by_market[mkt] = [np.nan, np.nan, np.nan]
                    continue
                mcs, mse, mae, theil = forecast_one_market_oneH_harci(data_rv_all=data_rv, P_corr=P_corr, P_kendall=P_kendall, P_partial=P_partial, P_tail=P_tail, market=mkt, fit_start_date=fit_start_date, H=H)
                if mcs is None:
                    tmp = pd.DataFrame(index=pd.DatetimeIndex([], name='date'))
                    tmp['RV_True'] = np.nan
                    tmp['RV_Predict'] = np.nan
                    tmp['MSE'] = np.nan
                    tmp['MAE'] = np.nan
                    model_outputs[str(mkt).strip().lstrip('.')] = tmp.copy()
                else:
                    model_outputs[str(mkt).strip().lstrip('.')] = mcs.copy()
                error_by_market[mkt] = [mse, mae, theil]
            ERROR = pd.DataFrame(error_by_market, index=['MSE', 'MAE', 'Theil'])
        return model_outputs

    def run_layer_partial():
        model_outputs = {}
        import os
        import numpy as np
        import pandas as pd
        import statsmodels.formula.api as smf
        from functools import partial
        from tqdm import tqdm as _tqdm
        tqdm = partial(_tqdm, bar_format="{bar} {percentage:3.0f}%", ncols=50)
        RET_PATH = RETURN_FILE
        RV_PATH = RV_FILE
        P_STORE_PATH = CENTRALITY_FILE
        PLOT = False
        fit_end_date = pd.to_datetime('2018-05-01')
        pred_start = pd.to_datetime('2018-05-02')
        pred_end = pd.to_datetime('2022-02-25')
        H_LIST = [22]
        FIT_START = 1000
        FIT_STEP = 100
        FIT_MAX = 1
        L_fit_list = [FIT_START + FIT_STEP * k for k in range(FIT_MAX)]
        NET_MIN = 1000
        NET_MAX = 1000
        NET_STEP = 50
        L_net_list = list(range(NET_MIN, NET_MAX + 1, NET_STEP))

        def _norm_date(x):
            return pd.to_datetime(x).tz_localize(None).normalize()

        def make_future_mean(rv_now: pd.Series, H: int) -> pd.Series:
            if H == 1:
                return rv_now
            ma = rv_now.rolling(window=H, center=False).mean()
            return ma.shift(-(H - 1))

        def read_returns(path):
            df = pd.read_csv(path)
            df = df.loc[:, ~df.columns.str.startswith('Unnamed')]
            df['date'] = pd.to_datetime(df['date'].astype(str).str.strip().str[:10], errors='coerce')
            df['date'] = df['date'].dt.tz_localize(None).dt.normalize()
            df = df.set_index('date').dropna(how='all')
            return df

        def read_rv_no_scale(path):
            data = pd.read_csv(path)
            if 'date' not in data.columns:
                raise RuntimeError("RV file must contain a 'date' column.")
            num_cols = [c for c in data.columns if c != 'date']
            data[num_cols] = data[num_cols].apply(pd.to_numeric, errors='coerce')
            data = data.dropna(how='all')
            data['date'] = pd.to_datetime(data['date'], utc=True, errors='coerce')
            data = data.dropna(subset=['date'])
            data['date'] = data['date'].dt.tz_localize(None)
            data['date'] = data['date'].dt.normalize()
            data = data.set_index('date')
            data = data.dropna(axis=1, how='all')
            return data
        df_ret = read_returns(RET_PATH)
        data_rv = read_rv_no_scale(RV_PATH)
        markets = [c for c in data_rv.columns if c in df_ret.columns]
        P_all = pd.read_csv(P_STORE_PATH)
        need_cols = ['date', 'market', 'ci_start_date', 'ci_end_date', 'L_net', 'P_corr', 'P_kendall', 'P_partial', 'P_tail']
        miss = [c for c in need_cols if c not in P_all.columns]
        if miss:
            raise RuntimeError(f'P.csv is missing required columns: {miss}; available columns={list(P_all.columns)}')
        P_all['date'] = pd.to_datetime(P_all['date'], errors='coerce').dt.tz_localize(None).dt.normalize()
        P_all['ci_start_date'] = P_all['ci_start_date'].apply(_norm_date)
        P_all['ci_end_date'] = P_all['ci_end_date'].apply(_norm_date)
        P_all['L_net'] = pd.to_numeric(P_all['L_net'], errors='coerce').astype('Int64')
        P_all = P_all.dropna(subset=['date', 'market', 'ci_start_date', 'ci_end_date', 'L_net']).copy()
        P_cache = {}

        def load_P_from_csv(L_net: int, ci_start_date, ci_end_date):
            key = (int(L_net), _norm_date(ci_start_date), _norm_date(ci_end_date))
            if key in P_cache:
                return P_cache[key]
            ln, cs, ce = key
            sub = P_all[(P_all['L_net'] == ln) & (P_all['ci_start_date'] == cs) & (P_all['ci_end_date'] == ce)]
            if sub.empty:
                P_cache[key] = (None, None, None, None)
                return P_cache[key]

            def _pivot(col):
                df = sub.pivot_table(index='date', columns='market', values=col, aggfunc='mean')
                return df.sort_index()
            P_corr = _pivot('P_corr')
            P_kendall = _pivot('P_kendall')
            P_partial = _pivot('P_partial')
            P_tail = _pivot('P_tail')
            P_cache[key] = (P_corr, P_kendall, P_partial, P_tail)
            return P_cache[key]

        def forecast_one_market_oneH_harci(data_rv_all: pd.DataFrame, P_corr: pd.DataFrame, P_kendall: pd.DataFrame, P_partial: pd.DataFrame, P_tail: pd.DataFrame, market: str, fit_start_date: pd.Timestamp, H: int):
            df = data_rv_all[market].to_frame('RV_now')
            df['RV_D'] = df['RV_now'].shift(1)
            df['RV_W'] = df['RV_D'].rolling(5).mean()
            df['RV_M'] = df['RV_D'].rolling(22).mean()
            df['Y'] = make_future_mean(df['RV_now'], H)
            ci = pd.DataFrame({'CI_corr': P_corr[market], 'CI_kendall': P_kendall[market], 'CI_partial': P_partial[market], 'CI_tail': P_tail[market]}).reindex(df.index)
            df = df.join(ci, how='inner').dropna()
            df['CI_corr_D'] = df['CI_corr'].shift(1)
            df['CI_kendall_D'] = df['CI_kendall'].shift(1)
            df['CI_partial_D'] = df['CI_partial'].shift(1)
            df['CI_tail_D'] = df['CI_tail'].shift(1)
            df = df.dropna()
            df = df.loc[fit_start_date:]
            data_test = df.loc[pred_start:pred_end].copy()
            if len(data_test) == 0:
                return (None, np.nan, np.nan, np.nan)
            window_len = len(df.loc[:pred_start]) - 1
            if window_len <= 0:
                return (None, np.nan, np.nan, np.nan)
            y_true = data_test['Y'].astype(float).copy()
            idx = y_true.index
            preds = []
            for i in tqdm(range(len(idx)), desc=f'HAR-CI {market} H{H}', leave=False):
                train = df.iloc[i:i + window_len].copy()
                train['_Y_'] = train['Y'].astype(float).values
                res = smf.ols('_Y_ ~ RV_D + RV_W + RV_M + CI_partial_D', train).fit()
                p = res.params
                pred = p.iloc[0] + p.iloc[1] * data_test['RV_D'].iloc[i] + p.iloc[2] * data_test['RV_W'].iloc[i] + p.iloc[3] * data_test['RV_M'].iloc[i] + p.iloc[4] * data_test['CI_partial_D'].iloc[i]
                preds.append(float(pred))
            y_pred = pd.Series(preds, index=idx)
            mcs = pd.concat([y_true.rename('RV_True'), y_pred.rename('RV_Predict')], axis=1)
            mcs['MSE'] = (mcs['RV_True'] - mcs['RV_Predict']) ** 2
            mcs['MAE'] = np.abs(mcs['RV_True'] - mcs['RV_Predict'])
            mse = float(np.nanmean(mcs['MSE'].values))
            mae = float(np.nanmean(mcs['MAE'].values))
            denom = np.sqrt(np.nanmean(mcs['RV_Predict'].values ** 2)) + np.sqrt(np.nanmean(mcs['RV_True'].values ** 2))
            theil = float(np.sqrt(np.nanmean((mcs['RV_True'].values - mcs['RV_Predict'].values) ** 2)) / denom) if denom != 0 else np.nan
            return (mcs, mse, mae, theil)
        L_fit = L_fit_list[0]
        L_net = L_net_list[0]
        for H in H_LIST:
            error_by_market = {}
            for mkt in markets:
                if fit_end_date not in data_rv.index:
                    error_by_market[mkt] = [np.nan, np.nan, np.nan]
                    continue
                fit_end_pos = data_rv.index.get_loc(fit_end_date)
                fit_start_pos = fit_end_pos - (L_fit - 1)
                if fit_start_pos < 0:
                    error_by_market[mkt] = [np.nan, np.nan, np.nan]
                    continue
                fit_start_date = data_rv.index[fit_start_pos]
                ret_pos_end = df_ret.index.searchsorted(fit_start_date) - 1
                if ret_pos_end < 0:
                    error_by_market[mkt] = [np.nan, np.nan, np.nan]
                    continue
                net_end_date = df_ret.index[ret_pos_end]
                ci_start_date = net_end_date
                ci_end_date = pred_end
                P_corr, P_kendall, P_partial, P_tail = load_P_from_csv(L_net=L_net, ci_start_date=ci_start_date, ci_end_date=ci_end_date)
                if P_corr is None:
                    error_by_market[mkt] = [np.nan, np.nan, np.nan]
                    continue
                if mkt not in P_corr.columns:
                    error_by_market[mkt] = [np.nan, np.nan, np.nan]
                    continue
                mcs, mse, mae, theil = forecast_one_market_oneH_harci(data_rv_all=data_rv, P_corr=P_corr, P_kendall=P_kendall, P_partial=P_partial, P_tail=P_tail, market=mkt, fit_start_date=fit_start_date, H=H)
                if mcs is None:
                    tmp = pd.DataFrame(index=pd.DatetimeIndex([], name='date'))
                    tmp['RV_True'] = np.nan
                    tmp['RV_Predict'] = np.nan
                    tmp['MSE'] = np.nan
                    tmp['MAE'] = np.nan
                    model_outputs[str(mkt).strip().lstrip('.')] = tmp.copy()
                else:
                    model_outputs[str(mkt).strip().lstrip('.')] = mcs.copy()
                error_by_market[mkt] = [mse, mae, theil]
            ERROR = pd.DataFrame(error_by_market, index=['MSE', 'MAE', 'Theil'])
        return model_outputs

    def run_layer_tail():
        model_outputs = {}
        import os
        import numpy as np
        import pandas as pd
        import statsmodels.formula.api as smf
        from functools import partial
        from tqdm import tqdm as _tqdm
        tqdm = partial(_tqdm, bar_format="{bar} {percentage:3.0f}%", ncols=50)
        RET_PATH = RETURN_FILE
        RV_PATH = RV_FILE
        P_STORE_PATH = CENTRALITY_FILE
        PLOT = False
        fit_end_date = pd.to_datetime('2018-05-01')
        pred_start = pd.to_datetime('2018-05-02')
        pred_end = pd.to_datetime('2022-02-25')
        H_LIST = [22]
        FIT_START = 1000
        FIT_STEP = 100
        FIT_MAX = 1
        L_fit_list = [FIT_START + FIT_STEP * k for k in range(FIT_MAX)]
        NET_MIN = 1000
        NET_MAX = 1000
        NET_STEP = 50
        L_net_list = list(range(NET_MIN, NET_MAX + 1, NET_STEP))

        def _norm_date(x):
            return pd.to_datetime(x).tz_localize(None).normalize()

        def make_future_mean(rv_now: pd.Series, H: int) -> pd.Series:
            if H == 1:
                return rv_now
            ma = rv_now.rolling(window=H, center=False).mean()
            return ma.shift(-(H - 1))

        def read_returns(path):
            df = pd.read_csv(path)
            df = df.loc[:, ~df.columns.str.startswith('Unnamed')]
            df['date'] = pd.to_datetime(df['date'].astype(str).str.strip().str[:10], errors='coerce')
            df['date'] = df['date'].dt.tz_localize(None).dt.normalize()
            df = df.set_index('date').dropna(how='all')
            return df

        def read_rv_no_scale(path):
            data = pd.read_csv(path)
            if 'date' not in data.columns:
                raise RuntimeError("RV file must contain a 'date' column.")
            num_cols = [c for c in data.columns if c != 'date']
            data[num_cols] = data[num_cols].apply(pd.to_numeric, errors='coerce')
            data = data.dropna(how='all')
            data['date'] = pd.to_datetime(data['date'], utc=True, errors='coerce')
            data = data.dropna(subset=['date'])
            data['date'] = data['date'].dt.tz_localize(None)
            data['date'] = data['date'].dt.normalize()
            data = data.set_index('date')
            data = data.dropna(axis=1, how='all')
            return data
        df_ret = read_returns(RET_PATH)
        data_rv = read_rv_no_scale(RV_PATH)
        markets = [c for c in data_rv.columns if c in df_ret.columns]
        P_all = pd.read_csv(P_STORE_PATH)
        need_cols = ['date', 'market', 'ci_start_date', 'ci_end_date', 'L_net', 'P_corr', 'P_kendall', 'P_partial', 'P_tail']
        miss = [c for c in need_cols if c not in P_all.columns]
        if miss:
            raise RuntimeError(f'P.csv is missing required columns: {miss}; available columns={list(P_all.columns)}')
        P_all['date'] = pd.to_datetime(P_all['date'], errors='coerce').dt.tz_localize(None).dt.normalize()
        P_all['ci_start_date'] = P_all['ci_start_date'].apply(_norm_date)
        P_all['ci_end_date'] = P_all['ci_end_date'].apply(_norm_date)
        P_all['L_net'] = pd.to_numeric(P_all['L_net'], errors='coerce').astype('Int64')
        P_all = P_all.dropna(subset=['date', 'market', 'ci_start_date', 'ci_end_date', 'L_net']).copy()
        P_cache = {}

        def load_P_from_csv(L_net: int, ci_start_date, ci_end_date):
            key = (int(L_net), _norm_date(ci_start_date), _norm_date(ci_end_date))
            if key in P_cache:
                return P_cache[key]
            ln, cs, ce = key
            sub = P_all[(P_all['L_net'] == ln) & (P_all['ci_start_date'] == cs) & (P_all['ci_end_date'] == ce)]
            if sub.empty:
                P_cache[key] = (None, None, None, None)
                return P_cache[key]

            def _pivot(col):
                df = sub.pivot_table(index='date', columns='market', values=col, aggfunc='mean')
                return df.sort_index()
            P_corr = _pivot('P_corr')
            P_kendall = _pivot('P_kendall')
            P_partial = _pivot('P_partial')
            P_tail = _pivot('P_tail')
            P_cache[key] = (P_corr, P_kendall, P_partial, P_tail)
            return P_cache[key]

        def forecast_one_market_oneH_harci(data_rv_all: pd.DataFrame, P_corr: pd.DataFrame, P_kendall: pd.DataFrame, P_partial: pd.DataFrame, P_tail: pd.DataFrame, market: str, fit_start_date: pd.Timestamp, H: int):
            df = data_rv_all[market].to_frame('RV_now')
            df['RV_D'] = df['RV_now'].shift(1)
            df['RV_W'] = df['RV_D'].rolling(5).mean()
            df['RV_M'] = df['RV_D'].rolling(22).mean()
            df['Y'] = make_future_mean(df['RV_now'], H)
            ci = pd.DataFrame({'CI_corr': P_corr[market], 'CI_kendall': P_kendall[market], 'CI_partial': P_partial[market], 'CI_tail': P_tail[market]}).reindex(df.index)
            df = df.join(ci, how='inner').dropna()
            df['CI_corr_D'] = df['CI_corr'].shift(1)
            df['CI_kendall_D'] = df['CI_kendall'].shift(1)
            df['CI_partial_D'] = df['CI_partial'].shift(1)
            df['CI_tail_D'] = df['CI_tail'].shift(1)
            df = df.dropna()
            df = df.loc[fit_start_date:]
            data_test = df.loc[pred_start:pred_end].copy()
            if len(data_test) == 0:
                return (None, np.nan, np.nan, np.nan)
            window_len = len(df.loc[:pred_start]) - 1
            if window_len <= 0:
                return (None, np.nan, np.nan, np.nan)
            y_true = data_test['Y'].astype(float).copy()
            idx = y_true.index
            preds = []
            for i in tqdm(range(len(idx)), desc=f'HAR-CI {market} H{H}', leave=False):
                train = df.iloc[i:i + window_len].copy()
                train['_Y_'] = train['Y'].astype(float).values
                res = smf.ols('_Y_ ~ RV_D + RV_W + RV_M + CI_tail_D', train).fit()
                p = res.params
                pred = p.iloc[0] + p.iloc[1] * data_test['RV_D'].iloc[i] + p.iloc[2] * data_test['RV_W'].iloc[i] + p.iloc[3] * data_test['RV_M'].iloc[i] + p.iloc[4] * data_test['CI_tail_D'].iloc[i]
                preds.append(float(pred))
            y_pred = pd.Series(preds, index=idx)
            mcs = pd.concat([y_true.rename('RV_True'), y_pred.rename('RV_Predict')], axis=1)
            mcs['MSE'] = (mcs['RV_True'] - mcs['RV_Predict']) ** 2
            mcs['MAE'] = np.abs(mcs['RV_True'] - mcs['RV_Predict'])
            mse = float(np.nanmean(mcs['MSE'].values))
            mae = float(np.nanmean(mcs['MAE'].values))
            denom = np.sqrt(np.nanmean(mcs['RV_Predict'].values ** 2)) + np.sqrt(np.nanmean(mcs['RV_True'].values ** 2))
            theil = float(np.sqrt(np.nanmean((mcs['RV_True'].values - mcs['RV_Predict'].values) ** 2)) / denom) if denom != 0 else np.nan
            return (mcs, mse, mae, theil)
        L_fit = L_fit_list[0]
        L_net = L_net_list[0]
        for H in H_LIST:
            error_by_market = {}
            for mkt in markets:
                if fit_end_date not in data_rv.index:
                    error_by_market[mkt] = [np.nan, np.nan, np.nan]
                    continue
                fit_end_pos = data_rv.index.get_loc(fit_end_date)
                fit_start_pos = fit_end_pos - (L_fit - 1)
                if fit_start_pos < 0:
                    error_by_market[mkt] = [np.nan, np.nan, np.nan]
                    continue
                fit_start_date = data_rv.index[fit_start_pos]
                ret_pos_end = df_ret.index.searchsorted(fit_start_date) - 1
                if ret_pos_end < 0:
                    error_by_market[mkt] = [np.nan, np.nan, np.nan]
                    continue
                net_end_date = df_ret.index[ret_pos_end]
                ci_start_date = net_end_date
                ci_end_date = pred_end
                P_corr, P_kendall, P_partial, P_tail = load_P_from_csv(L_net=L_net, ci_start_date=ci_start_date, ci_end_date=ci_end_date)
                if P_corr is None:
                    error_by_market[mkt] = [np.nan, np.nan, np.nan]
                    continue
                if mkt not in P_corr.columns:
                    error_by_market[mkt] = [np.nan, np.nan, np.nan]
                    continue
                mcs, mse, mae, theil = forecast_one_market_oneH_harci(data_rv_all=data_rv, P_corr=P_corr, P_kendall=P_kendall, P_partial=P_partial, P_tail=P_tail, market=mkt, fit_start_date=fit_start_date, H=H)
                if mcs is None:
                    tmp = pd.DataFrame(index=pd.DatetimeIndex([], name='date'))
                    tmp['RV_True'] = np.nan
                    tmp['RV_Predict'] = np.nan
                    tmp['MSE'] = np.nan
                    tmp['MAE'] = np.nan
                    model_outputs[str(mkt).strip().lstrip('.')] = tmp.copy()
                else:
                    model_outputs[str(mkt).strip().lstrip('.')] = mcs.copy()
                error_by_market[mkt] = [mse, mae, theil]
            ERROR = pd.DataFrame(error_by_market, index=['MSE', 'MAE', 'Theil'])
        return model_outputs

    def run_layer_corr_kendall_partial():
        model_outputs = {}
        import os
        import numpy as np
        import pandas as pd
        import statsmodels.formula.api as smf
        from functools import partial
        from tqdm import tqdm as _tqdm
        tqdm = partial(_tqdm, bar_format="{bar} {percentage:3.0f}%", ncols=50)
        RET_PATH = RETURN_FILE
        RV_PATH = RV_FILE
        P_STORE_PATH = CENTRALITY_FILE
        PLOT = False
        fit_end_date = pd.to_datetime('2018-05-01')
        pred_start = pd.to_datetime('2018-05-02')
        pred_end = pd.to_datetime('2022-02-25')
        H_LIST = [22]
        FIT_START = 1000
        FIT_STEP = 100
        FIT_MAX = 1
        L_fit_list = [FIT_START + FIT_STEP * k for k in range(FIT_MAX)]
        NET_MIN = 1000
        NET_MAX = 1000
        NET_STEP = 50
        L_net_list = list(range(NET_MIN, NET_MAX + 1, NET_STEP))

        def _norm_date(x):
            return pd.to_datetime(x).tz_localize(None).normalize()

        def make_future_mean(rv_now: pd.Series, H: int) -> pd.Series:
            if H == 1:
                return rv_now
            ma = rv_now.rolling(window=H, center=False).mean()
            return ma.shift(-(H - 1))

        def read_returns(path):
            df = pd.read_csv(path)
            df = df.loc[:, ~df.columns.str.startswith('Unnamed')]
            df['date'] = pd.to_datetime(df['date'].astype(str).str.strip().str[:10], errors='coerce')
            df['date'] = df['date'].dt.tz_localize(None).dt.normalize()
            df = df.set_index('date').dropna(how='all')
            return df

        def read_rv_no_scale(path):
            data = pd.read_csv(path)
            if 'date' not in data.columns:
                raise RuntimeError("RV file must contain a 'date' column.")
            num_cols = [c for c in data.columns if c != 'date']
            data[num_cols] = data[num_cols].apply(pd.to_numeric, errors='coerce')
            data = data.dropna(how='all')
            data['date'] = pd.to_datetime(data['date'], utc=True, errors='coerce')
            data = data.dropna(subset=['date'])
            data['date'] = data['date'].dt.tz_localize(None)
            data['date'] = data['date'].dt.normalize()
            data = data.set_index('date')
            data = data.dropna(axis=1, how='all')
            return data
        df_ret = read_returns(RET_PATH)
        data_rv = read_rv_no_scale(RV_PATH)
        markets = [c for c in data_rv.columns if c in df_ret.columns]
        P_all = pd.read_csv(P_STORE_PATH)
        need_cols = ['date', 'market', 'ci_start_date', 'ci_end_date', 'L_net', 'P_corr', 'P_kendall', 'P_partial', 'P_tail']
        miss = [c for c in need_cols if c not in P_all.columns]
        if miss:
            raise RuntimeError(f'P.csv is missing required columns: {miss}; available columns={list(P_all.columns)}')
        P_all['date'] = pd.to_datetime(P_all['date'], errors='coerce').dt.tz_localize(None).dt.normalize()
        P_all['ci_start_date'] = P_all['ci_start_date'].apply(_norm_date)
        P_all['ci_end_date'] = P_all['ci_end_date'].apply(_norm_date)
        P_all['L_net'] = pd.to_numeric(P_all['L_net'], errors='coerce').astype('Int64')
        P_all = P_all.dropna(subset=['date', 'market', 'ci_start_date', 'ci_end_date', 'L_net']).copy()
        P_cache = {}

        def load_P_from_csv(L_net: int, ci_start_date, ci_end_date):
            key = (int(L_net), _norm_date(ci_start_date), _norm_date(ci_end_date))
            if key in P_cache:
                return P_cache[key]
            ln, cs, ce = key
            sub = P_all[(P_all['L_net'] == ln) & (P_all['ci_start_date'] == cs) & (P_all['ci_end_date'] == ce)]
            if sub.empty:
                P_cache[key] = (None, None, None, None)
                return P_cache[key]

            def _pivot(col):
                df = sub.pivot_table(index='date', columns='market', values=col, aggfunc='mean')
                return df.sort_index()
            P_corr = _pivot('P_corr')
            P_kendall = _pivot('P_kendall')
            P_partial = _pivot('P_partial')
            P_tail = _pivot('P_tail')
            P_cache[key] = (P_corr, P_kendall, P_partial, P_tail)
            return P_cache[key]

        def forecast_one_market_oneH_harci(data_rv_all: pd.DataFrame, P_corr: pd.DataFrame, P_kendall: pd.DataFrame, P_partial: pd.DataFrame, P_tail: pd.DataFrame, market: str, fit_start_date: pd.Timestamp, H: int):
            df = data_rv_all[market].to_frame('RV_now')
            df['RV_D'] = df['RV_now'].shift(1)
            df['RV_W'] = df['RV_D'].rolling(5).mean()
            df['RV_M'] = df['RV_D'].rolling(22).mean()
            df['Y'] = make_future_mean(df['RV_now'], H)
            ci = pd.DataFrame({'CI_corr': P_corr[market], 'CI_kendall': P_kendall[market], 'CI_partial': P_partial[market], 'CI_tail': P_tail[market]}).reindex(df.index)
            df = df.join(ci, how='inner').dropna()
            df['CI_corr_D'] = df['CI_corr'].shift(1)
            df['CI_kendall_D'] = df['CI_kendall'].shift(1)
            df['CI_partial_D'] = df['CI_partial'].shift(1)
            df['CI_tail_D'] = df['CI_tail'].shift(1)
            df = df.dropna()
            df = df.loc[fit_start_date:]
            data_test = df.loc[pred_start:pred_end].copy()
            if len(data_test) == 0:
                return (None, np.nan, np.nan, np.nan)
            window_len = len(df.loc[:pred_start]) - 1
            if window_len <= 0:
                return (None, np.nan, np.nan, np.nan)
            y_true = data_test['Y'].astype(float).copy()
            idx = y_true.index
            preds = []
            for i in tqdm(range(len(idx)), desc=f'HAR-CI {market} H{H}', leave=False):
                train = df.iloc[i:i + window_len].copy()
                train['_Y_'] = train['Y'].astype(float).values
                res = smf.ols('_Y_ ~ RV_D + RV_W + RV_M + CI_corr_D + CI_kendall_D + CI_partial_D', train).fit()
                p = res.params
                pred = p.iloc[0] + p.iloc[1] * data_test['RV_D'].iloc[i] + p.iloc[2] * data_test['RV_W'].iloc[i] + p.iloc[3] * data_test['RV_M'].iloc[i] + p.iloc[4] * data_test['CI_corr_D'].iloc[i] + p.iloc[5] * data_test['CI_kendall_D'].iloc[i] + p.iloc[6] * data_test['CI_partial_D'].iloc[i]
                preds.append(float(pred))
            y_pred = pd.Series(preds, index=idx)
            mcs = pd.concat([y_true.rename('RV_True'), y_pred.rename('RV_Predict')], axis=1)
            mcs['MSE'] = (mcs['RV_True'] - mcs['RV_Predict']) ** 2
            mcs['MAE'] = np.abs(mcs['RV_True'] - mcs['RV_Predict'])
            mse = float(np.nanmean(mcs['MSE'].values))
            mae = float(np.nanmean(mcs['MAE'].values))
            denom = np.sqrt(np.nanmean(mcs['RV_Predict'].values ** 2)) + np.sqrt(np.nanmean(mcs['RV_True'].values ** 2))
            theil = float(np.sqrt(np.nanmean((mcs['RV_True'].values - mcs['RV_Predict'].values) ** 2)) / denom) if denom != 0 else np.nan
            return (mcs, mse, mae, theil)
        L_fit = L_fit_list[0]
        L_net = L_net_list[0]
        for H in H_LIST:
            error_by_market = {}
            for mkt in markets:
                if fit_end_date not in data_rv.index:
                    error_by_market[mkt] = [np.nan, np.nan, np.nan]
                    continue
                fit_end_pos = data_rv.index.get_loc(fit_end_date)
                fit_start_pos = fit_end_pos - (L_fit - 1)
                if fit_start_pos < 0:
                    error_by_market[mkt] = [np.nan, np.nan, np.nan]
                    continue
                fit_start_date = data_rv.index[fit_start_pos]
                ret_pos_end = df_ret.index.searchsorted(fit_start_date) - 1
                if ret_pos_end < 0:
                    error_by_market[mkt] = [np.nan, np.nan, np.nan]
                    continue
                net_end_date = df_ret.index[ret_pos_end]
                ci_start_date = net_end_date
                ci_end_date = pred_end
                P_corr, P_kendall, P_partial, P_tail = load_P_from_csv(L_net=L_net, ci_start_date=ci_start_date, ci_end_date=ci_end_date)
                if P_corr is None:
                    error_by_market[mkt] = [np.nan, np.nan, np.nan]
                    continue
                if mkt not in P_corr.columns:
                    error_by_market[mkt] = [np.nan, np.nan, np.nan]
                    continue
                mcs, mse, mae, theil = forecast_one_market_oneH_harci(data_rv_all=data_rv, P_corr=P_corr, P_kendall=P_kendall, P_partial=P_partial, P_tail=P_tail, market=mkt, fit_start_date=fit_start_date, H=H)
                if mcs is None:
                    tmp = pd.DataFrame(index=pd.DatetimeIndex([], name='date'))
                    tmp['RV_True'] = np.nan
                    tmp['RV_Predict'] = np.nan
                    tmp['MSE'] = np.nan
                    tmp['MAE'] = np.nan
                    model_outputs[str(mkt).strip().lstrip('.')] = tmp.copy()
                else:
                    model_outputs[str(mkt).strip().lstrip('.')] = mcs.copy()
                error_by_market[mkt] = [mse, mae, theil]
            ERROR = pd.DataFrame(error_by_market, index=['MSE', 'MAE', 'Theil'])
        return model_outputs

    def run_layer_corr_kendall_tail():
        model_outputs = {}
        import os
        import numpy as np
        import pandas as pd
        import statsmodels.formula.api as smf
        from functools import partial
        from tqdm import tqdm as _tqdm
        tqdm = partial(_tqdm, bar_format="{bar} {percentage:3.0f}%", ncols=50)
        RET_PATH = RETURN_FILE
        RV_PATH = RV_FILE
        P_STORE_PATH = CENTRALITY_FILE
        PLOT = False
        fit_end_date = pd.to_datetime('2018-05-01')
        pred_start = pd.to_datetime('2018-05-02')
        pred_end = pd.to_datetime('2022-02-25')
        H_LIST = [22]
        FIT_START = 1000
        FIT_STEP = 100
        FIT_MAX = 1
        L_fit_list = [FIT_START + FIT_STEP * k for k in range(FIT_MAX)]
        NET_MIN = 1000
        NET_MAX = 1000
        NET_STEP = 50
        L_net_list = list(range(NET_MIN, NET_MAX + 1, NET_STEP))

        def _norm_date(x):
            return pd.to_datetime(x).tz_localize(None).normalize()

        def make_future_mean(rv_now: pd.Series, H: int) -> pd.Series:
            if H == 1:
                return rv_now
            ma = rv_now.rolling(window=H, center=False).mean()
            return ma.shift(-(H - 1))

        def read_returns(path):
            df = pd.read_csv(path)
            df = df.loc[:, ~df.columns.str.startswith('Unnamed')]
            df['date'] = pd.to_datetime(df['date'].astype(str).str.strip().str[:10], errors='coerce')
            df['date'] = df['date'].dt.tz_localize(None).dt.normalize()
            df = df.set_index('date').dropna(how='all')
            return df

        def read_rv_no_scale(path):
            data = pd.read_csv(path)
            if 'date' not in data.columns:
                raise RuntimeError("RV file must contain a 'date' column.")
            num_cols = [c for c in data.columns if c != 'date']
            data[num_cols] = data[num_cols].apply(pd.to_numeric, errors='coerce')
            data = data.dropna(how='all')
            data['date'] = pd.to_datetime(data['date'], utc=True, errors='coerce')
            data = data.dropna(subset=['date'])
            data['date'] = data['date'].dt.tz_localize(None)
            data['date'] = data['date'].dt.normalize()
            data = data.set_index('date')
            data = data.dropna(axis=1, how='all')
            return data
        df_ret = read_returns(RET_PATH)
        data_rv = read_rv_no_scale(RV_PATH)
        markets = [c for c in data_rv.columns if c in df_ret.columns]
        P_all = pd.read_csv(P_STORE_PATH)
        need_cols = ['date', 'market', 'ci_start_date', 'ci_end_date', 'L_net', 'P_corr', 'P_kendall', 'P_partial', 'P_tail']
        miss = [c for c in need_cols if c not in P_all.columns]
        if miss:
            raise RuntimeError(f'P.csv is missing required columns: {miss}; available columns={list(P_all.columns)}')
        P_all['date'] = pd.to_datetime(P_all['date'], errors='coerce').dt.tz_localize(None).dt.normalize()
        P_all['ci_start_date'] = P_all['ci_start_date'].apply(_norm_date)
        P_all['ci_end_date'] = P_all['ci_end_date'].apply(_norm_date)
        P_all['L_net'] = pd.to_numeric(P_all['L_net'], errors='coerce').astype('Int64')
        P_all = P_all.dropna(subset=['date', 'market', 'ci_start_date', 'ci_end_date', 'L_net']).copy()
        P_cache = {}

        def load_P_from_csv(L_net: int, ci_start_date, ci_end_date):
            key = (int(L_net), _norm_date(ci_start_date), _norm_date(ci_end_date))
            if key in P_cache:
                return P_cache[key]
            ln, cs, ce = key
            sub = P_all[(P_all['L_net'] == ln) & (P_all['ci_start_date'] == cs) & (P_all['ci_end_date'] == ce)]
            if sub.empty:
                P_cache[key] = (None, None, None, None)
                return P_cache[key]

            def _pivot(col):
                df = sub.pivot_table(index='date', columns='market', values=col, aggfunc='mean')
                return df.sort_index()
            P_corr = _pivot('P_corr')
            P_kendall = _pivot('P_kendall')
            P_partial = _pivot('P_partial')
            P_tail = _pivot('P_tail')
            P_cache[key] = (P_corr, P_kendall, P_partial, P_tail)
            return P_cache[key]

        def forecast_one_market_oneH_harci(data_rv_all: pd.DataFrame, P_corr: pd.DataFrame, P_kendall: pd.DataFrame, P_partial: pd.DataFrame, P_tail: pd.DataFrame, market: str, fit_start_date: pd.Timestamp, H: int):
            df = data_rv_all[market].to_frame('RV_now')
            df['RV_D'] = df['RV_now'].shift(1)
            df['RV_W'] = df['RV_D'].rolling(5).mean()
            df['RV_M'] = df['RV_D'].rolling(22).mean()
            df['Y'] = make_future_mean(df['RV_now'], H)
            ci = pd.DataFrame({'CI_corr': P_corr[market], 'CI_kendall': P_kendall[market], 'CI_partial': P_partial[market], 'CI_tail': P_tail[market]}).reindex(df.index)
            df = df.join(ci, how='inner').dropna()
            df['CI_corr_D'] = df['CI_corr'].shift(1)
            df['CI_kendall_D'] = df['CI_kendall'].shift(1)
            df['CI_partial_D'] = df['CI_partial'].shift(1)
            df['CI_tail_D'] = df['CI_tail'].shift(1)
            df = df.dropna()
            df = df.loc[fit_start_date:]
            data_test = df.loc[pred_start:pred_end].copy()
            if len(data_test) == 0:
                return (None, np.nan, np.nan, np.nan)
            window_len = len(df.loc[:pred_start]) - 1
            if window_len <= 0:
                return (None, np.nan, np.nan, np.nan)
            y_true = data_test['Y'].astype(float).copy()
            idx = y_true.index
            preds = []
            for i in tqdm(range(len(idx)), desc=f'HAR-CI {market} H{H}', leave=False):
                train = df.iloc[i:i + window_len].copy()
                train['_Y_'] = train['Y'].astype(float).values
                res = smf.ols('_Y_ ~ RV_D + RV_W + RV_M + CI_corr_D + CI_kendall_D + CI_tail_D', train).fit()
                p = res.params
                pred = p.iloc[0] + p.iloc[1] * data_test['RV_D'].iloc[i] + p.iloc[2] * data_test['RV_W'].iloc[i] + p.iloc[3] * data_test['RV_M'].iloc[i] + p.iloc[4] * data_test['CI_corr_D'].iloc[i] + p.iloc[5] * data_test['CI_kendall_D'].iloc[i] + p.iloc[6] * data_test['CI_tail_D'].iloc[i]
                preds.append(float(pred))
            y_pred = pd.Series(preds, index=idx)
            mcs = pd.concat([y_true.rename('RV_True'), y_pred.rename('RV_Predict')], axis=1)
            mcs['MSE'] = (mcs['RV_True'] - mcs['RV_Predict']) ** 2
            mcs['MAE'] = np.abs(mcs['RV_True'] - mcs['RV_Predict'])
            mse = float(np.nanmean(mcs['MSE'].values))
            mae = float(np.nanmean(mcs['MAE'].values))
            denom = np.sqrt(np.nanmean(mcs['RV_Predict'].values ** 2)) + np.sqrt(np.nanmean(mcs['RV_True'].values ** 2))
            theil = float(np.sqrt(np.nanmean((mcs['RV_True'].values - mcs['RV_Predict'].values) ** 2)) / denom) if denom != 0 else np.nan
            return (mcs, mse, mae, theil)
        L_fit = L_fit_list[0]
        L_net = L_net_list[0]
        for H in H_LIST:
            error_by_market = {}
            for mkt in markets:
                if fit_end_date not in data_rv.index:
                    error_by_market[mkt] = [np.nan, np.nan, np.nan]
                    continue
                fit_end_pos = data_rv.index.get_loc(fit_end_date)
                fit_start_pos = fit_end_pos - (L_fit - 1)
                if fit_start_pos < 0:
                    error_by_market[mkt] = [np.nan, np.nan, np.nan]
                    continue
                fit_start_date = data_rv.index[fit_start_pos]
                ret_pos_end = df_ret.index.searchsorted(fit_start_date) - 1
                if ret_pos_end < 0:
                    error_by_market[mkt] = [np.nan, np.nan, np.nan]
                    continue
                net_end_date = df_ret.index[ret_pos_end]
                ci_start_date = net_end_date
                ci_end_date = pred_end
                P_corr, P_kendall, P_partial, P_tail = load_P_from_csv(L_net=L_net, ci_start_date=ci_start_date, ci_end_date=ci_end_date)
                if P_corr is None:
                    error_by_market[mkt] = [np.nan, np.nan, np.nan]
                    continue
                if mkt not in P_corr.columns:
                    error_by_market[mkt] = [np.nan, np.nan, np.nan]
                    continue
                mcs, mse, mae, theil = forecast_one_market_oneH_harci(data_rv_all=data_rv, P_corr=P_corr, P_kendall=P_kendall, P_partial=P_partial, P_tail=P_tail, market=mkt, fit_start_date=fit_start_date, H=H)
                if mcs is None:
                    tmp = pd.DataFrame(index=pd.DatetimeIndex([], name='date'))
                    tmp['RV_True'] = np.nan
                    tmp['RV_Predict'] = np.nan
                    tmp['MSE'] = np.nan
                    tmp['MAE'] = np.nan
                    model_outputs[str(mkt).strip().lstrip('.')] = tmp.copy()
                else:
                    model_outputs[str(mkt).strip().lstrip('.')] = mcs.copy()
                error_by_market[mkt] = [mse, mae, theil]
            ERROR = pd.DataFrame(error_by_market, index=['MSE', 'MAE', 'Theil'])
        return model_outputs

    def run_layer_corr_tail_partial():
        model_outputs = {}
        import os
        import numpy as np
        import pandas as pd
        import statsmodels.formula.api as smf
        from functools import partial
        from tqdm import tqdm as _tqdm
        tqdm = partial(_tqdm, bar_format="{bar} {percentage:3.0f}%", ncols=50)
        RET_PATH = RETURN_FILE
        RV_PATH = RV_FILE
        P_STORE_PATH = CENTRALITY_FILE
        PLOT = False
        fit_end_date = pd.to_datetime('2018-05-01')
        pred_start = pd.to_datetime('2018-05-02')
        pred_end = pd.to_datetime('2022-02-25')
        H_LIST = [22]
        FIT_START = 1000
        FIT_STEP = 100
        FIT_MAX = 1
        L_fit_list = [FIT_START + FIT_STEP * k for k in range(FIT_MAX)]
        NET_MIN = 1000
        NET_MAX = 1000
        NET_STEP = 50
        L_net_list = list(range(NET_MIN, NET_MAX + 1, NET_STEP))

        def _norm_date(x):
            return pd.to_datetime(x).tz_localize(None).normalize()

        def make_future_mean(rv_now: pd.Series, H: int) -> pd.Series:
            if H == 1:
                return rv_now
            ma = rv_now.rolling(window=H, center=False).mean()
            return ma.shift(-(H - 1))

        def read_returns(path):
            df = pd.read_csv(path)
            df = df.loc[:, ~df.columns.str.startswith('Unnamed')]
            df['date'] = pd.to_datetime(df['date'].astype(str).str.strip().str[:10], errors='coerce')
            df['date'] = df['date'].dt.tz_localize(None).dt.normalize()
            df = df.set_index('date').dropna(how='all')
            return df

        def read_rv_no_scale(path):
            data = pd.read_csv(path)
            if 'date' not in data.columns:
                raise RuntimeError("RV file must contain a 'date' column.")
            num_cols = [c for c in data.columns if c != 'date']
            data[num_cols] = data[num_cols].apply(pd.to_numeric, errors='coerce')
            data = data.dropna(how='all')
            data['date'] = pd.to_datetime(data['date'], utc=True, errors='coerce')
            data = data.dropna(subset=['date'])
            data['date'] = data['date'].dt.tz_localize(None)
            data['date'] = data['date'].dt.normalize()
            data = data.set_index('date')
            data = data.dropna(axis=1, how='all')
            return data
        df_ret = read_returns(RET_PATH)
        data_rv = read_rv_no_scale(RV_PATH)
        markets = [c for c in data_rv.columns if c in df_ret.columns]
        P_all = pd.read_csv(P_STORE_PATH)
        need_cols = ['date', 'market', 'ci_start_date', 'ci_end_date', 'L_net', 'P_corr', 'P_kendall', 'P_partial', 'P_tail']
        miss = [c for c in need_cols if c not in P_all.columns]
        if miss:
            raise RuntimeError(f'P.csv is missing required columns: {miss}; available columns={list(P_all.columns)}')
        P_all['date'] = pd.to_datetime(P_all['date'], errors='coerce').dt.tz_localize(None).dt.normalize()
        P_all['ci_start_date'] = P_all['ci_start_date'].apply(_norm_date)
        P_all['ci_end_date'] = P_all['ci_end_date'].apply(_norm_date)
        P_all['L_net'] = pd.to_numeric(P_all['L_net'], errors='coerce').astype('Int64')
        P_all = P_all.dropna(subset=['date', 'market', 'ci_start_date', 'ci_end_date', 'L_net']).copy()
        P_cache = {}

        def load_P_from_csv(L_net: int, ci_start_date, ci_end_date):
            key = (int(L_net), _norm_date(ci_start_date), _norm_date(ci_end_date))
            if key in P_cache:
                return P_cache[key]
            ln, cs, ce = key
            sub = P_all[(P_all['L_net'] == ln) & (P_all['ci_start_date'] == cs) & (P_all['ci_end_date'] == ce)]
            if sub.empty:
                P_cache[key] = (None, None, None, None)
                return P_cache[key]

            def _pivot(col):
                df = sub.pivot_table(index='date', columns='market', values=col, aggfunc='mean')
                return df.sort_index()
            P_corr = _pivot('P_corr')
            P_kendall = _pivot('P_kendall')
            P_partial = _pivot('P_partial')
            P_tail = _pivot('P_tail')
            P_cache[key] = (P_corr, P_kendall, P_partial, P_tail)
            return P_cache[key]

        def forecast_one_market_oneH_harci(data_rv_all: pd.DataFrame, P_corr: pd.DataFrame, P_kendall: pd.DataFrame, P_partial: pd.DataFrame, P_tail: pd.DataFrame, market: str, fit_start_date: pd.Timestamp, H: int):
            df = data_rv_all[market].to_frame('RV_now')
            df['RV_D'] = df['RV_now'].shift(1)
            df['RV_W'] = df['RV_D'].rolling(5).mean()
            df['RV_M'] = df['RV_D'].rolling(22).mean()
            df['Y'] = make_future_mean(df['RV_now'], H)
            ci = pd.DataFrame({'CI_corr': P_corr[market], 'CI_kendall': P_kendall[market], 'CI_partial': P_partial[market], 'CI_tail': P_tail[market]}).reindex(df.index)
            df = df.join(ci, how='inner').dropna()
            df['CI_corr_D'] = df['CI_corr'].shift(1)
            df['CI_kendall_D'] = df['CI_kendall'].shift(1)
            df['CI_partial_D'] = df['CI_partial'].shift(1)
            df['CI_tail_D'] = df['CI_tail'].shift(1)
            df = df.dropna()
            df = df.loc[fit_start_date:]
            data_test = df.loc[pred_start:pred_end].copy()
            if len(data_test) == 0:
                return (None, np.nan, np.nan, np.nan)
            window_len = len(df.loc[:pred_start]) - 1
            if window_len <= 0:
                return (None, np.nan, np.nan, np.nan)
            y_true = data_test['Y'].astype(float).copy()
            idx = y_true.index
            preds = []
            for i in tqdm(range(len(idx)), desc=f'HAR-CI {market} H{H}', leave=False):
                train = df.iloc[i:i + window_len].copy()
                train['_Y_'] = train['Y'].astype(float).values
                res = smf.ols('_Y_ ~ RV_D + RV_W + RV_M + CI_corr_D + CI_tail_D + CI_partial_D', train).fit()
                p = res.params
                pred = p.iloc[0] + p.iloc[1] * data_test['RV_D'].iloc[i] + p.iloc[2] * data_test['RV_W'].iloc[i] + p.iloc[3] * data_test['RV_M'].iloc[i] + p.iloc[4] * data_test['CI_corr_D'].iloc[i] + p.iloc[5] * data_test['CI_tail_D'].iloc[i] + p.iloc[6] * data_test['CI_partial_D'].iloc[i]
                preds.append(float(pred))
            y_pred = pd.Series(preds, index=idx)
            mcs = pd.concat([y_true.rename('RV_True'), y_pred.rename('RV_Predict')], axis=1)
            mcs['MSE'] = (mcs['RV_True'] - mcs['RV_Predict']) ** 2
            mcs['MAE'] = np.abs(mcs['RV_True'] - mcs['RV_Predict'])
            mse = float(np.nanmean(mcs['MSE'].values))
            mae = float(np.nanmean(mcs['MAE'].values))
            denom = np.sqrt(np.nanmean(mcs['RV_Predict'].values ** 2)) + np.sqrt(np.nanmean(mcs['RV_True'].values ** 2))
            theil = float(np.sqrt(np.nanmean((mcs['RV_True'].values - mcs['RV_Predict'].values) ** 2)) / denom) if denom != 0 else np.nan
            return (mcs, mse, mae, theil)
        L_fit = L_fit_list[0]
        L_net = L_net_list[0]
        for H in H_LIST:
            error_by_market = {}
            for mkt in markets:
                if fit_end_date not in data_rv.index:
                    error_by_market[mkt] = [np.nan, np.nan, np.nan]
                    continue
                fit_end_pos = data_rv.index.get_loc(fit_end_date)
                fit_start_pos = fit_end_pos - (L_fit - 1)
                if fit_start_pos < 0:
                    error_by_market[mkt] = [np.nan, np.nan, np.nan]
                    continue
                fit_start_date = data_rv.index[fit_start_pos]
                ret_pos_end = df_ret.index.searchsorted(fit_start_date) - 1
                if ret_pos_end < 0:
                    error_by_market[mkt] = [np.nan, np.nan, np.nan]
                    continue
                net_end_date = df_ret.index[ret_pos_end]
                ci_start_date = net_end_date
                ci_end_date = pred_end
                P_corr, P_kendall, P_partial, P_tail = load_P_from_csv(L_net=L_net, ci_start_date=ci_start_date, ci_end_date=ci_end_date)
                if P_corr is None:
                    error_by_market[mkt] = [np.nan, np.nan, np.nan]
                    continue
                if mkt not in P_corr.columns:
                    error_by_market[mkt] = [np.nan, np.nan, np.nan]
                    continue
                mcs, mse, mae, theil = forecast_one_market_oneH_harci(data_rv_all=data_rv, P_corr=P_corr, P_kendall=P_kendall, P_partial=P_partial, P_tail=P_tail, market=mkt, fit_start_date=fit_start_date, H=H)
                if mcs is None:
                    tmp = pd.DataFrame(index=pd.DatetimeIndex([], name='date'))
                    tmp['RV_True'] = np.nan
                    tmp['RV_Predict'] = np.nan
                    tmp['MSE'] = np.nan
                    tmp['MAE'] = np.nan
                    model_outputs[str(mkt).strip().lstrip('.')] = tmp.copy()
                else:
                    model_outputs[str(mkt).strip().lstrip('.')] = mcs.copy()
                error_by_market[mkt] = [mse, mae, theil]
            ERROR = pd.DataFrame(error_by_market, index=['MSE', 'MAE', 'Theil'])
        return model_outputs

    def run_layer_kendall_tail_partial():
        model_outputs = {}
        import os
        import numpy as np
        import pandas as pd
        import statsmodels.formula.api as smf
        from functools import partial
        from tqdm import tqdm as _tqdm
        tqdm = partial(_tqdm, bar_format="{bar} {percentage:3.0f}%", ncols=50)
        RET_PATH = RETURN_FILE
        RV_PATH = RV_FILE
        P_STORE_PATH = CENTRALITY_FILE
        PLOT = False
        fit_end_date = pd.to_datetime('2018-05-01')
        pred_start = pd.to_datetime('2018-05-02')
        pred_end = pd.to_datetime('2022-02-25')
        H_LIST = [22]
        FIT_START = 1000
        FIT_STEP = 100
        FIT_MAX = 1
        L_fit_list = [FIT_START + FIT_STEP * k for k in range(FIT_MAX)]
        NET_MIN = 1000
        NET_MAX = 1000
        NET_STEP = 50
        L_net_list = list(range(NET_MIN, NET_MAX + 1, NET_STEP))

        def _norm_date(x):
            return pd.to_datetime(x).tz_localize(None).normalize()

        def make_future_mean(rv_now: pd.Series, H: int) -> pd.Series:
            if H == 1:
                return rv_now
            ma = rv_now.rolling(window=H, center=False).mean()
            return ma.shift(-(H - 1))

        def read_returns(path):
            df = pd.read_csv(path)
            df = df.loc[:, ~df.columns.str.startswith('Unnamed')]
            df['date'] = pd.to_datetime(df['date'].astype(str).str.strip().str[:10], errors='coerce')
            df['date'] = df['date'].dt.tz_localize(None).dt.normalize()
            df = df.set_index('date').dropna(how='all')
            return df

        def read_rv_no_scale(path):
            data = pd.read_csv(path)
            if 'date' not in data.columns:
                raise RuntimeError("RV file must contain a 'date' column.")
            num_cols = [c for c in data.columns if c != 'date']
            data[num_cols] = data[num_cols].apply(pd.to_numeric, errors='coerce')
            data = data.dropna(how='all')
            data['date'] = pd.to_datetime(data['date'], utc=True, errors='coerce')
            data = data.dropna(subset=['date'])
            data['date'] = data['date'].dt.tz_localize(None)
            data['date'] = data['date'].dt.normalize()
            data = data.set_index('date')
            data = data.dropna(axis=1, how='all')
            return data
        df_ret = read_returns(RET_PATH)
        data_rv = read_rv_no_scale(RV_PATH)
        markets = [c for c in data_rv.columns if c in df_ret.columns]
        P_all = pd.read_csv(P_STORE_PATH)
        need_cols = ['date', 'market', 'ci_start_date', 'ci_end_date', 'L_net', 'P_corr', 'P_kendall', 'P_partial', 'P_tail']
        miss = [c for c in need_cols if c not in P_all.columns]
        if miss:
            raise RuntimeError(f'P.csv is missing required columns: {miss}; available columns={list(P_all.columns)}')
        P_all['date'] = pd.to_datetime(P_all['date'], errors='coerce').dt.tz_localize(None).dt.normalize()
        P_all['ci_start_date'] = P_all['ci_start_date'].apply(_norm_date)
        P_all['ci_end_date'] = P_all['ci_end_date'].apply(_norm_date)
        P_all['L_net'] = pd.to_numeric(P_all['L_net'], errors='coerce').astype('Int64')
        P_all = P_all.dropna(subset=['date', 'market', 'ci_start_date', 'ci_end_date', 'L_net']).copy()
        P_cache = {}

        def load_P_from_csv(L_net: int, ci_start_date, ci_end_date):
            key = (int(L_net), _norm_date(ci_start_date), _norm_date(ci_end_date))
            if key in P_cache:
                return P_cache[key]
            ln, cs, ce = key
            sub = P_all[(P_all['L_net'] == ln) & (P_all['ci_start_date'] == cs) & (P_all['ci_end_date'] == ce)]
            if sub.empty:
                P_cache[key] = (None, None, None, None)
                return P_cache[key]

            def _pivot(col):
                df = sub.pivot_table(index='date', columns='market', values=col, aggfunc='mean')
                return df.sort_index()
            P_corr = _pivot('P_corr')
            P_kendall = _pivot('P_kendall')
            P_partial = _pivot('P_partial')
            P_tail = _pivot('P_tail')
            P_cache[key] = (P_corr, P_kendall, P_partial, P_tail)
            return P_cache[key]

        def forecast_one_market_oneH_harci(data_rv_all: pd.DataFrame, P_corr: pd.DataFrame, P_kendall: pd.DataFrame, P_partial: pd.DataFrame, P_tail: pd.DataFrame, market: str, fit_start_date: pd.Timestamp, H: int):
            df = data_rv_all[market].to_frame('RV_now')
            df['RV_D'] = df['RV_now'].shift(1)
            df['RV_W'] = df['RV_D'].rolling(5).mean()
            df['RV_M'] = df['RV_D'].rolling(22).mean()
            df['Y'] = make_future_mean(df['RV_now'], H)
            ci = pd.DataFrame({'CI_corr': P_corr[market], 'CI_kendall': P_kendall[market], 'CI_partial': P_partial[market], 'CI_tail': P_tail[market]}).reindex(df.index)
            df = df.join(ci, how='inner').dropna()
            df['CI_corr_D'] = df['CI_corr'].shift(1)
            df['CI_kendall_D'] = df['CI_kendall'].shift(1)
            df['CI_partial_D'] = df['CI_partial'].shift(1)
            df['CI_tail_D'] = df['CI_tail'].shift(1)
            df = df.dropna()
            df = df.loc[fit_start_date:]
            data_test = df.loc[pred_start:pred_end].copy()
            if len(data_test) == 0:
                return (None, np.nan, np.nan, np.nan)
            window_len = len(df.loc[:pred_start]) - 1
            if window_len <= 0:
                return (None, np.nan, np.nan, np.nan)
            y_true = data_test['Y'].astype(float).copy()
            idx = y_true.index
            preds = []
            for i in tqdm(range(len(idx)), desc=f'HAR-CI {market} H{H}', leave=False):
                train = df.iloc[i:i + window_len].copy()
                train['_Y_'] = train['Y'].astype(float).values
                res = smf.ols('_Y_ ~ RV_D + RV_W + RV_M + CI_kendall_D + CI_tail_D + CI_partial_D', train).fit()
                p = res.params
                pred = p.iloc[0] + p.iloc[1] * data_test['RV_D'].iloc[i] + p.iloc[2] * data_test['RV_W'].iloc[i] + p.iloc[3] * data_test['RV_M'].iloc[i] + p.iloc[4] * data_test['CI_kendall_D'].iloc[i] + p.iloc[5] * data_test['CI_tail_D'].iloc[i] + p.iloc[6] * data_test['CI_partial_D'].iloc[i]
                preds.append(float(pred))
            y_pred = pd.Series(preds, index=idx)
            mcs = pd.concat([y_true.rename('RV_True'), y_pred.rename('RV_Predict')], axis=1)
            mcs['MSE'] = (mcs['RV_True'] - mcs['RV_Predict']) ** 2
            mcs['MAE'] = np.abs(mcs['RV_True'] - mcs['RV_Predict'])
            mse = float(np.nanmean(mcs['MSE'].values))
            mae = float(np.nanmean(mcs['MAE'].values))
            denom = np.sqrt(np.nanmean(mcs['RV_Predict'].values ** 2)) + np.sqrt(np.nanmean(mcs['RV_True'].values ** 2))
            theil = float(np.sqrt(np.nanmean((mcs['RV_True'].values - mcs['RV_Predict'].values) ** 2)) / denom) if denom != 0 else np.nan
            return (mcs, mse, mae, theil)
        L_fit = L_fit_list[0]
        L_net = L_net_list[0]
        for H in H_LIST:
            error_by_market = {}
            for mkt in markets:
                if fit_end_date not in data_rv.index:
                    error_by_market[mkt] = [np.nan, np.nan, np.nan]
                    continue
                fit_end_pos = data_rv.index.get_loc(fit_end_date)
                fit_start_pos = fit_end_pos - (L_fit - 1)
                if fit_start_pos < 0:
                    error_by_market[mkt] = [np.nan, np.nan, np.nan]
                    continue
                fit_start_date = data_rv.index[fit_start_pos]
                ret_pos_end = df_ret.index.searchsorted(fit_start_date) - 1
                if ret_pos_end < 0:
                    error_by_market[mkt] = [np.nan, np.nan, np.nan]
                    continue
                net_end_date = df_ret.index[ret_pos_end]
                ci_start_date = net_end_date
                ci_end_date = pred_end
                P_corr, P_kendall, P_partial, P_tail = load_P_from_csv(L_net=L_net, ci_start_date=ci_start_date, ci_end_date=ci_end_date)
                if P_corr is None:
                    error_by_market[mkt] = [np.nan, np.nan, np.nan]
                    continue
                if mkt not in P_corr.columns:
                    error_by_market[mkt] = [np.nan, np.nan, np.nan]
                    continue
                mcs, mse, mae, theil = forecast_one_market_oneH_harci(data_rv_all=data_rv, P_corr=P_corr, P_kendall=P_kendall, P_partial=P_partial, P_tail=P_tail, market=mkt, fit_start_date=fit_start_date, H=H)
                if mcs is None:
                    tmp = pd.DataFrame(index=pd.DatetimeIndex([], name='date'))
                    tmp['RV_True'] = np.nan
                    tmp['RV_Predict'] = np.nan
                    tmp['MSE'] = np.nan
                    tmp['MAE'] = np.nan
                    model_outputs[str(mkt).strip().lstrip('.')] = tmp.copy()
                else:
                    model_outputs[str(mkt).strip().lstrip('.')] = mcs.copy()
                error_by_market[mkt] = [mse, mae, theil]
            ERROR = pd.DataFrame(error_by_market, index=['MSE', 'MAE', 'Theil'])
        return model_outputs

    LAYER_RUNNERS = [
        ("HAR-CI-corr", run_layer_corr),
        ("HAR-CI-kendall", run_layer_kendall),
        ("HAR-CI-partial", run_layer_partial),
        ("HAR-CI-tail", run_layer_tail),
        ("HAR-CI-corr_kendall_partial", run_layer_corr_kendall_partial),
        ("HAR-CI-corr_kendall_tail", run_layer_corr_kendall_tail),
        ("HAR-CI-corr_tail_partial", run_layer_corr_tail_partial),
        ("HAR-CI-kendall_tail_partial", run_layer_kendall_tail_partial),
    ]

    return SimpleNamespace(**locals())


def build_table_05():
    from pathlib import Path

    import argparse

    import warnings

    import numpy as np

    import pandas as pd

    from functools import partial

    from tqdm import tqdm as _tqdm

    tqdm = partial(_tqdm, bar_format="{bar} {percentage:3.0f}%", ncols=50)

    from arch.bootstrap import MCS

    from openpyxl import Workbook

    from openpyxl.styles import Font, Alignment, Border, Side

    ROOT = Path(__file__).resolve().parents[1]

    RV_FILE = ROOT / "data/raw/RV.csv"

    RETURN_FILE = ROOT / "data/raw/Return.csv"

    LOG_CENTRALITY_FILE = ROOT / "data/processed/centrality/P_log.csv"

    OUTPUT_FILE = ROOT / "results/generated/Table_05_category2_forecast_performance.xlsx"

    THRESH = 1e-3

    SCALE_IF_MOST_GT_THRESH = 1e-4

    EPS = 1e-12

    RATIO_MIN = 1e-6

    RATIO_MAX = 1e6

    QLIKE_CAP = 100.0

    POST_SCALE_THRESHOLD = 1e4

    POST_SCALE_FACTOR = 1e-8

    def run_har():
        model_outputs = {}
        import pandas as pd
        import numpy as np
        from functools import partial
        from tqdm import tqdm as _tqdm
        tqdm = partial(_tqdm, bar_format="{bar} {percentage:3.0f}%", ncols=50)
        import statsmodels.formula.api as smf
        data = pd.read_csv(RV_FILE)
        num_cols = data.columns.drop('date')
        data[num_cols] = data[num_cols].apply(pd.to_numeric, errors='coerce')
        data = data.dropna()
        data['date'] = pd.to_datetime(data['date'], utc=True)
        data['date'] = data['date'].dt.tz_localize(None)
        data['date'] = data['date'].dt.normalize()
        data = data.set_index('date')
        for c in data.columns:
            data_RV = pd.DataFrame(data[c])
            data_RV.columns = ['RV_now']
            data_RV['RV_D'] = data_RV['RV_now'].shift(1)
            data_RV['RV_W'] = data_RV['RV_D'].rolling(5).mean()
            data_RV['RV_M'] = data_RV['RV_D'].rolling(22).mean()
            data_RV['RV_MA22'] = data_RV['RV_now'].rolling(window=22, center=False).mean().shift(-21)
            data_RV = data_RV.dropna().loc['2014-07-04':]
            time_start = '2018/05/02'
            time_stop = '2022-02-25'
            data_test = data_RV.loc[time_start:time_stop]
            data_predict = data_test['RV_MA22']
            train_len = len(data_RV.loc[:time_start]) - 1
            result = []
            for i in tqdm(range(len(data_test)), desc=f'HAR {str(c).strip().lstrip(".")}', leave=False):
                data_rolling = data_RV.iloc[i:i + train_len]
                params = smf.ols('RV_MA22 ~ RV_D + RV_W + RV_M', data_rolling).fit().params
                result.append(
                    params.iloc[0]
                    + params.iloc[1] * data_test['RV_D'].iloc[i]
                    + params.iloc[2] * data_test['RV_W'].iloc[i]
                    + params.iloc[3] * data_test['RV_M'].iloc[i]
                )
            result = pd.Series(result, index=data_predict.index)
            frame = pd.concat([data_predict, result], axis=1)
            frame.columns = ['RV_True', 'RV_Predict']
            frame.index = pd.to_datetime(data_test.index)
            frame['MSE'] = (frame['RV_True'] - frame['RV_Predict']) ** 2
            frame['MAE'] = np.abs(frame['RV_True'] - frame['RV_Predict'])
            frame['Qlike'] = frame['RV_True'] / frame['RV_Predict'] + np.log(frame['RV_Predict'])
            model_outputs[str(c).strip().lstrip('.')] = frame.copy()
        return model_outputs

    def run_loghar():
        model_outputs = {}
        import math
        import scipy as sc
        import time
        import os
        from scipy.stats import norm
        import pandas as pd
        import numpy as np
        from functools import partial
        from tqdm import tqdm as _tqdm
        tqdm = partial(_tqdm, bar_format="{bar} {percentage:3.0f}%", ncols=50)
        import statistics
        import statsmodels.formula.api as smf
        EPS_LOG = 1e-12
        EXP_CLIP_MAX = 50
        H_LIST = [22]
        SAMPLE_START = '2014-07-04'
        time_start = '2018/05/02'
        time_stop = '2022-02-25'

        def safe_log(s: pd.Series, eps: float=EPS_LOG) -> pd.Series:
            return np.log(s.clip(lower=eps))

        def safe_exp(x, clip_max: float=EXP_CLIP_MAX):
            return np.exp(np.clip(x, -clip_max, clip_max))

        def future_mean_level(rv_now: pd.Series, H: int) -> pd.Series:
            if H == 1:
                return rv_now
            ma = rv_now.rolling(window=H, center=False).mean()
            return ma.shift(-(H - 1))
        data = pd.read_csv(RV_FILE)
        num_cols = [c for c in data.columns if c != 'date']
        data[num_cols] = data[num_cols].apply(pd.to_numeric, errors='coerce')
        data = data.dropna()
        data['date'] = pd.to_datetime(data['date'], utc=True)
        data['date'] = data['date'].dt.tz_localize(None)
        data['date'] = data['date'].dt.normalize()
        data = data.set_index('date')
        for H in H_LIST:
            error = []
            for c in data.columns:
                data_RV = data[[c]].copy()
                data_RV.columns = ['RV_now']
                data_RV['RV_D_lvl'] = data_RV['RV_now'].shift(1)
                data_RV['RV_W_lvl'] = data_RV['RV_D_lvl'].rolling(5).mean()
                data_RV['RV_M_lvl'] = data_RV['RV_D_lvl'].rolling(22).mean()
                data_RV['RV_D'] = safe_log(data_RV['RV_D_lvl'])
                data_RV['RV_W'] = safe_log(data_RV['RV_W_lvl'])
                data_RV['RV_M'] = safe_log(data_RV['RV_M_lvl'])
                data_RV[f'RV_MA{H}_lvl'] = future_mean_level(data_RV['RV_now'], H)
                data_RV[f'RV_MA{H}'] = safe_log(data_RV[f'RV_MA{H}_lvl'])
                data_RV = data_RV.dropna()
                data_RV = data_RV.loc[SAMPLE_START:]
                data_test = data_RV.loc[time_start:time_stop].copy()
                if len(data_test) == 0:
                    error.append([np.nan, np.nan, np.nan])
                    tmp = pd.DataFrame(index=pd.DatetimeIndex([], name='date'))
                    tmp['RV_True'] = np.nan
                    tmp['RV_Predict'] = np.nan
                    tmp['MSE'] = np.nan
                    tmp['MAE'] = np.nan
                    tmp['Qlike'] = np.nan
                    model_outputs[str(c).strip().lstrip('.')] = tmp.copy()
                    continue
                y_true_log = data_test[f'RV_MA{H}'].copy()
                result_log = []
                n_test = len(data_RV.loc[time_start:time_stop])
                train_len = len(data_RV.loc[:time_start]) - 1
                for i in tqdm(range(n_test), desc=f'LogHAR {c} H{H}', leave=False):
                    data_rolling = data_RV.iloc[i:i + train_len].copy()
                    mod = smf.ols(f'RV_MA{H} ~ RV_D + RV_W + RV_M', data_rolling)
                    res = mod.fit()
                    params_new = res.params
                    r_predict_log = params_new.iloc[0] + params_new.iloc[1] * data_test['RV_D'].iloc[i] + params_new.iloc[2] * data_test['RV_W'].iloc[i] + params_new.iloc[3] * data_test['RV_M'].iloc[i]
                    result_log.append(r_predict_log)
                result_log = pd.Series(result_log, index=y_true_log.index)
                y_true = pd.Series(safe_exp(y_true_log.values), index=y_true_log.index)
                y_pred = pd.Series(safe_exp(result_log.values), index=result_log.index)
                SUM = pd.concat([y_true, y_pred], axis=1)
                SUM.columns = ['RV_True', 'RV_Predict']
                SUM_final = SUM.reset_index(drop=True)
                SUM_final.index = pd.to_datetime(data_test.index)
                SUM_final['MSE'] = (SUM_final['RV_True'] - SUM_final['RV_Predict']) ** 2
                SUM_final['MAE'] = np.abs(SUM_final['RV_True'] - SUM_final['RV_Predict'])
                pred_safe = SUM_final['RV_Predict'].clip(lower=EPS_LOG)
                true_safe = SUM_final['RV_True'].clip(lower=EPS_LOG)
                SUM_final['Qlike'] = true_safe / pred_safe + np.log(pred_safe)
                MSE = float(np.mean(SUM_final['MSE']))
                MAE = float(np.mean(SUM_final['MAE']))
                Theil = float(np.sqrt(np.mean((SUM_final['RV_True'] - SUM_final['RV_Predict']) ** 2)) / (np.sqrt(np.mean(SUM_final['RV_Predict'] ** 2)) + np.sqrt(np.mean(SUM_final['RV_True'] ** 2))))
                error.append([MSE, MAE, Theil])
                model_outputs[str(c).strip().lstrip('.')] = SUM_final.copy()
            ERROR = pd.DataFrame(error, columns=['MSE', 'MAE', 'Theil'], index=data.columns).T
        return model_outputs

    def run_loghar_ci():
        model_outputs = {}
        import os
        import numpy as np
        import pandas as pd
        import statsmodels.formula.api as smf
        from functools import partial
        from tqdm import tqdm as _tqdm
        tqdm = partial(_tqdm, bar_format="{bar} {percentage:3.0f}%", ncols=50)
        RET_PATH = RETURN_FILE
        RV_PATH = RV_FILE
        P_STORE_PATH = LOG_CENTRALITY_FILE
        fit_end_date = pd.to_datetime('2018-05-01')
        pred_start = pd.to_datetime('2018-05-02')
        pred_end = pd.to_datetime('2022-02-25')
        H_LIST = [22]
        FIT_START = 1000
        FIT_STEP = 100
        FIT_MAX = 1
        L_fit_list = [FIT_START + FIT_STEP * k for k in range(FIT_MAX)]
        NET_MIN = 700
        NET_MAX = 700
        NET_STEP = 50
        L_net_list = list(range(NET_MIN, NET_MAX + 1, NET_STEP))
        EPS_LOG = 1e-12
        EXP_CLIP_MAX = 50

        def _norm_date(x):
            return pd.to_datetime(x).tz_localize(None).normalize()

        def safe_log(s: pd.Series, eps: float=EPS_LOG) -> pd.Series:
            return np.log(s.clip(lower=eps))

        def safe_exp(x: pd.Series | np.ndarray, clip_max: float=EXP_CLIP_MAX):
            return np.exp(np.clip(x, -clip_max, clip_max))

        def make_future_mean(rv_now: pd.Series, H: int) -> pd.Series:
            if H == 1:
                return rv_now
            ma = rv_now.rolling(window=H, center=False).mean()
            return ma.shift(-(H - 1))

        def make_future_mean_log(rv_now: pd.Series, H: int) -> pd.Series:
            future_mean = make_future_mean(rv_now, H)
            return safe_log(future_mean)

        def read_returns(path):
            df = pd.read_csv(path)
            df = df.loc[:, ~df.columns.str.startswith('Unnamed')]
            df['date'] = pd.to_datetime(df['date'].astype(str).str.strip().str[:10], errors='coerce')
            df['date'] = df['date'].dt.tz_localize(None).dt.normalize()
            df = df.set_index('date').dropna(how='all')
            return df

        def read_rv_no_scale(path):
            data = pd.read_csv(path)
            if 'date' not in data.columns:
                raise RuntimeError("RV file must contain a 'date' column.")
            num_cols = [c for c in data.columns if c != 'date']
            data[num_cols] = data[num_cols].apply(pd.to_numeric, errors='coerce')
            data = data.dropna(how='all')
            data['date'] = pd.to_datetime(data['date'], utc=True, errors='coerce')
            data = data.dropna(subset=['date'])
            data['date'] = data['date'].dt.tz_localize(None)
            data['date'] = data['date'].dt.normalize()
            data = data.set_index('date')
            data = data.dropna(axis=1, how='all')
            return data
        df_ret = read_returns(RET_PATH)
        data_rv = read_rv_no_scale(RV_PATH)
        markets = [c for c in data_rv.columns if c in df_ret.columns]
        P_all = pd.read_csv(P_STORE_PATH)
        need_cols = ['date', 'market', 'ci_start_date', 'ci_end_date', 'L_net', 'P_corr', 'P_kendall', 'P_partial', 'P_tail']
        miss = [c for c in need_cols if c not in P_all.columns]
        if miss:
            raise RuntimeError(f'Model execution message{miss}Model execution message{list(P_all.columns)}')
        P_all['date'] = pd.to_datetime(P_all['date'], errors='coerce').dt.tz_localize(None).dt.normalize()
        P_all['ci_start_date'] = P_all['ci_start_date'].apply(_norm_date)
        P_all['ci_end_date'] = P_all['ci_end_date'].apply(_norm_date)
        P_all['L_net'] = pd.to_numeric(P_all['L_net'], errors='coerce').astype('Int64')
        P_all = P_all.dropna(subset=['date', 'market', 'ci_start_date', 'ci_end_date', 'L_net']).copy()
        P_cache = {}

        def load_P_from_csv(L_net: int, ci_start_date, ci_end_date):
            key = (int(L_net), _norm_date(ci_start_date), _norm_date(ci_end_date))
            if key in P_cache:
                return P_cache[key]
            (ln, cs, ce) = key
            sub = P_all[(P_all['L_net'] == ln) & (P_all['ci_start_date'] == cs) & (P_all['ci_end_date'] == ce)]
            if sub.empty:
                P_cache[key] = (None, None, None, None)
                return P_cache[key]

            def _pivot(col):
                df = sub.pivot_table(index='date', columns='market', values=col, aggfunc='mean')
                return df.sort_index()
            P_corr = _pivot('P_corr')
            P_kendall = _pivot('P_kendall')
            P_partial = _pivot('P_partial')
            P_tail = _pivot('P_tail')
            P_cache[key] = (P_corr, P_kendall, P_partial, P_tail)
            return P_cache[key]

        def forecast_one_market_oneH_logharci(data_rv_all: pd.DataFrame, P_corr: pd.DataFrame, P_kendall: pd.DataFrame, P_partial: pd.DataFrame, P_tail: pd.DataFrame, market: str, fit_start_date: pd.Timestamp, H: int):
            df = data_rv_all[market].to_frame('RV_now')
            df['RV_D_lvl'] = df['RV_now'].shift(1)
            df['RV_W_lvl'] = df['RV_D_lvl'].rolling(5).mean()
            df['RV_M_lvl'] = df['RV_D_lvl'].rolling(22).mean()
            df['RV_D'] = safe_log(df['RV_D_lvl'])
            df['RV_W'] = safe_log(df['RV_W_lvl'])
            df['RV_M'] = safe_log(df['RV_M_lvl'])
            df['Y'] = make_future_mean_log(df['RV_now'], H)
            ci = pd.DataFrame({'CI_corr': P_corr[market], 'CI_kendall': P_kendall[market], 'CI_partial': P_partial[market], 'CI_tail': P_tail[market]}).reindex(df.index)
            df = df.join(ci, how='inner').dropna()
            df['CI_corr_D'] = df['CI_corr'].shift(1)
            df['CI_kendall_D'] = df['CI_kendall'].shift(1)
            df['CI_partial_D'] = df['CI_partial'].shift(1)
            df['CI_tail_D'] = df['CI_tail'].shift(1)
            df = df.dropna()
            df = df.loc[fit_start_date:]
            data_test = df.loc[pred_start:pred_end].copy()
            if len(data_test) == 0:
                return (None, np.nan, np.nan, np.nan)
            window_len = len(df.loc[:pred_start]) - 1
            if window_len <= 0:
                return (None, np.nan, np.nan, np.nan)
            y_true_log = data_test['Y'].astype(float).copy()
            idx = y_true_log.index
            preds_log = []
            for i in tqdm(range(len(idx)), desc=f'LogHAR-CI {market} H{H}', leave=False):
                train = df.iloc[i:i + window_len].copy()
                train['_Y_'] = train['Y'].astype(float).values
                res = smf.ols('_Y_ ~ RV_D + RV_W + RV_M + CI_corr_D + CI_kendall_D + CI_partial_D + CI_tail_D', train).fit()
                p = res.params
                pred_log = p.iloc[0] + p.iloc[1] * data_test['RV_D'].iloc[i] + p.iloc[2] * data_test['RV_W'].iloc[i] + p.iloc[3] * data_test['RV_M'].iloc[i] + p.iloc[4] * data_test['CI_corr_D'].iloc[i] + p.iloc[5] * data_test['CI_kendall_D'].iloc[i] + p.iloc[6] * data_test['CI_partial_D'].iloc[i] + p.iloc[7] * data_test['CI_tail_D'].iloc[i]
                preds_log.append(float(pred_log))
            y_pred_log = pd.Series(preds_log, index=idx)
            y_true = pd.Series(safe_exp(y_true_log.values), index=idx)
            y_pred = pd.Series(safe_exp(y_pred_log.values), index=idx)
            mcs = pd.concat([y_true.rename('RV_True'), y_pred.rename('RV_Predict')], axis=1)
            mcs['MSE'] = (mcs['RV_True'] - mcs['RV_Predict']) ** 2
            mcs['MAE'] = np.abs(mcs['RV_True'] - mcs['RV_Predict'])
            mse = float(np.nanmean(mcs['MSE'].values))
            mae = float(np.nanmean(mcs['MAE'].values))
            denom = np.sqrt(np.nanmean(mcs['RV_Predict'].values ** 2)) + np.sqrt(np.nanmean(mcs['RV_True'].values ** 2))
            theil = float(np.sqrt(np.nanmean((mcs['RV_True'].values - mcs['RV_Predict'].values) ** 2)) / denom) if denom != 0 else np.nan
            return (mcs, mse, mae, theil)
        L_fit = L_fit_list[0]
        L_net = L_net_list[0]
        for H in H_LIST:
            error_by_market = {}
            for mkt in markets:
                if fit_end_date not in data_rv.index:
                    error_by_market[mkt] = [np.nan, np.nan, np.nan]
                    continue
                fit_end_pos = data_rv.index.get_loc(fit_end_date)
                fit_start_pos = fit_end_pos - (L_fit - 1)
                if fit_start_pos < 0:
                    error_by_market[mkt] = [np.nan, np.nan, np.nan]
                    continue
                fit_start_date = data_rv.index[fit_start_pos]
                ret_pos_end = df_ret.index.searchsorted(fit_start_date) - 1
                if ret_pos_end < 0:
                    error_by_market[mkt] = [np.nan, np.nan, np.nan]
                    continue
                net_end_date = df_ret.index[ret_pos_end]
                ci_start_date = net_end_date
                ci_end_date = pred_end
                (P_corr, P_kendall, P_partial, P_tail) = load_P_from_csv(L_net=L_net, ci_start_date=ci_start_date, ci_end_date=ci_end_date)
                if P_corr is None:
                    error_by_market[mkt] = [np.nan, np.nan, np.nan]
                    continue
                if mkt not in P_corr.columns:
                    error_by_market[mkt] = [np.nan, np.nan, np.nan]
                    continue
                (mcs, mse, mae, theil) = forecast_one_market_oneH_logharci(data_rv_all=data_rv, P_corr=P_corr, P_kendall=P_kendall, P_partial=P_partial, P_tail=P_tail, market=mkt, fit_start_date=fit_start_date, H=H)
                if mcs is None:
                    tmp = pd.DataFrame(index=pd.DatetimeIndex([], name='date'))
                    tmp['RV_True'] = np.nan
                    tmp['RV_Predict'] = np.nan
                    tmp['MSE'] = np.nan
                    tmp['MAE'] = np.nan
                    model_outputs[str(mkt).strip().lstrip('.')] = tmp.copy()
                else:
                    model_outputs[str(mkt).strip().lstrip('.')] = mcs.copy()
                error_by_market[mkt] = [mse, mae, theil]
            ERROR = pd.DataFrame(error_by_market, index=['MSE', 'MAE', 'Theil'])
        return model_outputs

    def run_hexpgl():
        model_outputs = {}
        import pandas as pd
        import numpy as np
        import statsmodels.formula.api as smf
        from functools import partial
        from tqdm import tqdm as _tqdm
        tqdm = partial(_tqdm, bar_format="{bar} {percentage:3.0f}%", ncols=50)
        import os

        def get_lambda(com):
            return np.log(1 + 1 / com)

        def calculate_exp_factor(series, com, window=500):
            lam = get_lambda(com)
            weights = np.exp(-np.arange(1, window + 1) * lam)
            weights /= weights.sum()
            return series.rolling(window).apply(lambda x: np.dot(x[::-1], weights), raw=True)
        data = pd.read_csv(RV_FILE)
        data['date'] = pd.to_datetime(data['date'])
        data = data.set_index('date')
        mean_rv_all = data.mean()
        global_signal = (data / mean_rv_all).mean(axis=1)
        exp_global_5 = calculate_exp_factor(global_signal, com=5)
        H_LIST = [22]
        com_list = [1, 5, 25, 125]
        time = '2018/05/02'
        time_stop = '2022-02-25'
        error_by_H = {H: [] for H in H_LIST}
        for c in data.columns:
            asset_mean = mean_rv_all[c]
            df_base = pd.DataFrame(index=data.index)
            df_base['RV_now'] = data[c]
            for j in com_list:
                df_base[f'ExpRV_{j}'] = calculate_exp_factor(df_base['RV_now'], com=j)
            df_base['ExpGlRV_5'] = exp_global_5 * asset_mean
            features = [f'ExpRV_{j}' for j in com_list] + ['ExpGlRV_5']
            for feat in features:
                df_base[f'{feat}_lag'] = df_base[feat].shift(1)
            reg_features = []
            for feat in features:
                col_name = f'diff_{feat}'
                df_base[col_name] = df_base[f'{feat}_lag'] - asset_mean
                reg_features.append(col_name)
            for H in H_LIST:
                df_model = df_base.copy()
                target_col = f'RV_MA{H}'
                if H == 1:
                    df_model[target_col] = df_model['RV_now']
                else:
                    df_model[target_col] = df_model['RV_now'].rolling(window=H).mean()
                    df_model[target_col] = df_model[target_col].shift(-H)
                df_model['target_diff'] = df_model[target_col] - asset_mean
                df_model = df_model.dropna()
                df_model = df_model.loc['2014-07-04':]
                data_test = df_model.loc[time:time_stop]
                data_predict = df_model.loc[time:time_stop][target_col]
                if len(data_test) == 0:
                    error_by_H[H].append([np.nan, np.nan, np.nan])
                    continue
                result = []
                for i in tqdm(range(len(df_model.loc[time:time_stop])), desc=f'{c}-H{H}'):
                    data_rolling = df_model.iloc[i:i + (len(df_model.loc[:time]) - 1)]
                    formula = 'target_diff ~ ' + ' + '.join(reg_features)
                    mod = smf.ols(formula, data_rolling)
                    res = mod.fit()
                    params_new = res.params
                    current_row = data_test.iloc[i]
                    pred_diff = params_new.iloc[0]
                    for (k, feat) in enumerate(reg_features, start=1):
                        pred_diff += params_new.iloc[k] * current_row[feat]
                    r_predict = pred_diff + asset_mean
                    result.append(r_predict)
                result = pd.Series(result, index=data_predict.index)
                SUM = pd.concat([data_predict, result], axis=1)
                SUM.columns = ['RV_True', 'RV_Predict']
                SUM_final = SUM
                SUM_final = SUM_final.reset_index(drop=True)
                SUM_final.index = pd.to_datetime(data_test.index)
                SUM_final['MSE'] = (SUM_final['RV_True'] - SUM_final['RV_Predict']) ** 2
                SUM_final['MAE'] = np.abs(SUM_final['RV_True'] - SUM_final['RV_Predict'])
                eps = 1e-12
                pred_clip = np.maximum(SUM_final['RV_Predict'], eps)
                SUM_final['Qlike'] = SUM_final['RV_True'] / pred_clip + np.log(pred_clip)
                MSE = np.mean((SUM_final['RV_True'] - SUM_final['RV_Predict']) ** 2)
                MAE = np.mean(np.abs(SUM_final['RV_True'] - SUM_final['RV_Predict']))
                Theil = np.sqrt(np.mean((SUM_final['RV_True'] - SUM_final['RV_Predict']) ** 2)) / (np.sqrt(np.mean(SUM_final['RV_Predict'] ** 2)) + np.sqrt(np.mean(SUM_final['RV_True'] ** 2)))
                error_by_H[H].append([MSE, MAE, Theil])
                model_outputs[str(c).strip().lstrip('.')] = SUM_final.copy()
        for H in H_LIST:
            ERROR = pd.DataFrame(error_by_H[H], columns=['MSE', 'MAE', 'Theil'], index=data.columns).T
        return model_outputs

    def run_mvf():
        model_outputs = {}
        import os
        import numpy as np
        import pandas as pd
        import statsmodels.formula.api as smf
        from functools import partial
        from tqdm import tqdm as _tqdm
        tqdm = partial(_tqdm, bar_format="{bar} {percentage:3.0f}%", ncols=50)
        H_LIST = [22]
        time_start = '2018/05/02'
        time_stop = '2022-02-25'
        EPS = 1e-12
        SCALE_10000 = False
        data = pd.read_csv(RV_FILE)
        if 'date' not in data.columns:
            raise RuntimeError("RV file must contain a 'date' column.")
        num_cols = data.columns.drop('date')
        data[num_cols] = data[num_cols].apply(pd.to_numeric, errors='coerce')
        if SCALE_10000:
            data[num_cols] = data[num_cols] * 10000
        data = data.dropna()
        data['date'] = pd.to_datetime(data['date'])
        data = data.set_index('date')
        data[data <= 0] = EPS
        crv_raw = data.mean(axis=1).clip(lower=EPS)
        xi_raw = data.divide(crv_raw, axis=0).clip(lower=EPS)

        def get_logmean_har_features(raw_series, h=22, eps=1e-12):
            raw = pd.Series(raw_series).astype(float).clip(lower=eps)
            df = pd.DataFrame(index=raw.index)
            df['D'] = np.log(raw.shift(1).clip(lower=eps))
            df['W'] = np.log(raw.shift(1).rolling(5).mean().clip(lower=eps))
            df['M'] = np.log(raw.shift(1).rolling(22).mean().clip(lower=eps))
            future_mean = raw.rolling(window=h).mean().shift(-(h - 1)).clip(lower=eps)
            df['Target'] = np.log(future_mean)
            return df.dropna()
        for H in H_LIST:
            df_log_crv = get_logmean_har_features(crv_raw, h=H, eps=EPS)
            df_log_crv = df_log_crv.loc['2014-07-04':]
            data_test_crv = df_log_crv.loc[time_start:time_stop]
            if data_test_crv.empty:
                raise RuntimeError(f'[H={H}] data_test_crv is empty. Check dates.')
            train_len_crv = len(df_log_crv.loc[:time_start]) - H
            if train_len_crv <= 10:
                raise RuntimeError(f'[H={H}] train_len_crv too small ({train_len_crv}). Check start date or reduce H.')
            crv_log_predictions = []
            for i in tqdm(range(len(data_test_crv)), desc=f'CRV-Log-HAR H={H}'):
                train_data = df_log_crv.iloc[i:i + train_len_crv]
                model = smf.ols('Target ~ D + W + M', train_data).fit()
                row = data_test_crv.iloc[i]
                pred = model.params['Intercept'] + model.params['D'] * row['D'] + model.params['W'] * row['W'] + model.params['M'] * row['M']
                crv_log_predictions.append(pred)
            pred_log_crv_series = pd.Series(crv_log_predictions, index=data_test_crv.index)
            error_rows = []
            for c in data.columns:
                df_log_xi = get_logmean_har_features(xi_raw[c], h=H, eps=EPS)
                df_log_xi = df_log_xi.loc['2014-07-04':]
                data_test_xi = df_log_xi.loc[time_start:time_stop]
                train_len_xi = len(df_log_xi.loc[:time_start]) - H
                if train_len_xi <= 10 or data_test_xi.empty:
                    error_rows.append([np.nan, np.nan, np.nan])
                    model_outputs[str(c).strip().lstrip('.')] = pd.DataFrame()
                    continue
                xi_log_predictions = []
                for i in tqdm(range(len(data_test_xi)), desc=f'{c}-xi-Log-HAR H={H}', leave=False):
                    train_data = df_log_xi.iloc[i:i + train_len_xi]
                    model = smf.ols('Target ~ D + W + M', train_data).fit()
                    row = data_test_xi.iloc[i]
                    pred = model.params['Intercept'] + model.params['D'] * row['D'] + model.params['W'] * row['W'] + model.params['M'] * row['M']
                    xi_log_predictions.append(pred)
                pred_log_xi_series = pd.Series(xi_log_predictions, index=data_test_xi.index)
                idx = pred_log_xi_series.index.intersection(pred_log_crv_series.index)
                if len(idx) == 0:
                    error_rows.append([np.nan, np.nan, np.nan])
                    model_outputs[str(c).strip().lstrip('.')] = pd.DataFrame()
                    continue
                pred_crv_mean = np.exp(pred_log_crv_series.loc[idx])
                pred_xi_mean = np.exp(pred_log_xi_series.loc[idx])
                rv_pred_mean = (pred_crv_mean * pred_xi_mean).clip(lower=EPS)
                rv_true_mean = data[c].rolling(H).mean().shift(-(H - 1)).loc[idx].clip(lower=EPS)
                out = pd.DataFrame({'RV_True': rv_true_mean, 'RV_Predict': rv_pred_mean}).dropna()
                if out.empty:
                    error_rows.append([np.nan, np.nan, np.nan])
                    model_outputs[str(c).strip().lstrip('.')] = out.copy()
                    continue
                out['MSE'] = (out['RV_True'] - out['RV_Predict']) ** 2
                out['MAE'] = np.abs(out['RV_True'] - out['RV_Predict'])
                out['Qlike'] = out['RV_True'] / out['RV_Predict'].clip(lower=EPS) + np.log(out['RV_Predict'].clip(lower=EPS))
                mse_val = float(out['MSE'].mean())
                mae_val = float(out['MAE'].mean())
                theil_val = float(np.sqrt(mse_val) / (np.sqrt((out['RV_Predict'] ** 2).mean()) + np.sqrt((out['RV_True'] ** 2).mean())))
                error_rows.append([mse_val, mae_val, theil_val])
                model_outputs[str(c).strip().lstrip('.')] = out.copy()
            ERROR_DF = pd.DataFrame(error_rows, columns=['MSE', 'MAE', 'Theil'], index=data.columns).T
        return model_outputs

    def run_multivariate():
        model_outputs = {}
        import os
        import numpy as np
        import pandas as pd
        from functools import partial
        from tqdm import tqdm as _tqdm
        tqdm = partial(_tqdm, bar_format="{bar} {percentage:3.0f}%", ncols=50)
        import statsmodels.api as sm
        EPS_LOG = 1e-12
        EXP_CLIP_MAX = 50
        H = 22
        SAMPLE_START = '2014-07-04'
        time_start = '2018/05/02'
        time_stop = '2022-02-25'

        def safe_log(x, eps: float=EPS_LOG):
            return np.log(np.clip(x, eps, None))

        def safe_exp(x, clip_max: float=EXP_CLIP_MAX):
            return np.exp(np.clip(x, -clip_max, clip_max))

        def future_mean_level(rv_now: pd.Series, H: int) -> pd.Series:
            if H == 1:
                return rv_now
            ma = rv_now.rolling(window=H, center=False).mean()
            return ma.shift(-(H - 1))
        data = pd.read_csv(RV_FILE)
        num_cols = [c for c in data.columns if c != 'date']
        data[num_cols] = data[num_cols].apply(pd.to_numeric, errors='coerce')
        data = data.dropna()
        data['date'] = pd.to_datetime(data['date'], utc=True)
        data['date'] = data['date'].dt.tz_localize(None)
        data['date'] = data['date'].dt.normalize()
        data = data.set_index('date').sort_index()
        data = data.fillna(method='ffill').dropna()
        data[data <= 0] = EPS_LOG
        markets = list(data.columns)
        rv_now = data.copy()
        rv_d_lvl = rv_now.shift(1)
        rv_w_lvl = rv_d_lvl.rolling(5).mean()
        rv_m_lvl = rv_d_lvl.rolling(22).mean()
        rv_d = safe_log(rv_d_lvl)
        rv_w = safe_log(rv_w_lvl)
        rv_m = safe_log(rv_m_lvl)
        X_all = pd.DataFrame(index=rv_now.index)
        for m in markets:
            X_all[f'RV_D_{m}'] = rv_d[m]
            X_all[f'RV_W_{m}'] = rv_w[m]
            X_all[f'RV_M_{m}'] = rv_m[m]
        X_cols = list(X_all.columns)
        error_rows = []
        time_start_dt = pd.to_datetime(time_start)
        time_stop_dt = pd.to_datetime(time_stop)
        for target in markets:
            y_lvl = future_mean_level(rv_now[target], H)
            y_log = pd.Series(safe_log(y_lvl), index=rv_now.index, name=f'RV_MA{H}')
            df = pd.concat([y_log, X_all], axis=1).dropna()
            df = df.loc[SAMPLE_START:].copy()
            data_test = df.loc[time_start_dt:time_stop_dt].copy()
            if len(data_test) == 0:
                tmp = pd.DataFrame(index=pd.DatetimeIndex([], name='date'))
                tmp['RV_True'] = np.nan
                tmp['RV_Predict'] = np.nan
                tmp['MSE'] = np.nan
                tmp['MAE'] = np.nan
                tmp['Qlike'] = np.nan
                model_outputs[str(target).strip().lstrip('.')] = tmp.copy()
                error_rows.append([np.nan, np.nan, np.nan])
                continue
            y_true_log = data_test[f'RV_MA{H}'].copy()
            test_index = data_test.index
            start_pos = df.index.get_loc(test_index[0])
            if start_pos <= H:
                tmp = pd.DataFrame(index=pd.DatetimeIndex([], name='date'))
                tmp['RV_True'] = np.nan
                tmp['RV_Predict'] = np.nan
                tmp['MSE'] = np.nan
                tmp['MAE'] = np.nan
                tmp['Qlike'] = np.nan
                model_outputs[str(target).strip().lstrip('.')] = tmp.copy()
                error_rows.append([np.nan, np.nan, np.nan])
                continue
            preds_log = []
            for i, forecast_date in enumerate(tqdm(test_index, desc=f'MV-LogHAR {target} H{H}', leave=False)):
                train_slice = df.iloc[:start_pos + i].copy()
                train_slice = train_slice.iloc[:-(H - 1)]
                y_tr = train_slice[f'RV_MA{H}'].values
                X_tr = train_slice[X_cols].values
                X_tr = sm.add_constant(X_tr, has_constant='add')
                res = sm.OLS(y_tr, X_tr).fit()
                x_row = data_test.loc[[forecast_date], X_cols].values
                x_row = sm.add_constant(x_row, has_constant='add')
                pred_i = float(res.predict(x_row)[0])
                preds_log.append(pred_i)
            preds_log = pd.Series(preds_log, index=y_true_log.index)
            y_true = pd.Series(safe_exp(y_true_log.values), index=y_true_log.index)
            y_pred = pd.Series(safe_exp(preds_log.values), index=preds_log.index)
            SUM = pd.concat([y_true, y_pred], axis=1)
            SUM.columns = ['RV_True', 'RV_Predict']
            SUM.index.name = 'date'
            SUM['MSE'] = (SUM['RV_True'] - SUM['RV_Predict']) ** 2
            SUM['MAE'] = np.abs(SUM['RV_True'] - SUM['RV_Predict'])
            pred_safe = SUM['RV_Predict'].clip(lower=EPS_LOG)
            true_safe = SUM['RV_True'].clip(lower=EPS_LOG)
            SUM['Qlike'] = true_safe / pred_safe + np.log(pred_safe)
            MSE = float(SUM['MSE'].mean())
            MAE = float(SUM['MAE'].mean())
            Theil = float(np.sqrt(np.mean((SUM['RV_True'] - SUM['RV_Predict']) ** 2)) / (np.sqrt(np.mean(SUM['RV_Predict'] ** 2)) + np.sqrt(np.mean(SUM['RV_True'] ** 2))))
            error_rows.append([MSE, MAE, Theil])
            model_outputs[str(target).strip().lstrip('.')] = SUM.copy()
        ERROR = pd.DataFrame(error_rows, columns=['MSE', 'MAE', 'Theil'], index=markets).T
        return model_outputs

    def calc_mse(y_true: np.ndarray, y_pred: np.ndarray) -> dict:
        y_true = np.asarray(y_true, dtype=float)
        y_pred = np.asarray(y_pred, dtype=float)
        mask = np.isfinite(y_true) & np.isfinite(y_pred)
        yt = y_true[mask]
        yp = y_pred[mask]
        if len(yt) == 0:
            return {'MSE': np.nan}
        err = yp - yt
        mse = float(np.mean(err ** 2))
        return {'MSE': mse}

    def qlike_with_shared_drop(y_true: pd.Series, y_pred: pd.Series, drop_dates: set) -> float:
        df = pd.concat([y_true.rename('y'), y_pred.rename('yhat')], axis=1).dropna()
        if df.empty:
            return np.nan
        if drop_dates:
            df = df.loc[~df.index.isin(drop_dates)]
        if df.empty:
            return np.nan
        y = df['y'].astype(float).values
        yhat = df['yhat'].astype(float).values
        ok_pos = (y > EPS) & (yhat > EPS)
        if not np.any(ok_pos):
            return np.nan
        ratio = y[ok_pos] / yhat[ok_pos]
        ok_ratio = np.isfinite(ratio) & (ratio >= RATIO_MIN) & (ratio <= RATIO_MAX)
        if not np.any(ok_ratio):
            return np.nan
        ratio = ratio[ok_ratio]
        q = ratio - np.log(ratio) - 1.0
        ok_q = np.isfinite(q) & (q >= 0.0) & (q <= QLIKE_CAP)
        if not np.any(ok_q):
            return np.nan
        return float(np.mean(q[ok_q]))

    def add_summary_rows(df: pd.DataFrame) -> pd.DataFrame:
        mean_row = df.mean(axis=0, skipna=True)
        q1_row = df.quantile(0.25, axis=0, interpolation='linear')
        q2_row = df.quantile(0.5, axis=0, interpolation='linear')
        q3_row = df.quantile(0.75, axis=0, interpolation='linear')
        df2 = df.copy()
        df2.loc['Mean'] = mean_row
        df2.loc['Q1'] = q1_row
        df2.loc['Q2'] = q2_row
        df2.loc['Q3'] = q3_row
        return df2

    def prepare_mcs_matrix_relaxed(sum_pre: pd.DataFrame, min_valid_per_model: int=10, min_rows_for_mcs: int=5):
        info = {'initial_shape': tuple(sum_pre.shape), 'dropped_low_valid_models': [], 'dropped_const_models': [], 'final_shape': None, 'reason': None}
        if sum_pre.empty or sum_pre.shape[1] == 0:
            info['reason'] = 'The loss matrix is empty'
            return (None, info)
        valid_counts = sum_pre.notna().sum(axis=0)
        keep_cols = valid_counts[valid_counts >= min_valid_per_model].index.tolist()
        info['dropped_low_valid_models'] = [c for c in sum_pre.columns if c not in keep_cols]
        if len(keep_cols) < 2:
            info['reason'] = f'Fewer than two models with at least {min_valid_per_model} valid observations'
            return (None, info)
        x = sum_pre[keep_cols].copy()
        x = x.dropna(axis=0, how='any')
        if x.empty:
            info['reason'] = 'No common valid observations'
            return (None, info)
        if x.shape[1] < 2:
            info['reason'] = 'Fewer than two models after dropping missing observations'
            return (None, info)
        if x.shape[0] < min_rows_for_mcs:
            info['reason'] = f'Insufficient common observations ({x.shape[0]} < {min_rows_for_mcs})'
            return (None, info)
        variances = x.var(axis=0, skipna=True)
        non_const_cols = variances[variances > 0].index.tolist()
        info['dropped_const_models'] = [c for c in x.columns if c not in non_const_cols]
        x = x[non_const_cols]
        if x.shape[1] < 2:
            info['reason'] = 'Fewer than two models after removing constant columns'
            return (None, info)
        if x.shape[0] < min_rows_for_mcs:
            info['reason'] = f'Insufficient common observations after removing constant columns ({x.shape[0]} < {min_rows_for_mcs})'
            return (None, info)
        info['final_shape'] = tuple(x.shape)
        return (x, info)

    def choose_one_model_to_drop_for_retry(raw_sum_pre: pd.DataFrame, current_cols: list):
        if len(current_cols) <= 2:
            return None
        sub = raw_sum_pre[current_cols].copy()
        valid_counts = sub.notna().sum(axis=0)
        min_valid = valid_counts.min()
        candidates = valid_counts[valid_counts == min_valid].index.tolist()
        if len(candidates) == 1:
            return candidates[0]
        var_map = {}
        for c in candidates:
            s = pd.to_numeric(sub[c], errors='coerce').dropna()
            if len(s) <= 1:
                var_map[c] = -1.0
            else:
                v = float(s.var())
                if np.isnan(v):
                    v = -1.0
                var_map[c] = v
        drop_col = sorted(var_map.items(), key=lambda kv: (kv[1], kv[0]))[0][0]
        return drop_col

    def run_mcs_with_retry(sum_pre: pd.DataFrame, country_code: str, min_valid_per_model: int, min_rows_for_mcs: int, max_retry_drop: int, debug_logs: list):
        details = {'country': country_code, 'retry_count': 0, 'dropped_in_retry': [], 'fail_reason': None, 'final_shape': None, 'final_models': None}
        if sum_pre.empty or sum_pre.shape[1] < 2:
            details['fail_reason'] = 'Fewer than two available models'
            return (False, None, None, None, details)
        current_cols = list(sum_pre.columns)
        for _ in range(max_retry_drop + 1):
            tmp_sum = sum_pre[current_cols].copy()
            (mcs_df, prep_info) = prepare_mcs_matrix_relaxed(tmp_sum, min_valid_per_model=min_valid_per_model, min_rows_for_mcs=min_rows_for_mcs)
            if mcs_df is None:
                if len(current_cols) <= 2:
                    details['fail_reason'] = f"Preparation failed with too few models to retry: {prep_info.get('reason')}"
                    return (False, None, None, None, details)
                drop_col = choose_one_model_to_drop_for_retry(sum_pre, current_cols)
                if drop_col is None:
                    details['fail_reason'] = f"Preparation failed and no model can be removed: {prep_info.get('reason')}"
                    return (False, None, None, None, details)
                current_cols.remove(drop_col)
                details['retry_count'] += 1
                details['dropped_in_retry'].append(drop_col)
                debug_logs.append(f"[{country_code}] Preparation failed; retry without model: {drop_col} | Reason: {prep_info.get('reason')}")
                continue
            try:
                mcs = MCS(losses=mcs_df, size=0.9, method='R', reps=20000, block_size=5)
                mcs.compute()
                details['final_shape'] = tuple(mcs_df.shape)
                details['final_models'] = list(mcs_df.columns)
                return (True, mcs.pvalues, mcs.included, mcs_df, details)
            except Exception as e:
                if len(current_cols) <= 2:
                    details['fail_reason'] = f'MCS failed with too few models to retry: {str(e)}'
                    return (False, None, None, None, details)
                drop_col = choose_one_model_to_drop_for_retry(sum_pre, current_cols)
                if drop_col is None:
                    details['fail_reason'] = f'MCS failed and no model can be removed: {str(e)}'
                    return (False, None, None, None, details)
                current_cols.remove(drop_col)
                details['retry_count'] += 1
                details['dropped_in_retry'].append(drop_col)
                debug_logs.append(f'[{country_code}] MCS failed; retry without model: {drop_col} | Error: {str(e)}')
                continue
        details['fail_reason'] = f'Maximum retry count exceeded: {max_retry_drop}'
        return (False, None, None, None, details)

    def qlike_series_with_rules(y_true: np.ndarray, y_pred: np.ndarray) -> np.ndarray:
        y_true = np.asarray(y_true, dtype=float)
        y_pred = np.asarray(y_pred, dtype=float)
        q = np.full(shape=len(y_true), fill_value=np.nan, dtype=float)
        mask = np.isfinite(y_true) & np.isfinite(y_pred)
        if not np.any(mask):
            return q
        yt = y_true[mask]
        yp = y_pred[mask]
        ok_pos = (yt > EPS) & (yp > EPS)
        ratio = np.full_like(yt, np.nan, dtype=float)
        ratio[ok_pos] = yt[ok_pos] / yp[ok_pos]
        ok_ratio = np.isfinite(ratio) & (ratio >= RATIO_MIN) & (ratio <= RATIO_MAX)
        q_sub = np.full_like(yt, np.nan, dtype=float)
        q_sub[ok_ratio] = ratio[ok_ratio] - np.log(ratio[ok_ratio]) - 1.0
        ok_q = np.isfinite(q_sub) & (q_sub >= 0.0) & (q_sub <= QLIKE_CAP)
        q_masked = np.full_like(yt, np.nan, dtype=float)
        q_masked[ok_q] = q_sub[ok_q]
        q[mask] = q_masked
        return q

    MODEL_RUNNERS = [
        ("LogHAR", "LogHAR", run_loghar),
        ("LogHAR-CI", "LogHAR-CI", run_loghar_ci),
        ("HExp-GlRV", "HExpGl", run_hexpgl),
        ("MVF-Log", "MVF", run_mvf),
        ("Multivariate-LogHAR", "Multivariate", run_multivariate),
    ]

    AUXILIARY_RUNNERS = [
        ("HAR-RV", "HAR", run_har),
    ]

    def load_prediction_pair(frame):
        dates = pd.to_datetime(frame.index, errors="coerce")
        y_true = pd.Series(pd.to_numeric(frame["RV_True"], errors="coerce").to_numpy(), index=dates)
        y_pred = pd.Series(pd.to_numeric(frame["RV_Predict"], errors="coerce").to_numpy(), index=dates)
        valid = dates.notna() & np.isfinite(y_true.to_numpy()) & np.isfinite(y_pred.to_numpy())
        y_true, y_pred = y_true[valid], y_pred[valid]
        if len(y_true) and np.mean(y_true.to_numpy() > THRESH) > 0.5:
            y_true = y_true * SCALE_IF_MOST_GT_THRESH
            y_pred = y_pred * SCALE_IF_MOST_GT_THRESH
        y_true = y_true[~y_true.index.duplicated(keep="last")].sort_index()
        y_pred = y_pred[~y_pred.index.duplicated(keep="last")].sort_index()
        return y_true, y_pred

    def summarize_losses(predictions, shared_drop_dates):
        models = [key for key, _, _ in MODEL_RUNNERS]
        all_models = [key for key, _, _ in AUXILIARY_RUNNERS] + models
        markets = sorted(predictions["LogHAR"])
        cache = {market: {model: load_prediction_pair(predictions[model][market]) for model in all_models} for market in markets}
        raw = {metric: pd.DataFrame(index=markets, columns=all_models, dtype=float) for metric in ("MSE", "QLIKE")}
        for market in markets:
            for model in all_models:
                y_true, y_pred = cache[market][model]
                raw["MSE"].loc[market, model] = calc_mse(y_true.to_numpy(), y_pred.to_numpy())["MSE"]
                raw["QLIKE"].loc[market, model] = qlike_with_shared_drop(y_true, y_pred, shared_drop_dates[market])
        summaries = {}
        for metric, values in raw.items():
            values = values.dropna(how="all").dropna(subset=["HAR-RV", "LogHAR"], how="any")
            har_relative = values.div(values["HAR-RV"].replace(0.0, np.nan), axis=0)
            if np.any(np.isfinite(har_relative.to_numpy()) & (np.abs(har_relative.to_numpy()) > POST_SCALE_THRESHOLD)):
                warnings.warn("The original relative-loss scaling rule was triggered for " + metric, RuntimeWarning)
                har_relative = har_relative.where(~(np.abs(har_relative) > POST_SCALE_THRESHOLD), har_relative * POST_SCALE_FACTOR)
            statistics = add_summary_rows(har_relative).loc[["Mean", "Q1", "Q2"]]
            normalized = statistics.div(statistics["LogHAR"].replace(0.0, np.nan), axis=0)
            summaries[metric] = normalized[models].T
        return summaries, shared_drop_dates

    def qlike_loss_series(frame, drop_dates):
        data = pd.DataFrame({
            "date": pd.to_datetime(frame.index, errors="coerce"),
            "y_true": pd.to_numeric(frame["RV_True"], errors="coerce").to_numpy(),
            "y_pred": pd.to_numeric(frame["RV_Predict"], errors="coerce").to_numpy(),
        }).dropna(subset=["date"])
        if data["date"].dt.tz is not None:
            data["date"] = data["date"].dt.tz_localize(None)
        data["date"] = data["date"].dt.normalize()
        data = data.sort_values("date").drop_duplicates(subset=["date"], keep="last").reset_index(drop=True)
        data = data.loc[~data["date"].isin(drop_dates)]
        return pd.Series(qlike_series_with_rules(data["y_true"].to_numpy(), data["y_pred"].to_numpy()), index=data["date"])

    def count_mcs_memberships(predictions, shared_drop_dates):
        models = [key for key, _, _ in MODEL_RUNNERS]
        markets = sorted(predictions["LogHAR"])
        counts = {metric: dict.fromkeys(models, 0) for metric in ("MSE", "QLIKE")}
        for metric in counts:
            for market in tqdm(markets, desc="MCS " + metric):
                losses = pd.DataFrame()
                for model in models:
                    frame = predictions[model][market]
                    if metric == "MSE":
                        series = pd.to_numeric(frame["MSE"], errors="coerce").reset_index(drop=True)
                    else:
                        series = qlike_loss_series(frame, shared_drop_dates[market])
                    if series.notna().sum():
                        losses[model] = series
                logs = []
                success, _, included, used, details = run_mcs_with_retry(losses, market, 10, 5, 50, logs)
                if not success:
                    raise RuntimeError(f"MCS {metric} failed for {market}: {details['fail_reason']}")
                omitted = [model for model in models if model not in used.columns]
                if omitted:
                    warnings.warn(f"MCS {metric}, {market}: excluded from the test: {', '.join(omitted)}", RuntimeWarning)
                for model in included:
                    counts[metric][model] += 1
        return counts

    def save_table(summaries, counts, output_path=OUTPUT_FILE, bold_mode="loghar-ci"):
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "Table 5"
        sheet.merge_cells("A1:E1")
        sheet.merge_cells("G1:K1")
        sheet["A1"] = "MSE"
        sheet["G1"] = "QLIKE"
        for start in (1, 7):
            for offset, label in enumerate(("Model", "Mean", "Q1", "Median", "MCS(#)")):
                sheet.cell(2, start + offset, label)
        for row, (key, label, _) in enumerate(MODEL_RUNNERS, start=3):
            for start, metric in ((1, "MSE"), (7, "QLIKE")):
                sheet.cell(row, start, label)
                for offset, statistic in enumerate(("Mean", "Q1", "Q2"), start=1):
                    value = float(summaries[metric].loc[key, statistic])
                    if not np.isfinite(value):
                        raise ValueError(f"Cannot export a missing {metric} {statistic} for {key}")
                    sheet.cell(row, start + offset, value).number_format = "0.000"
                sheet.cell(row, start + 4, int(counts[metric][key])).number_format = "0"
        thin = Side(style="thin", color="000000")
        last_row = 2 + len(MODEL_RUNNERS)
        for row in sheet:
            for cell in row:
                cell.font = Font(name="Times New Roman", size=11, bold=False)
                cell.alignment = Alignment(horizontal="left" if cell.column in (1, 7) and cell.row > 1 else "center", vertical="center")
                if cell.column != 6:
                    if cell.row == 1:
                        cell.border = Border(top=thin, bottom=thin)
                    elif cell.row in (2, last_row):
                        cell.border = Border(bottom=thin)
        for start, metric in ((1, "MSE"), (7, "QLIKE")):
            for row, (key, _, _) in enumerate(MODEL_RUNNERS, start=3):
                for offset, stat in enumerate(("Mean", "Q1", "Q2", "MCS(#)"), start=1):
                    bold = bold_mode == "loghar-ci" and key == "LogHAR-CI"
                    if bold_mode == "best":
                        bold = counts[metric][key] == max(counts[metric].values()) if offset == 4 else np.isclose(summaries[metric].loc[key, stat], summaries[metric][stat].min(), rtol=0, atol=1e-12)
                    if bold:
                        sheet.cell(row, start + offset).font = Font(name="Times New Roman", size=11, bold=True)
        for column in ("A", "G"):
            sheet.column_dimensions[column].width = 22
        for column in ("B", "C", "D", "E", "H", "I", "J", "K"):
            sheet.column_dimensions[column].width = 12
        sheet.column_dimensions["F"].width = 3
        sheet.freeze_panes = "B3"
        sheet.sheet_view.showGridLines = False
        sheet.print_options.horizontalCentered = True
        sheet.page_setup.orientation = "landscape"
        sheet.page_setup.paperSize = sheet.PAPERSIZE_A4
        sheet.page_setup.fitToWidth = 1
        sheet.page_setup.fitToHeight = 1
        sheet.sheet_properties.pageSetUpPr.fitToPage = True
        sheet.print_area = f"A1:K{last_row}"
        output_path = Path(output_path)
        output_path.parent.mkdir(parents=True, exist_ok=True)
        workbook.save(output_path)

    def check_inputs():
        for path in (RV_FILE, RETURN_FILE, LOG_CENTRALITY_FILE):
            if not path.is_file():
                raise FileNotFoundError(f"Required input not found: {path}")
        markets = [str(c).strip().lstrip(".") for c in pd.read_csv(RV_FILE, nrows=0).columns if c != "date"]
        if len(markets) != 18 or len(set(markets)) != 18:
            raise ValueError("Table 5 requires 18 distinct markets in RV.csv")
        centrality = pd.read_csv(LOG_CENTRALITY_FILE, usecols=["market", "L_net"])
        centrality["market"] = centrality["market"].astype(str).str.strip().str.lstrip(".")
        windows = set(pd.to_numeric(centrality["L_net"], errors="coerce").dropna().astype(int))
        if 700 not in windows:
            raise ValueError("P_log.csv does not contain the 700-day network window required by LogHAR-CI")
        if set(centrality["market"].dropna()) != set(markets):
            raise ValueError("P_log.csv markets do not match RV.csv")
        return set(markets)

    return SimpleNamespace(**locals())


def build_table_06():
    from pathlib import Path

    import argparse

    import warnings

    import numpy as np

    import pandas as pd

    from functools import partial

    from tqdm import tqdm as _tqdm

    tqdm = partial(_tqdm, bar_format="{bar} {percentage:3.0f}%", ncols=50)

    from arch.bootstrap import MCS

    from openpyxl import Workbook

    from openpyxl.styles import Font, Alignment, Border, Side

    ROOT = Path(__file__).resolve().parents[1]

    RV_FILE = ROOT / "data/raw/RV.csv"

    RETURN_FILE = ROOT / "data/raw/Return.csv"

    CENTRALITY_FILE = ROOT / "data/processed/centrality/P.csv"

    OUTPUT_FILE = ROOT / "results/generated/Table_06_category3_forecast_performance.xlsx"

    THRESH = 1e-3

    SCALE_IF_MOST_GT_THRESH = 1e-4

    EPS = 1e-12

    RATIO_MIN = 1e-6

    RATIO_MAX = 1e6

    QLIKE_CAP = 100.0

    POST_SCALE_THRESHOLD = 1e4

    POST_SCALE_FACTOR = 1e-8

    def run_har():
        model_outputs = {}
        import math
        import scipy as sc
        import time
        import os
        from scipy.stats import norm
        import pandas as pd
        import numpy as np
        from functools import partial
        from tqdm import tqdm as _tqdm
        tqdm = partial(_tqdm, bar_format="{bar} {percentage:3.0f}%", ncols=50)
        import statistics
        import statsmodels.formula.api as smf
        import os
        data = pd.read_csv(RV_FILE)
        num_cols = data.columns.drop('date')
        data[num_cols] = data[num_cols].apply(pd.to_numeric, errors='coerce')
        data = data.dropna()
        data['date'] = pd.to_datetime(data['date'], utc=True)
        data['date'] = data['date'].dt.tz_localize(None)
        data['date'] = data['date'].dt.normalize()
        data = data.set_index('date')
        error = []
        H = 22
        for c in data.columns:
            data_RV = data['%s' % c]
            data_RV = pd.DataFrame(data_RV)
            data_RV.columns = ['RV_now']
            data_RV['RV_D'] = data_RV['RV_now'].shift(1)
            data_RV['RV_W'] = data_RV['RV_D'].rolling(5).mean()
            data_RV['RV_M'] = data_RV['RV_D'].rolling(22).mean()
            data_RV['RV_MA22'] = data_RV['RV_now'].rolling(window=22, center=False).mean()
            data_RV['RV_MA22'] = data_RV['RV_MA22'].shift(-21)
            data_RV = data_RV.dropna()
            data_RV = data_RV.loc['2014-07-04':]
            time = '2018/05/02'
            time_stop = '2022-02-25'
            Loss_ALL = []
            data_test = data_RV.loc[time:time_stop]
            data_predict = data_RV.loc[time:time_stop]['RV_MA22']
            result = []
            for i in tqdm(range(len(data_RV.loc[time:time_stop]))):
                data_rolling = data_RV.iloc[i:i + (len(data_RV.loc[:time]) - 1)]
                mod = smf.ols('RV_MA22 ~ RV_D + RV_W + RV_M', data_rolling)
                res = mod.fit()
                params_new = res.params
                r_predict = params_new.iloc[0] + params_new.iloc[1] * data_test['RV_D'].iloc[i] + params_new.iloc[2] * data_test['RV_W'].iloc[i] + params_new.iloc[3] * data_test['RV_M'].iloc[i]
                result.append(r_predict)
            result = pd.Series(result, index=data_predict.index)
            SUM = pd.concat([data_predict, result], axis=1)
            SUM.columns = ['RV_True', 'RV_Predict']
            SUM_final = SUM
            SUM_final = SUM_final.reset_index(drop=True)
            SUM_final.index = pd.to_datetime(data_test.index)
            SUM_final['MSE'] = (SUM_final['RV_True'] - SUM_final['RV_Predict']) ** 2
            SUM_final['MAE'] = np.abs(SUM_final['RV_True'] - SUM_final['RV_Predict'])
            SUM_final['Qlike'] = SUM_final['RV_True'] / SUM_final['RV_Predict'] + np.log(SUM_final['RV_Predict'])
            MSE = np.mean((SUM_final['RV_True'] - SUM_final['RV_Predict']) ** 2)
            MAE = np.mean(np.abs(SUM_final['RV_True'] - SUM_final['RV_Predict']))
            Theil = np.sqrt(np.mean((SUM_final['RV_True'] - SUM_final['RV_Predict']) ** 2)) / (np.sqrt(np.mean(SUM_final['RV_Predict'] ** 2)) + np.sqrt(np.mean(SUM_final['RV_True'] ** 2)))
            Loss_ALL = [MSE, MAE, Theil]
            error.append(Loss_ALL)
            model_outputs[str(c).strip().lstrip('.')] = SUM_final.copy()
        ERROR = pd.DataFrame(error, columns=['MSE', 'MAE', 'Theil'], index=data.columns).T
        return model_outputs

    def run_har_ci():
        model_outputs = {}
        import os
        import numpy as np
        import pandas as pd
        import statsmodels.formula.api as smf
        from functools import partial
        from tqdm import tqdm as _tqdm
        tqdm = partial(_tqdm, bar_format="{bar} {percentage:3.0f}%", ncols=50)
        RET_PATH = RETURN_FILE
        RV_PATH = RV_FILE
        P_STORE_PATH = CENTRALITY_FILE
        PLOT = False
        fit_end_date = pd.to_datetime('2018-05-01')
        pred_start = pd.to_datetime('2018-05-02')
        pred_end = pd.to_datetime('2022-02-25')
        H_LIST = [22]
        FIT_START = 1000
        FIT_STEP = 100
        FIT_MAX = 1
        L_fit_list = [FIT_START + FIT_STEP * k for k in range(FIT_MAX)]
        NET_MIN = 1000
        NET_MAX = 1000
        NET_STEP = 50
        L_net_list = list(range(NET_MIN, NET_MAX + 1, NET_STEP))

        def _norm_date(x):
            return pd.to_datetime(x).tz_localize(None).normalize()

        def make_future_mean(rv_now: pd.Series, H: int) -> pd.Series:
            if H == 1:
                return rv_now
            ma = rv_now.rolling(window=H, center=False).mean()
            return ma.shift(-(H - 1))

        def read_returns(path):
            df = pd.read_csv(path)
            df = df.loc[:, ~df.columns.str.startswith('Unnamed')]
            df['date'] = pd.to_datetime(df['date'].astype(str).str.strip().str[:10], errors='coerce')
            df['date'] = df['date'].dt.tz_localize(None).dt.normalize()
            df = df.set_index('date').dropna(how='all')
            return df

        def read_rv_no_scale(path):
            data = pd.read_csv(path)
            if 'date' not in data.columns:
                raise RuntimeError("RV file must contain a 'date' column.")
            num_cols = [c for c in data.columns if c != 'date']
            data[num_cols] = data[num_cols].apply(pd.to_numeric, errors='coerce')
            data = data.dropna(how='all')
            data['date'] = pd.to_datetime(data['date'], utc=True, errors='coerce')
            data = data.dropna(subset=['date'])
            data['date'] = data['date'].dt.tz_localize(None)
            data['date'] = data['date'].dt.normalize()
            data = data.set_index('date')
            data = data.dropna(axis=1, how='all')
            return data
        df_ret = read_returns(RET_PATH)
        data_rv = read_rv_no_scale(RV_PATH)
        markets = [c for c in data_rv.columns if c in df_ret.columns]
        P_all = pd.read_csv(P_STORE_PATH)
        need_cols = ['date', 'market', 'ci_start_date', 'ci_end_date', 'L_net', 'P_corr', 'P_kendall', 'P_partial', 'P_tail']
        miss = [c for c in need_cols if c not in P_all.columns]
        if miss:
            raise RuntimeError(f'P.csv is missing required columns: {miss}; available columns={list(P_all.columns)}')
        P_all['date'] = pd.to_datetime(P_all['date'], errors='coerce').dt.tz_localize(None).dt.normalize()
        P_all['ci_start_date'] = P_all['ci_start_date'].apply(_norm_date)
        P_all['ci_end_date'] = P_all['ci_end_date'].apply(_norm_date)
        P_all['L_net'] = pd.to_numeric(P_all['L_net'], errors='coerce').astype('Int64')
        P_all = P_all.dropna(subset=['date', 'market', 'ci_start_date', 'ci_end_date', 'L_net']).copy()
        P_cache = {}

        def load_P_from_csv(L_net: int, ci_start_date, ci_end_date):
            key = (int(L_net), _norm_date(ci_start_date), _norm_date(ci_end_date))
            if key in P_cache:
                return P_cache[key]
            (ln, cs, ce) = key
            sub = P_all[(P_all['L_net'] == ln) & (P_all['ci_start_date'] == cs) & (P_all['ci_end_date'] == ce)]
            if sub.empty:
                P_cache[key] = (None, None, None, None)
                return P_cache[key]

            def _pivot(col):
                df = sub.pivot_table(index='date', columns='market', values=col, aggfunc='mean')
                return df.sort_index()
            P_corr = _pivot('P_corr')
            P_kendall = _pivot('P_kendall')
            P_partial = _pivot('P_partial')
            P_tail = _pivot('P_tail')
            P_cache[key] = (P_corr, P_kendall, P_partial, P_tail)
            return P_cache[key]

        def forecast_one_market_oneH_harci(data_rv_all: pd.DataFrame, P_corr: pd.DataFrame, P_kendall: pd.DataFrame, P_partial: pd.DataFrame, P_tail: pd.DataFrame, market: str, fit_start_date: pd.Timestamp, H: int):
            df = data_rv_all[market].to_frame('RV_now')
            df['RV_D'] = df['RV_now'].shift(1)
            df['RV_W'] = df['RV_D'].rolling(5).mean()
            df['RV_M'] = df['RV_D'].rolling(22).mean()
            df['Y'] = make_future_mean(df['RV_now'], H)
            ci = pd.DataFrame({'CI_corr': P_corr[market], 'CI_kendall': P_kendall[market], 'CI_partial': P_partial[market], 'CI_tail': P_tail[market]}).reindex(df.index)
            df = df.join(ci, how='inner').dropna()
            df['CI_corr_D'] = df['CI_corr'].shift(1)
            df['CI_kendall_D'] = df['CI_kendall'].shift(1)
            df['CI_partial_D'] = df['CI_partial'].shift(1)
            df['CI_tail_D'] = df['CI_tail'].shift(1)
            df = df.dropna()
            df = df.loc[fit_start_date:]
            data_test = df.loc[pred_start:pred_end].copy()
            if len(data_test) == 0:
                return (None, np.nan, np.nan, np.nan)
            window_len = len(df.loc[:pred_start]) - 1
            if window_len <= 0:
                return (None, np.nan, np.nan, np.nan)
            y_true = data_test['Y'].astype(float).copy()
            idx = y_true.index
            preds = []
            for i in tqdm(range(len(idx)), desc=f'HAR-CI {market} H{H}', leave=False):
                train = df.iloc[i:i + window_len].copy()
                train['_Y_'] = train['Y'].astype(float).values
                res = smf.ols('_Y_ ~ RV_D + RV_W + RV_M + CI_corr_D + CI_kendall_D + CI_partial_D + CI_tail_D', train).fit()
                p = res.params
                pred = p.iloc[0] + p.iloc[1] * data_test['RV_D'].iloc[i] + p.iloc[2] * data_test['RV_W'].iloc[i] + p.iloc[3] * data_test['RV_M'].iloc[i] + p.iloc[4] * data_test['CI_corr_D'].iloc[i] + p.iloc[5] * data_test['CI_kendall_D'].iloc[i] + p.iloc[6] * data_test['CI_partial_D'].iloc[i] + p.iloc[7] * data_test['CI_tail_D'].iloc[i]
                preds.append(float(pred))
            y_pred = pd.Series(preds, index=idx)
            mcs = pd.concat([y_true.rename('RV_True'), y_pred.rename('RV_Predict')], axis=1)
            mcs['MSE'] = (mcs['RV_True'] - mcs['RV_Predict']) ** 2
            mcs['MAE'] = np.abs(mcs['RV_True'] - mcs['RV_Predict'])
            mse = float(np.nanmean(mcs['MSE'].values))
            mae = float(np.nanmean(mcs['MAE'].values))
            denom = np.sqrt(np.nanmean(mcs['RV_Predict'].values ** 2)) + np.sqrt(np.nanmean(mcs['RV_True'].values ** 2))
            theil = float(np.sqrt(np.nanmean((mcs['RV_True'].values - mcs['RV_Predict'].values) ** 2)) / denom) if denom != 0 else np.nan
            return (mcs, mse, mae, theil)
        L_fit = L_fit_list[0]
        L_net = L_net_list[0]
        for H in H_LIST:
            error_by_market = {}
            for mkt in markets:
                if fit_end_date not in data_rv.index:
                    error_by_market[mkt] = [np.nan, np.nan, np.nan]
                    continue
                fit_end_pos = data_rv.index.get_loc(fit_end_date)
                fit_start_pos = fit_end_pos - (L_fit - 1)
                if fit_start_pos < 0:
                    error_by_market[mkt] = [np.nan, np.nan, np.nan]
                    continue
                fit_start_date = data_rv.index[fit_start_pos]
                ret_pos_end = df_ret.index.searchsorted(fit_start_date) - 1
                if ret_pos_end < 0:
                    error_by_market[mkt] = [np.nan, np.nan, np.nan]
                    continue
                net_end_date = df_ret.index[ret_pos_end]
                ci_start_date = net_end_date
                ci_end_date = pred_end
                (P_corr, P_kendall, P_partial, P_tail) = load_P_from_csv(L_net=L_net, ci_start_date=ci_start_date, ci_end_date=ci_end_date)
                if P_corr is None:
                    error_by_market[mkt] = [np.nan, np.nan, np.nan]
                    continue
                if mkt not in P_corr.columns:
                    error_by_market[mkt] = [np.nan, np.nan, np.nan]
                    continue
                (mcs, mse, mae, theil) = forecast_one_market_oneH_harci(data_rv_all=data_rv, P_corr=P_corr, P_kendall=P_kendall, P_partial=P_partial, P_tail=P_tail, market=mkt, fit_start_date=fit_start_date, H=H)
                if mcs is None:
                    tmp = pd.DataFrame(index=pd.DatetimeIndex([], name='date'))
                    tmp['RV_True'] = np.nan
                    tmp['RV_Predict'] = np.nan
                    tmp['MSE'] = np.nan
                    tmp['MAE'] = np.nan
                    model_outputs[str(mkt).strip().lstrip('.')] = tmp.copy()
                else:
                    model_outputs[str(mkt).strip().lstrip('.')] = mcs.copy()
                error_by_market[mkt] = [mse, mae, theil]
            ERROR = pd.DataFrame(error_by_market, index=['MSE', 'MAE', 'Theil'])
        return model_outputs

    def run_bagging():
        model_outputs = {}
        from sklearn import metrics
        import pandas as pd
        from functools import partial
        from tqdm import tqdm as _tqdm
        tqdm = partial(_tqdm, bar_format="{bar} {percentage:3.0f}%", ncols=50)
        import numpy as np
        from sklearn.preprocessing import StandardScaler
        import warnings
        warnings.filterwarnings('ignore')
        from sklearn.metrics import mean_squared_error
        from sklearn.metrics import mean_absolute_error
        from sklearn.tree import DecisionTreeRegressor
        from sklearn.tree import ExtraTreeRegressor
        from sklearn.preprocessing import MinMaxScaler
        from sklearn.ensemble import RandomForestRegressor
        from sklearn.ensemble import AdaBoostRegressor
        from sklearn.ensemble import GradientBoostingRegressor
        from sklearn.ensemble import BaggingRegressor
        from sklearn.metrics import mean_absolute_percentage_error
        import os
        data = pd.read_csv(RV_FILE)
        data = data.dropna()
        data['date'] = pd.to_datetime(data['date'], utc=True)
        data['date'] = data['date'].dt.tz_localize(None)
        data['date'] = data['date'].dt.normalize()
        data = data.fillna(method='ffill')
        data = data.set_index('date')
        P_STORE_PATH = CENTRALITY_FILE
        P_all = pd.read_csv(P_STORE_PATH)
        need_cols = ['date', 'market', 'P_corr', 'P_kendall', 'P_partial', 'P_tail']
        miss = [c for c in need_cols if c not in P_all.columns]
        if miss:
            raise RuntimeError(f'Model execution message{miss}Model execution message{list(P_all.columns)}')
        P_all['date'] = pd.to_datetime(P_all['date'], errors='coerce').dt.tz_localize(None).dt.normalize()
        P_all = P_all.dropna(subset=['date', 'market']).copy()
        P_ci = P_all.groupby(['date', 'market'])[['P_corr', 'P_kendall', 'P_partial', 'P_tail']].mean().reset_index()
        error = []
        for country in tqdm(data.columns):
            data_RV = data[[country]].copy()
            data_RV['RV_now'] = data_RV[country]
            data_RV = data_RV.drop(columns=[country])
            data_RV['RV_MA22'] = data_RV['RV_now'].rolling(window=22, center=False).mean()
            data_RV['RV_MA22'] = data_RV['RV_MA22'].shift(-21)
            data_RV['RV_D'] = data_RV['RV_now'].shift(1)
            data_RV['RV_W'] = data_RV['RV_D'].rolling(5).mean()
            data_RV['RV_M'] = data_RV['RV_D'].rolling(22).mean()
            ci_one = P_ci[P_ci['market'] == country].copy()
            if ci_one.empty:
                tmp = pd.DataFrame(index=pd.DatetimeIndex([], name='date'))
                tmp['RV_True'] = np.nan
                tmp['RV_Predict'] = np.nan
                tmp['MSE'] = np.nan
                tmp['MAE'] = np.nan
                tmp['MAPE'] = np.nan
                model_outputs[str(country).strip().lstrip('.')] = tmp.copy()
                error.append([np.nan, np.nan, np.nan])
                continue
            ci_one = ci_one.set_index('date')[['P_corr', 'P_kendall', 'P_partial', 'P_tail']]
            ci_one = ci_one.rename(columns={'P_corr': 'CI_corr', 'P_kendall': 'CI_kendall', 'P_partial': 'CI_partial', 'P_tail': 'CI_tail'})
            data_RV = data_RV.join(ci_one, how='left')
            data_RV['CI_corr_D'] = data_RV['CI_corr'].shift(1)
            data_RV['CI_kendall_D'] = data_RV['CI_kendall'].shift(1)
            data_RV['CI_partial_D'] = data_RV['CI_partial'].shift(1)
            data_RV['CI_tail_D'] = data_RV['CI_tail'].shift(1)
            data_RV = data_RV.dropna()
            data_RV = data_RV.loc['2014-07-04':]
            time = '2018-05-02'
            time_stop = '2022-02-25'
            X_all = data_RV[['RV_D', 'RV_W', 'RV_M', 'CI_corr_D', 'CI_kendall_D', 'CI_partial_D', 'CI_tail_D']]
            y_all = data_RV['RV_MA22']
            train_mask = X_all.index < pd.to_datetime(time)
            test_mask = (X_all.index >= pd.to_datetime(time)) & (X_all.index <= pd.to_datetime(time_stop))
            (X_train, y_train) = (X_all[train_mask], y_all[train_mask])
            (X_test, y_test) = (X_all[test_mask], y_all[test_mask])
            if X_test.dropna().empty or y_test.dropna().empty or X_train.dropna().empty or y_train.dropna().empty:
                tmp = pd.DataFrame(index=pd.DatetimeIndex([], name='date'))
                tmp['RV_True'] = np.nan
                tmp['RV_Predict'] = np.nan
                tmp['MSE'] = np.nan
                tmp['MAE'] = np.nan
                tmp['MAPE'] = np.nan
                model_outputs[str(country).strip().lstrip('.')] = tmp.copy()
                error.append([np.nan, np.nan, np.nan])
                continue
            base_estimator = DecisionTreeRegressor(max_depth=10, random_state=42)
            try:
                Bagging = BaggingRegressor(estimator=base_estimator, n_estimators=1000, max_samples=0.8, max_features=0.8, random_state=42, n_jobs=-1)
            except TypeError:
                Bagging = BaggingRegressor(base_estimator=base_estimator, n_estimators=1000, max_samples=0.8, max_features=0.8, random_state=42, n_jobs=-1)
            Bagging.fit(X_train.dropna(), y_train.dropna())
            predictions = Bagging.predict(X_test.dropna())
            valid_idx = X_test.dropna().index
            result = pd.Series(predictions, index=valid_idx, name='RV_Predict')
            SUM = pd.concat([y_test.loc[valid_idx], result], axis=1)
            SUM.columns = ['RV_True', 'RV_Predict']
            SUM_final = SUM.dropna()
            SUM_final['MSE'] = (SUM_final['RV_True'] - SUM_final['RV_Predict']) ** 2
            SUM_final['MAE'] = np.abs(SUM_final['RV_True'] - SUM_final['RV_Predict'])
            SUM_final['MAPE'] = np.abs((SUM_final['RV_True'] - SUM_final['RV_Predict']) / SUM_final['RV_True'])
            MSE = np.mean((SUM_final['RV_True'] - SUM_final['RV_Predict']) ** 2)
            MAE = np.mean(np.abs(SUM_final['RV_True'] - SUM_final['RV_Predict']))
            Theil = np.sqrt(np.mean((SUM_final['RV_True'] - SUM_final['RV_Predict']) ** 2)) / (np.sqrt(np.mean(SUM_final['RV_Predict'] ** 2)) + np.sqrt(np.mean(SUM_final['RV_True'] ** 2)))
            error.append([MSE, MAE, Theil])
            model_outputs[str(country).strip().lstrip('.')] = SUM_final.copy()
        ERROR = pd.DataFrame(error, columns=['MSE', 'MAE', 'Theil'], index=data.columns).T
        return model_outputs

    def run_gradient():
        model_outputs = {}
        from sklearn import metrics
        import pandas as pd
        from functools import partial
        from tqdm import tqdm as _tqdm
        tqdm = partial(_tqdm, bar_format="{bar} {percentage:3.0f}%", ncols=50)
        import numpy as np
        from sklearn.preprocessing import StandardScaler
        import warnings
        warnings.filterwarnings('ignore')
        from sklearn.metrics import mean_squared_error
        from sklearn.metrics import mean_absolute_error
        from sklearn.tree import DecisionTreeRegressor
        from sklearn.tree import ExtraTreeRegressor
        from sklearn.preprocessing import MinMaxScaler
        from sklearn.ensemble import RandomForestRegressor
        from sklearn.ensemble import AdaBoostRegressor
        from sklearn.ensemble import GradientBoostingRegressor
        from sklearn.ensemble import BaggingRegressor
        from sklearn.metrics import mean_absolute_percentage_error
        import os
        data = pd.read_csv(RV_FILE)
        data = data.dropna()
        data['date'] = pd.to_datetime(data['date'], utc=True)
        data['date'] = data['date'].dt.tz_localize(None)
        data['date'] = data['date'].dt.normalize()
        data = data.fillna(method='ffill')
        data = data.set_index('date')
        P_STORE_PATH = CENTRALITY_FILE
        P_all = pd.read_csv(P_STORE_PATH)
        need_cols = ['date', 'market', 'P_corr', 'P_kendall', 'P_partial', 'P_tail']
        miss = [c for c in need_cols if c not in P_all.columns]
        if miss:
            raise RuntimeError(f'Model execution message{miss}Model execution message{list(P_all.columns)}')
        P_all['date'] = pd.to_datetime(P_all['date'], errors='coerce').dt.tz_localize(None).dt.normalize()
        P_all = P_all.dropna(subset=['date', 'market']).copy()
        P_ci = P_all.groupby(['date', 'market'])[['P_corr', 'P_kendall', 'P_partial', 'P_tail']].mean().reset_index()
        error = []
        for country in tqdm(data.columns):
            data_RV = data[[country]].copy()
            data_RV['RV_now'] = data_RV[country]
            data_RV = data_RV.drop(columns=[country])
            data_RV['RV_D'] = data_RV['RV_now'].shift(1)
            data_RV['RV_W'] = data_RV['RV_D'].rolling(5).mean()
            data_RV['RV_M'] = data_RV['RV_D'].rolling(22).mean()
            data_RV['RV_MA22'] = data_RV['RV_now'].rolling(window=22, center=False).mean()
            data_RV['RV_MA22'] = data_RV['RV_MA22'].shift(-21)
            ci_one = P_ci[P_ci['market'] == country].copy()
            if ci_one.empty:
                tmp = pd.DataFrame(index=pd.DatetimeIndex([], name='date'))
                tmp['RV_True'] = np.nan
                tmp['RV_Predict'] = np.nan
                tmp['MSE'] = np.nan
                tmp['MAE'] = np.nan
                tmp['MAPE'] = np.nan
                model_outputs[str(country).strip().lstrip('.')] = tmp.copy()
                error.append([np.nan, np.nan, np.nan])
                continue
            ci_one = ci_one.set_index('date')[['P_corr', 'P_kendall', 'P_partial', 'P_tail']]
            ci_one = ci_one.rename(columns={'P_corr': 'CI_corr', 'P_kendall': 'CI_kendall', 'P_partial': 'CI_partial', 'P_tail': 'CI_tail'})
            data_RV = data_RV.join(ci_one, how='left')
            data_RV['CI_corr_D'] = data_RV['CI_corr'].shift(1)
            data_RV['CI_kendall_D'] = data_RV['CI_kendall'].shift(1)
            data_RV['CI_partial_D'] = data_RV['CI_partial'].shift(1)
            data_RV['CI_tail_D'] = data_RV['CI_tail'].shift(1)
            data_RV = data_RV.dropna()
            data_RV = data_RV.loc['2014-07-04':]
            time = '2018-05-02'
            time_stop = '2022-02-25'
            X_all = data_RV[['RV_D', 'RV_W', 'RV_M', 'CI_corr_D', 'CI_kendall_D', 'CI_partial_D', 'CI_tail_D']]
            y_all = data_RV['RV_MA22']
            train_mask = X_all.index < pd.to_datetime(time)
            test_mask = (X_all.index >= pd.to_datetime(time)) & (X_all.index <= pd.to_datetime(time_stop))
            (X_train, y_train) = (X_all[train_mask], y_all[train_mask])
            (X_test, y_test) = (X_all[test_mask], y_all[test_mask])
            if X_test.dropna().empty or y_test.dropna().empty or X_train.dropna().empty or y_train.dropna().empty:
                tmp = pd.DataFrame(index=pd.DatetimeIndex([], name='date'))
                tmp['RV_True'] = np.nan
                tmp['RV_Predict'] = np.nan
                tmp['MSE'] = np.nan
                tmp['MAE'] = np.nan
                tmp['MAPE'] = np.nan
                model_outputs[str(country).strip().lstrip('.')] = tmp.copy()
                error.append([np.nan, np.nan, np.nan])
                continue
            gb = GradientBoostingRegressor(n_estimators=1000, learning_rate=0.02, max_depth=10, subsample=0.8, min_samples_split=50, random_state=32, validation_fraction=0.1, n_iter_no_change=10)
            gb.fit(X_train.dropna(), y_train.dropna())
            predictions = gb.predict(X_test.dropna())
            valid_idx = X_test.dropna().index
            result = pd.Series(predictions, index=valid_idx, name='RV_Predict')
            SUM = pd.concat([y_test.loc[valid_idx], result], axis=1)
            SUM.columns = ['RV_True', 'RV_Predict']
            SUM_final = SUM.dropna()
            SUM_final['MSE'] = (SUM_final['RV_True'] - SUM_final['RV_Predict']) ** 2
            SUM_final['MAE'] = np.abs(SUM_final['RV_True'] - SUM_final['RV_Predict'])
            SUM_final['MAPE'] = np.abs((SUM_final['RV_True'] - SUM_final['RV_Predict']) / SUM_final['RV_True'])
            MSE = np.mean((SUM_final['RV_True'] - SUM_final['RV_Predict']) ** 2)
            MAE = np.mean(np.abs(SUM_final['RV_True'] - SUM_final['RV_Predict']))
            Theil = np.sqrt(np.mean((SUM_final['RV_True'] - SUM_final['RV_Predict']) ** 2)) / (np.sqrt(np.mean(SUM_final['RV_Predict'] ** 2)) + np.sqrt(np.mean(SUM_final['RV_True'] ** 2)))
            error.append([MSE, MAE, Theil])
            model_outputs[str(country).strip().lstrip('.')] = SUM_final.copy()
        ERROR = pd.DataFrame(error, columns=['MSE', 'MAE', 'Theil'], index=data.columns).T
        return model_outputs

    def run_random_forest():
        model_outputs = {}
        from sklearn import metrics
        import pandas as pd
        from functools import partial
        from tqdm import tqdm as _tqdm
        tqdm = partial(_tqdm, bar_format="{bar} {percentage:3.0f}%", ncols=50)
        import numpy as np
        from sklearn.preprocessing import StandardScaler
        import warnings
        warnings.filterwarnings('ignore')
        from sklearn.metrics import mean_squared_error
        from sklearn.metrics import mean_absolute_error
        from sklearn.tree import DecisionTreeRegressor
        from sklearn.tree import ExtraTreeRegressor
        from sklearn.preprocessing import MinMaxScaler
        from sklearn.ensemble import RandomForestRegressor
        from sklearn.ensemble import AdaBoostRegressor
        from sklearn.ensemble import GradientBoostingRegressor
        from sklearn.ensemble import BaggingRegressor
        from sklearn.metrics import mean_absolute_percentage_error
        import os
        H = 22
        data = pd.read_csv(RV_FILE)
        data = data.dropna()
        data['date'] = pd.to_datetime(data['date'], utc=True)
        data['date'] = data['date'].dt.tz_localize(None)
        data['date'] = data['date'].dt.normalize()
        data = data.fillna(method='ffill')
        data = data.set_index('date')
        P_STORE_PATH = CENTRALITY_FILE
        P_all = pd.read_csv(P_STORE_PATH)
        need_cols = ['date', 'market', 'P_corr', 'P_kendall', 'P_partial', 'P_tail']
        miss = [c for c in need_cols if c not in P_all.columns]
        if miss:
            raise RuntimeError(f'Model execution message{miss}Model execution message{list(P_all.columns)}')
        P_all['date'] = pd.to_datetime(P_all['date'], errors='coerce').dt.tz_localize(None).dt.normalize()
        P_all = P_all.dropna(subset=['date', 'market']).copy()
        P_ci = P_all.groupby(['date', 'market'])[['P_corr', 'P_kendall', 'P_partial', 'P_tail']].mean().reset_index()
        error = []
        for country in data.columns:
            data_RV = data[[country]].copy()
            data_RV['RV_now'] = data_RV[country]
            data_RV = data_RV.drop(columns=[country])
            data_RV['RV_D'] = data_RV['RV_now'].shift(1)
            data_RV['RV_W'] = data_RV['RV_D'].rolling(5).mean()
            data_RV['RV_M'] = data_RV['RV_D'].rolling(22).mean()
            data_RV['RV_MA22'] = data_RV['RV_now'].rolling(window=H, center=False).mean()
            data_RV['RV_MA22'] = data_RV['RV_MA22'].shift(-(H - 1))
            ci_one = P_ci[P_ci['market'] == country].copy()
            if ci_one.empty:
                tmp = pd.DataFrame(index=pd.DatetimeIndex([], name='date'))
                tmp['RV_True'] = np.nan
                tmp['RV_Predict'] = np.nan
                tmp['MSE'] = np.nan
                tmp['MAE'] = np.nan
                model_outputs[str(country).strip().lstrip('.')] = tmp.copy()
                error.append([np.nan, np.nan, np.nan])
                continue
            ci_one = ci_one.set_index('date')[['P_corr', 'P_kendall', 'P_partial', 'P_tail']]
            ci_one = ci_one.rename(columns={'P_corr': 'CI_corr', 'P_kendall': 'CI_kendall', 'P_partial': 'CI_partial', 'P_tail': 'CI_tail'})
            data_RV = data_RV.join(ci_one, how='left')
            data_RV['CI_corr_D'] = data_RV['CI_corr'].shift(1)
            data_RV['CI_kendall_D'] = data_RV['CI_kendall'].shift(1)
            data_RV['CI_partial_D'] = data_RV['CI_partial'].shift(1)
            data_RV['CI_tail_D'] = data_RV['CI_tail'].shift(1)
            data_RV = data_RV.dropna()
            data_RV = data_RV.loc['2014-07-04':]
            time = '2018-05-02'
            time_stop = '2022-02-25'
            X_all = data_RV[['RV_D', 'RV_W', 'RV_M', 'CI_corr_D', 'CI_kendall_D', 'CI_partial_D', 'CI_tail_D']].copy()
            y_all = data_RV['RV_MA22'].copy()
            time_ts = pd.to_datetime(time)
            time_stop_ts = pd.to_datetime(time_stop)
            if len(data_RV.loc[time:time_stop]) == 0:
                tmp = pd.DataFrame(index=pd.DatetimeIndex([], name='date'))
                tmp['RV_True'] = np.nan
                tmp['RV_Predict'] = np.nan
                tmp['MSE'] = np.nan
                tmp['MAE'] = np.nan
                model_outputs[str(country).strip().lstrip('.')] = tmp.copy()
                error.append([np.nan, np.nan, np.nan])
                continue
            if time_ts in data_RV.index:
                train_idx = int(data_RV.index.get_loc(time_ts))
            else:
                train_idx = int(data_RV.index.searchsorted(time_ts))
            test_len = len(data_RV.loc[time:time_stop])
            X_test = X_all.iloc[train_idx:train_idx + test_len]
            y_test = y_all.iloc[train_idx:train_idx + test_len]
            if X_test.empty or y_test.empty:
                tmp = pd.DataFrame(index=pd.DatetimeIndex([], name='date'))
                tmp['RV_True'] = np.nan
                tmp['RV_Predict'] = np.nan
                tmp['MSE'] = np.nan
                tmp['MAE'] = np.nan
                model_outputs[str(country).strip().lstrip('.')] = tmp.copy()
                error.append([np.nan, np.nan, np.nan])
                continue
            rf = RandomForestRegressor(n_estimators=200, max_depth=1, random_state=42, n_jobs=-1)
            predictions = []
            for i in tqdm(range(len(X_test)), desc=f'RF {country}', leave=False):
                pos = train_idx + i
                train_end_exclusive = pos - H + 1
                train_start = i
                if train_end_exclusive <= train_start:
                    predictions.append(np.nan)
                    continue
                current_train_X = X_all.iloc[train_start:train_end_exclusive].copy()
                current_train_y = y_all.iloc[train_start:train_end_exclusive].copy()
                current_train_X = current_train_X.dropna()
                current_train_y = current_train_y.loc[current_train_X.index].dropna()
                current_train_X = current_train_X.loc[current_train_y.index]
                if current_train_X.empty or current_train_y.empty:
                    predictions.append(np.nan)
                    continue
                rf.fit(current_train_X, current_train_y)
                x_one = X_test.iloc[[i]].dropna()
                if x_one.empty:
                    predictions.append(np.nan)
                    continue
                pred = rf.predict(x_one)
                predictions.append(float(pred[0]) if len(pred) > 0 else np.nan)
            result = pd.Series(predictions, index=y_test.index, name='RV_Predict')
            SUM = pd.concat([y_test.rename('RV_True'), result], axis=1)
            SUM_final = SUM.dropna()
            if len(SUM_final) == 0:
                tmp = pd.DataFrame(index=pd.DatetimeIndex([], name='date'))
                tmp['RV_True'] = np.nan
                tmp['RV_Predict'] = np.nan
                tmp['MSE'] = np.nan
                tmp['MAE'] = np.nan
                model_outputs[str(country).strip().lstrip('.')] = tmp.copy()
                error.append([np.nan, np.nan, np.nan])
                continue
            SUM_final['MSE'] = (SUM_final['RV_True'] - SUM_final['RV_Predict']) ** 2
            SUM_final['MAE'] = np.abs(SUM_final['RV_True'] - SUM_final['RV_Predict'])
            MSE = float(np.mean((SUM_final['RV_True'] - SUM_final['RV_Predict']) ** 2))
            MAE = float(np.mean(np.abs(SUM_final['RV_True'] - SUM_final['RV_Predict'])))
            denom = np.sqrt(np.mean(SUM_final['RV_Predict'] ** 2)) + np.sqrt(np.mean(SUM_final['RV_True'] ** 2))
            Theil = float(np.sqrt(np.mean((SUM_final['RV_True'] - SUM_final['RV_Predict']) ** 2)) / denom) if denom != 0 else np.nan
            error.append([MSE, MAE, Theil])
            model_outputs[str(country).strip().lstrip('.')] = SUM_final.copy()
        ERROR = pd.DataFrame(error, columns=['MSE', 'MAE', 'Theil'], index=data.columns).T
        return model_outputs

    def run_deep_learning():
        deep_outputs = {'LSTM': {}, 'CNN': {}, 'Transformer': {}}
        import os
        import warnings
        import numpy as np
        import pandas as pd
        import tensorflow as tf
        from functools import partial
        from tqdm import tqdm as _tqdm
        tqdm = partial(_tqdm, bar_format="{bar} {percentage:3.0f}%", ncols=50)
        from sklearn.preprocessing import MinMaxScaler
        from sklearn.metrics import mean_squared_error, mean_absolute_error
        from tensorflow.keras.models import Sequential
        from tensorflow.keras.layers import LSTM, Dense, Conv1D, MaxPooling1D, Flatten, MultiHeadAttention, Dropout, Input, GlobalAveragePooling1D
        from tensorflow.keras.callbacks import EarlyStopping
        warnings.filterwarnings('ignore')
        np.random.seed(42)
        tf.random.set_seed(42)
        RV_PATH = RV_FILE
        P_STORE_PATH = CENTRALITY_FILE
        OUT_LSTM = 'LSTM'
        OUT_CNN = 'CNN'
        OUT_TR = 'Transformer'
        PRED_START = pd.to_datetime('2018-05-02')
        PRED_END = pd.to_datetime('2022-02-25')
        FIT_START = pd.to_datetime('2014-07-04')
        n_steps = 1
        data = pd.read_csv(RV_PATH)
        data = data.dropna()
        data['date'] = pd.to_datetime(data['date'], utc=True)
        data['date'] = data['date'].dt.tz_localize(None)
        data['date'] = data['date'].dt.normalize()
        data = data.fillna(method='ffill')
        data = data.set_index('date')
        data = data.loc[FIT_START:]
        P_all = pd.read_csv(P_STORE_PATH)
        need_cols = ['date', 'market', 'P_corr', 'P_kendall', 'P_partial', 'P_tail']
        miss = [c for c in need_cols if c not in P_all.columns]
        if miss:
            raise RuntimeError(f'Model execution message{miss}Model execution message{list(P_all.columns)}')
        P_all['date'] = pd.to_datetime(P_all['date'], errors='coerce').dt.tz_localize(None).dt.normalize()
        P_all = P_all.dropna(subset=['date', 'market']).copy()
        P_ci = P_all.groupby(['date', 'market'])[['P_corr', 'P_kendall', 'P_partial', 'P_tail']].mean().reset_index()

        def build_lstm_model(input_shape):
            model = Sequential([LSTM(64, activation='relu', input_shape=input_shape), Dense(1)])
            model.compile(optimizer='adam', loss='mse')
            return model

        def build_cnn_model(input_shape):
            model = Sequential([Conv1D(filters=64, kernel_size=1, activation='relu', input_shape=input_shape), MaxPooling1D(pool_size=1), Flatten(), Dense(50, activation='relu'), Dense(1)])
            model.compile(optimizer='adam', loss='mse')
            return model

        def transformer_encoder(inputs, head_size, num_heads, ff_dim, dropout=0.0):
            x = MultiHeadAttention(key_dim=head_size, num_heads=num_heads, dropout=dropout)(inputs, inputs)
            x = Dropout(dropout)(x)
            res = x + inputs
            x = Dense(ff_dim, activation='relu')(res)
            x = Dense(inputs.shape[-1])(x)
            return x + res

        def build_transformer_model(input_shape):
            inputs = Input(shape=input_shape)
            x = transformer_encoder(inputs, head_size=32, num_heads=4, ff_dim=64, dropout=0.0)
            x = GlobalAveragePooling1D()(x)
            outputs = Dense(1)(x)
            model = tf.keras.Model(inputs, outputs)
            model.compile(optimizer='adam', loss='mse')
            return model

        def make_h22_target(rv_now: pd.Series) -> pd.Series:
            y = rv_now.rolling(window=22, center=False).mean().shift(-21)
            return y

        def create_dataset_with_dates(X_df: pd.DataFrame, y_s: pd.Series, time_steps=1):
            (Xs, ys, y_dates) = ([], [], [])
            for i in range(len(X_df) - time_steps):
                Xs.append(X_df.iloc[i:i + time_steps].values)
                ys.append(y_s.iloc[i + time_steps])
                y_dates.append(y_s.index[i + time_steps])
            return (np.array(Xs), np.array(ys), pd.DatetimeIndex(y_dates))

        def calc_theil(y_true, y_pred):
            y_true = np.asarray(y_true).reshape(-1, 1)
            y_pred = np.asarray(y_pred).reshape(-1, 1)
            denom = np.sqrt(np.mean(y_pred ** 2)) + np.sqrt(np.mean(y_true ** 2))
            if denom == 0:
                return np.nan
            return float(np.sqrt(np.mean((y_true - y_pred) ** 2)) / denom)

        def save_mcs_and_metrics(model, X_test, y_test_scaled, scaler_y, date_index, out_dir, country):
            y_pred_scaled = model.predict(X_test, verbose=0)
            y_pred = scaler_y.inverse_transform(y_pred_scaled).reshape(-1)
            y_true = scaler_y.inverse_transform(y_test_scaled.reshape(-1, 1)).reshape(-1)
            mcs = pd.DataFrame({'RV_True': y_true, 'RV_Predict': y_pred}, index=pd.DatetimeIndex(date_index, name='date'))
            mcs = mcs.dropna()
            if len(mcs) == 0:
                tmp = pd.DataFrame(index=pd.DatetimeIndex([], name='date'))
                tmp['RV_True'] = np.nan
                tmp['RV_Predict'] = np.nan
                tmp['MSE'] = np.nan
                tmp['MAE'] = np.nan
                deep_outputs[out_dir][str(country).strip().lstrip('.')] = tmp.copy()
                return [np.nan, np.nan, np.nan]
            mcs['MSE'] = (mcs['RV_True'] - mcs['RV_Predict']) ** 2
            mcs['MAE'] = np.abs(mcs['RV_True'] - mcs['RV_Predict'])
            eps = 1e-12
            pred_safe = np.clip(mcs['RV_Predict'].values, eps, None)
            true_safe = np.clip(mcs['RV_True'].values, eps, None)
            mcs['Qlike'] = true_safe / pred_safe + np.log(pred_safe)
            deep_outputs[out_dir][str(country).strip().lstrip('.')] = mcs.copy()
            mse = float(mean_squared_error(mcs['RV_True'], mcs['RV_Predict']))
            mae = float(mean_absolute_error(mcs['RV_True'], mcs['RV_Predict']))
            theil = calc_theil(mcs['RV_True'].values, mcs['RV_Predict'].values)
            return [mse, mae, theil]
        error_lstm = []
        error_cnn = []
        error_transformer = []
        for country in tqdm(data.columns, desc='Markets'):
            data_RV = data[[country]].copy()
            data_RV['RV_now'] = data_RV[country]
            data_RV = data_RV.drop(columns=[country])
            data_RV['RV_D'] = data_RV['RV_now'].shift(1)
            data_RV['RV_W'] = data_RV['RV_D'].rolling(5).mean()
            data_RV['RV_M'] = data_RV['RV_D'].rolling(22).mean()
            data_RV['RV_MA22'] = make_h22_target(data_RV['RV_now'])
            ci_one = P_ci[P_ci['market'] == country].copy()
            if ci_one.empty:
                for out_dir in [OUT_LSTM, OUT_CNN, OUT_TR]:
                    tmp = pd.DataFrame(index=pd.DatetimeIndex([], name='date'))
                    tmp['RV_True'] = np.nan
                    tmp['RV_Predict'] = np.nan
                    tmp['MSE'] = np.nan
                    tmp['MAE'] = np.nan
                    tmp['Qlike'] = np.nan
                    deep_outputs[out_dir][str(country).strip().lstrip('.')] = tmp.copy()
                error_lstm.append([np.nan, np.nan, np.nan])
                error_cnn.append([np.nan, np.nan, np.nan])
                error_transformer.append([np.nan, np.nan, np.nan])
                continue
            ci_one = ci_one.set_index('date')[['P_corr', 'P_kendall', 'P_partial', 'P_tail']]
            ci_one = ci_one.rename(columns={'P_corr': 'CI_corr', 'P_kendall': 'CI_kendall', 'P_partial': 'CI_partial', 'P_tail': 'CI_tail'})
            data_RV = data_RV.join(ci_one, how='left')
            data_RV['CI_corr_D'] = data_RV['CI_corr'].shift(1)
            data_RV['CI_kendall_D'] = data_RV['CI_kendall'].shift(1)
            data_RV['CI_partial_D'] = data_RV['CI_partial'].shift(1)
            data_RV['CI_tail_D'] = data_RV['CI_tail'].shift(1)
            data_RV = data_RV.dropna()
            data_RV = data_RV.loc[FIT_START:]
            X = data_RV[['RV_D', 'RV_W', 'RV_M', 'CI_corr_D', 'CI_kendall_D', 'CI_partial_D', 'CI_tail_D']].copy()
            y = data_RV['RV_MA22'].copy()
            if len(X) <= n_steps + 10:
                for out_dir in [OUT_LSTM, OUT_CNN, OUT_TR]:
                    tmp = pd.DataFrame(index=pd.DatetimeIndex([], name='date'))
                    tmp['RV_True'] = np.nan
                    tmp['RV_Predict'] = np.nan
                    tmp['MSE'] = np.nan
                    tmp['MAE'] = np.nan
                    tmp['Qlike'] = np.nan
                    deep_outputs[out_dir][str(country).strip().lstrip('.')] = tmp.copy()
                error_lstm.append([np.nan, np.nan, np.nan])
                error_cnn.append([np.nan, np.nan, np.nan])
                error_transformer.append([np.nan, np.nan, np.nan])
                continue
            train_mask = X.index < PRED_START
            test_mask = (X.index >= PRED_START) & (X.index <= PRED_END)
            X_train_df = X.loc[train_mask].copy()
            y_train_s = y.loc[train_mask].copy()
            X_test_df = X.loc[test_mask].copy()
            y_test_s = y.loc[test_mask].copy()
            if len(X_test_df) == 0 or len(X_train_df) == 0:
                for out_dir in [OUT_LSTM, OUT_CNN, OUT_TR]:
                    tmp = pd.DataFrame(index=pd.DatetimeIndex([], name='date'))
                    tmp['RV_True'] = np.nan
                    tmp['RV_Predict'] = np.nan
                    tmp['MSE'] = np.nan
                    tmp['MAE'] = np.nan
                    tmp['Qlike'] = np.nan
                    deep_outputs[out_dir][str(country).strip().lstrip('.')] = tmp.copy()
                error_lstm.append([np.nan, np.nan, np.nan])
                error_cnn.append([np.nan, np.nan, np.nan])
                error_transformer.append([np.nan, np.nan, np.nan])
                continue
            X_context = X.iloc[max(0, X.index.get_loc(X_test_df.index[0]) - n_steps):X.index.get_loc(X_test_df.index[0])]
            y_context = y.loc[X_context.index]
            X_test_for_seq = pd.concat([X_context, X_test_df], axis=0)
            y_test_for_seq = pd.concat([y_context, y_test_s], axis=0)
            scaler_X = MinMaxScaler()
            scaler_y = MinMaxScaler()
            X_train_scaled = pd.DataFrame(scaler_X.fit_transform(X_train_df), index=X_train_df.index, columns=X_train_df.columns)
            y_train_scaled = pd.Series(scaler_y.fit_transform(y_train_s.values.reshape(-1, 1)).reshape(-1), index=y_train_s.index)
            X_test_scaled = pd.DataFrame(scaler_X.transform(X_test_for_seq), index=X_test_for_seq.index, columns=X_test_for_seq.columns)
            y_test_scaled = pd.Series(scaler_y.transform(y_test_for_seq.values.reshape(-1, 1)).reshape(-1), index=y_test_for_seq.index)
            (X_train_seq, y_train_seq, train_dates) = create_dataset_with_dates(X_train_scaled, y_train_scaled, n_steps)
            (X_test_seq, y_test_seq, test_dates) = create_dataset_with_dates(X_test_scaled, y_test_scaled, n_steps)
            keep_test = (test_dates >= PRED_START) & (test_dates <= PRED_END)
            X_test_seq = X_test_seq[keep_test]
            y_test_seq = y_test_seq[keep_test]
            test_dates = test_dates[keep_test]
            if len(X_train_seq) == 0 or len(X_test_seq) == 0:
                for out_dir in [OUT_LSTM, OUT_CNN, OUT_TR]:
                    tmp = pd.DataFrame(index=pd.DatetimeIndex([], name='date'))
                    tmp['RV_True'] = np.nan
                    tmp['RV_Predict'] = np.nan
                    tmp['MSE'] = np.nan
                    tmp['MAE'] = np.nan
                    tmp['Qlike'] = np.nan
                    deep_outputs[out_dir][str(country).strip().lstrip('.')] = tmp.copy()
                error_lstm.append([np.nan, np.nan, np.nan])
                error_cnn.append([np.nan, np.nan, np.nan])
                error_transformer.append([np.nan, np.nan, np.nan])
                continue
            early_stopping = EarlyStopping(monitor='val_loss', patience=10, verbose=0, restore_best_weights=True)
            lstm_model = build_lstm_model((n_steps, X_train_seq.shape[2]))
            lstm_model.fit(X_train_seq, y_train_seq, epochs=200, batch_size=32, validation_split=0.1, callbacks=[early_stopping], verbose=0)
            lstm_pred_scaled = lstm_model.predict(X_test_seq, verbose=0)
            lstm_metrics_real = save_mcs_and_metrics(lstm_model, X_test_seq, y_test_seq, scaler_y, test_dates, OUT_LSTM, country)
            cnn_model = build_cnn_model((n_steps, X_train_seq.shape[2]))
            cnn_model.fit(X_train_seq, y_train_seq, epochs=200, batch_size=32, validation_split=0.1, callbacks=[early_stopping], verbose=0)
            cnn_pred_scaled = cnn_model.predict(X_test_seq, verbose=0)
            cnn_metrics_real = save_mcs_and_metrics(cnn_model, X_test_seq, y_test_seq, scaler_y, test_dates, OUT_CNN, country)
            transformer_model = build_transformer_model((n_steps, X_train_seq.shape[2]))
            transformer_model.fit(X_train_seq, y_train_seq, epochs=200, batch_size=32, validation_split=0.1, callbacks=[early_stopping], verbose=0)
            tr_pred_scaled = transformer_model.predict(X_test_seq, verbose=0)
            transformer_metrics_real = save_mcs_and_metrics(transformer_model, X_test_seq, y_test_seq, scaler_y, test_dates, OUT_TR, country)
            error_lstm.append(lstm_metrics_real)
            error_cnn.append(cnn_metrics_real)
            error_transformer.append(transformer_metrics_real)
        ERROR_lstm = pd.DataFrame(error_lstm, columns=['MSE', 'MAE', 'Theil'], index=data.columns).T
        ERROR_cnn = pd.DataFrame(error_cnn, columns=['MSE', 'MAE', 'Theil'], index=data.columns).T
        ERROR_transformer = pd.DataFrame(error_transformer, columns=['MSE', 'MAE', 'Theil'], index=data.columns).T
        return deep_outputs

    def calc_mse(y_true: np.ndarray, y_pred: np.ndarray) -> dict:
        y_true = np.asarray(y_true, dtype=float)
        y_pred = np.asarray(y_pred, dtype=float)
        mask = np.isfinite(y_true) & np.isfinite(y_pred)
        yt = y_true[mask]
        yp = y_pred[mask]
        if len(yt) == 0:
            return {'MSE': np.nan}
        err = yp - yt
        mse = float(np.mean(err ** 2))
        return {'MSE': mse}

    def qlike_with_shared_drop(y_true: pd.Series, y_pred: pd.Series, drop_dates: set) -> float:
        df = pd.concat([y_true.rename('y'), y_pred.rename('yhat')], axis=1).dropna()
        if df.empty:
            return np.nan
        if drop_dates:
            df = df.loc[~df.index.isin(drop_dates)]
        if df.empty:
            return np.nan
        y = df['y'].astype(float).values
        yhat = df['yhat'].astype(float).values
        ok_pos = (y > EPS) & (yhat > EPS)
        if not np.any(ok_pos):
            return np.nan
        ratio = y[ok_pos] / yhat[ok_pos]
        ok_ratio = np.isfinite(ratio) & (ratio >= RATIO_MIN) & (ratio <= RATIO_MAX)
        if not np.any(ok_ratio):
            return np.nan
        ratio = ratio[ok_ratio]
        q = ratio - np.log(ratio) - 1.0
        ok_q = np.isfinite(q) & (q >= 0.0) & (q <= QLIKE_CAP)
        if not np.any(ok_q):
            return np.nan
        return float(np.mean(q[ok_q]))

    def add_summary_rows(df: pd.DataFrame) -> pd.DataFrame:
        mean_row = df.mean(axis=0, skipna=True)
        q1_row = df.quantile(0.25, axis=0, interpolation='linear')
        q2_row = df.quantile(0.5, axis=0, interpolation='linear')
        q3_row = df.quantile(0.75, axis=0, interpolation='linear')
        df2 = df.copy()
        df2.loc['Mean'] = mean_row
        df2.loc['Q1'] = q1_row
        df2.loc['Q2'] = q2_row
        df2.loc['Q3'] = q3_row
        return df2

    def prepare_mcs_matrix_relaxed(sum_pre: pd.DataFrame, min_valid_per_model: int=10, min_rows_for_mcs: int=5):
        info = {'initial_shape': tuple(sum_pre.shape), 'dropped_low_valid_models': [], 'dropped_const_models': [], 'final_shape': None, 'reason': None}
        if sum_pre.empty or sum_pre.shape[1] == 0:
            info['reason'] = 'The loss matrix is empty'
            return (None, info)
        valid_counts = sum_pre.notna().sum(axis=0)
        keep_cols = valid_counts[valid_counts >= min_valid_per_model].index.tolist()
        info['dropped_low_valid_models'] = [c for c in sum_pre.columns if c not in keep_cols]
        if len(keep_cols) < 2:
            info['reason'] = f'Fewer than two models with at least {min_valid_per_model} valid observations'
            return (None, info)
        x = sum_pre[keep_cols].copy()
        x = x.dropna(axis=0, how='any')
        if x.empty:
            info['reason'] = 'No common valid observations'
            return (None, info)
        if x.shape[1] < 2:
            info['reason'] = 'Fewer than two models after dropping missing observations'
            return (None, info)
        if x.shape[0] < min_rows_for_mcs:
            info['reason'] = f'Insufficient common observations ({x.shape[0]} < {min_rows_for_mcs})'
            return (None, info)
        variances = x.var(axis=0, skipna=True)
        non_const_cols = variances[variances > 0].index.tolist()
        info['dropped_const_models'] = [c for c in x.columns if c not in non_const_cols]
        x = x[non_const_cols]
        if x.shape[1] < 2:
            info['reason'] = 'Fewer than two models after removing constant columns'
            return (None, info)
        if x.shape[0] < min_rows_for_mcs:
            info['reason'] = f'Insufficient common observations after removing constant columns ({x.shape[0]} < {min_rows_for_mcs})'
            return (None, info)
        info['final_shape'] = tuple(x.shape)
        return (x, info)

    def choose_one_model_to_drop_for_retry(raw_sum_pre: pd.DataFrame, current_cols: list):
        if len(current_cols) <= 2:
            return None
        sub = raw_sum_pre[current_cols].copy()
        valid_counts = sub.notna().sum(axis=0)
        min_valid = valid_counts.min()
        candidates = valid_counts[valid_counts == min_valid].index.tolist()
        if len(candidates) == 1:
            return candidates[0]
        var_map = {}
        for c in candidates:
            s = pd.to_numeric(sub[c], errors='coerce').dropna()
            if len(s) <= 1:
                var_map[c] = -1.0
            else:
                v = float(s.var())
                if np.isnan(v):
                    v = -1.0
                var_map[c] = v
        drop_col = sorted(var_map.items(), key=lambda kv: (kv[1], kv[0]))[0][0]
        return drop_col

    def run_mcs_with_retry(sum_pre: pd.DataFrame, country_code: str, min_valid_per_model: int, min_rows_for_mcs: int, max_retry_drop: int, debug_logs: list):
        details = {'country': country_code, 'retry_count': 0, 'dropped_in_retry': [], 'fail_reason': None, 'final_shape': None, 'final_models': None}
        if sum_pre.empty or sum_pre.shape[1] < 2:
            details['fail_reason'] = 'Fewer than two available models'
            return (False, None, None, None, details)
        current_cols = list(sum_pre.columns)
        for _ in range(max_retry_drop + 1):
            tmp_sum = sum_pre[current_cols].copy()
            (mcs_df, prep_info) = prepare_mcs_matrix_relaxed(tmp_sum, min_valid_per_model=min_valid_per_model, min_rows_for_mcs=min_rows_for_mcs)
            if mcs_df is None:
                if len(current_cols) <= 2:
                    details['fail_reason'] = f"Preparation failed with too few models to retry: {prep_info.get('reason')}"
                    return (False, None, None, None, details)
                drop_col = choose_one_model_to_drop_for_retry(sum_pre, current_cols)
                if drop_col is None:
                    details['fail_reason'] = f"Preparation failed and no model can be removed: {prep_info.get('reason')}"
                    return (False, None, None, None, details)
                current_cols.remove(drop_col)
                details['retry_count'] += 1
                details['dropped_in_retry'].append(drop_col)
                debug_logs.append(f"[{country_code}] Preparation failed; retry without model: {drop_col} | Reason: {prep_info.get('reason')}")
                continue
            try:
                mcs = MCS(losses=mcs_df, size=0.9, method='R', reps=20000, block_size=5)
                mcs.compute()
                details['final_shape'] = tuple(mcs_df.shape)
                details['final_models'] = list(mcs_df.columns)
                return (True, mcs.pvalues, mcs.included, mcs_df, details)
            except Exception as e:
                if len(current_cols) <= 2:
                    details['fail_reason'] = f'MCS failed with too few models to retry: {str(e)}'
                    return (False, None, None, None, details)
                drop_col = choose_one_model_to_drop_for_retry(sum_pre, current_cols)
                if drop_col is None:
                    details['fail_reason'] = f'MCS failed and no model can be removed: {str(e)}'
                    return (False, None, None, None, details)
                current_cols.remove(drop_col)
                details['retry_count'] += 1
                details['dropped_in_retry'].append(drop_col)
                debug_logs.append(f'[{country_code}] MCS failed; retry without model: {drop_col} | Error: {str(e)}')
                continue
        details['fail_reason'] = f'Maximum retry count exceeded: {max_retry_drop}'
        return (False, None, None, None, details)

    def qlike_series_with_rules(y_true: np.ndarray, y_pred: np.ndarray) -> np.ndarray:
        y_true = np.asarray(y_true, dtype=float)
        y_pred = np.asarray(y_pred, dtype=float)
        q = np.full(shape=len(y_true), fill_value=np.nan, dtype=float)
        mask = np.isfinite(y_true) & np.isfinite(y_pred)
        if not np.any(mask):
            return q
        yt = y_true[mask]
        yp = y_pred[mask]
        ok_pos = (yt > EPS) & (yp > EPS)
        ratio = np.full_like(yt, np.nan, dtype=float)
        ratio[ok_pos] = yt[ok_pos] / yp[ok_pos]
        ok_ratio = np.isfinite(ratio) & (ratio >= RATIO_MIN) & (ratio <= RATIO_MAX)
        q_sub = np.full_like(yt, np.nan, dtype=float)
        q_sub[ok_ratio] = ratio[ok_ratio] - np.log(ratio[ok_ratio]) - 1.0
        ok_q = np.isfinite(q_sub) & (q_sub >= 0.0) & (q_sub <= QLIKE_CAP)
        q_masked = np.full_like(yt, np.nan, dtype=float)
        q_masked[ok_q] = q_sub[ok_q]
        q[mask] = q_masked
        return q

    MODEL_LABELS = [
        ("HAR-RV", "HAR"),
        ("HAR-CI", "HAR-CI"),
        ("bagging", "Bagging"),
        ("gradient", "Gradient"),
        ("RF", "RF"),
        ("CNN", "CNN"),
        ("LSTM", "LSTM"),
        ("Transformer", "Transformer"),
    ]

    STANDARD_RUNNERS = [
        ("HAR-RV", "HAR", run_har),
        ("HAR-CI", "HAR-CI", run_har_ci),
        ("bagging", "Bagging", run_bagging),
        ("gradient", "Gradient", run_gradient),
        ("RF", "RF", run_random_forest),
    ]

    def load_prediction_pair(frame):
        dates = pd.to_datetime(frame.index, errors="coerce")
        y_true = pd.Series(pd.to_numeric(frame["RV_True"], errors="coerce").to_numpy(), index=dates)
        y_pred = pd.Series(pd.to_numeric(frame["RV_Predict"], errors="coerce").to_numpy(), index=dates)
        valid = dates.notna() & np.isfinite(y_true.to_numpy()) & np.isfinite(y_pred.to_numpy())
        y_true, y_pred = y_true[valid], y_pred[valid]
        if len(y_true) and np.mean(y_true.to_numpy() > THRESH) > 0.5:
            y_true = y_true * SCALE_IF_MOST_GT_THRESH
            y_pred = y_pred * SCALE_IF_MOST_GT_THRESH
        y_true = y_true[~y_true.index.duplicated(keep="last")].sort_index()
        y_pred = y_pred[~y_pred.index.duplicated(keep="last")].sort_index()
        return y_true, y_pred

    def summarize_losses(predictions, shared_drop_dates):
        models = [key for key, _ in MODEL_LABELS]
        markets = sorted(predictions["HAR-RV"])
        cache = {market: {model: load_prediction_pair(predictions[model][market]) for model in models} for market in markets}
        raw = {metric: pd.DataFrame(index=markets, columns=models, dtype=float) for metric in ("MSE", "QLIKE")}
        for market in markets:
            for model in models:
                y_true, y_pred = cache[market][model]
                raw["MSE"].loc[market, model] = calc_mse(y_true.to_numpy(), y_pred.to_numpy())["MSE"]
                raw["QLIKE"].loc[market, model] = qlike_with_shared_drop(y_true, y_pred, shared_drop_dates[market])
        summaries = {}
        for metric, values in raw.items():
            values = values.dropna(how="all").dropna(subset=["HAR-RV"], how="all")
            relative = values.div(values["HAR-RV"].replace(0.0, np.nan), axis=0)
            if np.any(np.isfinite(relative.to_numpy()) & (np.abs(relative.to_numpy()) > POST_SCALE_THRESHOLD)):
                warnings.warn("The original relative-loss scaling rule was triggered for " + metric, RuntimeWarning)
                relative = relative.where(~(np.abs(relative) > POST_SCALE_THRESHOLD), relative * POST_SCALE_FACTOR)
            summaries[metric] = add_summary_rows(relative).loc[["Mean", "Q1", "Q2"]].T
        return summaries, shared_drop_dates

    def qlike_loss_series(frame, drop_dates):
        data = pd.DataFrame({
            "date": pd.to_datetime(frame.index, errors="coerce"),
            "y_true": pd.to_numeric(frame["RV_True"], errors="coerce").to_numpy(),
            "y_pred": pd.to_numeric(frame["RV_Predict"], errors="coerce").to_numpy(),
        }).dropna(subset=["date"])
        if data["date"].dt.tz is not None:
            data["date"] = data["date"].dt.tz_localize(None)
        data["date"] = data["date"].dt.normalize()
        data = data.sort_values("date").drop_duplicates(subset=["date"], keep="last").reset_index(drop=True)
        data = data.loc[~data["date"].isin(drop_dates)]
        return pd.Series(qlike_series_with_rules(data["y_true"].to_numpy(), data["y_pred"].to_numpy()), index=data["date"])

    def count_mcs_memberships(predictions, shared_drop_dates):
        models = [key for key, _ in MODEL_LABELS]
        markets = sorted(predictions["HAR-RV"])
        counts = {metric: dict.fromkeys(models, 0) for metric in ("MSE", "QLIKE")}
        for metric in counts:
            for market in tqdm(markets, desc="MCS " + metric):
                losses = pd.DataFrame()
                for model in models:
                    frame = predictions[model][market]
                    if metric == "MSE":
                        series = pd.to_numeric(frame["MSE"], errors="coerce").reset_index(drop=True)
                    else:
                        series = qlike_loss_series(frame, shared_drop_dates[market])
                    if series.notna().sum():
                        losses[model] = series
                logs = []
                success, _, included, used, details = run_mcs_with_retry(losses, market, 10, 5, 50, logs)
                if not success:
                    raise RuntimeError(f"MCS {metric} failed for {market}: {details['fail_reason']}")
                omitted = [model for model in models if model not in used.columns]
                if omitted:
                    warnings.warn(f"MCS {metric}, {market}: excluded from the test: {', '.join(omitted)}", RuntimeWarning)
                for model in included:
                    counts[metric][model] += 1
        return counts

    def save_table(summaries, counts, output_path=OUTPUT_FILE, bold_mode="har-ci"):
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "Table 6"
        sheet.merge_cells("A1:E1")
        sheet.merge_cells("G1:K1")
        sheet["A1"] = "MSE"
        sheet["G1"] = "QLIKE"
        for start in (1, 7):
            for offset, label in enumerate(("Model", "Mean", "Q1", "Median", "MCS(#)")):
                sheet.cell(2, start + offset, label)
        for row, (key, label) in enumerate(MODEL_LABELS, start=3):
            for start, metric in ((1, "MSE"), (7, "QLIKE")):
                sheet.cell(row, start, label)
                for offset, statistic in enumerate(("Mean", "Q1", "Q2"), start=1):
                    value = float(summaries[metric].loc[key, statistic])
                    if not np.isfinite(value):
                        raise ValueError(f"Cannot export a missing {metric} {statistic} for {key}")
                    sheet.cell(row, start + offset, value).number_format = "0.000"
                sheet.cell(row, start + 4, int(counts[metric][key])).number_format = "0"
        thin = Side(style="thin", color="000000")
        last_row = 2 + len(MODEL_LABELS)
        for row in sheet:
            for cell in row:
                cell.font = Font(name="Times New Roman", size=11, bold=False)
                cell.alignment = Alignment(horizontal="left" if cell.column in (1, 7) and cell.row > 1 else "center", vertical="center")
                if cell.column != 6:
                    if cell.row == 1:
                        cell.border = Border(top=thin, bottom=thin)
                    elif cell.row in (2, last_row):
                        cell.border = Border(bottom=thin)
        for start, metric in ((1, "MSE"), (7, "QLIKE")):
            for row, (key, _) in enumerate(MODEL_LABELS, start=3):
                for offset, stat in enumerate(("Mean", "Q1", "Q2", "MCS(#)"), start=1):
                    bold = bold_mode == "har-ci" and key == "HAR-CI"
                    if bold_mode == "best":
                        bold = counts[metric][key] == max(counts[metric].values()) if offset == 4 else np.isclose(summaries[metric].loc[key, stat], summaries[metric][stat].min(), rtol=0, atol=1e-12)
                    if bold:
                        sheet.cell(row, start + offset).font = Font(name="Times New Roman", size=11, bold=True)
        for column in ("A", "G"):
            sheet.column_dimensions[column].width = 18
        for column in ("B", "C", "D", "E", "H", "I", "J", "K"):
            sheet.column_dimensions[column].width = 12
        sheet.column_dimensions["F"].width = 3
        sheet.freeze_panes = "B3"
        sheet.sheet_view.showGridLines = False
        sheet.print_options.horizontalCentered = True
        sheet.page_setup.orientation = "landscape"
        sheet.page_setup.paperSize = sheet.PAPERSIZE_A4
        sheet.page_setup.fitToWidth = 1
        sheet.page_setup.fitToHeight = 1
        sheet.sheet_properties.pageSetUpPr.fitToPage = True
        sheet.print_area = f"A1:K{last_row}"
        output_path = Path(output_path)
        output_path.parent.mkdir(parents=True, exist_ok=True)
        workbook.save(output_path)

    def check_inputs():
        for path in (RV_FILE, RETURN_FILE, CENTRALITY_FILE):
            if not path.is_file():
                raise FileNotFoundError(f"Required input not found: {path}")
        markets = [str(c).strip().lstrip(".") for c in pd.read_csv(RV_FILE, nrows=0).columns if c != "date"]
        if len(markets) != 18 or len(set(markets)) != 18:
            raise ValueError("Table 6 requires 18 distinct markets in RV.csv")
        centrality = pd.read_csv(CENTRALITY_FILE, usecols=["market", "L_net"])
        centrality["market"] = centrality["market"].astype(str).str.strip().str.lstrip(".")
        windows = set(pd.to_numeric(centrality["L_net"], errors="coerce").dropna().astype(int))
        if 1000 not in windows:
            raise ValueError("P.csv does not contain the 1000-day network window required by Table 6")
        if set(centrality["market"].dropna()) != set(markets):
            raise ValueError("P.csv markets do not match RV.csv")
        return set(markets)

    def validate_model_output(label, result, markets):
        if set(result) != markets:
            raise RuntimeError(f"{label}: missing or unexpected markets: {sorted(markets.symmetric_difference(result))}")
        for market, frame in result.items():
            if frame.empty or not {"RV_True", "RV_Predict", "MSE"}.issubset(frame.columns):
                raise RuntimeError(f"{label}: no usable forecast output for {market}")

    return SimpleNamespace(**locals())


def validate_forecasts(label, result, markets):
    if set(result) != set(markets):
        raise RuntimeError(f"{label}: forecast markets do not match the inputs")
    for market, frame in result.items():
        if frame.empty or not {"RV_True", "RV_Predict", "MSE"}.issubset(frame.columns):
            raise RuntimeError(f"{label}: no usable forecast output for {market}")


def compute_shared_drop_dates(table4, predictions4, predictions5, predictions6, layer_predictions):
    markets = sorted(predictions4["HAR-RV"])
    shared = {market: set() for market in markets}
    sources = list(predictions4.items())
    sources += [(key, result) for key, result in predictions5.items() if key != "HAR-RV"]
    sources += [(key, predictions6[key]) for key in ("bagging", "gradient", "RF")]
    sources += list(layer_predictions.items())
    for label, results in sources:
        validate_forecasts(label, results, markets)
        for market, frame in results.items():
            y_true, y_pred = table4.load_prediction_pair(frame)
            losses = table4.qlike_series_with_rules(y_true.to_numpy(), y_pred.to_numpy())
            shared[market].update(y_true.index[~np.isfinite(losses)].normalize())
    return shared


def main():
    table4, table5, table6 = build_table_04(), build_table_05(), build_table_06()
    markets = table4.check_inputs()
    if table5.check_inputs() != markets or table6.check_inputs() != markets:
        raise ValueError("The three tables must use the same markets")
    predictions4, predictions5, predictions6, layer_predictions = {}, {}, {}, {}
    for table, runners, target in (
        (table4, table4.MODEL_RUNNERS, predictions4),
        (table5, table5.AUXILIARY_RUNNERS + table5.MODEL_RUNNERS, predictions5),
        (table6, table6.STANDARD_RUNNERS, predictions6),
    ):
        for key, label, runner in tqdm(runners):
            result = runner()
            validate_forecasts(label, result, markets)
            target[key] = result
    deep_outputs = table6.run_deep_learning()
    for key in ("CNN", "LSTM", "Transformer"):
        validate_forecasts(key, deep_outputs[key], markets)
        predictions6[key] = deep_outputs[key]
    for key, runner in tqdm(table4.LAYER_RUNNERS):
        result = runner()
        validate_forecasts(key, result, markets)
        layer_predictions[key] = result
    shared_drop_dates = compute_shared_drop_dates(table4, predictions4, predictions5, predictions6, layer_predictions)
    for table, predictions in ((table4, predictions4), (table5, predictions5), (table6, predictions6)):
        summaries, _ = table.summarize_losses(predictions, shared_drop_dates)
        counts = table.count_mcs_memberships(predictions, shared_drop_dates)
        table.save_table(summaries, counts)


if __name__ == "__main__":
    main()

