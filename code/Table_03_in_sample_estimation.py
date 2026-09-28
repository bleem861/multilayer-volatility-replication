from copy import copy
from pathlib import Path
import argparse
import os
import numpy as np
import pandas as pd
from openpyxl.styles import PatternFill
from openpyxl.utils import get_column_letter
import statsmodels.formula.api as smf

PROJECT_ROOT = Path(__file__).resolve().parents[1]

RET_PATH = PROJECT_ROOT / 'data' / 'raw' / 'Return.csv'
RV_PATH = PROJECT_ROOT / 'data' / 'raw' / 'RV.csv'
parser = argparse.ArgumentParser()
parser.add_argument('--recomputed-centrality', action='store_true', help='Use recomputed centrality indices.')
args = parser.parse_args()
P_STORE_PATH = (
    PROJECT_ROOT / 'results' / 'generated' / 'centrality' / 'P.csv'
    if args.recomputed_centrality
    else PROJECT_ROOT / 'data' / 'processed' / 'centrality' / 'P.csv'
)

H = 22

IN_START = pd.to_datetime('2014-07-03')
IN_END   = pd.to_datetime('2018-05-01')

L_net = 1000

BASE_OUT_ROOT = PROJECT_ROOT / 'results' / 'generated'
OUT_DIR = PROJECT_ROOT / 'results' / 'generated'
os.makedirs(OUT_DIR, exist_ok=True)

def _norm_date(x):
    return pd.to_datetime(x).tz_localize(None).normalize()

def make_future_mean(rv_now: pd.Series, H: int) -> pd.Series:
    ma = rv_now.rolling(window=H, center=False).mean()
    return ma.shift(-(H - 1))

def star_from_p(pval: float) -> str:
    if pd.isna(pval):
        return ""
    if pval <= 0.01:
        return "***"
    if pval <= 0.05:
        return "**"
    if pval <= 0.10:
        return "*"
    return ""

def fmt_coef(coef: float, pval: float, ndigits: int = 3) -> str:
    if pd.isna(coef):
        return ""
    return f"{coef:.{ndigits}f}{star_from_p(pval)}"

def read_rv_scale_10000(path):
    data = pd.read_csv(path)
    if 'date' not in data.columns:
        raise RuntimeError("RV file must contain a 'date' column.")
    num_cols = [c for c in data.columns if c != 'date']

    data[num_cols] = data[num_cols].apply(pd.to_numeric, errors='coerce')

    data[num_cols] = data[num_cols] * 10000

    data = data.dropna(how='all')
    data['date'] = pd.to_datetime(data['date'], utc=True, errors='coerce')
    data = data.dropna(subset=['date'])
    data['date'] = data['date'].dt.tz_localize(None).dt.normalize()
    data = data.set_index('date')

    data = data.dropna(axis=1, how='all')

    data.columns = [str(c).replace('.', '') for c in data.columns]

    return data

data_rv = read_rv_scale_10000(RV_PATH)

markets = [c for c in data_rv.columns]

P_all = pd.read_csv(P_STORE_PATH)

need_cols = ['date', 'market', 'L_net', 'P_corr', 'P_kendall', 'P_partial', 'P_tail']
miss = [c for c in need_cols if c not in P_all.columns]
if miss:
    raise RuntimeError(f"P.csv is missing required columns: {miss}. Available columns: {list(P_all.columns)}")

P_all['date'] = pd.to_datetime(P_all['date'], errors='coerce').dt.tz_localize(None).dt.normalize()

P_all['market'] = P_all['market'].astype(str).str.replace('.', '', regex=False).str.strip()

P_all['L_net'] = pd.to_numeric(P_all['L_net'], errors='coerce').astype('Int64')
P_all = P_all.dropna(subset=['date', 'market', 'L_net']).copy()

def load_P_by_lnet_and_date_range(L_net: int, start_date: pd.Timestamp, end_date: pd.Timestamp):
    start_date = _norm_date(start_date)
    end_date   = _norm_date(end_date)

    sub = P_all[
        (P_all['L_net'] == int(L_net)) &
        (P_all['date'] >= start_date) &
        (P_all['date'] <= end_date)
    ].copy()

    if sub.empty:
        return None, None, None, None

    def _pivot(col):
        df = sub.pivot_table(index='date', columns='market', values=col, aggfunc='mean')
        return df.sort_index()

    return _pivot('P_corr'), _pivot('P_kendall'), _pivot('P_partial'), _pivot('P_tail')

P_corr, P_kendall, P_partial, P_tail = load_P_by_lnet_and_date_range(L_net, IN_START, IN_END)
if P_corr is None:
    raise RuntimeError(
        f"No centrality data found for L_net={L_net} between {IN_START.date()} and {IN_END.date()}."
    )

rv_set = set(data_rv.columns)
p_set = set(P_corr.columns)

def fit_in_sample_one_market_H22(market: str):
    if market not in data_rv.columns:
        return None
    if market not in P_corr.columns:

        return None
    if market not in P_kendall.columns or market not in P_partial.columns or market not in P_tail.columns:
        return None

    df = data_rv[market].to_frame('RV_now')

    df['RV_D'] = df['RV_now'].shift(1)
    df['RV_W'] = df['RV_D'].rolling(5).mean()
    df['RV_M'] = df['RV_D'].rolling(22).mean()

    df['Y'] = make_future_mean(df['RV_now'], H)

    ci = pd.DataFrame({
        'CI_corr':    P_corr[market],
        'CI_kendall': P_kendall[market],
        'CI_partial': P_partial[market],
        'CI_tail':    P_tail[market],
    }).reindex(df.index)

    df = df.join(ci, how='inner')

    df['CI_corr_D']    = df['CI_corr'].shift(1)
    df['CI_kendall_D'] = df['CI_kendall'].shift(1)
    df['CI_partial_D'] = df['CI_partial'].shift(1)
    df['CI_tail_D']    = df['CI_tail'].shift(1)

    need_common = [
        'Y', 'RV_D', 'RV_W', 'RV_M',
        'CI_corr_D', 'CI_kendall_D', 'CI_partial_D', 'CI_tail_D'
    ]
    df = df.dropna(subset=need_common)

    df_in = df.loc[IN_START:IN_END].copy()
    if len(df_in) < 30:
        return None

    df_in['_Y_'] = df_in['Y'].astype(float).values

    res = smf.ols(
        '_Y_ ~ RV_D + RV_W + RV_M + CI_corr_D + CI_kendall_D + CI_partial_D + CI_tail_D',
        df_in
    ).fit()

    res_drop_corr = smf.ols(
        '_Y_ ~ RV_D + RV_W + RV_M + CI_kendall_D + CI_partial_D + CI_tail_D',
        df_in
    ).fit()

    res_drop_kendall = smf.ols(
        '_Y_ ~ RV_D + RV_W + RV_M + CI_corr_D + CI_partial_D + CI_tail_D',
        df_in
    ).fit()

    res_drop_partial = smf.ols(
        '_Y_ ~ RV_D + RV_W + RV_M + CI_corr_D + CI_kendall_D + CI_tail_D',
        df_in
    ).fit()

    res_drop_tail = smf.ols(
        '_Y_ ~ RV_D + RV_W + RV_M + CI_corr_D + CI_kendall_D + CI_partial_D',
        df_in
    ).fit()

    out = {
        'beta_corr':    (res.params.get('CI_corr_D', np.nan),    res.pvalues.get('CI_corr_D', np.nan)),
        'beta_kendall': (res.params.get('CI_kendall_D', np.nan), res.pvalues.get('CI_kendall_D', np.nan)),
        'beta_partial': (res.params.get('CI_partial_D', np.nan), res.pvalues.get('CI_partial_D', np.nan)),
        'beta_tail':    (res.params.get('CI_tail_D', np.nan),    res.pvalues.get('CI_tail_D', np.nan)),

        'R2': res.rsquared,

        'R2_drop_corr':    res_drop_corr.rsquared,
        'R2_drop_kendall': res_drop_kendall.rsquared,
        'R2_drop_partial': res_drop_partial.rsquared,
        'R2_drop_tail':    res_drop_tail.rsquared,
    }
    return out

coef_table = pd.DataFrame(index=[
    'beta_corr', 'beta_kendall', 'beta_partial', 'beta_tail',
    'R2',
    'R2_drop_corr', 'R2_drop_kendall', 'R2_drop_partial', 'R2_drop_tail'
])

kept_markets = []
for mkt in markets:
    ins = fit_in_sample_one_market_H22(mkt)
    if ins is None:
        continue

    kept_markets.append(mkt)

    coef_table.loc['beta_corr',    mkt] = fmt_coef(ins['beta_corr'][0],    ins['beta_corr'][1])
    coef_table.loc['beta_kendall', mkt] = fmt_coef(ins['beta_kendall'][0], ins['beta_kendall'][1])
    coef_table.loc['beta_partial', mkt] = fmt_coef(ins['beta_partial'][0], ins['beta_partial'][1])
    coef_table.loc['beta_tail',    mkt] = fmt_coef(ins['beta_tail'][0],    ins['beta_tail'][1])

    coef_table.loc['R2', mkt] = f"{ins['R2']:.3f}"

    coef_table.loc['R2_drop_corr',    mkt] = f"{ins['R2_drop_corr']:.3f}"
    coef_table.loc['R2_drop_kendall', mkt] = f"{ins['R2_drop_kendall']:.3f}"
    coef_table.loc['R2_drop_partial', mkt] = f"{ins['R2_drop_partial']:.3f}"
    coef_table.loc['R2_drop_tail',    mkt] = f"{ins['R2_drop_tail']:.3f}"

coef_table = coef_table.T

coef_xlsx = os.path.join(OUT_DIR, "Table_03_in_sample_estimation.xlsx")
r2_columns = ['R2_drop_corr', 'R2_drop_kendall', 'R2_drop_partial', 'R2_drop_tail']
excel_table = coef_table.copy()
excel_table[['R2'] + r2_columns] = excel_table[['R2'] + r2_columns].apply(pd.to_numeric)
with pd.ExcelWriter(coef_xlsx, engine='openpyxl') as writer:
    excel_table.to_excel(writer, sheet_name='Table 3', index_label='Market')
    worksheet = writer.sheets['Table 3']
    for cell in list(worksheet[1]) + list(worksheet['A']):
        font = copy(cell.font)
        font.bold = False
        cell.font = font
    gray_fill = PatternFill(fill_type='solid', fgColor='D9D9D9')
    for row_number, (_, row) in enumerate(excel_table.iterrows(), start=2):
        minimum = row[r2_columns].min()
        for column in ['R2'] + r2_columns:
            column_number = excel_table.columns.get_loc(column) + 2
            cell = worksheet.cell(row=row_number, column=column_number)
            cell.number_format = '0.000'
            if column in r2_columns and pd.notna(row[column]) and row[column] == minimum:
                cell.fill = gray_fill
    worksheet.column_dimensions['A'].width = 14
    for column_number, column in enumerate(excel_table.columns, start=2):
        worksheet.column_dimensions[get_column_letter(column_number)].width = max(16, len(column) + 2)
